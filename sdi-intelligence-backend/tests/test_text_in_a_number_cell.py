"""Text in a cell typed as a number does not stop the AM CRM reading the tracker.

14 October 2026, the 07:50 morning review:
    ValueError: could not convert string to float: 'Completed'
Nick's tracker had been rewritten outside Excel, leaving text in cells typed as
numbers. Excel Online refuses that file (FileCorruptTryRepair) and openpyxl
raised on the first such cell, so the morning email, the spoken briefing and
the read-only copy all failed together.
"""
import io
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

import morning_review  # noqa: E402
import repair_tracker  # noqa: E402
import voicecrm_excel  # noqa: E402

HEADERS = ["Project code", "ACCOUNT", "OVERVIEW / DESCRIPTION / DELIVERABLES", "AM Owner",
           "Status", "NEXT STEPS", "Commercial Status", "BUDGET COST"]
KEY = ["AM Owner", "Status", "NEXT STEPS", "Commercial Status", "BUDGET COST"]


def _broken_tracker() -> bytes:
    """A tracker as the outside tool left it: "Completed" in a cell typed as a number."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ACCOUNT PROJECT TRACKER"
    ws.append(["NG tracker 2026"])
    ws.append([])
    ws.append(HEADERS)
    ws.append([None, "TPS", "H2 forecast", "NG", 999999, "chase", "Ordered", 15000])
    ws.append([None, "Tesco", "SDA tables & ends", "NG", "In production", "", "Ordered", 30359.96])
    buf = io.BytesIO()
    wb.save(buf)
    src, out = zipfile.ZipFile(io.BytesIO(buf.getvalue())), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            raw = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                assert b"<v>999999</v>" in raw
                raw = raw.replace(b"<v>999999</v>", b"<v>Completed</v>")
            dst.writestr(info, raw)
    return out.getvalue()


def _read_with_stock_openpyxl(data: bytes, tmp_path: Path) -> subprocess.CompletedProcess:
    """In a fresh process, so nothing this suite imported has taught openpyxl anything."""
    path = tmp_path / "book.xlsx"
    path.write_bytes(data)
    code = ("import openpyxl,sys; wb=openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True); "
            "print([c for r in wb.active.iter_rows(values_only=True) for c in r if c])")
    return subprocess.run([sys.executable, "-c", code, str(path)], capture_output=True, text=True)


def test_the_file_reproduces_this_mornings_failure(tmp_path):
    done = _read_with_stock_openpyxl(_broken_tracker(), tmp_path)
    assert done.returncode != 0
    assert "could not convert string to float: 'Completed'" in done.stderr


def test_the_morning_review_reads_it(tmp_path):
    rows, headers, sheet = morning_review.rows_from_xlsx(_broken_tracker(), KEY)
    assert sheet == "ACCOUNT PROJECT TRACKER" and len(rows) == 2
    assert rows[0]["Status"] == "Completed" and rows[0]["BUDGET COST"] == 15000
    assert rows[1]["Status"] == "In production" and rows[1]["BUDGET COST"] == 30359.96


def test_the_read_only_copy_reads_it():
    store = voicecrm_excel.ExcelStore.__new__(voicecrm_excel.ExcelStore)
    store.key_columns, store.sheet, store._chosen_sheet = KEY, "", ""
    store.download = lambda token: (_broken_tracker(), None)
    refused = {"state": "graph_error", "status": 501, "detail": "FileCorruptTryRepair"}
    out = store._rows_from_copy("token", refused)
    assert out["state"] == "ok" and out["source"] == "copy", out
    statuses = [item["fields"]["Status"] for item in out["items"]]
    assert statuses == ["Completed", "In production"], statuses


def test_the_repair_keeps_every_value_and_reads_cleanly_anywhere(tmp_path):
    broken = _broken_tracker()
    fixed, report = repair_tracker.repair(broken)
    assert report == [("ACCOUNT PROJECT TRACKER", "E4", "Completed")]
    done = _read_with_stock_openpyxl(fixed, tmp_path)
    assert done.returncode == 0, done.stderr
    assert "'Completed'" in done.stdout and "15000" in done.stdout and "30359.96" in done.stdout
    # Every other part of the file is byte for byte what it was.
    before, after = zipfile.ZipFile(io.BytesIO(broken)), zipfile.ZipFile(io.BytesIO(fixed))
    assert before.namelist() == after.namelist()
    for name in before.namelist():
        if name != "xl/worksheets/sheet1.xml":
            assert before.read(name) == after.read(name), name


def test_a_sound_file_is_left_exactly_as_it_is():
    wb = openpyxl.Workbook()
    wb.active.append(["Status", 12, 3.5, -1e3, "Completed"])
    buf = io.BytesIO()
    wb.save(buf)
    fixed, report = repair_tracker.repair(buf.getvalue())
    assert report == [] and fixed == buf.getvalue()


@pytest.mark.parametrize("cell, expect", [
    ('<c r="A1" t="n"><v>Completed</v></c>',
     '<c r="A1" t="inlineStr"><is><t xml:space="preserve">Completed</t></is></c>'),
    ('<c r="B2" s="3"><v>M&amp;S</v></c>',
     '<c r="B2" s="3" t="inlineStr"><is><t xml:space="preserve">M&amp;S</t></is></c>'),
    ('<c r="C3"><f>A1</f><v>Completed</v></c>', '<c r="C3" t="str"><f>A1</f><v>Completed</v></c>'),
    ('<c r="D4" t="s"><v>7</v></c>', None),
    ('<c r="E5"><v>1.5E3</v></c>', None),
    ('<c r="F6" s="2"/>', None),
])
def test_each_kind_of_cell(cell, expect):
    xml = f"<sheetData><row r=\"1\">{cell}</row></sheetData>"
    fixed, changed = repair_tracker.fix_sheet_xml(xml)
    assert fixed == (xml if expect is None else xml.replace(cell, expect))
    assert bool(changed) == (expect is not None)
