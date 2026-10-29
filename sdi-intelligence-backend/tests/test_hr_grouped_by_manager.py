"""Simon's lists grouped by BrightHR manager: who reports to whom, end to end.

BrightHR's field name was not in any documentation to hand, so the usual
spellings are tried and HR_MANAGER_FIELD can name the real one. Fictional
people throughout.
"""
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import hr_brighthr as bh                # noqa: E402
import hr_config as cfg                 # noqa: E402
import hr_rollcall                      # noqa: E402


@pytest.mark.parametrize("raw, expected", [
    ({"lineManager": {"id": "BH-7", "name": {"givenName": "Max", "familyName": "Boss"}}},
     ("BH-7", "Max Boss", "lineManager")),
    ({"manager": {"firstName": "Max", "lastName": "Boss"}}, ("", "Max Boss", "manager")),
    ({"managerId": "BH-7"}, ("BH-7", "", "managerId")),
    ({"approvers": [{}, {"id": "BH-8", "displayName": "Ann Approver"}]},
     ("BH-8", "Ann Approver", "approvers")),
    ({"reportsTo": "Max Boss"}, ("", "Max Boss", "reportsTo")),
    ({"email": "x@wearesdi.com"}, ("", "", "")),
])
def test_the_usual_shapes_of_a_manager_field_are_read(raw, expected, monkeypatch):
    monkeypatch.setattr(cfg, "HR_MANAGER_FIELD", "")
    got = bh.manager_of(raw)
    assert (got["manager_id"], got["manager_name"], got["manager_field"]) == expected


def test_the_field_can_be_named_once_the_probe_has_shown_it(monkeypatch):
    raw = {"manager": {"id": "BH-1"}, "holidayApprover": {"id": "BH-2"}}
    monkeypatch.setattr(cfg, "HR_MANAGER_FIELD", "holidayApprover.id")
    assert bh.manager_of(raw)["manager_id"] == "BH-2"


def test_a_manager_given_only_by_id_gets_their_name_from_the_staff_list(monkeypatch):
    monkeypatch.setattr(cfg, "HR_MANAGER_FIELD", "")
    staff = bh.name_managers([bh.normalise_employee(r) for r in (
        {"id": "BH-7", "firstName": "Max", "lastName": "Boss"},
        {"id": "BH-1", "firstName": "Jo", "lastName": "Office", "lineManagerId": "BH-7"},
        {"id": "BH-2", "firstName": "Sam", "lastName": "Floor", "lineManagerId": "BH-404"},
    )])
    by_id = {e["id"]: e["manager_name"] for e in staff}
    assert by_id == {"BH-7": "", "BH-1": "Max Boss", "BH-2": ""}


def test_a_check_carries_each_persons_manager_into_the_snapshot(tmp_path, monkeypatch):
    import hr_blip
    monkeypatch.setattr(cfg, "HR_SNAPSHOT_DIR", str(tmp_path))
    monkeypatch.setattr(cfg, "HR_OUTPUT_DIR", "")
    monkeypatch.setattr(cfg, "HR_MANAGER_FIELD", "")
    monkeypatch.setattr(hr_blip, "_get_token", lambda: "t")
    monkeypatch.setattr(hr_blip, "STAFF_LIST_PAUSE", 0)
    monkeypatch.setattr(bh, "fetch_employees", lambda token: [
        {"id": "BH-7", "firstName": "Max", "lastName": "Boss"},
        {"id": "BH-1", "firstName": "Jo", "lastName": "Office", "lineManager": {"id": "BH-7"}},
        {"id": "BH-2", "firstName": "Sam", "lastName": "Floor", "lineManager": {"id": "BH-7"}},
    ])
    monkeypatch.setattr(hr_blip, "_query_one", lambda token, eid, from_dt, throttle=None:
                        ([{"start": "2026-10-29T07:00:00Z"}], True) if eid == "BH-1" else ([], True))
    summary = hr_blip.run_blip()
    snap = json.loads((tmp_path / "blip_latest.json").read_text(encoding="utf-8"))
    assert [(p["first_name"], p["manager"]) for p in snap["on_site"]] == [("Jo", "Max Boss")]
    assert snap["managers"] == {"jo office": "Max Boss", "sam floor": "Max Boss"}
    assert summary["staff_with_manager"] == 2 and summary["manager_fields"] == ["lineManager"]


def test_the_roll_call_groups_reception_sign_ins_by_manager_too():
    """Someone InVentry has on site who is not clocked in on BrightHR still
    belongs to a manager: the snapshot's staff list says which."""
    snap = {"on_site": [{"first_name": "Jo", "surname": "Office", "manager": "Max Boss",
                         "clocking": {"start": "2026-10-29T07:00:00Z"}}],
            "managers": {"sam floor": "Ann Approver"}}
    import datetime
    people, _ = hr_rollcall.build([{"name": "Sam Floor", "since": "", "by": "reception"},
                                   {"name": "Vic Visitor", "since": "", "by": "reception"}], snap,
                                  now=datetime.datetime(2026, 10, 29, 9, tzinfo=datetime.timezone.utc),
                                  max_age_hours=20)
    assert {p["name"]: p["manager"] for p in people} == {
        "Jo Office": "Max Boss", "Sam Floor": "Ann Approver", "Vic Visitor": ""}


def test_the_page_has_the_button_and_remembers_the_choice():
    page = (Path(__file__).resolve().parent.parent / "appportal" / "hr-inventry.html").read_text(encoding="utf-8")
    assert "data-group" in page and "Group by manager" in page
    assert "No manager listed" in page
    assert "localStorage.setItem('hrGroupByManager'" in page


# ── the managers list HR keeps, because BrightHR does not hold managers ──────

import hr_managers                      # noqa: E402


def write_sheet(path, rows):
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(list(hr_managers.HEADERS))
    for r in rows:
        book.active.append(r)
    book.save(path)


@pytest.mark.parametrize("title, department", [
    ("Bench Joiner", "Joinery"), ("Sprayer/Bench Joiner", "Joinery"),
    ("Team Leader acrylics Machine", "Acrylics"), ("Afternoon Shift Paintline Operative", "Paintline"),
    ("Press Breke Operator", "Press brake, laser and folding"), ("Project manger", "Project management"),
    ("Quailty Supervisor", "Quality"), ("", "No job title in BrightHR"), ("Approval", "Other"),
])
def test_the_sheet_groups_by_a_department_guessed_from_the_job_title(title, department):
    assert hr_managers.department_of(title) == department


def test_a_person_is_matched_by_id_then_email_then_name(tmp_path):
    sheet = tmp_path / "managers.xlsx"
    write_sheet(sheet, [["Joinery", "Bench Joiner", "Jo Office", "", "Max Boss", "BH-1"],
                        ["Joinery", "Bench Joiner", "Old Name", "sam@wearesdi.com", "Ann Lead", ""],
                        ["Joinery", "Bench Joiner", "Kim Weld", "", "Ann Lead", ""],
                        ["Joinery", "Bench Joiner", "Not Filled", "", "", "BH-9"]])
    staff = hr_managers.apply([
        {"id": "BH-1", "first_name": "Jo", "surname": "Office", "email": ""},
        {"id": "BH-2", "first_name": "Sam", "surname": "Floor", "email": "SAM@wearesdi.com"},
        {"id": "BH-3", "first_name": "Kim", "surname": "Weld", "email": ""},
        {"id": "BH-9", "first_name": "Not", "surname": "Filled", "email": ""},
    ], path=sheet)
    assert [e.get("manager_name", "") for e in staff] == ["Max Boss", "Ann Lead", "Ann Lead", ""]
    assert hr_managers.LAST_PROBLEM == ""


def test_a_csv_works_too(tmp_path):
    sheet = tmp_path / "managers.csv"
    sheet.write_text("Name,Manager\nJo Office,Max Boss\n", encoding="utf-8")
    (e,) = hr_managers.apply([{"id": "x", "first_name": "Jo", "surname": "Office"}], path=sheet)
    assert e["manager_name"] == "Max Boss"


def test_a_missing_list_does_not_stop_a_check_and_says_why(tmp_path):
    staff = hr_managers.apply([{"id": "x", "first_name": "Jo", "surname": "Office"}],
                              path=tmp_path / "nowhere.xlsx")
    assert "not found" in hr_managers.LAST_PROBLEM and "manager_name" not in staff[0]


def test_the_sheet_tool_lists_current_staff_and_keeps_what_was_typed(tmp_path, monkeypatch):
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
    import hr_managers_sheet as tool
    monkeypatch.setattr(cfg, "HR_MANAGER_FIELD", "")
    sheet = tmp_path / "HR managers.xlsx"
    write_sheet(sheet, [["Joinery", "Bench Joiner", "Jo Office", "", "Max Boss", "BH-1"]])
    monkeypatch.setattr(bh, "get_access_token", lambda: "t")
    monkeypatch.setattr(bh, "fetch_employees", lambda token: [
        {"id": "BH-1", "name": {"givenName": "Jo", "familyName": "Office"},
         "employment": {"jobTitle": "Bench Joiner"}, "_metadata": {"isTerminated": False}},
        {"id": "BH-2", "name": {"givenName": "New", "familyName": "Starter"},
         "employment": {"jobTitle": "Welder"}, "_metadata": {"isTerminated": False}},
        {"id": "BH-3", "name": {"givenName": "Lee", "familyName": "Leaver"},
         "employment": {"jobTitle": "Welder"}, "_metadata": {"isTerminated": True}},
    ])
    monkeypatch.setattr(sys, "argv", ["x", str(sheet)])
    assert tool.main() == 0
    from openpyxl import load_workbook
    rows = [[c or "" for c in r] for r in load_workbook(sheet).active.iter_rows(values_only=True)]
    assert rows[0] == list(hr_managers.HEADERS)
    assert sorted((r[2], r[4], r[0]) for r in rows[1:]) == [
        ("Jo Office", "Max Boss", "Joinery"), ("New Starter", "", "Welding and fabrication")]
