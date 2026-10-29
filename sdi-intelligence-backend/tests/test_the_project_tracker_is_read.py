"""Liam's project tracker: any layout, read as the person, drawings opened.

The tracker's columns were not seen when this was written, so nothing assumes
them: a heading row is found, every row kept as heading -> value, and any path
offered as a file. Fictional rows throughout.
"""
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import project_tracker as pt            # noqa: E402


def workbook():
    from openpyxl import Workbook
    book = Workbook()
    ws = book.active
    ws.title = "Items"
    ws.append(["MILWAUKEE AMSTERDAM CONFERENCE 2027 - PC 402426"])           # a title, not headings
    ws.append([])
    ws.append(["Item", "Status", "Due", "Drawing"])
    ws.append(["Header unit", "In production", datetime(2027, 1, 12), r"W:\Design\Customers April 08 Onwards\TTi\402426\GA-01.pdf"])
    ws.append(["Plinth", "Awaiting approval", None, "W:/Prototype/TTi/402426 Plinth"])
    ws.append(["Signage", "Done", None, "see drawing"])
    ws["D6"].hyperlink = r"file:///W:/Production/TTi/402426/Signage.pdf"
    hidden = book.create_sheet("Old")
    hidden.sheet_state = "hidden"
    hidden.append(["x", "y"])
    data = BytesIO()
    book.save(data)
    return data.getvalue()


def test_the_heading_row_is_found_below_a_title():
    (sheet,) = pt.parse(workbook(), "TRACKER.xlsx")                         # hidden sheet left out
    assert sheet["headings"] == ["Item", "Status", "Due", "Drawing"]
    assert [r["cells"]["Item"] for r in sheet["rows"]] == ["Header unit", "Plinth", "Signage"]
    assert sheet["rows"][0]["cells"]["Due"] == "2027-01-12" and sheet["rows"][0]["row"] == 4


def test_drawing_paths_are_found_in_cells_and_behind_links():
    (sheet,) = pt.parse(workbook())
    paths = [[p["path"] for p in r["paths"]] for r in sheet["rows"]]
    assert paths == [[r"W:\Design\Customers April 08 Onwards\TTi\402426\GA-01.pdf"],
                     [r"W:\Prototype\TTi\402426 Plinth"],
                     [r"W:\Production\TTi\402426\Signage.pdf"]]
    assert [r["paths"][0]["is_file"] for r in sheet["rows"]] == [True, False, True]


def test_search_finds_rows_with_every_word_and_says_where():
    read_ = [{"label": "Amsterdam", "files": [{"name": "TRACKER.xlsx", "sheets": pt.parse(workbook())}]}]
    answer = pt.search(read_, "awaiting plinth")
    assert answer["match_count"] == 1
    assert answer["matches"][0]["source"] == "TRACKER.xlsx › Items › row 5"
    assert pt.search(read_, "")["match_count"] == 3


class Graph:
    """SharePoint as the person: a site, a search by name, a download."""
    def __init__(self, status=200):
        self.status, self.calls = status, []

    def get(self, url, headers=None, params=None, follow_redirects=False):
        self.calls.append(url)
        import httpx
        if self.status != 200:
            return httpx.Response(self.status, request=httpx.Request("GET", url))
        if url.endswith("/sites/sdidisplays.sharepoint.com:/sites/TTiProjects2022"):
            return httpx.Response(200, json={"id": "SITE"}, request=httpx.Request("GET", url))
        if "/search(" in url:
            return httpx.Response(200, request=httpx.Request("GET", url), json={"value": [
                {"id": "F1", "name": "TRACKER - MILWAUKEE AMSTERDAM CONFERENCE 2027.xlsx", "file": {},
                 "size": 100, "eTag": "E1", "lastModifiedDateTime": "2026-10-29T10:00:00Z",
                 "webUrl": "https://example.test/tracker"},
                {"id": "F2", "name": "TRACKER - MILWAUKEE AMSTERDAM CONFERENCE 2027 old.xlsx", "file": {}}]})
        if url.endswith("/items/F1/content"):
            return httpx.Response(200, content=workbook(), request=httpx.Request("GET", url))
        return httpx.Response(404, request=httpx.Request("GET", url))


@pytest.fixture(autouse=True)
def fresh(monkeypatch):
    pt._sites.clear()
    pt._parsed.clear()
    monkeypatch.delenv("SDI_PROJECT_TRACKERS", raising=False)
    monkeypatch.delenv("SDI_PROJECT_TRACKER_SITE", raising=False)


def test_the_tracker_is_found_by_its_name_on_the_site_and_read_as_the_person():
    graph = Graph()
    (t,) = pt.read_all("token-for-liam", client=graph)
    assert t["label"] == "Amsterdam Conference 2027"
    (f,) = t["files"]                                                        # the exact name only
    assert f["name"] == "TRACKER - MILWAUKEE AMSTERDAM CONFERENCE 2027.xlsx"
    assert f["sheets"][0]["rows"][0]["cells"]["Item"] == "Header unit"
    pt.read_all("token-for-liam", client=graph)                              # unchanged eTag: no second download
    assert sum(u.endswith("/content") for u in graph.calls) == 1


def test_someone_who_cannot_open_it_is_told_it_is_their_access():
    for status in (403, 404):
        pt._sites.clear()
        with pytest.raises(pt.Unavailable) as refused:
            pt.read_all("token-for-sue", client=Graph(status=status))
        assert refused.value.reason == "access"


def test_a_folder_or_a_path_can_be_named_instead(monkeypatch):
    monkeypatch.setenv("SDI_PROJECT_TRACKERS", "A=Milwaukee 2021/Projects 2027|TRACKER.xlsx")
    assert pt.trackers() == [{"label": "A", "what": "Milwaukee 2021/Projects 2027"},
                             {"label": "TRACKER.xlsx", "what": "TRACKER.xlsx"}]


# ── over HTTP, through the same TTI list ─────────────────────────────────────

from test_the_tti_views_end_to_end import SOMEBODY_ELSE, client  # noqa: E402,F401


def test_the_tab_is_behind_the_tti_list_and_reads_as_the_person(client, monkeypatch):
    seen = {}
    monkeypatch.setattr(pt, "read_all", lambda token: seen.setdefault("token", token) and [])
    assert client.get("/api/tti/despatch/projects").status_code == 200
    assert seen["token"] == "graph-token-for-sue"
    client.who["user"] = SOMEBODY_ELSE
    assert client.get("/api/tti/despatch/projects").status_code == 403


def test_no_access_to_the_tracker_is_marked_for_the_page(client, monkeypatch):
    def refuse(token):
        raise pt.Unavailable("No spreadsheet called … your account can open.", "access")
    monkeypatch.setattr(pt, "read_all", refuse)
    r = client.get("/api/tti/despatch/projects")
    assert r.status_code == 503 and r.headers["x-tti-reason"] == "access"
