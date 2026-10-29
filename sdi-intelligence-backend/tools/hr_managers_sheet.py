"""Write, or bring up to date, the managers list HR keeps. Reads BrightHR only.

    python tools/hr_managers_sheet.py              (to HR_MANAGERS_FILE)
    python tools/hr_managers_sheet.py "\\\\server\\share\\HR managers.xlsx"

Every current BrightHR employee gets a row, grouped by a department guessed
from their job title, with an empty Manager column for HR to fill in. Run it
again after new starters: anything already typed in Manager is kept, new
starters are added, and leavers are dropped. Close the file in Excel first.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import hr_brighthr as bh         # noqa: E402
import hr_config as cfg          # noqa: E402
import hr_managers               # noqa: E402


def main() -> int:
    target = Path(sys.argv[1] if len(sys.argv) > 1 else cfg.HR_MANAGERS_FILE or "")
    if not str(target) or str(target) == ".":
        print("Say where to write it: HR_MANAGERS_FILE in .env, or a path after the command.")
        return 2
    records = bh.fetch_employees(bh.get_access_token())
    kept = hr_managers.read(target) if target.is_file() else {"id": {}, "email": {}, "name": {}}
    rows = []
    for raw in records:
        e = bh.normalise_employee(raw)
        if e["terminated"] or not e["id"]:
            continue
        title = ((raw.get("employment") or {}).get("jobTitle") or "").strip()
        name = f"{e['first_name']} {e['surname']}".strip()
        manager = (kept["id"].get(hr_managers._key(e["id"])) or kept["email"].get(hr_managers._key(e["email"]))
                   or kept["name"].get(hr_managers._key(name)) or "")
        rows.append([hr_managers.department_of(title), title, name, e["email"], manager, e["id"]])
    rows.sort(key=lambda r: (r[0] == "No job title in BrightHR", r[0], r[1].lower(), r[2].lower()))

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = "Managers"
    sheet.append(list(hr_managers.HEADERS))
    for r in rows:
        sheet.append(r)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for cell in sheet["E"][1:]:
        cell.fill = PatternFill("solid", fgColor="FFF2CC")          # the column to fill in
    for col, width in zip("ABCDEF", (30, 32, 26, 34, 26, 38)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    notes = book.create_sheet("How to use")
    for line in ("Type each person's manager (or approver) in the yellow Manager column, as their name.",
                 "Filter by Department to fill in a whole team at once; it is only a guess from the job title.",
                 "HRInventry Anywhere groups people by this column after the next 'Check who's clocked in'.",
                 "Keep the Name, Email and BrightHR id columns as they are: they are how people are matched.",
                 "New starters: ask James to re-run the sheet. What is already typed is kept."):
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.stem + ".writing.xlsx")
    book.save(tmp)
    tmp.replace(target)
    filled = sum(1 for r in rows if r[4])
    print(f"Wrote {target}: {len(rows)} current staff, {filled} with a manager already, "
          f"{len(rows) - filled} to fill in.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
