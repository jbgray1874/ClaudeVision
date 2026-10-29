"""Who manages whom, from a list HR keeps - because BrightHR does not say.

Simon Foister asked for HRInventry Anywhere's lists grouped by each person's
manager or approver. BrightHR's Customer API was checked on 29 October 2026
(tools/probe_brighthr_fields.py): its employee records carry name, email,
job title, payroll number, start and end dates and a reference, and no
manager; Blip clockings carry no approver either. Job titles are no stand-in:
102 of them for 203 people, 27 blank, the same job spelt three ways.

So the list is a spreadsheet HR keeps: everyone's name and their manager.
tools/hr_managers_sheet.py writes it, pre-filled with every current employee
and a department guessed from their job title so it fills in quickly, and
adds new starters each time it is run again without losing what is typed.
Every check reads it (HR_MANAGERS_FILE), so a change shows on the next check.

It only reads the sheet. A row is matched by BrightHR id, then email, then
name; a person not in it, or with no manager typed, is "No manager listed".
"""
import csv
import re
from pathlib import Path

import hr_config as cfg

LAST_PROBLEM = ""          # why the list could not be read, for the check's notes

HEADERS = ("Department (guessed from job title)", "Job title", "Name", "Email",
           "Manager", "BrightHR id")

# Only to group the sheet so whole teams can be filled in at once. It is a
# guess from a free-text job title, said so in the column heading, and nothing
# in the app groups by it.
_DEPARTMENTS = (
    ("Joinery", r"joiner"),
    ("Acrylics", r"acrylic"),
    ("Paintline", r"paint|powder|spray"),
    ("Welding and fabrication", r"weld|fabricat|fettler"),
    ("Press brake, laser and folding", r"press|brake|breke|fold|laser|punch"),
    ("Assembly and packing", r"assembl|packing"),
    ("Prototype and toolroom", r"prototype|toolroom"),
    ("Warehouse and logistics", r"warehouse|goods in|dispatch|despatch|forklift|logistics"),
    ("Customer Services", r"customer serv"),
    ("Finance", r"finance|financial|payroll|treasury|accountant|accounts"),
    ("Account management", r"account (manager|director)|senior account"),
    ("Project management", r"project (manag|mang|supervis)|vm project"),
    ("Design", r"design|creative|creavtive"),
    ("Estimating and buying", r"estimat|buyer|buying|purchas"),
    ("Quality", r"qua(l|il)ity|quailty"),
    ("Health and safety", r"health"),
    ("Installation", r"install"),
    ("Directors and operations", r"director|head of|production (manager|planner)|technical manager|operations"),
)


def department_of(job_title: str) -> str:
    title = " ".join(str(job_title or "").lower().split())
    if not title:
        return "No job title in BrightHR"
    for name, pattern in _DEPARTMENTS:
        if re.search(pattern, title):
            return name
    return "Other"


def _key(text) -> str:
    return " ".join(str(text or "").lower().split())


def _rows(path: Path) -> list:
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [list(r) for r in csv.reader(f)]
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [["" if v is None else str(v) for v in r]
                for r in book.worksheets[0].iter_rows(values_only=True)]
    finally:
        book.close()


def read(path=None) -> dict:
    """{"id": {...}, "email": {...}, "name": {...}} -> manager, from the sheet."""
    path = Path(path or cfg.HR_MANAGERS_FILE)
    rows = _rows(path)
    if not rows:
        return {"id": {}, "email": {}, "name": {}}
    head = [_key(h) for h in rows[0]]

    def col(*words):
        for i, h in enumerate(head):
            if any(w in h for w in words):
                return i
        return None
    c_mgr, c_id = col("manager", "approver"), col("brighthr id", "brighthr_id")
    c_mail, c_name = col("email"), next((i for i, h in enumerate(head) if h == "name"), None)
    if c_mgr is None:
        raise ValueError(f"{path.name} has no Manager column in its first row.")
    out = {"id": {}, "email": {}, "name": {}}
    for r in rows[1:]:
        cell = lambda i: (r[i] if i is not None and i < len(r) else "") or ""
        manager = " ".join(cell(c_mgr).split())
        if not manager:
            continue
        for kind, i in (("id", c_id), ("email", c_mail), ("name", c_name)):
            if cell(i).strip():
                out[kind][_key(cell(i))] = manager
    return out


def apply(staff: list, path=None) -> list:
    """Give each person their manager from the sheet. Never fails a check: a
    sheet that cannot be read leaves managers blank and says why."""
    global LAST_PROBLEM
    LAST_PROBLEM = ""
    where = path or cfg.HR_MANAGERS_FILE
    if not where:
        return staff
    try:
        listed = read(where)
    except FileNotFoundError:
        LAST_PROBLEM = f"The managers list was not found at {where}, so nobody is grouped by manager."
        return staff
    except Exception as exc:                                        # noqa: BLE001
        LAST_PROBLEM = f"The managers list at {where} could not be read ({exc})."
        return staff
    for e in staff:
        name = f"{e.get('first_name', '')} {e.get('surname', '')}"
        manager = (listed["id"].get(_key(e.get("id"))) or listed["email"].get(_key(e.get("email")))
                   or listed["name"].get(_key(name)))
        if manager:
            e["manager_name"], e["manager_field"] = manager, "managers list"
    return staff
