"""Read a saved TTI schedule workbook without modifying it or any share.

The deployment adapter must first obtain an authorised, version-pinned copy.
This loader supplies source coordinates and a SHA-256 snapshot identity.
"""

from __future__ import annotations

import hashlib
import io
from pathlib import Path

from openpyxl import load_workbook

from sue_despatch import SourceRow, _norm


REQUIRED = {"sdiorderno", "partcode", "lorryfromsdi"}
MAX_BYTES = 30 * 1024 * 1024
MAX_ROWS = 20000


def load_schedule_snapshot(path: str | Path) -> tuple[str, list[SourceRow]]:
    """Return (SHA-256, rows) from a trusted local .xlsx snapshot, never live K:."""
    source = Path(path)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError("An existing .xlsx snapshot is required")
    if source.stat().st_size > MAX_BYTES:
        raise ValueError("Workbook exceeds the review limit")
    return read_schedule_bytes(source.read_bytes())


def read_schedule_bytes(data: bytes) -> tuple[str, list[SourceRow]]:
    """The same reader, over bytes - so a workbook fetched from SharePoint and
    one saved on a share are parsed by one piece of code rather than two.

    Nothing is written to disk on the way through: the copy exists for the life
    of the request and no longer, which is the right lifetime for somebody
    else's commercial tracker.
    """
    if len(data) > MAX_BYTES:
        raise ValueError("Workbook exceeds the review limit")
    digest = hashlib.sha256(data).hexdigest()
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as problem:                                 # noqa: BLE001
        raise ValueError(f"This does not open as a workbook "
                         f"({type(problem).__name__})") from problem
    result: list[SourceRow] = []
    try:
        for sheet in book.worksheets:
            header = None
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_number > MAX_ROWS:
                    raise ValueError("Sheet exceeds the review row limit")
                if header is None:
                    candidate = [str(v).strip() if v is not None else "" for v in values]
                    if REQUIRED <= {_norm(v) for v in candidate}:
                        header = candidate
                    elif row_number >= 20:
                        break
                    continue
                cells = {name: value for name, value in zip(header, values) if name}
                if any(value not in (None, "") for value in cells.values()):
                    result.append(SourceRow(sheet.title, row_number, cells))
    finally:
        book.close()
    if not result:
        raise ValueError("No TTI schedule rows found in this snapshot")
    return digest, result
