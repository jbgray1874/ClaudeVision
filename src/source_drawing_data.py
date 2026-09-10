"""Everything the engine read out of every file, one row at a time, in the estimating folder.

WHY THIS EXISTS. "When I open up DXF files they contain only images." That is what the
viewer shows; it is not what the file holds. Every flat export in this corpus is vector
geometry — SDI's own 117620202M is 55 LINEs, 2 CIRCLEs and 2 ARCs across two layers, SLD-0
carrying the profile and BENDLINES carrying the folds. Its blank measures 1009.49 x 363.91
and its holes are Ø5.0, and not one of those figures is a printed dimension: the circle IS
its diameter. The GA export alongside it carries 64 MTEXT and 20 DIMENSION entities, and its
text includes a stated weight of 384g.

So the question that actually matters is not what the files contain. It is which of it we
READ, which part we attached it to, and whether it reached the price. Three different
failures hide in that gap and they look identical from outside:

    the file did not have it          -> nothing anyone can do
    we did not read it                -> a reader to fix
    we read it and dropped it later   -> a pipeline join to fix, and the worst of the three,
                                         because the evidence was in the building

This workbook makes those distinguishable. One row per fact, every file in the pack listed
one after another, with where it came from and where it ended up. It prices nothing and
decides nothing; it is the audit that says which of the three we are looking at, and it is
the thing to read before writing another rule on a guess.

Written to the estimating folder as source_drawing_data.xlsx, per job.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence

SHEETS = ("Files", "Pages", "DXF file vs engine", "Facts", "BOM rows",
          "Operations", "Not extracted")


def _text(value: Any, limit: int = 300) -> str:
    if value is None:
        return ""
    out = str(value)
    return out if len(out) <= limit else out[: limit - 1] + "…"


def _num(value: Any) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _parts(summary: Mapping[str, Any]) -> List[Mapping[str, Any]]:
    for holder in (summary.get("estimate_summary") or {}, summary):
        if isinstance(holder, Mapping):
            for key in ("canonical_part_estimates", "part_estimates"):
                rows = holder.get(key)
                if isinstance(rows, list) and rows:
                    return [r for r in rows if isinstance(r, Mapping)]
    writeup = (summary.get("manufacturing_writeup") or {}).get("parts")
    return [r for r in (writeup or []) if isinstance(r, Mapping)]


def _costed_fields(part: Mapping[str, Any]) -> Dict[str, Any]:
    """The fields that actually move money, so a fact can be marked used or not."""
    material = part.get("material_estimate") if isinstance(
        part.get("material_estimate"), Mapping) else {}
    return {
        "normalized_material": part.get("normalized_material"),
        "normalized_thickness_mm": part.get("normalized_thickness_mm"),
        "blank_length_mm": part.get("blank_length_mm") or material.get("blank_length_mm"),
        "blank_width_mm": part.get("blank_width_mm") or material.get("blank_width_mm"),
        "quantity": part.get("quantity"),
        "wire_gauge_mm": part.get("wire_gauge_mm"),
        "wire_length_mm": part.get("wire_length_mm"),
        "stated_weight_kg": part.get("stated_weight_kg"),
    }


def files_rows(summary: Mapping[str, Any]) -> List[Dict[str, Any]]:
    """Every source file the job saw, and what came out of it."""
    rows: List[Dict[str, Any]] = []
    seen: set = set()

    pages_by_file: Dict[str, int] = {}
    for page in (summary.get("pages") or []):
        if isinstance(page, Mapping):
            name = str(page.get("source_pdf_name") or page.get("source_file") or "").strip()
            if name:
                pages_by_file[name] = pages_by_file.get(name, 0) + 1
    for name, count in sorted(pages_by_file.items()):
        seen.add(name.lower())
        rows.append({"file": name, "kind": "PDF", "read": "yes",
                     "pages_or_entities": count, "yielded": f"{count} page(s) analysed"})

    # DXF and model files, wherever the readers recorded them on the parts.
    for part in _parts(summary):
        for key, kind in (("dxf_file", "DXF"), ("dxf_path", "DXF"),
                          ("solidworks_file", "SolidWorks"),
                          ("source_file", "source")):
            name = str(part.get(key) or "").strip()
            if not name or name.lower() in seen:
                continue
            seen.add(name.lower())
            rows.append({"file": Path(name).name, "kind": kind, "read": "yes",
                         "pages_or_entities": "",
                         "yielded": f"geometry for {part.get('part_number')}"})

    # Files the pack contained and NOTHING read — the ones worth the most attention.
    for issue in ((summary.get("document_analysis") or {}).get("pack_issues") or []):
        if not isinstance(issue, Mapping):
            continue
        for name in (issue.get("files") or []):
            base = Path(str(name)).name
            if base.lower() in seen:
                continue
            seen.add(base.lower())
            rows.append({"file": base, "kind": Path(base).suffix.upper().lstrip(".") or "?",
                         "read": "NO", "pages_or_entities": "",
                         "yielded": _text(issue.get("message") or "not read")})
    return rows


def _match_part(summary: Mapping[str, Any], dxf_name: str) -> tuple:
    """(part, how it was matched, ambiguous?) — the file-to-part association, not a guess.

    An earlier version substring-matched a part code into the filename and took the longest
    hit silently. Attribution IS the audit: a fact credited to the wrong part is worse than a
    fact nobody credited, because it reads as evidence. So the association the pipeline
    ITSELF recorded is used first, and where that is absent the filename fallback must be
    UNIQUE — two candidates are reported as ambiguous rather than resolved by length.

    THE FULL PATH IS MATCHED BEFORE THE BASENAME, AND A BASENAME-ONLY HIT SAYS SO. A pack
    holding revA/117620202M.dxf and revB/117620202M.dxf produced two rows that were identical
    in every visible column — same file name, same part — one agreeing with the engine and one
    a red NO. The pipeline had recorded revB; revA was a different file that merely shared a
    name, and nothing on the page said so. Pass the full path here (not probe["file"]) and an
    exact path match wins outright; a basename match against a DIFFERENT recorded path is
    still offered, because it is usually right, but it is marked ambiguous and names the path
    the pipeline actually used.

    AND AN EXACT PATH MATCH HAS TO BE UNIQUE TOO. Returning on the first hit made uniqueness an
    assumption rather than a check: where two part records named the same DXF the first won
    silently and the row was reported as unambiguous. That is not a contrived case at SDI — a
    handed pair is routinely cut from one flat export, so both hands reference the same file.
    Every claimant is collected before the decision, and more than one is reported as ambiguous
    with all of them named. It may well be that the fact is the same for both; that is a
    judgement for whoever reads the row, and the audit's job is to say that the choice exists.
    """
    target = str(dxf_name).strip().lower()
    target_name = Path(target).name
    exact: List[Mapping[str, Any]] = []
    basename_only: Optional[tuple] = None
    for part in _parts(summary):
        for key in ("dxf_file", "dxf_path", "flat_pattern_file"):
            recorded = str(part.get(key) or "").strip().lower()
            if not recorded:
                continue
            if recorded == target:
                if part not in exact:
                    exact.append(part)
                continue
            if Path(recorded).name == target_name and basename_only is None:
                # The path is compared case-folded but REPORTED as recorded: a lowercased
                # path is not the path anyone can go and look at.
                basename_only = (part, str(part.get(key)).strip())
    if len(exact) == 1:
        return exact[0], "the pipeline's own file-to-part association", False
    if len(exact) > 1:
        named = ", ".join(_text(p.get("part_number") or "(unnamed)") for p in exact[:4])
        more = f" and {len(exact) - 4} more" if len(exact) > 4 else ""
        return None, (f"ambiguous: {len(exact)} part records name this exact file "
                      f"({named}{more}), so no single part owns these figures"), True
    if basename_only is not None:
        part, recorded = basename_only
        if Path(target).parent == Path(""):
            # Only a bare name was supplied, so there is no path to disagree with.
            return part, "the pipeline's own file-to-part association", False
        return part, ("file name only: the pipeline recorded a DIFFERENT path for this part "
                      f"({recorded}), so figures here may come from another file of the same "
                      "name"), True
    stem = Path(dxf_name).stem.upper().replace("_", "-")
    hits = [p for p in _parts(summary)
            if str(p.get("part_number") or "").strip().upper().replace("_", "-")
            and str(p.get("part_number")).strip().upper().replace("_", "-") in stem]
    if len(hits) == 1:
        return hits[0], "the part code in the filename", False
    if len(hits) > 1:
        return None, "ambiguous: " + ", ".join(str(h.get("part_number")) for h in hits[:4]), True
    return None, "no part matched this file", False


def _display_names(paths: Sequence[Any]) -> Dict[str, str]:
    """str(path) -> the shortest label that still tells these files apart.

    The basename alone is the right label almost always, and is what an estimator recognises.
    But when two files in one pack share a basename, showing both as "117620202M.dxf" makes
    the page lie twice over: the rows look like duplicates, and a disagreement on one reads as
    a disagreement on the other. Where the names collide, enough parent directory is prefixed
    to separate them.
    """
    by_name: Dict[str, List[str]] = {}
    for path in paths:
        by_name.setdefault(Path(str(path)).name, []).append(str(path))
    labels: Dict[str, str] = {}
    for name, group in by_name.items():
        if len(group) == 1:
            labels[group[0]] = name
            continue
        for key in group:
            parent = Path(key).parent.name
            labels[key] = f"{parent}/{name}" if parent else key
        # Still colliding (same parent name too) — fall back to the whole path, which is ugly
        # but never ambiguous. A pretty label that merges two files is the worse trade.
        if len({labels[k] for k in group}) != len(group):
            for key in group:
                labels[key] = key
    return labels


def dxf_comparison_rows(summary: Mapping[str, Any],
                        dxf_paths: Optional[Sequence[Any]] = None) -> List[Dict[str, Any]]:
    """AVAILABLE IN THE FILE -> EXTRACTED -> ASSIGNED TO A PART -> USED IN COSTING.

    The other sheets can only report what the pipeline recorded, which blinds them to the
    failure worth the most: a fact that was in the file and reached nothing. This one reads
    the DXFs with ezdxf, independently of the production readers, and puts the file's own
    inventory beside the engine's figures.

    IT COMPARES LIKE WITH LIKE, AND SAYS WHEN IT CANNOT. A count of circles is not a count of
    holes — a Ø24 disc's outline is a circle and the part has no hole at all. A count of lines
    on a bend layer is not a count of bends. Those are shown side by side and marked NOT
    COMPARABLE, because an audit that manufactures disagreements is noise.
    """
    try:
        from dxf_probe import probe_dxf
    except Exception:                                                    # noqa: BLE001
        return []

    paths: List[Any] = list(dxf_paths or [])
    if not paths:
        seen: set = set()
        for part in _parts(summary):
            for key in ("dxf_file", "dxf_path", "flat_pattern_file"):
                value = part.get(key)
                if value and str(value).lower() not in seen:
                    seen.add(str(value).lower())
                    paths.append(value)
    if not paths:
        return []

    labels = _display_names(paths)
    rows: List[Dict[str, Any]] = []
    for path in paths:
        shown = labels.get(str(path)) or Path(str(path)).name
        try:
            probe = _probe_once(path)
        except Exception as err:                                         # noqa: BLE001
            rows.append({"file": shown, "part": "", "fact": "could not be read",
                         "in_the_file": "", "engine_has": "", "comparable": "",
                         "agrees": "", "note": f"{type(err).__name__}: {err}"})
            continue
        if not probe.get("readable"):
            rows.append({"file": shown, "part": "", "fact": "not readable",
                         "in_the_file": "", "engine_has": "", "comparable": "", "agrees": "",
                         "note": probe.get("error") or "ezdxf could not open this file"})
            continue
        if probe.get("entities_are_raster_only"):
            rows.append({"file": shown, "part": "", "fact": "content",
                         "in_the_file": "a raster image only", "engine_has": "",
                         "comparable": "n/a", "agrees": "n/a",
                         "note": "no geometry in this file to extract — no software reveals "
                                 "what is not there. Ask for a vector export"})
            continue
        if probe.get("raster_only_undetermined"):
            # DO NOT SEND ANYONE TO ASK FOR A VECTOR EXPORT OF A FILE THAT MAY ALREADY BE ONE.
            # An image with no reachable geometry beside it looks raster-only, but a block we
            # could not open may hold the entire profile. The row says what we know and what
            # we do not, and then carries on to the comparisons rather than stopping here.
            rows.append({"file": shown, "part": "", "fact": "content",
                         "in_the_file": "an image, and geometry we could not reach",
                         "engine_has": "", "comparable": "n/a", "agrees": "n/a",
                         "note": "this file MAY hold nothing but a raster image, but a block "
                                 "reference could not be resolved ("
                                 + "; ".join(probe.get("unresolved_blocks") or [])
                                 + ") and may contain the profile. Not called image-only"})

        part, how, ambiguous = _match_part(summary, path)
        pn = _text(part.get("part_number")) if part else ""
        geometry = (part.get("geometry_rollup") or {}) if part else {}
        features = (part.get("manufacturing_features") or {}) if part else {}

        def _engine(*keys: str) -> Any:
            for source in (part or {}, geometry, features):
                for key in keys:
                    value = (source or {}).get(key)
                    if value not in (None, "", []):
                        return value
            return None

        unit_note = "" if probe.get("units_known") else             f"units {probe.get('units')} — figures are unitless, not millimetres"
        partial = " (partial: " + ", ".join(probe.get("unsupported") or []) + ")"             if probe.get("unsupported") else ""

        # WHAT STOPS A FIGURE BEING SCORED AT ALL. A red "NO" on this page sends somebody to
        # look for an engine defect, so it must never be produced by the probe's own limits.
        #
        #   units unknown  — the file declares no $INSUNITS, so the probe's figure is in file
        #                    units and the engine's is in millimetres. One unitless drawing
        #                    read 30000 against the engine's 762 mm and scored NO; that is the
        #                    inch-to-mm factor, not a disagreement.
        #   partial total   — a sum taken while something was skipped (an unresolvable block, a
        #                    curve ezdxf could not flatten) is knowingly short of the whole. It
        #                    cannot be held against a complete figure: 300 vs 420 scored NO on
        #                    a file whose own `unsupported` list said geometry was missing.
        #   attribution only
        #     by file name   — the engine's figures for this part came from a file at a
        #                    DIFFERENT path that happens to share this basename. Then the two
        #                    numbers are not one fact measured twice; they are two files.
        #
        # All three are disclosed, not hidden — the row still shows both numbers and says why
        # they are not scored, so a reader can take it further if they want to.
        units_block = ("" if probe.get("units_known") else
                       f"NOT SCORED: the file declares no units ($INSUNITS {probe.get('units')}), "
                       f"so this figure is in file units and the engine's is in millimetres")
        partial_block = ("" if not probe.get("outline_length_partial") else
                         "NOT SCORED: this total was summed while geometry was skipped, so it "
                         "is knowingly short of the whole file")
        attribution_block = ("" if not (ambiguous and part) else
                             "NOT SCORED: this file is matched to the part by name only")
        # COMPLETENESS IS PER MEASUREMENT. Marking only the outline partial left a blank of
        # 100 x 50 scored as definitive on a file with a block we could not open — geometry we
        # never reached may lie outside that extent — and let a fold-axis count be scored when
        # the bend layer inside that block was never seen.
        unreached = ("" if probe.get("geometry_is_complete", True) else
                     "NOT SCORED: geometry in this file could not be reached ("
                     + "; ".join(probe.get("unresolved_blocks") or []) + ")")
        blank_block = attribution_block or units_block or (
            unreached if probe.get("blank_is_partial") else "")
        axes_block = attribution_block or units_block or (
            unreached if probe.get("candidate_fold_axes_partial") else "")
        units_block = attribution_block or units_block
        partial_block = attribution_block or partial_block or unreached

        # comparable facts: same thing measured two ways
        checks = [
            ("blank length", probe.get("blank_length_mm"), _engine("blank_length_mm"), True, "",
             blank_block),
            ("blank width", probe.get("blank_width_mm"), _engine("blank_width_mm"), True, "",
             blank_block),
            ("outline length", probe.get("outline_length"),
             _engine("cut_length_mm", "dxf_measured_cut_length"), True,
             ("the file's total profile length" + partial),
             units_block or partial_block),
            # inventory vs interpretation: shown together, never scored
            ("circles in the file", probe.get("circle_count") or None,
             _engine("hole_count", "estimated_hole_count"), False,
             "a circle is not necessarily a hole — a disc's OUTLINE is a circle. Deciding "
             "which are holes needs the part's role and the drawing's instructions", ""),
            # AN AXIS IS NOT A BEND, SO THE VERDICT MUST SAY SO TOO. This was scored against
            # bend_count while the note beside it admitted that nothing here knows which
            # profile owns which segment — the row therefore graded a comparison its own
            # explanation said could not be made. It is the same category error as circles
            # versus holes, which this table has always refused to score: the geometry states
            # fold LINES, and how many manufacturing bends those represent needs the part's
            # role. Two tabs folding on one line read as one axis; a nested pair of parts read
            # as shared axes. Both figures are shown side by side — a difference is worth
            # looking at, and the note says what would explain one — but neither is graded
            # right or wrong against the other.
            ("candidate fold axes", probe.get("candidate_fold_axes") or None,
             _engine("bend_count", "bend_count_dxf", "fold_count"), False,
             (f"{probe.get('bend_layer_line_count')} segment(s) on the bend layer collapse "
              f"to {probe.get('candidate_fold_axes')} fold axis/axes. An AXIS is not proven "
              f"to be one manufacturing bend: nothing here knows which profile owns which "
              f"segment, so two tabs folding on one line would read as one"
              + ("; and geometry in this file could not be reached, so the bend layer may be "
                 "incomplete" if probe.get("candidate_fold_axes_partial") else ""))
             if probe.get("bend_layer_line_count") else "",
             axes_block),
        ]
        for label, available, extracted, comparable, note, blocked in checks:
            if available is None and extracted in (None, "", []):
                continue
            verdict = ""
            if not comparable:
                verdict = "NOT COMPARABLE"
            elif blocked and available is not None:
                verdict = "NOT COMPARABLE"
            elif available is not None and extracted not in (None, "", []):
                a, b = _num(available), _num(extracted)
                if a is not None and b is not None:
                    verdict = "yes" if abs(a - b) <= max(0.5, abs(a) * 0.02) else "NO"
            elif available is not None:
                # DELIBERATELY NOT "NOT EXTRACTED". This audit inspects a handful of fields;
                # their absence is not proof the engine never read or used the value.
                verdict = "not in the fields checked"
            rows.append({
                "file": shown,
                "part": pn or ("(ambiguous)" if ambiguous else "(no part matched)"),
                "fact": label,
                "in_the_file": available,
                "engine_has": extracted,
                "comparable": "yes" if (comparable and not (blocked and available is not None))
                              else "no",
                "agrees": verdict,
                "note": "; ".join(x for x in ((blocked if available is not None else ""),
                                              note, unit_note, probe.get("extent_is") or "",
                                              "" if not ambiguous else f"attribution {how}")
                                  if x),
            })
    return rows


_PROBE_CACHE: Dict[str, Dict[str, Any]] = {}


def _probe_once(path: Any) -> Dict[str, Any]:
    """One read per VERSION of a file. The comparison table and the detail section both want
    the same inventory, and reading each DXF twice is both slower and a way for the two halves
    of one page to disagree.

    THE KEY IS THE FILE'S IDENTITY PLUS ITS STATE, NOT JUST ITS PATH. Keyed on the path alone,
    this cache is a correctness hazard in any process that outlives one job: a DXF rewritten
    between two audits — a revision dropped into the same folder under the same name, which is
    exactly how a drawing office issues one — would be reported from the old read, and the page
    would describe geometry the file no longer has while naming it as current. Adding the
    modification time and size makes a rewritten file a different key, so it is read again. A
    file whose stat cannot be taken is not cached at all: better a second read than a stale
    answer."""
    from dxf_probe import probe_dxf
    try:
        stat = Path(str(path)).stat()
        key = f"{path}|{stat.st_mtime_ns}|{stat.st_size}"
    except Exception:                                                    # noqa: BLE001
        return probe_dxf(path)
    if key not in _PROBE_CACHE:
        _PROBE_CACHE[key] = probe_dxf(path)
    return _PROBE_CACHE[key]


def dxf_detail_blocks(summary: Mapping[str, Any],
                      dxf_paths: Optional[Sequence[Any]] = None) -> List[Dict[str, Any]]:
    """The COMPLETE inventory of each DXF — everything read, nothing summarised away.

    The comparison sheet answers "does the engine agree?". This answers the prior question:
    what is actually IN this file? Every entity type and its count, every layer, every text
    string, every distinct circle diameter, the units as declared, and anything the reader
    could not measure. Nothing is elided: a section that shows the interesting rows and hides
    the rest is how a fact goes missing without anyone deciding to drop it.
    """
    paths: List[Any] = list(dxf_paths or [])
    if not paths:
        seen: set = set()
        for part in _parts(summary):
            for key in ("dxf_file", "dxf_path", "flat_pattern_file"):
                value = part.get(key)
                if value and str(value).lower() not in seen:
                    seen.add(str(value).lower())
                    paths.append(value)
    labels = _display_names(paths)
    out: List[Dict[str, Any]] = []
    for path in paths:
        try:
            probe = _probe_once(path)
        except Exception:                                                # noqa: BLE001
            continue
        # The FULL path, so two files sharing a basename are not credited to one part.
        part, how, ambiguous = _match_part(summary, path)
        out.append({"probe": probe,
                    "shown_as": labels.get(str(path)) or Path(str(path)).name,
                    "part": _text((part or {}).get("part_number")),
                    "attribution": how, "ambiguous": ambiguous})
    return out


def page_rows(summary: Mapping[str, Any]) -> List[Dict[str, Any]]:
    """Every page of every PDF and everything read off it. One row per page, always.

    A page that yielded nothing gets a row saying so — otherwise the pages that failed are
    simply absent, which reads as a shorter drawing pack.

    THE WORDING IS EXACT ON PURPOSE. An earlier version checked four fields and, when none
    was populated, printed "nothing was read from this page" — on a page whose process notes
    HAD been extracted. That is a false statement about the pack, made by a document whose
    only job is to be true about the pack. It now lists everything it finds and, when it
    finds none of them, says which fields it looked at.
    """
    checked = [
        ("overall dimensions", lambda d, a: (
            f"{d.get('overall_length_mm')} x {d.get('overall_width_mm')}"
            if d.get("overall_length_mm") or d.get("overall_width_mm") else None)),
        ("dimension figures", lambda d, a: (
            f"{len(d['all_dimensions_mm'])}" if d.get("all_dimensions_mm") else None)),
        ("hole sizes", lambda d, a: (
            f"{len(d['hole_sizes_mm'])}" if d.get("hole_sizes_mm") else None)),
        ("materials", lambda d, a: ", ".join(str(m) for m in a["materials"][:4])
            if a.get("materials") else None),
        ("finishes", lambda d, a: ", ".join(str(m) for m in a["surface_finishes"][:4])
            if a.get("surface_finishes") else None),
        ("part codes", lambda d, a: f"{len(a['part_numbers'])}"
            if a.get("part_numbers") else None),
        ("process notes", lambda d, a: f"{len(a['process_notes'])}"
            if a.get("process_notes") else None),
        ("textual operations", lambda d, a: ", ".join(str(o) for o in a["textual_operations"][:6])
            if a.get("textual_operations") else None),
        ("drawing numbers", lambda d, a: ", ".join(str(o) for o in a["drawing_numbers"][:4])
            if a.get("drawing_numbers") else None),
        ("revision", lambda d, a: str(a.get("revision")) if a.get("revision") else None),
        ("scales", lambda d, a: ", ".join(str(o) for o in a["scales"][:3])
            if a.get("scales") else None),
        ("BOM table", lambda d, a: f"{len(a['bom_rows'])} row(s)" if a.get("bom_rows") else None),
    ]
    rows: List[Dict[str, Any]] = []
    for page in (summary.get("pages") or []):
        if not isinstance(page, Mapping):
            continue
        analysis = page.get("page_analysis") or {}
        dims = analysis.get("dimensions") or {}
        role = (page.get("page_role") or {}).get("primary_role") or page.get("page_roles") or ""
        found = []
        for label, reader in checked:
            try:
                value = reader(dims, analysis)
            except Exception:                                            # noqa: BLE001
                value = None
            if value:
                found.append(f"{label}: {value}")
        rows.append({
            "file": _text(page.get("source_pdf_name") or page.get("source_file") or ""),
            "page": page.get("source_page_number") or page.get("page_number"),
            "role": _text(role),
            "read_from_it": "; ".join(found) if found else (
                "none of the fields checked were populated ("
                + ", ".join(label for label, _ in checked) + ")"),
        })
    return rows


def fact_rows(summary: Mapping[str, Any]) -> List[Dict[str, Any]]:
    """One row per extracted datum: what, on which part, from where, and did it get used.

    The `used_in_costing` column is the point of the sheet. A value that was read, attached
    to the right part and then beaten by a stronger source is CORRECT behaviour and says so;
    a value that was read and simply never referenced again is a join to fix.
    """
    rows: List[Dict[str, Any]] = []
    for part in _parts(summary):
        pn = _text(part.get("part_number"))
        live = _costed_fields(part)
        displaced = part.get("_displaced") if isinstance(part.get("_displaced"), Mapping) else {}

        for field, value in live.items():
            if value in (None, ""):
                continue
            rows.append({
                "part": pn,
                "field": field,
                "value": value,
                "source": _text(part.get(f"{field.replace('normalized_', '')}_source")
                                or part.get(f"{field}_source") or ""),
                "outcome": "USED — this is the figure the price is built on",
                "file_or_page": _text(part.get("drawing_files") or part.get("pages") or ""),
            })

        for field, entries in (displaced or {}).items():
            for entry in (entries or []):
                if not isinstance(entry, Mapping):
                    continue
                applied = entry.get("applied")
                rows.append({
                    "part": pn,
                    "field": field,
                    "value": entry.get("value"),
                    "source": _text(entry.get("source")),
                    "outcome": ("USED" if applied else
                                f"read, then beaten by a stronger source"
                                + (f" ({_text(entry.get('displaced_by'))})"
                                   if entry.get("displaced_by") else "")),
                    "file_or_page": _text(entry.get("page") or entry.get("file") or ""),
                })
    return rows


def bom_rows(summary: Mapping[str, Any]) -> List[Dict[str, Any]]:
    """Every parts-list row the readers returned, with the page it was read from."""
    out: List[Dict[str, Any]] = []
    for row in ((summary.get("document_analysis") or {}).get("bom_rows") or []):
        if not isinstance(row, Mapping):
            continue
        out.append({
            "file_or_page": _text(row.get("source_page") or row.get("page") or ""),
            "item": _text(row.get("item") or row.get("item_no") or ""),
            "part_number": _text(row.get("part_number")),
            "description": _text(row.get("description")),
            "quantity": row.get("quantity"),
            "material_as_printed": _text(row.get("material_text")),
            "thickness_mm": row.get("thickness_mm"),
            "stated_weight_kg": row.get("stated_weight_kg"),
            "read_by": _text(row.get("source") or row.get("reader") or ""),
        })
    return out


def operation_rows(summary: Mapping[str, Any]) -> List[Dict[str, Any]]:
    """Every route decision, its target and why it stands or does not."""
    out: List[Dict[str, Any]] = []
    route = (summary.get("estimate_summary") or {}).get("canonical_route") or {}
    for decision in (route.get("decisions") or []):
        if not isinstance(decision, Mapping):
            continue
        out.append({
            "operation": _text(decision.get("operation")),
            "target": _text(decision.get("part_number") or decision.get("target")),
            "status": _text(decision.get("status")),
            "decided_by": _text(decision.get("decided_by")),
            "reason": _text(decision.get("reason")),
        })
    return out


def not_extracted_rows(summary: Mapping[str, Any]) -> List[Dict[str, Any]]:
    """What the pack held that nothing read. The shortest sheet worth the most attention."""
    out: List[Dict[str, Any]] = []
    for issue in ((summary.get("document_analysis") or {}).get("pack_issues") or []):
        if isinstance(issue, Mapping):
            out.append({"what": _text(issue.get("code") or "pack issue"),
                        "detail": _text(issue.get("message"), 600)})
    for part in _parts(summary):
        for flag in (part.get("review_flags") or []):
            text = str(flag)
            if any(cue in text for cue in ("UNRESOLVED", "not in the material lexicon",
                                           "NOT read as a diameter", "is not known",
                                           "could not be taken")):
                out.append({"what": _text(part.get("part_number")), "detail": _text(text, 600)})
    return out


def build_tables(summary: Mapping[str, Any],
                 dxf_paths: Optional[Sequence[Any]] = None) -> Dict[str, List[Dict[str, Any]]]:
    return {
        "Files": files_rows(summary),
        "Pages": page_rows(summary),
        "DXF file vs engine": dxf_comparison_rows(summary, dxf_paths),
        "Facts": fact_rows(summary),
        "BOM rows": bom_rows(summary),
        "Operations": operation_rows(summary),
        "Not extracted": not_extracted_rows(summary),
    }


def write_source_drawing_data(summary: Mapping[str, Any], out_dir: Any,
                              job: str = "",
                              dxf_paths: Optional[Sequence[Any]] = None,
                              tables: Optional[Mapping[str, List[Dict[str, Any]]]] = None,
                              ) -> Optional[Path]:
    """Write source_drawing_data.xlsx into the estimating folder. Returns the path.

    Pass `tables` to write from a snapshot already built — see the note on the HTML writer for
    why the caller should build once and hand the same dict to both.

    Never raises into the run: an audit that breaks the job it audits is worse than no audit.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except Exception:                                                    # noqa: BLE001
        return None
    tables = dict(tables) if tables is not None else build_tables(summary, dxf_paths)
    book = Workbook()
    book.remove(book.active)
    for name in SHEETS:
        rows = tables.get(name) or []
        sheet = book.create_sheet(name[:31])
        if not rows:
            sheet.cell(row=1, column=1, value=f"No {name.lower()} recorded for this job.")
            continue
        headers = list(rows[0].keys())
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = Font(bold=True)
        for index, row in enumerate(rows, start=2):
            for column, header in enumerate(headers, start=1):
                sheet.cell(row=index, column=column, value=row.get(header))
        for column, header in enumerate(headers, start=1):
            width = max(len(str(header)),
                        *(len(str(r.get(header) or "")) for r in rows[:200]))
            sheet.column_dimensions[
                sheet.cell(row=1, column=column).column_letter].width = min(max(width + 2, 10), 80)
        sheet.freeze_panes = "A2"

    target = Path(out_dir)
    target.mkdir(parents=True, exist_ok=True)
    path = target / (f"{job}_source_drawing_data.xlsx" if job else "source_drawing_data.xlsx")
    book.save(path)
    return path


# ── The same tables, as a page ────────────────────────────────────────────────────────

_CSS = """
:root{--bg:#fbfbfa;--surface:#fff;--surface-2:#f6f6f4;--ink:#17181a;--dim:#5c6066;
--faint:#8b9096;--line:#e4e4e1;--ok:#1a7f52;--bad:#b4342a;--warn:#a8641a;--info:#2b5fa8;
--disp:'Inter',-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;
--mono:ui-monospace,'SF Mono',Menlo,Consolas,monospace}
@media(prefers-color-scheme:dark){:root{--bg:#131416;--surface:#1a1c1f;--surface-2:#212429;
--ink:#e8e9ea;--dim:#a4a9b0;--faint:#767b83;--line:#2c3037;--ok:#4ec08a;--bad:#e8776c;
--warn:#dda257;--info:#7aa8e8}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.55 var(--disp);
-webkit-font-smoothing:antialiased}
.wrap{max-width:1180px;margin:0 auto;padding:40px 20px 72px}
header{border-bottom:1px solid var(--line);padding-bottom:22px;margin-bottom:30px}
.kicker{font-size:10px;letter-spacing:2.4px;text-transform:uppercase;color:var(--faint);
font-weight:700;margin-bottom:8px}
h1{font-size:27px;line-height:1.2;margin:0 0 8px;font-weight:750;letter-spacing:-.4px}
.sub{color:var(--dim);margin:0;max-width:68ch;font-size:14px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;
margin:26px 0 34px}
.card{background:var(--surface);border:1px solid var(--line);border-radius:8px;padding:14px 16px}
.card .n{font-size:25px;font-weight:750;letter-spacing:-.5px;line-height:1.1}
.card .l{font-size:11px;color:var(--faint);text-transform:uppercase;letter-spacing:1.1px;
margin-top:5px;font-weight:600}
.card.alert .n{color:var(--bad)}
h2{font-size:17px;margin:34px 0 4px;font-weight:700;letter-spacing:-.2px}
h2 .cnt{color:var(--faint);font-weight:500;font-size:13px;letter-spacing:0}
.why{color:var(--dim);margin:0 0 14px;font-size:13px;max-width:78ch}
.panel{background:var(--surface);border:1px solid var(--line);border-radius:8px;
overflow:auto;max-height:none}
table{border-collapse:collapse;width:100%;font-size:12.5px}
th{position:sticky;top:0;background:var(--surface-2);text-align:left;padding:9px 12px;
font-size:10px;letter-spacing:1.1px;text-transform:uppercase;color:var(--faint);
font-weight:700;border-bottom:1px solid var(--line);white-space:nowrap}
td{padding:8px 12px;border-bottom:1px solid var(--line);vertical-align:top}
tr:last-child td{border-bottom:0}
td.num{font-family:var(--mono);white-space:nowrap;text-align:right}
td.k{font-family:var(--mono);font-size:11.5px;color:var(--dim);white-space:nowrap}
.pill{display:inline-block;padding:1.5px 8px;border-radius:99px;font-size:10.5px;
font-weight:700;letter-spacing:.3px;white-space:nowrap}
.pill.ok{background:color-mix(in srgb,var(--ok) 14%,transparent);color:var(--ok)}
.pill.bad{background:color-mix(in srgb,var(--bad) 14%,transparent);color:var(--bad)}
.pill.warn{background:color-mix(in srgb,var(--warn) 16%,transparent);color:var(--warn)}
.pill.mute{background:var(--surface-2);color:var(--faint)}
.note{color:var(--faint);font-size:11.5px;max-width:52ch}
.empty{padding:22px;color:var(--faint);font-size:13px}
.fhead{font-family:var(--mono);font-size:13px;margin-bottom:10px;padding-bottom:8px;
border-bottom:1px solid var(--line);color:var(--ink)}
table.kv{font-size:12.5px}
table.kv th{position:static;background:none;width:190px;text-transform:none;letter-spacing:0;
font-size:11.5px;color:var(--faint);font-weight:600;border-bottom:1px solid var(--line);
vertical-align:top;padding:6px 12px 6px 0;white-space:nowrap}
table.kv td{padding:6px 0;font-family:var(--mono);font-size:11.5px;word-break:break-word}
footer{margin-top:44px;padding-top:18px;border-top:1px solid var(--line);
color:var(--faint);font-size:11.5px;max-width:76ch}
@media(max-width:640px){.wrap{padding:24px 14px 48px}h1{font-size:22px}table{font-size:12px}}
@media print{body{background:#fff}.panel{break-inside:avoid}th{position:static}}
"""

_PILL = {
    "yes": "ok", "NO": "bad", "NOT COMPARABLE": "mute",
    "not in the fields checked": "warn", "n/a": "mute", "NO ": "bad",
}


def _esc(value: Any) -> str:
    import html
    return html.escape("" if value is None else str(value), quote=True)


def _cell(header: str, value: Any) -> str:
    """A verdict becomes a pill, a number right-aligns, a code gets the mono face."""
    if header == "agrees" and value:
        return f'<td><span class="pill {_PILL.get(str(value), "mute")}">{_esc(value)}</span></td>'
    if header == "read" and str(value).upper() == "NO":
        return '<td><span class="pill bad">not read</span></td>'
    if header in ("note", "detail", "yielded", "reason", "assumption"):
        return f'<td class="note">{_esc(value)}</td>'
    if header in ("file", "part", "part_number", "field", "fact", "operation", "target"):
        return f'<td class="k">{_esc(value)}</td>'
    if isinstance(value, (int, float)):
        return f'<td class="num">{_esc(value)}</td>'
    return f"<td>{_esc(value)}</td>"


_WHY = {
    "Pages": "Every page of every PDF and what was read off it. A page that yielded "\
             "nothing gets a row saying so — otherwise the pages that failed are simply "\
             "absent, which reads as a shorter drawing pack.",
    "Files": "Every file the job saw and whether anything read it. A CAD file nobody opened "
             "appears nowhere else — the estimate simply prices what it has.",
    "DXF file vs engine": "The DXFs opened independently with ezdxf, their own measurements "
                          "beside the engine's. A reader cannot be checked against its own "
                          "output. Facts that are not the same thing — a circle is not a "
                          "hole — are shown together and marked NOT COMPARABLE, never scored.",
    "Facts": "One row per datum: where it came from, and whether it is the figure the price "
             "rests on or was read and then beaten by something stronger.",
    "BOM rows": "The parts list as the drawing office typed it — material cell verbatim.",
    "Operations": "Every route decision, its target and why it stands or does not.",
    "Not extracted": "What the pack held that nothing read. The shortest section, and usually "
                     "the one worth the most.",
}


def write_source_drawing_html(summary: Mapping[str, Any], out_dir: Any, job: str = "",
                              dxf_paths: Optional[Sequence[Any]] = None,
                              tables: Optional[Mapping[str, List[Dict[str, Any]]]] = None,
                              ) -> Optional[Path]:
    """The same tables as a page: estimating works from the spreadsheet, management reads this.

    PASS `tables` AND THE TWO OUTPUTS ARE THE SAME SNAPSHOT. This docstring used to claim the
    workbook and the page "cannot drift apart" because both were "built from build_tables()" —
    but calling the same FUNCTION twice is not the same DATA. Two builds read the DXFs again,
    re-walk the summary, and nothing held them to the same answer: a file rewritten between the
    two calls, or a summary mutated by anything running in between, and the spreadsheet
    estimating works from would disagree with the page management reads, with no indication
    which was right. When the caller builds once and hands the same dict to both writers, the
    claim is true by construction rather than by coincidence. Built here only if no snapshot is
    supplied, so a lone call still works.
    """
    tables = dict(tables) if tables is not None else build_tables(summary, dxf_paths)
    files = tables.get("Files") or []
    comparisons = tables.get("DXF file vs engine") or []
    unread = sum(1 for f in files if str(f.get("read", "")).upper() == "NO")
    disagree = sum(1 for r in comparisons if r.get("agrees") == "NO")
    unchecked = sum(1 for r in comparisons
                    if r.get("agrees") == "not in the fields checked")

    parts: List[str] = [
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width,initial-scale=1'>",
        f"<title>Source drawing data{(' · ' + _esc(job)) if job else ''}</title>",
        f"<style>{_CSS}</style></head><body><div class='wrap'>",
        "<header><div class='kicker'>SDI Estimating Intelligence</div>",
        f"<h1>Source drawing data{(' — ' + _esc(job)) if job else ''}</h1>",
        "<p class='sub'>Everything the drawing pack gave us, file by file, and what the "
        "engine did with it. This page prices nothing and decides nothing — it exists so "
        "that <em>the file did not have it</em>, <em>we did not read it</em> and <em>we read "
        "it and dropped it before the price</em> can be told apart.</p></header>",
        "<div class='cards'>",
        f"<div class='card'><div class='n'>{len(files)}</div><div class='l'>Files seen</div></div>",
        f"<div class='card{' alert' if unread else ''}'><div class='n'>{unread}</div>"
        "<div class='l'>Not read</div></div>",
        f"<div class='card'><div class='n'>{len(tables.get('BOM rows') or [])}</div>"
        "<div class='l'>BOM rows</div></div>",
        f"<div class='card{' alert' if disagree else ''}'><div class='n'>{disagree}</div>"
        "<div class='l'>File vs engine differ</div></div>",
        f"<div class='card'><div class='n'>{unchecked}</div>"
        "<div class='l'>Not in fields checked</div></div>",
        "</div>",
    ]

    for name in SHEETS:
        rows = tables.get(name) or []
        parts.append(f"<h2>{_esc(name)} <span class='cnt'>{len(rows)}</span></h2>")
        parts.append(f"<p class='why'>{_esc(_WHY.get(name, ''))}</p>")
        if not rows:
            parts.append(f"<div class='panel'><div class='empty'>Nothing recorded for this "
                         f"job.</div></div>")
            continue
        headers = list(rows[0].keys())
        parts.append("<div class='panel'><table><thead><tr>"
                     + "".join(f"<th>{_esc(h.replace('_', ' '))}</th>" for h in headers)
                     + "</tr></thead><tbody>")
        for row in rows:
            parts.append("<tr>" + "".join(_cell(h, row.get(h)) for h in headers) + "</tr>")
        parts.append("</tbody></table></div>")

    # ── The complete inventory of every DXF, nothing elided ──────────────────────────
    blocks = dxf_detail_blocks(summary, dxf_paths)
    if blocks:
        parts.append("<h2>Every DXF in full <span class='cnt'>"
                     f"{len(blocks)}</span></h2>")
        parts.append("<p class='why'>The complete inventory of each file — every entity type, "
                     "every layer, every text string, every distinct circle. Nothing is "
                     "summarised away: a section that shows the interesting rows and hides the "
                     "rest is how a fact goes missing without anyone deciding to drop it.</p>")
        for entry in blocks:
            probe = entry["probe"]
            # The disambiguated label, not the bare basename: two files of the same name in
            # one pack must not head two sections identically.
            head = _esc(entry.get("shown_as") or probe.get("file"))
            who = f" &middot; {_esc(entry['part'])}" if entry.get("part") else ""
            parts.append(f"<div class='panel' style='padding:16px 18px;margin-bottom:12px'>")
            parts.append(f"<div class='fhead'><b>{head}</b>{who}</div>")
            if not probe.get("readable"):
                parts.append(f"<p class='note'>Not readable — {_esc(probe.get('error'))}</p></div>")
                continue
            kind = ("flat pattern export" if probe.get("looks_like_flat_export")
                    else "drawing export")
            # EVERY VALUE HERE COMES OUT OF A FILE, so every value is escaped and only the
            # separators are markup. The first version inserted strings raw on the reasoning
            # that they were "just layer names" — layer names, entity types, attribution text
            # and error messages are all source-derived, and a table that renders one of them
            # as markup is a defect whether or not today's files happen to exploit it.
            _E = _esc

            def _list(values: Sequence[Any], sep: str = ", ", empty: str = "none") -> str:
                items = [_E(v) for v in (values or [])]
                return sep.join(items) if items else empty

            facts = [
                ("read with", _E(probe.get("reader"))),
                ("classified as", _E(kind)),
                ("units declared", _E(probe.get("units"))
                 + ("" if probe.get("units_known") else " &mdash; figures are unitless")),
                ("extent", f"{_E(probe.get('extent_length'))} &times; "
                           f"{_E(probe.get('extent_width'))}"),
                ("blank published",
                 ((f"{_E(probe.get('blank_length_mm'))} &times; "
                   f"{_E(probe.get('blank_width_mm'))}")
                  + (" (partial &mdash; geometry in this file could not be reached, so the "
                     "real blank may be larger)" if probe.get("blank_is_partial") else ""))
                 if probe.get("blank_length_mm")
                 else "none &mdash; " + (_E(probe.get("extent_is")) or "not a flat export")),
                ("profile length", _E(probe.get("outline_length"))
                 + (" (partial)" if probe.get("outline_length_partial") else "")),
                ("circles", _E(probe.get("circle_count"))
                 + ((" &mdash; &Oslash; " + _list(probe.get("circle_diameters")))
                    if probe.get("circle_diameters") else "")),
                ("candidate fold axes",
                 f"{_E(probe.get('candidate_fold_axes'))} from "
                 f"{_E(probe.get('bend_layer_line_count'))} segment(s) on the bend layer"
                 + (" (partial &mdash; the bend layer may continue inside geometry we could "
                    "not reach)" if probe.get("candidate_fold_axes_partial") else "")),
                ("dimension entities", _E(probe.get("dimension_entities"))),
                ("layers", _list(probe.get("layers"))),
                ("entities", _list([f"{k} \u00d7{v}" for k, v in
                                    (probe.get("entity_counts") or {}).items()])),
                # HOW DEEP THE BLOCKS WENT, because the counts above depend on it. A block
                # reference is not geometry, and a block inside a block inside a block is the
                # ordinary shape of an assembly export; a file whose profile was two levels
                # down once reported one INSERT and nothing else.
                ("block nesting resolved",
                 (f"{_E(probe.get('block_nesting_depth'))} level(s) deep"
                  if probe.get("block_nesting_depth") else "no block references")
                 + ((" &mdash; NOT resolved: " + _list(probe.get("unresolved_blocks")))
                    if probe.get("unresolved_blocks") else "")),
                ("not measured", _list(probe.get("unsupported"), empty="nothing")),
                ("text in the file", _list(probe.get("text_values"), sep=" &vert; ")
                 + (f" ({_E(probe.get('text_count'))} string(s))"
                    if probe.get("text_count") else "")),
                ("attributed to a part by", _E(entry.get("attribution"))),
            ]
            parts.append("<table class='kv'>")
            for label, value in facts:
                parts.append(f"<tr><th>{_E(label)}</th><td>{value}</td></tr>")
            parts.append("</table></div>")

    parts.append(
        "<footer>Geometry proves shape, not method: a modelled hole does not prove drilling, "
        "and a line on a bend layer is not a bend. Measurements are read with ezdxf and "
        "converted from the units the file declares; a file declaring none publishes no "
        "blank. Anything the reader could not measure is named rather than folded into a "
        "total.</footer></div></body></html>")

    target = Path(out_dir)
    target.mkdir(parents=True, exist_ok=True)
    path = target / (f"{job}_source_drawing_data.html" if job else "source_drawing_data.html")
    path.write_text("\n".join(parts), encoding="utf-8")
    return path
