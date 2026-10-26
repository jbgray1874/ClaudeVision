"""The AI BOX CALCULATOR sheet itself, read and calculated where the estimators keep it (D-542).

James Gray, 15 Oct 2026: "can we change it ... so that it's reading from the s/sheet and using
the s/sheet for the calculation in case estimating tweak it ... we should make the file a .env
variable."

SDI_BOX_CALCULATOR_XLSX names the file (config.BOX_CALCULATOR_SHEET["path"]). This module does
three things with it and nothing else:

  read_sheet()      the sheet's own figures and formulas, read from the file as saved — the
                    board, sheet size and price, set-up, running rate, panel oversize, nest gap
                    and trim, scrap, and which pairs its TOTAL adds. Checked against the
                    sheet's own labels first: a sheet laid out differently is refused by name,
                    never read from the wrong cells.
  calculate()       on a machine with Excel, the sheet calculating: a COPY of the file is
                    opened, the box size and quantity are entered into its input cells, Excel
                    calculates, and the TOTAL, running, set-up share and each pair's nest and
                    cost are read back. The estimators' file is never written to.
  the engine copy   where Excel is not available, box_costing runs its copy of the sheet's
                    formulas on the figures read_sheet() found — and says so.

Every answer carries where it came from (file, saved time, a short checksum) so the e-mail can
say exactly which version of the sheet priced the box.
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    import config
except ImportError:                                                  # pragma: no cover
    from . import config                                             # type: ignore

ENV_NAME = "SDI_BOX_CALCULATOR_XLSX"

_READ_CACHE: Dict[Tuple[str, float], Dict[str, Any]] = {}
_CALC_CACHE: Dict[Tuple[str, float, float, float, float, int], Dict[str, Any]] = {}


def settings() -> Dict[str, Any]:
    return dict(getattr(config, "BOX_CALCULATOR_SHEET", {}) or {})


def configured_path() -> Optional[str]:
    p = settings().get("path")
    return str(p).strip() if p and str(p).strip() else None


def _num(v: Any) -> Optional[float]:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f == f else None


def _cell_ref(cell: str) -> Tuple[str, int]:
    m = re.fullmatch(r"([A-Z]+)(\d+)", str(cell).strip().upper())
    if not m:
        raise ValueError(f"not a cell reference: {cell!r}")
    return m.group(1), int(m.group(2))


def read_sheet(path: Optional[str] = None) -> Dict[str, Any]:
    """The sheet's figures and formulas as saved, or why they cannot be used.

    {ok, reason, path, file, modified, sha256, sheet_name, figures, formulas, notes,
     pairs_in_total, engine_copy_ok}. `figures` uses BOX_COSTING's own keys, so box_costing
    prices from it exactly as it prices from config."""
    path = path or configured_path()
    out: Dict[str, Any] = {"ok": False, "path": path, "env": ENV_NAME, "notes": []}
    if not path:
        out["reason"] = f"{ENV_NAME} is not set"
        return out
    p = Path(path)
    if not p.is_file():
        out["reason"] = f"{ENV_NAME} names {path}, which is not there (or not reachable)"
        return out
    try:
        mtime = p.stat().st_mtime
    except OSError as exc:
        out["reason"] = f"cannot read {path}: {exc}"
        return out
    key = (str(p.resolve()), mtime)
    if key in _READ_CACHE:
        return dict(_READ_CACHE[key])
    try:
        data = p.read_bytes()
        import openpyxl
        from io import BytesIO
        wb_f = openpyxl.load_workbook(BytesIO(data))
        wb_v = openpyxl.load_workbook(BytesIO(data), data_only=True)
    except Exception as exc:                                         # noqa: BLE001
        out["reason"] = f"cannot open {p.name}: {type(exc).__name__}: {exc}"
        return out
    out.update({"file": p.name,
                "modified": _dt.datetime.fromtimestamp(mtime).strftime("%d %b %Y %H:%M"),
                "sha256": hashlib.sha256(data).hexdigest()[:12]})
    s = settings()
    labels = dict(s.get("labels") or {})
    # THE SHEET IS FOUND BY ITS LABELS, NOT ITS TAB NAME. The tab is called what the
    # estimator called it; the input cells are where the labels say they are.
    ws_f = ws_v = None
    for name in wb_f.sheetnames:
        cand = wb_f[name]
        if all(str(cand[c].value or "").strip().lower() == str(t).strip().lower()
               for c, t in labels.items()):
            ws_f, ws_v = cand, wb_v[name]
            break
    if ws_f is None:
        first = wb_f[wb_f.sheetnames[0]]
        wrong = [f"{c} reads {str(first[c].value or '')!r}, expected {t!r}"
                 for c, t in labels.items()
                 if str(first[c].value or "").strip().lower() != str(t).strip().lower()]
        out["reason"] = ("the sheet is not laid out as the engine expects — "
                         + "; ".join(wrong[:4])
                         + " (update config.BOX_CALCULATOR_SHEET if the inputs have moved)")
        return out
    out["sheet_name"] = ws_f.title

    def f(cell: str) -> Any:            # the formula (or typed value) in a cell
        return ws_f[cell].value

    def v(cell: str) -> Any:            # the value Excel last saved for it
        return ws_v[cell].value

    formulas: Dict[str, str] = {}
    for row in ws_f.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                formulas[c.coordinate] = c.value
    out["formulas"] = formulas

    figs: Dict[str, Any] = {}
    problems: List[str] = []
    fig_cells = dict(s.get("figures") or {})
    figs["stock_board"] = str(v(fig_cells.get("stock_board", "A3")) or "").strip() or None
    for k in ("sheet_cost_gbp", "setup_cost_gbp"):
        n = _num(v(fig_cells[k]))
        if n is None:
            problems.append(f"{fig_cells[k]} ({k.replace('_gbp', '').replace('_', ' ')}) "
                            f"is not a number")
        figs[k] = n
    m = re.search(r"\*\s*([0-9]*\.?[0-9]+)\s*$", str(f(s.get("running", "E3")) or ""))
    if m:
        figs["running_rate_gbp_per_mm"] = float(m.group(1))
    else:
        problems.append(f"{s.get('running', 'E3')} is no longer (L+B+H) x a rate: "
                        f"{f(s.get('running', 'E3'))!r}")
    rows = list(s.get("pair_rows") or (7, 8, 9))
    r0 = rows[0]
    m = re.search(r"\+\s*([0-9]*\.?[0-9]+)\s*$", str(f(f"C{r0}") or ""))
    if m:
        figs["panel_oversize_mm"] = float(m.group(1))
    else:
        problems.append(f"C{r0} is no longer a size + an oversize: {f(f'C{r0}')!r}")
    m = re.search(r"\(G\d+\s*-\s*([0-9]*\.?[0-9]+)\)\s*/\s*\(D\d+\s*\+\s*\(\s*([0-9]*\.?[0-9]+)"
                  r"\s*\*\s*2\s*\)\)", str(f(f"H{r0}") or ""))
    if m:
        figs["sheet_edge_trim_mm"] = float(m.group(1))
        figs["nest_gap_mm"] = float(m.group(2))
    else:
        problems.append(f"H{r0}'s nest formula has changed: {str(f(f'H{r0}'))[:80]!r}")
    for k, col in (("sheet_length_mm", "F"), ("sheet_width_mm", "G"), ("scrap_fraction", "I")):
        n = _num(v(f"{col}{r0}"))
        if n is None:
            problems.append(f"{col}{r0} ({k}) is not a number")
        figs[k] = n
        for r in rows[1:]:
            if _num(v(f"{col}{r}")) is not None and _num(v(f"{col}{r}")) != n:
                out["notes"].append(f"{col}{r} reads {v(f'{col}{r}')} where {col}{r0} reads "
                                    f"{n:g} — the engine's copy uses {col}{r0} for every pair")
    # A PAIR'S SHEET COST TYPED, NOT LINKED. J7 and J8 read =C3; on the 14 Oct sheet J9 is a
    # typed 8 — change the sheet price in C3 and the ends pair does not follow. Said, not fixed:
    # the sheet is the estimators'.
    sc_cell = fig_cells.get("sheet_cost_gbp", "C3")
    for r in rows:
        jf = f(f"J{r}")
        if not (isinstance(jf, str) and jf.replace("$", "").upper() == f"={sc_cell}"):
            out["notes"].append(
                f"J{r} (the {str(v(f'A{r}') or f'row {r}')} pair's sheet cost) is "
                f"{'typed as ' + str(jf) if not isinstance(jf, str) else jf}, not ={sc_cell} — "
                f"a change to the sheet price in {sc_cell} does not reach that pair")
    m = re.fullmatch(r"=\s*SUM\(\s*K(\d+)\s*:\s*K(\d+)\s*\)\s*",
                     str(f(s.get("pairs_total", "K10")) or ""), re.IGNORECASE)
    if m and int(m.group(1)) == r0:
        out["pairs_in_total"] = int(m.group(2)) - r0 + 1
    else:
        problems.append(f"{s.get('pairs_total', 'K10')} is no longer SUM(K{r0}:K..): "
                        f"{f(s.get('pairs_total', 'K10'))!r}")
    tot = str(f(s.get("total", "I2")) or "").replace(" ", "").upper()
    want = {s.get("running", "E3"), s.get("setup_share", "H3"), s.get("pairs_total", "K10")}
    if not (tot.startswith("=") and set(re.findall(r"[A-Z]+\d+", tot)) == want):
        problems.append(f"the TOTAL ({s.get('total', 'I2')}) is no longer running + set-up "
                        f"share + the pairs: {f(s.get('total', 'I2'))!r}")
    setup_f = str(f(s.get("setup_share", "H3")) or "").replace(" ", "").upper()
    if setup_f != f"={fig_cells.get('setup_cost_gbp', 'D3')}/{s['inputs']['qty']}":
        problems.append(f"the set-up share ({s.get('setup_share', 'H3')}) is no longer "
                        f"set-up / quantity: {f(s.get('setup_share', 'H3'))!r}")
    figs["stated_by"] = f"{p.name} ({path})"
    figs["stated_on"] = out["modified"]
    out["figures"] = {k: val for k, val in figs.items() if val is not None}
    # THE ENGINE'S COPY needs every figure; Excel needs only the labels to be right.
    out["engine_copy_ok"] = not problems
    out["problems"] = problems
    out["ok"] = True
    out["reason"] = ""
    _READ_CACHE[key] = dict(out)
    return out


def excel_available() -> Tuple[bool, str]:
    if not settings().get("use_excel", True):
        return False, "SDI_BOX_CALCULATOR_EXCEL is off"
    if sys.platform != "win32":
        return False, "this machine has no Excel (not Windows)"
    try:
        import win32com.client  # type: ignore  # noqa: F401
    except Exception:                                                # noqa: BLE001
        return False, "pywin32 is not installed, so Excel cannot be driven"
    return True, ""


def calculate(read: Dict[str, Any], length_mm: float, breadth_mm: float, height_mm: float,
              quantities: Iterable[int]) -> Tuple[Dict[int, Dict[str, Any]], str]:
    """The sheet calculating the box at each quantity: {qty: {total, running, setup_share,
    pairs: [{per_sheet, cost}]}}, and '' — or {}, and why Excel could not do it.

    One Excel session for all the quantities. A COPY of the file is calculated; the
    estimators' own file is opened by nobody but them."""
    path = read.get("path")
    if not (read.get("ok") and path):
        return {}, read.get("reason") or "the sheet was not read"
    ok, why = excel_available()
    if not ok:
        return {}, why
    try:
        mtime = Path(path).stat().st_mtime
    except OSError as exc:
        return {}, f"cannot read {path}: {exc}"
    key0 = (str(Path(path).resolve()), mtime, float(length_mm), float(breadth_mm),
            float(height_mm))
    qs = sorted({max(1, int(q)) for q in quantities})
    have = {q: _CALC_CACHE[key0 + (q,)] for q in qs if key0 + (q,) in _CALC_CACHE}
    todo = [q for q in qs if q not in have]
    if not todo:
        return have, ""
    s = settings()
    rows = list(s.get("pair_rows") or (7, 8, 9))
    tmpdir = tempfile.mkdtemp(prefix="sdi_boxcalc_")
    copy = Path(tmpdir) / Path(path).name
    try:
        shutil.copy2(path, copy)
        from wep_readback_from_xlsx import _close_excel, _open_xlsx_excel_com, \
            with_busy_retry

        def _run() -> Dict[int, Dict[str, Any]]:
            excel, com_wb = _open_xlsx_excel_com(copy, prime_sheet=read.get("sheet_name"),
                                                 read_only=True)
            try:
                ws = com_wb.Worksheets(read.get("sheet_name"))
                got: Dict[int, Dict[str, Any]] = {}
                ins = s["inputs"]
                ws.Range(ins["length"]).Value = float(length_mm)
                ws.Range(ins["breadth"]).Value = float(breadth_mm)
                ws.Range(ins["height"]).Value = float(height_mm)
                for q in todo:
                    ws.Range(ins["qty"]).Value = int(q)
                    try:
                        excel.Calculate()
                    except Exception:                                # noqa: BLE001
                        pass
                    got[q] = {
                        "total": _num(ws.Range(s["total"]).Value),
                        "running": _num(ws.Range(s["running"]).Value),
                        "setup_share": _num(ws.Range(s["setup_share"]).Value),
                        "pairs": [{"per_sheet": _num(ws.Range(f"H{r}").Value),
                                   "cost": _num(ws.Range(f"K{r}").Value)} for r in rows],
                    }
                return got
            finally:
                _close_excel(excel, com_wb, save=False)

        got, _, _ = with_busy_retry(_run, "box calculator sheet")
    except Exception as exc:                                         # noqa: BLE001
        return {}, f"Excel could not calculate the sheet ({type(exc).__name__}: {exc})"
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    for q, res in got.items():
        if res.get("total") is None:
            return {}, (f"the sheet's TOTAL came back empty at {q} off — a cell error or a "
                        f"panel that does not fit its sheet")
        _CALC_CACHE[key0 + (q,)] = res
        have[q] = res
    return have, ""
