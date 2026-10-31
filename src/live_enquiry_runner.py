#!/usr/bin/env python3
"""
The Live Enquiry runner (D-420): the estimators' Live Enquiry workbook drives AI estimates,
one after another.

    python src\\live_enquiry_runner.py scan                  what each row would do; runs nothing
    python src\\live_enquiry_runner.py run                   queue every ready row, one at a time
    python src\\live_enquiry_runner.py run --max 1           just the next one
    python src\\live_enquiry_runner.py watch                 run, wait, read the sheet again
    python src\\live_enquiry_runner.py list                  what has run, and what it came to
    python src\\live_enquiry_runner.py qty 8188-08 1,5,10,50 quantities the sheet does not give
    python src\\live_enquiry_runner.py retry 8188-08         let a failed job try again
    python src\\live_enquiry_runner.py baseline              jobs run by hand before: leave them

A ROW IS A JOB. Customer names a folder under the Live Enquiry root, Drawing No. names a
folder under that, and the drawings directly in it are the pack — never its sub-folders.

    <root>\\M&S\\8188-08\\0348503_8188-08-GA_..._REV_H.PDF

The older layout is found too: a pack at the root named by its exact drawing number,
"12633-10-GA-AvantiConsumableHolderandChillerDisplay". The customer's folder is looked in first.

THE PORTAL DOES THE ESTIMATE. Every ready row is queued through POST /api/estimate exactly as a
person queues it from the page — same staging, same runner, same deliverables, filed to the
same place — and the next row is queued only when that one has finished. Nothing here prices
anything; if this file has an opinion about costing, that is a bug.

THE WORKBOOK IS READ, AND WRITTEN IN ONE PLACE ONLY. It is the estimators' working file and is
often open. What ran, when, on which pack and at what figure is kept in this runner's own
ledger, and every scan writes a status table beside it (live_enquiry_status.csv) saying, for
every row, what it did and why. The one write (D-608a): AI Estimate Self-Running sets the AI
Live Enquiry sheet's AICompleted to YES on a row whose run finished, and only while the
workbook is not open in Excel (mark_completed).

A ROW THAT CANNOT RUN SAYS WHY, AND WHO MOVES IT:
    not asked   AI CHECK is not YES
    held        a person has to act: a note on the row (WAIT FOR DRAWINGS), no single drawing
                number, two folders that both fit, no quantity stated
    waiting     nothing to do yet: the customer or drawing folder is not on the share, or has
                no drawings in it. Picked up by itself on the scan after it appears.
    done        estimated on this pack at these quantities; a new revision dropped into the
                folder, or new quantities, makes it ready again
    failed      the last run failed; `retry` lets it go again

WHERE THE DRAWINGS COME FROM is a list of sources, tried in order. Today it is the Live Enquiry
share. The design area, through the endpoint James will provide (SDI_DESIGN_AREA_ENDPOINT),
becomes the second: a row whose folder is missing asks it for the pack. Until that endpoint
exists the source says so and finds nothing. (src/live_enquiry_collector.py is the earlier,
provisional search of W:\\Production; it copies drawings and runs nothing.)

Settings: src/config.py, LIVE_ENQUIRY_* — the workbook path (SDI_LIVE_ENQUIRY_WORKBOOK), the
root, the vocabularies. The portal is SDI_SERVER and SDI_API_KEY, the runner's own names.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config  # noqa: E402

try:
    from enquiry import DRAWING_EXTENSIONS  # the engine's own idea of a drawing
except Exception:                            # noqa: BLE001
    DRAWING_EXTENSIONS = {".pdf", ".dxf", ".dwg"}

LEDGER_SCHEMA = "live_enquiry_ledger.v1"
_IGNORED_FILES = re.compile(r"^(?:~\$|\.|thumbs\.db$|desktop\.ini$)", re.IGNORECASE)


# ── THE SHEET ────────────────────────────────────────────────────────────────────────
@dataclass
class EnquiryRow:
    row: int                                  # the sheet's own row number
    customer: str
    drawing: str
    description: str = ""
    received: Optional[date] = None
    due: Optional[date] = None
    ai_check: str = ""
    estimator: str = ""
    account_manager: str = ""
    quantities: str = ""
    notes: str = ""
    # THE AI LIVE ENQUIRY SHEET (D-608): its AICompleted answer, and whether the sheet has that
    # column at all. A sheet with it runs by its own rule; one without keeps D-420's.
    ai_completed: str = ""
    completed_column: bool = False


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _as_date(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _open_workbook_bytes(path: Path) -> io.BytesIO:
    """The workbook's bytes, decrypted when it carries a password.

    Read into memory first: a workbook being saved on the share while it is read comes back
    half-written, and a copy taken in one read is at least one consistent version of it.
    """
    data = Path(path).read_bytes()
    if data[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":       # a plain .xlsx is a zip
        return io.BytesIO(data)
    # AN OLE CONTAINER: an encrypted .xlsx, or an old .xls. Either way openpyxl cannot read it.
    password = os.getenv("SDI_LIVE_ENQUIRY_PASSWORD", "")
    if not password:
        raise RuntimeError(
            f"{path} is password-protected (or is an old .xls). Set SDI_LIVE_ENQUIRY_PASSWORD "
            f"in .env, or point SDI_LIVE_ENQUIRY_WORKBOOK at an unprotected .xlsx copy.")
    try:
        import msoffcrypto  # type: ignore
    except ImportError as exc:
        raise RuntimeError("The workbook is encrypted and msoffcrypto-tool is not installed: "
                           "pip install msoffcrypto-tool") from exc
    out = io.BytesIO()
    office = msoffcrypto.OfficeFile(io.BytesIO(data))
    office.load_key(password=password)
    office.decrypt(out)
    out.seek(0)
    return out


def _columns(header: Sequence[Any]) -> Dict[str, int]:
    found: Dict[str, int] = {}
    for idx, cell in enumerate(header):
        text = _text(cell).upper()
        if not text:
            continue
        for name, pattern in config.LIVE_ENQUIRY_COLUMNS:
            if name not in found and re.search(pattern, text):
                found[name] = idx
                break
    return found


def _header(table: Sequence[Sequence[Any]]) -> Tuple[int, Dict[str, int]]:
    """(index of the header row, column of each field): the first row naming both a Customer
    and a Drawing column. One function, so reading the sheet and writing AICompleted back
    find the same columns."""
    for i, line in enumerate(table[:15]):
        cols = _columns(line)
        if "customer" in cols and "drawing" in cols:
            return i, cols
    raise ValueError("No header row naming both a Customer and a Drawing No. column.")


def rows_from_table(table: Sequence[Sequence[Any]]) -> List[EnquiryRow]:
    """Enquiry rows from a sheet's cells. The header is the first row naming both a Customer
    and a Drawing column; a cell under no header is a note on its row."""
    header_at, cols = _header(table)
    header = table[header_at]
    headed = {i for i, cell in enumerate(header) if _text(cell)}
    out: List[EnquiryRow] = []
    for offset, line in enumerate(table[header_at + 1:], start=header_at + 2):
        def cell(name: str) -> Any:
            i = cols.get(name)
            return line[i] if i is not None and i < len(line) else None
        customer, drawing = _text(cell("customer")), _text(cell("drawing"))
        if not customer and not drawing:
            continue
        notes = " | ".join(_text(v) for i, v in enumerate(line)
                           if i not in headed and _text(v))
        out.append(EnquiryRow(
            row=offset, customer=customer, drawing=drawing,
            description=_text(cell("description")), received=_as_date(cell("received")),
            due=_as_date(cell("due")), ai_check=_text(cell("ai_check")),
            estimator=_text(cell("estimator")), account_manager=_text(cell("account_manager")),
            quantities=_text(cell("quantities")), notes=notes,
            ai_completed=_text(cell("ai_completed")), completed_column="ai_completed" in cols))
    return out


def completed_word(text: str) -> str:
    """An AICompleted answer as one word: "Re-Run", "re run" and "RERUN" are all RERUN."""
    return re.sub(r"[\s\-_]+", "", str(text or "").upper())


def pick_sheet(names: Sequence[str], wanted: str = "") -> str:
    """The sheet to read (D-608a): the one asked for, else the first of
    config.LIVE_ENQUIRY_SHEET_NAMES the workbook has, else its only sheet. Never its first tab
    by default: the AI Live Enquiry workbook's first tab is "Archived"."""
    by_fold = {str(n).casefold(): str(n) for n in names}
    wanted = (wanted or config.LIVE_ENQUIRY_SHEET or "").strip()
    if wanted:
        if wanted.casefold() in by_fold:
            return by_fold[wanted.casefold()]
        raise ValueError(f"The workbook has no sheet '{wanted}'; its sheets are "
                         f"{', '.join(names)}. Set SDI_LIVE_ENQUIRY_SHEET to one of them.")
    for name in config.LIVE_ENQUIRY_SHEET_NAMES:
        if name.casefold() in by_fold:
            return by_fold[name.casefold()]
    if len(names) == 1:
        return str(names[0])
    raise ValueError(f"The workbook has sheets {', '.join(names)} and none is called "
                     f"{' or '.join(config.LIVE_ENQUIRY_SHEET_NAMES)}. Set SDI_LIVE_ENQUIRY_SHEET "
                     f"to the sheet the jobs are on.")


def read_sheet(path: Path, sheet: str = "") -> List[EnquiryRow]:
    import openpyxl
    wb = openpyxl.load_workbook(_open_workbook_bytes(path), read_only=True, data_only=True)
    try:
        ws = wb[pick_sheet(wb.sheetnames, sheet)]
        return rows_from_table([list(r) for r in ws.iter_rows(values_only=True)])
    finally:
        wb.close()


def mark_completed(path: Path, row: "EnquiryRow", sheet: str = "") -> Tuple[bool, str]:
    """Write config.LIVE_ENQUIRY_COMPLETED_MARK into the row's AICompleted cell (D-608a): JG,
    31 Oct 2026, "when we complete the run for that row, we populate the AICompleted column J
    with YES, so it is not re-ran unless re-set". (True, what was written) or (False, why not).

    THE ONE WRITE TO THE WORKBOOK, and it is careful about whose file it is:
      * not while the workbook is open in Excel (its ~$ lock file is beside it): Excel would
        save over the YES, or the person would be editing a file changed under them. The
        caller tries again on a later read;
      * only on the row that still holds this customer and drawing. A row moved by a sort or
        an insert is found again by its customer and drawing; two rows holding the job, or
        none, are left alone and said;
      * formulas, formats and every other cell are kept (openpyxl loads the formulas, not
        their values), and the file is replaced in one step, never half-written.
    A password-protected workbook cannot be written back, and says so."""
    import openpyxl
    path = Path(path)
    lock = path.with_name("~$" + path.name)
    if lock.exists():
        return False, (f"the workbook is open in Excel ({lock.name} beside it); AICompleted is "
                       f"written on a later read, once it is closed")
    try:
        if path.read_bytes()[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            return False, "the workbook is password-protected, so AICompleted cannot be written"
        wb = openpyxl.load_workbook(path)
    except OSError as exc:
        return False, f"the workbook could not be opened to write AICompleted ({exc})"
    try:
        ws = wb[pick_sheet(wb.sheetnames, sheet)]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        header_at, cols = _header(table)
        if "ai_completed" not in cols:
            return False, "the sheet has no AICompleted column"
        want = (customer_key(row.customer), drawing_key(row.drawing))

        def holds(n: int) -> bool:
            line = table[n - 1] if 0 < n <= len(table) else []
            def cell(name: str) -> Any:
                i = cols.get(name)
                return line[i] if i is not None and i < len(line) else None
            return (customer_key(_text(cell("customer"))),
                    drawing_key(_text(cell("drawing")))) == want

        target = row.row if holds(row.row) else None
        if target is None:
            found = [n for n in range(header_at + 2, len(table) + 1) if holds(n)]
            if len(found) != 1:
                return False, (f"row {row.row} no longer holds {row.drawing} for {row.customer}, "
                               f"and {len(found) or 'no'} other row(s) do; AICompleted not "
                               f"written")
            target = found[0]
        ws.cell(row=target, column=cols["ai_completed"] + 1,
                value=config.LIVE_ENQUIRY_COMPLETED_MARK)
        tmp = path.with_name(path.stem + ".sdi-writing" + path.suffix)
        wb.save(tmp)
    except (OSError, ValueError) as exc:
        return False, f"AICompleted could not be written ({exc})"
    finally:
        wb.close()
    try:
        os.replace(tmp, path)
    except OSError as exc:
        try:
            tmp.unlink()
        except OSError:
            pass
        return False, (f"the workbook could not be replaced ({exc}); AICompleted is written on "
                       f"a later read")
    return True, (f"AICompleted on row {target} set to {config.LIVE_ENQUIRY_COMPLETED_MARK}")


# ── NAMES ON THE SHARE ───────────────────────────────────────────────────────────────
def customer_key(name: str) -> str:
    """M&S, m & s and "M&S " are one customer; punctuation other than & is not part of it."""
    return re.sub(r"[^0-9a-z&]+", "", str(name or "").casefold())


def drawing_key(name: str) -> str:
    """A drawing number as the folder and the sheet can both spell it: case, spaces and the
    underscore/hyphen difference removed, and the sheet-type tail (GA, SA2) taken off."""
    text = re.sub(r"\s+", " ", str(name or "").strip().upper().replace("_", "-"))
    text = re.sub(config.LIVE_ENQUIRY_SHEET_SUFFIX, "", text)
    return text.replace(" ", "")


def _lead_token(folder_name: str) -> str:
    """The drawing number a folder is named by: "12633-01-GA Wine lifter" -> "12633-01-GA"."""
    return re.split(r"\s+-\s+|\s+|\(", folder_name.strip(), maxsplit=1)[0]


def names_the_drawing(folder_name: str, key: str) -> bool:
    """A folder named by this drawing: its number, then nothing, a sheet tail or words.
    "12633-10-GA-AvantiConsumableHolder" names 12633-10-GA; "12633-10-02-GA" does not — a
    digit after the number is another drawing, never a description."""
    if not key:
        return False
    if drawing_key(_lead_token(folder_name)) == key or drawing_key(folder_name) == key:
        return True
    flat = re.sub(r"\s+", "", str(folder_name or "").upper().replace("_", "-"))
    if not flat.startswith(key):
        return False
    rest = flat[len(key):]
    return rest == "" or bool(re.match(r"^-(?:(?:GA|SA)\d*)?(?:-?[A-Z(]|$)", rest))


def names_the_drawing_inside(folder_name: str, key: str) -> bool:
    """The drawing number anywhere in the folder's name, as a whole number (D-608): JG's
    "pattern match for the Drawing No. in the name of the folder", for a folder named
    "0359974 Editions Footwear Podium 12678-01-GA". Not inside a longer number: the character
    before it is not a letter or digit, and what follows obeys names_the_drawing's rule, so
    "12678-01-02-GA" still does not name 12678-01."""
    if not key:
        return False
    text = re.sub(r"\s+", " ", str(folder_name or "").upper().replace("_", "-"))
    for m in re.finditer(re.escape(key), text):
        if m.start() and text[m.start() - 1].isalnum():
            continue
        rest = text[m.end():]
        if (rest == "" or rest[0] in " (" or
                re.match(r"^-(?:(?:GA|SA)\d*)?(?:-?[A-Z( ]|$)", rest)):
            return True
    return False


def _customer_words(name: str) -> List[str]:
    """A customer's name as words that survive spelling: case, punctuation, AND/&, a plural
    s and the company words (LTD, UK, PLC — config.LIVE_ENQUIRY_CUSTOMER_NOISE_WORDS) gone."""
    noise = {w.upper() for w in getattr(config, "LIVE_ENQUIRY_CUSTOMER_NOISE_WORDS", ())}
    words = []
    for w in re.findall(r"[0-9A-Z&]+", str(name or "").upper().replace(" AND ", " & ")):
        if w in noise:
            continue
        words.append(w[:-1] if len(w) > 3 and w.endswith("S") else w)
    return words


def _one_leads_the_other(a: List[str], b: List[str]) -> bool:
    """FANATICS PARIS and Fanatics, TTI and TTi Milwaukee, M & S and M&S Food: the shorter
    name, run together, is the longer's opening words run together."""
    if not a or not b:
        return False
    short, long_ = (a, b) if len("".join(a)) <= len("".join(b)) else (b, a)
    joined, ends, run = "".join(short), set(), ""
    for w in long_:
        run += w
        ends.add(len(run))
    return "".join(long_).startswith(joined) and len(joined) in ends


def customer_folders(folders: Sequence[Path], customer: str) -> List[Path]:
    """The folder(s) the sheet's Customer names — by name, no list of aliases (D-423).

    The same name however it is spelt wins outright. Otherwise every folder whose name leads,
    or is led by, the customer's: "FANATICS PARIS" finds Fanatics. A folder carrying a drawing
    number is a pack, never a customer. More than one is returned as found — the caller holds
    the row and names them; it never picks."""
    want = _customer_words(customer)
    if not want:
        return []
    named = [f for f in folders if not re.search(r"\d{4,}", f.name)]
    same = [f for f in named if "".join(_customer_words(f.name)) == "".join(want)]
    if same:
        return same
    return [f for f in named if _one_leads_the_other(want, _customer_words(f.name))]


def _subfolders(folder: Path) -> List[Path]:
    try:
        return sorted(p for p in folder.iterdir() if p.is_dir())
    except OSError:
        return []


def pack_files(folder: Path) -> List[Path]:
    """Every file directly in the drawing folder — the pack. Sub-folders are never read."""
    try:
        return sorted(p for p in folder.iterdir()
                      if p.is_file() and not _IGNORED_FILES.match(p.name))
    except OSError:
        return []


def fingerprint(files: Sequence[Path]) -> str:
    """Names, sizes and times of the pack: a new revision dropped in changes it."""
    parts = []
    for p in files:
        try:
            st = p.stat()
            parts.append(f"{p.name.casefold()}:{st.st_size}:{int(st.st_mtime)}")
        except OSError:
            parts.append(f"{p.name.casefold()}:?")
    import hashlib
    return hashlib.sha1("|".join(parts).encode("utf-8")).hexdigest()[:12]


@dataclass
class Located:
    status: str                    # found | waiting | held
    reason: str = ""
    folder: Optional[Path] = None
    client_folder_name: str = ""
    files: List[Path] = field(default_factory=list)
    source: str = ""


class LiveEnquiryShare:
    """The packs the estimators drop: <root>\\<customer>\\<drawing>."""
    name = "Live Enquiry share"

    def __init__(self, root: Path):
        self.root = Path(root)

    def find(self, row: EnquiryRow) -> Located:
        if not self.root.is_dir():
            return Located("held", f"the Live Enquiry root {self.root} cannot be read from this "
                                   f"machine")
        clients = _subfolders(self.root)
        client = customer_folders(clients, row.customer)
        if len(client) > 1:
            return Located("held", f"{len(client)} folders could be '{row.customer}': "
                                   f"{', '.join(c.name for c in client)} — leave one, or "
                                   f"spell the Customer as its folder", source=self.name)
        key = drawing_key(row.drawing)
        client_dir = client[0] if client else None
        hits = ([d for d in _subfolders(client_dir) if names_the_drawing(d.name, key)]
                if client_dir else [])
        if not hits and client_dir is not None:
            # A FOLDER THAT CARRIES THE NUMBER INSIDE ITS NAME (D-608), looked for only under
            # the customer's folder and only when no folder is named by the number: the
            # stricter match wins wherever it finds one.
            hits = [d for d in _subfolders(client_dir) if names_the_drawing_inside(d.name, key)]
        at_root = False
        if not hits:
            # THE OLDER LAYOUT: a pack at the root named by its drawing number,
            # "12633-10-GA-AvantiConsumableHolderandChillerDisplay". The number must be exact.
            hits = [d for d in clients if names_the_drawing(d.name, key)]
            at_root = bool(hits)
        if not hits:
            if client_dir is None:
                return Located("waiting", f"no '{row.customer}' folder under the Live Enquiry "
                                          f"root, and no {row.drawing} pack at the root, yet",
                               source=self.name)
            return Located("waiting", f"no {row.drawing} folder under {client_dir.name} yet",
                           client_folder_name=client_dir.name, source=self.name)
        where = "the root" if at_root else client_dir.name
        client_name = "" if at_root else client_dir.name
        if len(hits) > 1:
            return Located("held", f"{len(hits)} folders under {where} fit {row.drawing}: "
                                   f"{', '.join(h.name for h in hits)} — leave one",
                           client_folder_name=client_name, source=self.name)
        folder = hits[0]
        files = pack_files(folder)
        drawings = [p for p in files if p.suffix.lower() in DRAWING_EXTENSIONS]
        old_layout = (f" (at the root, not under a '{row.customer}' folder)" if at_root else "")
        if not drawings:
            return Located("waiting", f"{folder.name}{old_layout} has no drawings in it yet "
                                      f"(only files directly in the folder are read)",
                           folder=folder, client_folder_name=client_name, files=files,
                           source=self.name)
        return Located("found", f"{len(drawings)} drawing(s) in {folder.name}{old_layout}",
                       folder=folder, client_folder_name=client_name, files=files,
                       source=self.name)


class DesignArea:
    """The design area, through the endpoint James will provide.

    Its job, when it exists: given a row's customer and drawing number, put that drawing's
    pack into <root>\\<customer>\\<drawing> and say so — after which the share finds it like
    any other. Until the endpoint is given it finds nothing and says why.
    """
    name = "design area"

    def __init__(self, endpoint: str = ""):
        self.endpoint = (endpoint or "").strip()

    def find(self, row: EnquiryRow) -> Optional[Located]:
        if not self.endpoint:
            return None
        return Located("waiting", f"SDI_DESIGN_AREA_ENDPOINT is set ({self.endpoint}) but this "
                                  f"build does not call it yet", source=self.name)


def locate(row: EnquiryRow, sources: Sequence[Any]) -> Located:
    """The first source that has the pack; otherwise the most useful reason not."""
    first: Optional[Located] = None
    for source in sources:
        got = source.find(row)
        if got is None:
            continue
        if got.status == "found":
            return got
        if first is None:
            first = got
        elif got.status == "held" and first.status != "held":
            first = got
    return first or Located("waiting", "no source of drawings is configured")


# ── QUANTITIES ───────────────────────────────────────────────────────────────────────
def parse_quantities(text: str) -> List[int]:
    """"1, 5, 10, 50", "1/5/10/50 off", "50" -> [1, 5, 10, 50]. Order kept, repeats dropped;
    the first is the run quantity and the rest are breaks, as on the page."""
    seen: List[int] = []
    for part in re.split(r"[,;/\s]+", re.sub(r"(?i)\b(?:off|pcs|units|qty\.?)\b|:", " ",
                                             str(text or ""))):
        if part.isdigit() and int(part) >= 1 and int(part) not in seen:
            seen.append(int(part))
    return seen


def quantities_for(row: EnquiryRow, located: Located,
                   console: Dict[str, List[int]]) -> Tuple[List[int], str]:
    """The quantities and where they were stated. Nothing stated is nothing: setup is spread
    over the run quantity, so a guessed one moves every labour line, and the row is held."""
    if row.completed_column:
        # THE AI LIVE ENQUIRY SHEET'S QUANTITY IS ITS QTY COLUMN (D-608), and nothing else: not
        # the console, not a file in the folder, not the description.
        q = parse_quantities(row.quantities)
        return (q, "the sheet's QTY column") if q else ([], "")
    given = console.get(drawing_key(row.drawing))
    if given:
        return list(given), "given at the console (qty)"
    if row.quantities and parse_quantities(row.quantities):
        return parse_quantities(row.quantities), "the sheet's quantity column"
    if located.folder is not None:
        for name in config.LIVE_ENQUIRY_QTY_FILES:
            for p in located.files:
                if p.name.casefold() == name.casefold():
                    try:
                        q = parse_quantities(p.read_text(encoding="utf-8", errors="replace"))
                    except OSError:
                        q = []
                    if q:
                        return q, f"{p.name} in the drawing folder"
    for pattern in config.LIVE_ENQUIRY_QTY_IN_TEXT:
        found = [int(m) for m in re.findall(pattern, row.description.upper()) if int(m) >= 1]
        if found:
            return list(dict.fromkeys(found)), "the job description"
    return [], ""


# ── THE LEDGER ───────────────────────────────────────────────────────────────────────
def job_key(row: EnquiryRow) -> str:
    return f"{customer_key(row.customer)}|{drawing_key(row.drawing)}"


def load_ledger(path: Path) -> Dict[str, Any]:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        if isinstance(data, dict) and data.get("schema") == LEDGER_SCHEMA:
            data.setdefault("jobs", {})
            data.setdefault("quantities", {})
            return data
    except (OSError, ValueError):
        pass
    return {"schema": LEDGER_SCHEMA, "jobs": {}, "quantities": {}}


def save_ledger(path: Path, ledger: Dict[str, Any]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(ledger, indent=2, ensure_ascii=False, default=str),
                   encoding="utf-8")
    os.replace(tmp, path)               # never a half-written ledger


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


# ── THE DECISION ─────────────────────────────────────────────────────────────────────
@dataclass
class Decision:
    row: EnquiryRow
    status: str                       # ready | done | running | failed | held | waiting | not_asked
    reason: str
    located: Optional[Located] = None
    quantities: List[int] = field(default_factory=list)
    quantity_source: str = ""
    fingerprint: str = ""
    last: Dict[str, Any] = field(default_factory=dict)
    rerun: bool = False               # AICompleted says Re-Run (D-608)


def _hold_note(notes: str) -> str:
    for pattern in config.LIVE_ENQUIRY_HOLD_NOTES:
        if re.search(pattern, notes.upper()):
            return notes
    return ""


def _not_one_drawing(drawing: str) -> str:
    text = drawing.strip().upper()
    if not text or text in {w.upper() for w in config.LIVE_ENQUIRY_NOT_A_DRAWING}:
        return f"Drawing No. is '{drawing or 'blank'}' — not one drawing; the pack is made up " \
               f"by hand"
    if re.search(r"[+,;&]|\bAND\b", text):
        return f"Drawing No. '{drawing}' names more than one drawing — one row per drawing"
    return ""


def decide(row: EnquiryRow, sources: Sequence[Any], ledger: Dict[str, Any],
           seen: Optional[Dict[str, int]] = None) -> Decision:
    if row.ai_check.strip().upper() not in {w.upper() for w in config.LIVE_ENQUIRY_ASKED}:
        return Decision(row, "not_asked", f"AI CHECK is '{row.ai_check or 'blank'}'")
    # THE AI LIVE ENQUIRY SHEET'S RULE (D-608): YES, a QTY, and AICompleted NO, N or Re-Run.
    rerun = False
    if row.completed_column:
        said = completed_word(row.ai_completed)
        if said not in config.LIVE_ENQUIRY_RUN_WHEN_COMPLETED:
            return Decision(row, "not_asked",
                            f"AICompleted is '{row.ai_completed or 'blank'}' — NO, N or Re-Run "
                            f"runs it")
        if not parse_quantities(row.quantities):
            return Decision(row, "held", "no QTY on the sheet — put the quantities in the QTY "
                                         "column")
        rerun = said in config.LIVE_ENQUIRY_RERUN
    note = _hold_note(row.notes)
    if note:
        return Decision(row, "held", f"the row's note says '{note}'")
    bad = _not_one_drawing(row.drawing)
    if bad:
        return Decision(row, "held", bad)
    if not row.customer:
        return Decision(row, "held", "no Customer on the row")
    key = job_key(row)
    if seen is not None:
        if key in seen:
            return Decision(row, "held", f"the same job as row {seen[key]}")
        seen[key] = row.row
    located = locate(row, sources)
    job = ledger["jobs"].get(key) or {}
    attempts = job.get("attempts") or []
    if not attempts and job.get("baseline") == "*":
        return Decision(row, "done", f"run by hand before the runner (baselined "
                                     f"{job.get('baselined_at')})", located=located)
    if located.status != "found":
        return Decision(row, located.status, located.reason, located=located)
    if not attempts and job.get("baseline") == fingerprint(located.files):
        return Decision(row, "done", f"run by hand before the runner (baselined "
                                     f"{job.get('baselined_at')})", located=located)
    qty, qsource = quantities_for(row, located, ledger.get("quantities", {}))
    if not qty:
        return Decision(row, "held",
                        f"no quantity stated — add it to the sheet, put "
                        f"{config.LIVE_ENQUIRY_QTY_FILES[0]} in {located.folder.name}, or "
                        f"`qty {row.drawing} 1,5,10`", located=located)
    fp = fingerprint(located.files)
    last = attempts[-1] if attempts else {}
    d = Decision(row, "ready", "first run", located=located, quantities=qty,
                 quantity_source=qsource, fingerprint=fp, last=last, rerun=rerun)
    # RE-RUN RUNS ONCE PER MARKING (D-608). Until AICompleted is written back the cell still
    # says Re-Run after the run; the ledger remembers it was taken, and forgets when the cell
    # is seen saying anything else (plan()), so marking it Re-Run again runs it again.
    rerun_due = rerun and not job.get("rerun_taken")
    if not last:
        if job.get("baseline"):
            d.reason = "the pack has changed since it was baselined"
        return d
    same = last.get("fingerprint") == fp and list(last.get("quantities") or []) == qty
    if last.get("status") in {"queued", "running"}:
        d.status, d.reason = "running", f"queued as run {last.get('run_id')} at " \
                                        f"{last.get('queued_at')}"
        return d
    if rerun_due and last.get("status") in {"done", "error"}:
        d.reason = "Re-Run asked for on the sheet"
        return d
    if same and last.get("status") == "done":
        if row.completed_column and last.get("sheet_marked"):
            # RESET BY A PERSON (D-608a): this run's YES reached the sheet, and the cell now
            # asks for a run again, so somebody cleared or changed it on purpose.
            d.reason = (f"AICompleted reset since this run was marked "
                        f"{config.LIVE_ENQUIRY_COMPLETED_MARK} at {last.get('sheet_marked')}")
            return d
        d.status, d.reason = "done", f"estimated {last.get('finished_at')} — " \
                                     f"{last.get('output_path') or 'no output path'}"
        if row.completed_column:
            d.reason += (f" (AICompleted not yet written: "
                         f"{last.get('sheet_mark_note') or 'waiting for the next read'})")
        return d
    if not same:
        d.reason = ("the pack has changed since the last run"
                    if last.get("fingerprint") != fp else "the quantities have changed")
        return d
    tries = sum(1 for a in attempts if a.get("fingerprint") == fp
                and list(a.get("quantities") or []) == qty and a.get("status") != "done")
    if job.get("retry") or tries < config.LIVE_ENQUIRY_MAX_ATTEMPTS:
        d.reason = "retry asked for" if job.get("retry") else f"attempt {tries + 1}"
        return d
    d.status, d.reason = "failed", f"last run {last.get('status')}: " \
                                   f"{(last.get('error') or '').strip()[:200]} — `retry " \
                                   f"{row.drawing}` to run it again"
    return d


def baseline(rows: Sequence[EnquiryRow], sources: Sequence[Any], ledger: Dict[str, Any],
             only: Sequence[str] = ()) -> List[str]:
    """Mark the jobs already on the sheet as run, without running them.

    The sheet carries jobs estimated by hand before this runner existed; started cold it would
    queue every one of them again. A baselined job stays done until its pack changes. Only
    rows asked for (AI CHECK YES) whose pack is on the share are marked, and with `only`, just
    those drawings — which are marked even where their pack is not found."""
    want = {drawing_key(d) for d in only}
    marked: List[str] = []
    for row in rows:
        if want and drawing_key(row.drawing) not in want:
            continue
        if row.ai_check.strip().upper() not in {w.upper() for w in config.LIVE_ENQUIRY_ASKED}:
            continue
        if _not_one_drawing(row.drawing) or not row.customer:
            continue
        located = locate(row, sources)
        if located.status != "found" and not want:
            continue
        job = ledger["jobs"].setdefault(job_key(row), {
            "customer": row.customer, "drawing": row.drawing, "attempts": []})
        if job.get("attempts"):
            continue                      # it has really run; its own record stands
        # NAMED, IT IS DONE WHEREVER ITS PACK IS: "baseline 12633-01-GA" for a job estimated by
        # hand from a folder the share does not lay out as customer\drawing. "*" = any pack.
        job["baseline"] = fingerprint(located.files) if located.status == "found" else "*"
        job["baselined_at"] = _now()
        marked.append(row.drawing)
    return marked


def plan(rows: Sequence[EnquiryRow], sources: Sequence[Any],
         ledger: Dict[str, Any]) -> List[Decision]:
    """Every row's decision, ready rows first by the date the estimate is wanted, or, on the
    AI Live Enquiry sheet, in the sheet's own order from the top."""
    seen: Dict[str, int] = {}
    # A RE-RUN TAKEN IS FORGOTTEN once the cell says something else (D-608), so the next
    # Re-Run on the same row is a new request.
    for r in rows:
        if r.completed_column and completed_word(r.ai_completed) not in config.LIVE_ENQUIRY_RERUN:
            (ledger.get("jobs", {}).get(job_key(r)) or {}).pop("rerun_taken", None)
    decisions = [decide(r, sources, ledger, seen) for r in rows]
    if any(r.completed_column for r in rows):
        # THE AI LIVE ENQUIRY SHEET RUNS FROM THE TOP (D-608): ready rows in the sheet's order.
        return sorted(decisions, key=lambda d: (d.status != "ready", d.row.row))
    far = date.max
    return sorted(decisions, key=lambda d: (d.status != "ready", d.row.due or far, d.row.row))


# ── THE PORTAL ───────────────────────────────────────────────────────────────────────
class PortalError(RuntimeError):
    def __init__(self, status: int, message: str):
        super().__init__(f"{status}: {message}")
        self.status, self.message = status, message


class Portal:
    """POST /api/estimate and GET /api/estimate/{id}, as the page uses them."""

    def __init__(self, server: str, api_key: str = "", timeout: float = 60.0):
        self.server, self.api_key, self.timeout = server.rstrip("/"), api_key, timeout

    def _call(self, method: str, path: str, body: Optional[Dict[str, Any]] = None) -> Any:
        import urllib.error
        import urllib.request
        data = json.dumps(body).encode("utf-8") if body is not None else None
        req = urllib.request.Request(self.server + path, data=data, method=method)
        req.add_header("Content-Type", "application/json")
        if self.api_key:
            req.add_header("X-SDI-Key", self.api_key)
        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                return json.loads(resp.read().decode("utf-8") or "null")
        except urllib.error.HTTPError as exc:
            text = exc.read().decode("utf-8", errors="replace")
            try:
                text = json.loads(text).get("detail", text)
            except (ValueError, AttributeError):
                pass
            raise PortalError(exc.code, str(text)) from None
        except (urllib.error.URLError, OSError) as exc:
            raise PortalError(0, f"the portal at {self.server} did not answer ({exc})") from None

    def runners(self) -> Any:
        return self._call("GET", "/api/estimate/runners")

    def start(self, body: Dict[str, Any]) -> Dict[str, Any]:
        return self._call("POST", "/api/estimate", body)

    def status(self, run_id: str) -> Dict[str, Any]:
        return self._call("GET", f"/api/estimate/{run_id}")

    def self_running(self, report: Dict[str, Any]) -> Dict[str, Any]:
        """Report to the portal's AI Estimate Self-Running switch and hear whether it is on
        (D-588). The answer is the only thing that lets the self-running watcher queue a job."""
        return self._call("POST", "/api/estimate/self-running/heartbeat", report)


def request_body(d: Decision) -> Dict[str, Any]:
    """What the page sends for a pack estimate. The client is the share's folder name — the
    spelling the output tree is filed under — and the drawing number is the row's."""
    return {
        "client": d.located.client_folder_name or d.row.customer,
        "drawing_number": d.row.drawing,
        "units": d.quantities[0],
        "quantity_breaks": d.quantities[1:],
        "job_folder": str(d.located.folder),
        "files": [],
        "deliverables": True,
        "method": "both",
        "email_to": config.LIVE_ENQUIRY_EMAIL_TO or None,
        "email_quote": False,
    }


def _attempt_from(snapshot: Dict[str, Any], attempt: Dict[str, Any]) -> Dict[str, Any]:
    attempt.update({
        "status": snapshot.get("status") or attempt.get("status"),
        "error": snapshot.get("error") or "",
        "output_path": snapshot.get("output_path") or attempt.get("output_path"),
        "engine_price_gbp": snapshot.get("engine_price_gbp"),
        "seconds": snapshot.get("seconds"),
        "deliverables": [x.get("name") for x in (snapshot.get("deliverables") or [])
                         if isinstance(x, dict)],
    })
    if attempt["status"] in {"done", "error"}:
        attempt.setdefault("finished_at", _now())
    return attempt


def refresh_open_runs(ledger: Dict[str, Any], portal: Portal) -> None:
    """A run left queued or running by an earlier cycle (or an earlier process) is asked about
    before anything new is decided, so the ledger never acts on a stale 'running'."""
    for job in ledger["jobs"].values():
        for attempt in job.get("attempts") or []:
            if attempt.get("status") not in {"queued", "running"}:
                continue
            try:
                _attempt_from(portal.status(attempt["run_id"]), attempt)
            except PortalError as exc:
                if exc.status == 404:
                    attempt.update(status="error", finished_at=_now(),
                                   error="the portal no longer knows this run (restarted?)")


def run_one(d: Decision, portal: Portal, ledger: Dict[str, Any], ledger_path: Path,
            say: Callable[[str], None] = print, poll_seconds: float = 30.0,
            sleep: Callable[[float], None] = time.sleep,
            clock: Callable[[], float] = time.time) -> Dict[str, Any]:
    """Queue one row and wait for it to finish. The attempt is in the ledger from the moment it
    is queued, so a runner stopped mid-wait finds it again on the next start."""
    attempt: Dict[str, Any] = {
        "queued_at": _now(), "status": "queued", "row": d.row.row,
        "folder": str(d.located.folder), "fingerprint": d.fingerprint,
        "quantities": list(d.quantities), "quantity_source": d.quantity_source,
        # A ROW ON THE AI LIVE ENQUIRY SHEET HAS AICompleted WRITTEN BACK when it finishes
        # (D-608a); kept on the attempt so a run that finishes while nobody waits on it is
        # still written on a later read.
        "marks_sheet": bool(d.row.completed_column),
    }
    started = portal.start(request_body(d))           # refused: nothing is recorded here
    job = ledger["jobs"].setdefault(job_key(d.row), {
        "customer": d.row.customer, "drawing": d.row.drawing, "attempts": []})
    attempt["run_id"] = started.get("run_id")
    attempt["output_path"] = started.get("output_path")
    job["attempts"].append(attempt)
    job.pop("retry", None)
    if d.rerun:
        job["rerun_taken"] = True                 # this Re-Run has been answered (D-608)
    save_ledger(ledger_path, ledger)
    say(f"   queued {d.row.drawing} ({d.row.customer}) at {'/'.join(map(str, d.quantities))} "
        f"— run {attempt['run_id']}")
    deadline = clock() + config.LIVE_ENQUIRY_RUN_TIMEOUT_MINUTES * 60
    while True:
        sleep(poll_seconds)
        try:
            snap = portal.status(attempt["run_id"])
        except PortalError as exc:
            if exc.status == 404:
                attempt.update(status="error", finished_at=_now(),
                               error="the portal no longer knows this run (restarted?)")
                save_ledger(ledger_path, ledger)
                return attempt
            say(f"   (portal did not answer: {exc.message}; still waiting)")
            snap = {}
        if snap:
            _attempt_from(snap, attempt)
            save_ledger(ledger_path, ledger)
        if attempt["status"] in {"done", "error"}:
            return attempt
        _touch_lock(ledger_path)
        if clock() > deadline:
            say(f"   {d.row.drawing} has run past {config.LIVE_ENQUIRY_RUN_TIMEOUT_MINUTES} "
                f"min; not waiting on it any longer this cycle (the portal keeps it)")
            return attempt


def run_ready(decisions: Sequence[Decision], portal: Portal, ledger: Dict[str, Any],
              ledger_path: Path, limit: Optional[int] = None,
              say: Callable[[str], None] = print,
              before_each: Optional[Callable[[Decision], bool]] = None,
              after_each: Optional[Callable[[Decision, Dict[str, Any]], None]] = None,
              **wait: Any) -> List[Dict[str, Any]]:
    """Every ready row, one after another. Stops when no runner is connected or a run is
    still going at its time limit — a queue built behind work that is not moving helps no
    one. `before_each` is asked before each row is queued (D-588): the self-running watcher
    reports the row it is about to run and answers False once the portal's switch is off, so
    the run in progress finishes and nothing more is queued. `after_each` is told each
    attempt as it ends (D-608a): the watcher writes AICompleted back for a finished one."""
    done: List[Dict[str, Any]] = []
    for d in decisions:
        if d.status != "ready":
            continue
        if limit is not None and len(done) >= limit:
            break
        if before_each is not None and not before_each(d):
            say("   stopped: AI Estimate Self-Running was switched off on the portal")
            break
        try:
            attempt = run_one(d, portal, ledger, ledger_path, say=say, **wait)
        except PortalError as exc:
            if exc.status in (0, 503):
                say(f"   stopped: {exc.message}")
                break
            if exc.status == 409:
                say(f"   {d.row.drawing}: already queued or running on the portal — "
                    f"left to it")
                continue
            job = ledger["jobs"].setdefault(job_key(d.row), {
                "customer": d.row.customer, "drawing": d.row.drawing, "attempts": []})
            job["attempts"].append({
                "queued_at": _now(), "finished_at": _now(), "status": "error",
                "row": d.row.row, "folder": str(d.located.folder),
                "fingerprint": d.fingerprint, "quantities": list(d.quantities),
                "error": f"the portal refused it: {exc.message}"})
            save_ledger(ledger_path, ledger)
            say(f"   {d.row.drawing}: the portal refused it — {exc.message}")
            continue
        done.append(attempt)
        price = attempt.get("engine_price_gbp")
        say(f"   {d.row.drawing}: {attempt['status']}"
            + (f" — £{price:,.2f} a unit at {d.quantities[0]} off" if price else "")
            + (f" — {attempt.get('error')}" if attempt.get("error") else ""))
        if after_each is not None:
            after_each(d, attempt)
        if attempt["status"] not in {"done", "error"}:
            break
    return done


# ── WHAT IT SAYS ─────────────────────────────────────────────────────────────────────
STATUS_COLUMNS = ("row", "customer", "drawing", "description", "estimator", "due", "status",
                  "reason", "folder", "quantities", "quantity_source", "last_run",
                  "last_status", "unit_gbp", "output")


def status_rows(decisions: Sequence[Decision]) -> List[Dict[str, Any]]:
    out = []
    for d in sorted(decisions, key=lambda x: x.row.row):
        out.append({
            "row": d.row.row, "customer": d.row.customer, "drawing": d.row.drawing,
            "description": d.row.description, "estimator": d.row.estimator,
            "due": d.row.due.isoformat() if d.row.due else "", "status": d.status,
            "reason": d.reason,
            "folder": str(d.located.folder) if d.located and d.located.folder else "",
            "quantities": "/".join(map(str, d.quantities)),
            "quantity_source": d.quantity_source,
            "last_run": d.last.get("finished_at") or d.last.get("queued_at") or "",
            "last_status": d.last.get("status") or "",
            "unit_gbp": d.last.get("engine_price_gbp") or "",
            "output": d.last.get("output_path") or "",
        })
    return out


def write_status(decisions: Sequence[Decision], ledger_path: Path) -> Path:
    path = Path(ledger_path).with_name("live_enquiry_status.csv")
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=STATUS_COLUMNS)
            w.writeheader()
            w.writerows(status_rows(decisions))
    except PermissionError:
        # Open in Excel. The scan is still right; only the table could not be refreshed.
        print(f"   [live] {path.name} is open elsewhere; not refreshed this scan", flush=True)
    return path


def print_plan(decisions: Sequence[Decision], say: Callable[[str], None] = print) -> None:
    order = ("ready", "running", "failed", "held", "waiting", "done", "not_asked")
    for status in order:
        group = [d for d in decisions if d.status == status]
        if not group:
            continue
        say(f"\n{status.upper().replace('_', ' ')} ({len(group)})")
        for d in group:
            qty = f" @ {'/'.join(map(str, d.quantities))}" if d.quantities else ""
            say(f"  row {d.row.row:>3}  {d.row.customer:<16.16} {d.row.drawing:<18.18}"
                f"{qty:<14} {d.reason}")


# ── ONE AT A TIME ────────────────────────────────────────────────────────────────────
def _lock_path(ledger_path: Path) -> Path:
    return Path(ledger_path).with_name("live_enquiry.lock")


def _touch_lock(ledger_path: Path) -> None:
    try:
        _lock_path(ledger_path).write_text(f"{os.getpid()} {time.time():.0f}", encoding="utf-8")
    except OSError:
        pass


def take_lock(ledger_path: Path, stale_after: float = 600.0) -> bool:
    """Two runners reading one sheet would queue every job twice. A lock not touched for ten
    minutes belongs to a process that has gone, and is taken over."""
    path = _lock_path(ledger_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        pid, at = path.read_text(encoding="utf-8").split()
        if int(pid) != os.getpid() and time.time() - float(at) < stale_after:
            return False
    except (OSError, ValueError):
        pass
    _touch_lock(ledger_path)
    return True


def release_lock(ledger_path: Path) -> None:
    try:
        path = _lock_path(ledger_path)
        if path.read_text(encoding="utf-8").split()[0] == str(os.getpid()):
            path.unlink()
    except (OSError, IndexError):
        pass


# ── THE COMMAND LINE ─────────────────────────────────────────────────────────────────
def _settings(args: argparse.Namespace) -> Dict[str, Any]:
    return {
        "workbook": Path(args.workbook or config.LIVE_ENQUIRY_WORKBOOK)
        if (args.workbook or config.LIVE_ENQUIRY_WORKBOOK) else None,
        "root": Path(args.root or config.LIVE_ENQUIRY_ROOT),
        "ledger": Path(args.ledger or config.LIVE_ENQUIRY_LEDGER),
    }


def _sources(root: Path) -> List[Any]:
    return [LiveEnquiryShare(root), DesignArea(config.DESIGN_AREA_ENDPOINT)]


def scan(s: Dict[str, Any], ledger: Dict[str, Any]) -> List[Decision]:
    if not s["workbook"]:
        raise SystemExit("Set SDI_LIVE_ENQUIRY_WORKBOOK (in .env) to the Live Enquiry "
                         "workbook, or pass --workbook.")
    rows = read_sheet(s["workbook"])
    decisions = plan(rows, _sources(s["root"]), ledger)
    write_status(decisions, s["ledger"])
    return decisions


def _cycle(args: argparse.Namespace, s: Dict[str, Any], portal: Portal) -> None:
    ledger = load_ledger(s["ledger"])
    refresh_open_runs(ledger, portal)
    save_ledger(s["ledger"], ledger)
    decisions = scan(s, ledger)
    print(f"\n[{_now()}] {s['workbook']}")
    print_plan(decisions)
    ready = [d for d in decisions if d.status == "ready"]
    if not ready:
        print("\nNothing ready to run.")
        return
    print(f"\nRunning {min(len(ready), args.max or len(ready))} of {len(ready)} ready, "
          f"one after another:")
    run_ready(decisions, portal, ledger, s["ledger"], limit=args.max)
    write_status(plan(read_sheet(s["workbook"]), _sources(s["root"]), ledger), s["ledger"])


def _log_to(path: str) -> None:
    """Everything printed goes to `path`, line-buffered and appended — the scheduled task
    (D-425) runs under pythonw with no console, where print() would go nowhere."""
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fh = open(path, "a", encoding="utf-8", buffering=1)
    sys.stdout = fh
    sys.stderr = fh
    print(f"\n===== {_now()} live enquiry runner, pid {os.getpid()} =====", flush=True)


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Drive AI estimates from the Live Enquiry sheet.")
    ap.add_argument("command",
                    choices=("scan", "run", "watch", "list", "qty", "retry", "baseline"))
    ap.add_argument("args", nargs="*")
    ap.add_argument("--workbook", default="")
    ap.add_argument("--root", default="")
    ap.add_argument("--ledger", default="")
    ap.add_argument("--server", default=os.getenv("SDI_SERVER", "http://10.0.0.5:8071"))
    ap.add_argument("--api-key", default=os.getenv("SDI_API_KEY", ""))
    ap.add_argument("--max", type=int, default=None, help="run at most this many this cycle")
    ap.add_argument("--log", default="",
                    help="append everything printed to this file (the scheduled task runs "
                         "windowless, so this is the only place its output goes)")
    args = ap.parse_args(argv)
    if args.log:
        _log_to(args.log)
    s = _settings(args)

    if args.command == "qty":
        if len(args.args) != 2 or not parse_quantities(args.args[1]):
            raise SystemExit("usage: qty <drawing> 1,5,10,50")
        ledger = load_ledger(s["ledger"])
        ledger["quantities"][drawing_key(args.args[0])] = parse_quantities(args.args[1])
        save_ledger(s["ledger"], ledger)
        print(f"{args.args[0]}: {'/'.join(map(str, parse_quantities(args.args[1])))} — used "
              f"before the sheet's own column on the next scan")
        return 0
    if args.command == "retry":
        if len(args.args) != 1:
            raise SystemExit("usage: retry <drawing>")
        ledger = load_ledger(s["ledger"])
        hits = [j for k, j in ledger["jobs"].items()
                if k.split("|", 1)[1] == drawing_key(args.args[0])]
        for j in hits:
            j["retry"] = True
        save_ledger(s["ledger"], ledger)
        print(f"{args.args[0]}: " + ("will run on the next cycle" if hits
                                     else "not in the ledger — it has never run"))
        return 0 if hits else 1
    if args.command == "baseline":
        if not s["workbook"]:
            raise SystemExit("Set SDI_LIVE_ENQUIRY_WORKBOOK, or pass --workbook.")
        ledger = load_ledger(s["ledger"])
        marked = baseline(read_sheet(s["workbook"]), _sources(s["root"]), ledger, args.args)
        save_ledger(s["ledger"], ledger)
        print(f"Baselined {len(marked)} job(s) as already run: {', '.join(marked) or 'none'}")
        return 0
    if args.command == "list":
        ledger = load_ledger(s["ledger"])
        for job in ledger["jobs"].values():
            if job.get("baseline") and not job.get("attempts"):
                print(f"{job.get('baselined_at', ''):<20} {job['customer']:<16.16} "
                      f"{job['drawing']:<18.18} {'':<12} baselined (run by hand before the "
                      f"runner)")
            for a in job.get("attempts") or []:
                price = a.get("engine_price_gbp")
                print(f"{a.get('queued_at', ''):<20} {job['customer']:<16.16} "
                      f"{job['drawing']:<18.18} {'/'.join(map(str, a.get('quantities') or [])):<12}"
                      f" {a.get('status', ''):<8} {('£%.2f' % price) if price else '':>10}  "
                      f"{a.get('output_path') or a.get('error') or ''}")
        return 0
    if args.command == "scan":
        ledger = load_ledger(s["ledger"])
        decisions = scan(s, ledger)
        print_plan(decisions)
        print(f"\nStatus table: {Path(s['ledger']).with_name('live_enquiry_status.csv')}")
        return 0

    portal = Portal(args.server, args.api_key)
    if not take_lock(s["ledger"]):
        raise SystemExit(f"Another Live Enquiry runner holds {_lock_path(s['ledger'])}. Stop it "
                         f"first; two would queue every job twice.")
    try:
        while True:
            try:
                _cycle(args, s, portal)
            except SystemExit:
                raise
            except Exception as exc:                          # noqa: BLE001
                # One bad read of the sheet (saved mid-read, share blipped) must not end the
                # watch; it is said, and the next cycle reads it again.
                print(f"\n[{_now()}] cycle failed: {type(exc).__name__}: {exc}", flush=True)
                if args.command == "run":
                    return 1
            if args.command == "run":
                return 0
            wait = config.LIVE_ENQUIRY_POLL_MINUTES * 60
            print(f"\nNext read of the sheet in {config.LIVE_ENQUIRY_POLL_MINUTES} min.",
                  flush=True)
            while wait > 0:
                _touch_lock(s["ledger"])
                time.sleep(min(60, wait))
                wait -= 60
    finally:
        release_lock(s["ledger"])


if __name__ == "__main__":
    sys.exit(main())
