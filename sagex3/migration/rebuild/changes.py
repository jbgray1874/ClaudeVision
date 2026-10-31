"""What changed between two runs - the write-up that goes with every rebuild.

Compares the data tab of each mapped spreadsheet in the previous run's folder with the same file
in the new run's folder, by the template's key, and writes:
  * WHAT_CHANGED.md   - a plain summary per file: rows before / after, added, removed, changed,
                        and the fields that changed most, with a few examples each
  * What_changed.xlsx - every added, removed and changed row, one tab per file

Usage:
    python changes.py <previous run folder> <new run folder>
"""
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

# file name stem -> (data sheet, first data row, key columns (0-based), what one row is)
# None as the key compares whole rows (files whose rows have no natural key).
FILES = {
    "ZSDIITM_products_MAPPED": ("Data with Headings", 4, (1,), "product"),
    "YITC_product_costs_MAPPED": ("YITC", 3, (2,), "product cost"),
    "YA00408BOM_boms_MAPPED": ("Data with Headings", 3, (0, 1, 7), "BOM line"),
    "ZSDIROU2_routes_MAPPED": ("ZSDROU2", 4, (0, 1, 6), "route operation"),
    "ZSDIITMSUP_product_suppliers_MAPPED": ("Template", 3, (0,), "product-supplier"),
    "YA004BPS_suppliers_MAPPED": ("Data with Headings", 3, (1,), "supplier"),
    "YA004BPSC_bank_details_MAPPED": ("Data with Headings", 3, (0,), "supplier bank"),
    "YA004BPC_customers_MAPPED": ("Data with Headings", 3, (1, 17), "customer address"),
    "YA004BPCDL_delivery_addresses_MAPPED": ("Data with headings", 3, (0, 1), "delivery address"),
    "YA004SMR_guide_MAPPED_clean_run": ("Data with Headings", 3, None, "stock lot"),
    "YA004POH_guide_MAPPED_clean_run": ("Data with Headings", 3, None, "PO line"),
    "YA004SOHPJ_guide_MAPPED_clean_run": ("Data with Headings", 3, None, "sales order line"),
    "YA004MFG_works_orders_MAPPED_clean_run": ("Data With Headings", 3, None, "works order"),
    "YA004BICY_guide_MAPPED": ("Data with Headings", 2, (1,), "customer open item"),
    "YA004BISY_guide_MAPPED": ("Data with Headings", 2, (1,), "supplier open item"),
}


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def rows_of(path, sheet, first):
    wb = load_workbook(path, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    # the headings are the row above the data; the X3 field codes, where there is one, above that
    head = [norm(v) for v in next(ws.iter_rows(min_row=first - 1, max_row=first - 1, values_only=True))]
    if first > 2:
        codes = [norm(v) for v in next(ws.iter_rows(min_row=first - 2, max_row=first - 2, values_only=True))]
        head = [f"{h} ({c})" if c and h and c != h else h or c for h, c in zip(head, codes)]
    out = []
    for r in ws.iter_rows(min_row=first, values_only=True):
        r = [norm(v) for v in r]
        while r and r[-1] == "":
            r.pop()
        if r:
            out.append(tuple(r))
    return head, out


def compare(prev, new, key):
    """-> added, removed, changed [(key, before, after, [fields])]"""
    if key is None:
        a, b = Counter(prev), Counter(new)
        return list((b - a).elements()), list((a - b).elements()), []
    k = lambda r: tuple(r[i] if i < len(r) else "" for i in key)
    pa, pb = {k(r): r for r in prev}, {k(r): r for r in new}
    added = [pb[x] for x in pb if x not in pa]
    removed = [pa[x] for x in pa if x not in pb]
    changed = []
    for x in pa.keys() & pb.keys():
        a, b = pa[x], pb[x]
        n = max(len(a), len(b))
        a, b = a + ("",) * (n - len(a)), b + ("",) * (n - len(b))
        diff = [i for i in range(n) if a[i] != b[i]]
        if diff:
            changed.append((x, a, b, diff))
    return added, removed, changed


def main(prev_dir, new_dir):
    prev_dir, new_dir = Path(prev_dir), Path(new_dir)
    md = [f"# What changed\n\nPrevious run: `{prev_dir}`  \nThis run: `{new_dir}`\n",
          "| File | Before | After | Added | Removed | Changed |", "|---|---|---|---|---|---|"]
    detail = []
    wb = Workbook()
    wb.remove(wb.active)
    hf, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
    for stem, (sheet, first, key, what) in FILES.items():
        p, n = prev_dir / f"{stem}.xlsx", new_dir / f"{stem}.xlsx"
        if not n.exists():
            continue
        head, new = rows_of(n, sheet, first)
        if not p.exists():
            md.append(f"| {stem} | - | {len(new):,} | new file | | |")
            continue
        _, prev = rows_of(p, sheet, first)
        added, removed, changed = compare(prev, new, key)
        md.append(f"| {stem} | {len(prev):,} | {len(new):,} | {len(added):,} | {len(removed):,} | {len(changed):,} |")
        if added or removed or changed:
            fields = Counter(head[i] if i < len(head) else f"col {i + 1}" for _, _, _, d in changed for i in d)
            lines = [f"\n## {stem}\n", f"- {len(added):,} {what}s added, {len(removed):,} removed, {len(changed):,} changed."]
            if fields:
                lines.append("- Fields that changed: " + ", ".join(f"{f} ({c:,})" for f, c in fields.most_common(8)) + ".")
            for label, rows in (("Added", added), ("Removed", removed)):
                if rows:
                    lines.append(f"- {label}, e.g.: " + "; ".join(" / ".join(r[i] for i in (key or (0, 1, 2)) if i < len(r))
                                                                  for r in rows[:5]) + ".")
            for x, a, b, d in changed[:5]:
                lines.append("- " + " / ".join(x) + ": " + ", ".join(
                    f"{head[i] if i < len(head) else i} {a[i] or '(blank)'} → {b[i] or '(blank)'}" for i in d[:4]))
            detail += lines
            ws = wb.create_sheet(stem[:31])
            ws.append(["Change"] + head + ["Fields changed"])
            for c in ws[1]:
                c.font, c.fill = hf, fill
            for r in added:
                ws.append(["Added"] + list(r))
            for r in removed:
                ws.append(["Removed"] + list(r))
            for x, a, b, d in changed:
                ws.append(["Before"] + list(a) + [", ".join(head[i] if i < len(head) else str(i) for i in d)])
                ws.append(["After"] + list(b))
            ws.freeze_panes = "B2"
    (new_dir / "WHAT_CHANGED.md").write_text("\n".join(md + detail) + "\n", encoding="utf-8")
    if wb.sheetnames:
        wb.save(new_dir / "What_changed.xlsx")
    print("\n".join(md))


if __name__ == "__main__":
    main(*sys.argv[1:3])
