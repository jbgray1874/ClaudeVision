"""Shared rules for the repeatable rebuild - one copy of each, used by every builder.

The rules are the ones the runbook records (sections 4 and 5); the earlier one-off scripts each
carried their own copy. Change a rule here and every file follows.
"""
import csv
import re
import unicodedata
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENC = "cp850"                       # Access exports and X3 load files are DOS cp850


def read_semi(path):
    """A ;-delimited load file, as lists of strings."""
    with open(path, newline="", encoding=ENC, errors="replace") as f:
        return [r for r in csv.reader(f, delimiter=";") if r and r[0].strip()]


def write_semi(path, rows):
    with open(path, "w", newline="", encoding=ENC, errors="replace") as f:
        csv.writer(f, delimiter=";", lineterminator="\r\n").writerows(rows)


def decode(raw):
    """SSMS saves results as UTF-16 (with a byte-order mark), UTF-8 or the Windows code page,
    depending on its settings. Read whichever it is, so a £ or an accent isn't lost."""
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    if raw[:3] == b"\xef\xbb\xbf":
        return raw[3:].decode("utf-8")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def read_grids(path):
    """A SQL Server 'Results To File' export: grids separated by blank lines, tab-delimited.
    Returns {label: rows} for grids whose first column is a label (the extract SQL puts one
    on every grid), and a list of unlabelled grids in order."""
    text = decode(Path(path).read_bytes()).replace("\r\n", "\n")
    labelled, plain = {}, []
    for block in text.split("\n\n"):
        rows = [line.split("\t") for line in block.strip("\n").split("\n") if line.strip()]
        # SSMS can be set to print column headings, a line of dashes and "(n rows affected)"
        rows = [r for r in rows if r[0] != "grid" and not set(r[0]) <= set("- ")
                and not re.fullmatch(r"\(\d+ rows? affected\)", r[0].strip())]
        if not rows:
            continue
        label = rows[0][0]
        if label.isupper() and all(r[0] == label for r in rows) and len(rows[0]) > 1:
            labelled.setdefault(label, []).extend(r[1:] for r in rows)
        else:
            plain.append(rows)
    return labelled, plain


# ---- product codes (runbook 4.1, the V7 rule) ------------------------------------------
def product_code(code, slash="-"):
    """Upper case; spaces, &, ( ) and CHAR(160) removed; . / _ become hyphens; any other symbol
    removed. The same rule in every file, so a product code always matches."""
    s = code.replace("\xa0", "").strip().upper()
    s = re.sub(r"[ &()]", "", s).replace(".", "-").replace("/", slash).replace("_", "-")
    return re.sub(r"[^A-Z0-9\-]", "", s)


def product_codes(raw_codes):
    """Map every raw Access code to its X3 code. Where a / turned into - would land on a code
    that already exists, the / is dropped instead (the V7 rule). Returns {raw: x3} and the
    codes still shared by two raw products after that (runbook: the 33 shared codes)."""
    raw_codes = list(dict.fromkeys(raw_codes))
    first = {c: product_code(c) for c in raw_codes}
    taken = {}
    for c, n in first.items():
        taken.setdefault(n, []).append(c)
    out = {}
    for c, n in first.items():
        if len(taken[n]) > 1 and "/" in c and product_code(c, "") not in taken:
            n = product_code(c, "")
        out[c] = n
    shared = {}
    for c, n in out.items():
        shared.setdefault(n, []).append(c)
    return out, {n: cs for n, cs in shared.items() if len(cs) > 1}


# ---- Ethan's character rule (runbook 5) ------------------------------------------------
def x3_text(v):
    """Letters, numbers, spaces, - _ . only. & -> AND, / -> -, accents dropped, the lost
    apostrophe in '£s' removed, £ -> GBP, the rest removed. Not for emails, phones, websites."""
    v = v or ""
    for a, b in (("Ø", "DIA "), ("ø", "dia "), ("°", " deg"), ("½", "1/2"), ("¼", "1/4"),
                 ("–", "-"), ("—", "-"), ("’", "'")):
        v = v.replace(a, b)
    v = re.sub("(?<=[A-Za-z])£(?=s\\b)", "", v)
    v = re.sub("£s\\b", " GBP", v).replace("£", " GBP ")
    v = v.replace(" ? ", " - ")
    v = unicodedata.normalize("NFKD", v)
    v = "".join(ch for ch in v if not unicodedata.combining(ch))
    v = v.replace("&", " AND ").replace("/", "-")
    return " ".join(re.sub(r"[^A-Za-z0-9._\- ]", "", v).split())


def clip(v, n):
    return (v or "")[:n].rstrip()


def ensure_dir(p):
    Path(p).mkdir(parents=True, exist_ok=True)
    return Path(p)


# ---- spreadsheet tabs, styled the same in every file -----------------------------------
HEAD_FONT, HEAD_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")


def add_sheet(wb, title, head, rows, widths=None, index=None):
    ws = wb.create_sheet(title, index)
    ws.append(list(head))
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths or [18] * len(head), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def notes(wb, lines, index=None):
    ws = wb.create_sheet("Notes", index)
    for line in lines:
        ws.append(list(line))
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws
