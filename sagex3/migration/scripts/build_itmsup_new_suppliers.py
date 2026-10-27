"""ZSDIITMSUP with the products bought from the new suppliers.

Input is grid 2 of sql/product_suppliers_new_only.sql: every product ever bought from one of the
suppliers that were missing from X3, with all of its suppliers ranked (priority 1 = last bought
from). For each of those products that is in the clean-run product file, the priority-1 supplier
becomes its main supplier - added where the product had none, replaced where the old file named
an older supplier only because the newer one had no code. Everything else stays as it was.

Product codes go through the V7 code rule, which reproduces the existing file exactly.

Usage:
    python build_itmsup_new_suppliers.py <grid 2 .txt>
"""
import csv
import re
import sys
from collections import Counter

from openpyxl import load_workbook
from openpyxl.styles import Font

BASE = "/tmp/claude-0/-home-user-ClaudeVision/e9ffe9d9-c7dd-5e34-8a16-6cb3c65dadea/scratchpad"
V8, OUT, T = f"{BASE}/out_v8", f"{BASE}/out_sheets", f"{BASE}/tmpl_t2"


def read(p):
    with open(p, newline="", encoding="cp850") as f:
        return [r for r in csv.reader(f, delimiter=";") if r and r[0].strip()]


def fill(ws, first_row, rows):
    """As build_sheets.py: clear the template's data rows, then write ours."""
    if ws.max_row >= first_row:
        ws.delete_rows(first_row, ws.max_row)
    for r in rows:
        ws.append([v if v != "" else None for v in r])
    ws.freeze_panes = ws.cell(first_row, 1)


def check(path, sheet, first_row, rows):
    got = [["" if v is None else str(v) for v in r][:3] for r in
           load_workbook(path, read_only=True)[sheet].iter_rows(min_row=first_row, values_only=True)]
    got = [g for g in got if any(g)]
    assert got == rows, path
    return len(got)


def notes(wb, lines):
    ws = wb.create_sheet("Notes")
    for l in lines:
        ws.append(l)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 100


def clean(c):
    """The V7 product code rule."""
    return re.sub(r"[^A-Z0-9\-]", "", c.strip().upper().replace("/", "-").replace("_", "-"))


def main(grid):
    products = {r[1] for r in read(f"{V8}/ZSDIITM_load8_products.txt")}
    suppliers = {r[1] for r in read(f"{BASE}/out_sue/YA004BPS_load7_suppliers.txt")}
    rows = read(f"{V8}/ZSDIITMSUP_load8_product_suppliers.txt")
    at = {r[0]: i for i, r in enumerate(rows)}
    first = {}
    for r in (l.split("\t") for l in open(grid, encoding="utf-8").read().splitlines() if l.count("\t") == 15):
        if r[2] == "1":
            first[clean(r[0])] = r[1]
    log, changed = Counter(), []
    for p, s in sorted(first.items()):
        if p not in products:
            log["product not in the product file - not loaded"] += 1
            continue
        assert s in suppliers, (p, s)
        if p not in at:
            at[p] = len(rows)
            rows.append([p, s, p])
            log["added - had no supplier"] += 1
        elif rows[at[p]][1] != s:
            changed.append([p, rows[at[p]][1], s])
            rows[at[p]][1] = s
            log["main supplier changed to the one last bought from"] += 1
        else:
            log["unchanged"] += 1
    rows.sort(key=lambda r: r[0])
    assert len({r[0] for r in rows}) == len(rows)
    assert all(r[0] in products and r[1] in suppliers for r in rows)

    path = f"{OUT}/ZSDIITMSUP_product_suppliers_MAPPED.xlsx"
    wb = load_workbook(f"{T}/ZSDIITMSUP.xlsx")
    fill(wb["Template"], 3, rows)
    notes(wb, [["ZSDIITMSUP - product-suppliers, clean run"], [], ["Rows", len(rows)],
               ["Built from", "Access purchase orders since September 2022: the supplier each product was last bought from"],
               ["New suppliers", f"{log['added - had no supplier']} products added that are bought only from the suppliers "
                                 f"that were missing from X3; {log['main supplier changed to the one last bought from']} "
                                 "changed to the new supplier they were last bought from (Changed tab)"],
               ["Not loaded", "3,150 alternative suppliers - the template has no priority column. Ethan to add PIO if wanted"]])
    ws = wb.create_sheet("Changed")
    ws.append(["Product", "Was", "Now"])
    for r in changed:
        ws.append(r)
    wb.save(path)
    print(path.split("/")[-1], check(path, "Template", 3, rows), "rows checked;", dict(log))


if __name__ == "__main__":
    main(sys.argv[1])
