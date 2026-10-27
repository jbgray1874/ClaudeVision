"""Customer and supplier mapped spreadsheets for the clean run.

The original template workbooks are no longer here: only their first 14 column codes and
headings were recorded. Data positions are exactly those of the files that loaded, so
Marc's macro (which skips the two heading rows) writes the same file. Columns past N carry
the X3 code only where the data and Ethan's notes fix it - column R is the address code and
AG the address default flag, as his import notes say - and a plain label otherwise.
"""
import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "/tmp/claude-0/-home-user-ClaudeVision/e9ffe9d9-c7dd-5e34-8a16-6cb3c65dadea/scratchpad"
OUT = f"{BASE}/out_sheets"
HF, FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
GREY = PatternFill("solid", fgColor="F2F2F2")


def read(p):
    with open(p, newline="", encoding="cp850") as f:
        return [r for r in csv.reader(f, delimiter=";") if r and r[0].strip()]


CUSTOMER = [  # (X3 code or "", heading, where the heading comes from)
    ("BCGCOD", "Category", "template"), ("BPCNUM", "Customer", "template"), ("BPRNAM", "Company Name", "template"),
    ("BPRNAM(1)", "Company Name", "template"), ("BPRSHO", "Short title", "template"), ("BPAINV", "Address", "template"),
    ("BPAPYR", "Address", "template"), ("BPAADD", "Default Address", "template"),
    ("BPDADD", "Ship-to Customer Address", "template"), ("CUR", "Currency", "template"),
    ("CRN", "Co. Reg. No.", "template"), ("NAF", "SIC code", "template"), ("EECNUM", "Intraco VAT no", "template"),
    ("BPCINV", "Bill-to", "template"),
    ("VACBPR", "Tax Rule", "data: DOM / ROWC, the X3 tax rule codes"), ("PTE", "Payment Terms", "data: TFR_30_EOM ..."),
    ("ACCCOD", "Accounting code", "data: STD"), ("BPAADD", "Address", "Ethan's notes: column R, the address code"),
    ("BPADES", "Description", "data: the address name"), ("BPAADDLIG(0)", "Address Line", "data"),
    ("BPAADDLIG(1)", "Address Line", "data"), ("BPAADDLIG(2)", "Address Line", "data"),
    ("POSCOD", "Postal Code", "data"), ("", "County / region", "data: SURREY, FRANCE ..."),
    ("CRY", "Country", "data: GB, IE ..."), ("", "Telephone", "data"), ("", "Telephone 2", "data"),
    ("", "Email", "data"), ("", "Email", "data"), ("", "(blank in the loaded file)", "data"),
    ("", "(blank in the loaded file)", "data"), ("", "(blank in the loaded file)", "data"),
    ("BPAADDFLG", "Default address flag", "Ethan's notes: column AG, 2 for address 001, else 1"),
    ("", "Town / county", "data: LONDON, TYNE AND WEAR ..."),
]
SUPPLIER = [
    ("BSGCOD", "Category", "template"), ("BPSNUM", "Supplier", "template"), ("BPSNAM", "Company Name", "template"),
    ("CRY", "Country", "template"), ("CUR", "Currency", "template"), ("CRN", "Co. Reg. No.", "template"),
    ("EECNUM", "Intraco VAT no", "template"), ("VACBPR", "Tax Rule", "template"), ("PTE", "Payment Terms", "template"),
    ("ACCCOD", "Accounting code", "template"), ("BPAADD", "Address", "template"), ("BPADES", "Description", "template"),
    ("BPAADDLIG(0)", "Address Line", "template"), ("BPAADDLIG(1)", "Address Line", "template"),
    ("BPAADDLIG(2)", "Address Line", "data"), ("", "County / region", "data: LEICESTERSHIRE, CHINA ..."),
    ("", "(blank in the loaded file)", "data"), ("POSCOD", "Postal Code", "data"), ("CRY", "Country", "data"),
    ("", "Telephone", "data"), ("", "Telephone 2", "data"), ("", "Website / email", "data"), ("", "Email", "data"),
] + [("", "(blank in the loaded file)", "data")] * 9 + [
    ("BPAADDFLG", "Default address flag", "Ethan's notes: column AG, 2 for every supplier"),
]


def build(name, title, layout, rows, before, changes):
    assert all(len(r) == len(layout) for r in rows), (name, set(len(r) for r in rows), len(layout))
    wb = Workbook()
    ws = wb.active
    ws.title = "Data with Headings"
    ws.append([c for c, _, _ in layout])
    ws.append([h for _, h, _ in layout])
    for i, (_, _, src) in enumerate(layout, 1):
        ws.cell(2, i).font = Font(bold=True)
        if src != "template":
            ws.cell(1, i).fill = ws.cell(2, i).fill = GREY
    for r in rows:
        ws.append([v if v != "" else None for v in r])
    for i in range(1, len(layout) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "C3"

    g = wb.create_sheet("Guide")
    g.append(["Column", "X3 field", "Heading", "Heading from"])
    for c in g[1]:
        c.font, c.fill = HF, FILL
    for i, (c, h, src) in enumerate(layout, 1):
        g.append([get_column_letter(i), c, h, src])
    for col, w in zip("ABCD", (8, 14, 28, 60)):
        g.column_dimensions[col].width = w

    ch = wb.create_sheet("Changes")
    ch.append(["Code", "Address", "Column", "Field", "Was", "Now"])
    for c in ch[1]:
        c.font, c.fill = HF, FILL
    n = 0
    for a, b in zip(before, rows):
        for i, (x, y) in enumerate(zip(a, b)):
            if x != y:
                ch.append([b[1], b[17] if name == "YA004BPC" else b[10], get_column_letter(i + 1),
                           layout[i][0] or layout[i][1], x, y])
                n += 1
    for col, w in zip("ABCDEF", (10, 8, 8, 16, 40, 40)):
        ch.column_dimensions[col].width = w
    ch.freeze_panes = "A2"

    nt = wb.create_sheet("Notes")
    for l in [[f"{name} - {title}, clean run"], [], ["Rows", len(rows)], ["Changes from the file loaded in UAT", n]] \
            + changes + [[], ["Headings",
                              "Columns A to N carry the X3 codes and headings of your template. Grey columns are "
                              "labelled from the data and Ethan's notes - the template workbook is not to hand. Every "
                              "value is in the same column as the file that loaded, so the macro writes the same file"]]:
        nt.append(l)
    nt["A1"].font = Font(bold=True, size=13)
    nt.column_dimensions["A"].width = 40
    nt.column_dimensions["B"].width = 100
    path = f"{OUT}/{name}_{title.lower()}_MAPPED.xlsx"
    wb.save(path)

    # read back: every cell must equal the load data
    from openpyxl import load_workbook
    got = [["" if v is None else str(v) for v in r]
           for r in load_workbook(path, read_only=True)["Data with Headings"].iter_rows(min_row=3, values_only=True)]
    got = [g + [""] * (len(layout) - len(g)) for g in got]
    assert got == rows, name
    print(path.split("/")[-1], len(rows), "rows checked,", n, "changes listed")


bps = read(f"{BASE}/out_v8/YA004BPS_load6_suppliers.txt")
build("YA004BPS", "Suppliers", SUPPLIER, bps, read(f"{BASE}/upl2/bps.txt"), [
    ["VAT numbers", "Dummy numbers removed - the PL numbers copied across 117 suppliers, TEMPJAF placeholders, GB "
                    "numbers on overseas companies. 185 suppliers. UK suppliers load with a blank"],
    ["Tax rule", "28 suppliers on EU set to ROWS - X3 has no EU tax rule. Finance to confirm"],
    ["Names and addresses", "Ethan's character rule. Emails, phones and websites untouched"]])
bpc = read(f"{BASE}/out_v8/YA004BPC_load6_customers.txt")
build("YA004BPC", "Customers", CUSTOMER, bpc, read(f"{BASE}/upl2/bpc.txt"), [
    ["Tax rule", "10 EU customers (56 address rows) set to ROW / ROWC, as UNI100 which loaded. Finance to confirm"],
    ["VAT number", "ROL100's GB76361992 is 8 digits - blanked. Finance to supply the real one"],
    ["TEC100", "Postcodes moved out of the address lines (28823, 4785-598) for the POSCOD error"],
    ["Two rows", "A second email after a ';' had pushed every later field along by one - first email kept"],
    ["Names and addresses", "Ethan's character rule. Emails and phones untouched"]])
