"""YA004MFG - open works orders, from works_orders_extract.sql grid 2.

One line per open Access works order with something still to make:
  ITMREF  the product, on the cleaned codes (UAT codes now, V7 codes for the clean run)
  STRDAT  the date the works order was created (no open order has a first date)
  ENDDAT  the scheduled date
  EXTQTY  ordered less made (the amended quantity has never been used)
  MFGNUM  blank: X3 numbers the orders, as Ethan's example numbers MFG-S01A000001 suggest
Held back, with the reason: nothing left to make, product not in the product file, or
product with no BOM or route on code 10.
"""
import csv
import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "/tmp/claude-0/-home-user-ClaudeVision/e9ffe9d9-c7dd-5e34-8a16-6cb3c65dadea/scratchpad"
OUT = f"{BASE}/out_mfg"
ENC = "cp850"


def read(p):
    with open(p, newline="", encoding=ENC) as f:
        return [r for r in csv.reader(f, delimiter=";") if r and r[0].strip()]


def write(p, rows):
    with open(p, "w", newline="", encoding=ENC) as f:
        csv.writer(f, delimiter=";", lineterminator="\r\n").writerows(rows)


def key(s):
    """The normalised key the V6 standardisation matched codes on."""
    return (str(s or "").strip().replace(" ", "").replace("&", "")
            .replace("(", "").replace(")", "").replace(".", "-").upper())


lines = [l.split("\t") for l in open(f"{BASE}/wor/out.txt", encoding="utf-8").read().splitlines()]
wor = [r for r in lines if len(r) == 25]
HEAD = ["unique", "works_order", "parent_order", "product", "rev", "sales_order", "customer", "order_type",
        "orig_qty", "amd_qty", "good_qty", "rejected_qty", "scrap_qty", "created", "first_date", "mrp_date",
        "scheduled", "last_mat_due", "done_date", "mat_stat", "booked_mat", "booked_lab", "booked_subcon",
        "dest_wh", "dest_bin"]
wor = [dict(zip(HEAD, r)) for r in wor]

v6 = read(f"{BASE}/out_std6/ZSDIITM_load6_products.txt")
v7 = read(f"{BASE}/out_v7/ZSDIITM_load7_products.txt")
to_v7 = {a[1]: b[1] for a, b in zip(v6, v7)}
cat = {r[1]: r[0] for r in v6}
by_key = {}
for r in v6:
    by_key.setdefault(key(r[1]), r[1])
uat = {r[1] for r in read(f"{BASE}/uploads_final/ZSDIITM_Cleansing.txt")}
bom = {r[0] for r in read(f"{BASE}/out_std6/YA00408BOM_load6_loadable.txt")}
rou = {r[0] for r in read(f"{BASE}/out_std6/ZSDIROU2_load6_routes.txt")}


def num(s):
    v = float(s)
    return str(int(v)) if v == int(v) else f"{v:g}"


load, held, why = [], [], Counter()
for w in wor:
    code = by_key.get(key(w["product"]))
    todo = float(w["orig_qty"]) - float(w["good_qty"])
    reason = ""
    if todo <= 0:
        reason = "nothing left to make"
    elif code is None:
        reason = "product not in the product file"
    elif code not in bom:
        reason = "product has no BOM on code 10"
    elif code not in rou:
        reason = "product has no route on code 10"
    start = w["created"]
    end = w["scheduled"] if w["scheduled"] not in ("", "NULL") else start
    row = {"w": w, "code": code, "todo": todo, "start": start, "end": end}
    if reason:
        why[reason] += 1
        held.append((row, reason))
    else:
        load.append(row)


def line(r, code):
    return ["S01A", "10", code, r["start"], max(r["end"], r["start"]), num(r["todo"]), "", code, "10", "C", "1"]


now = [line(r, r["code"]) for r in load if r["code"] in uat]
not_in_uat = [r for r in load if r["code"] not in uat]
clean = [line(r, to_v7[r["code"]]) for r in load]
write(f"{OUT}/YA004MFG_load1_UAT_now.txt", now)
write(f"{OUT}/YA004MFG_load7_clean_run.txt", clean)

# ---- the guide -------------------------------------------------------------------
FIELDS = [("MFGFCY", "Production Site", "S01A", "X3 Site", '"S01A"', "Hardcoded"),
          ("BOMALT", "BOM Code", "10", "Hardcoded to 10", '"10" - every BOM in the BOM file is on 10', "Hardcoded"),
          ("ITMREF", "Product", "TES-11106-01-01", "Parent Product code for Works Order",
           "WOR_ITEM_ONLY, on the cleaned product codes", "Dynamic"),
          ("STRDAT", "Start date", "20261106", "Start Date YYYYMMDD",
           "WOR_CREATED_DATE - WOR_FIRST_DATE is blank on every open order", "Dynamic"),
          ("ENDDAT", "End date", "20261106", "End Date YYYYMMDD", "WOR_DATE_SCHED; never before the start", "Dynamic"),
          ("EXTQTY", "Planned", "117", "Qty of parent product to be made",
           "WOR_ORIG_QTY less WOR_GOOD_QTY - still to make. WOR_AMD_QTY is never used", "Dynamic"),
          ("MFGNUM", "Order", "67", "", "Blank - X3 numbers the orders. Access number kept on the Access reference tab",
           "To confirm"),
          ("*1", "Routing", "TES-11106-01-01", "Same as parent product code", "= ITMREF", "Dynamic"),
          ("*2", "Code", "10", "Hardcode to 10", '"10" - every route in the route file is on 10', "Hardcoded"),
          ("*3", "Update code", "C", "Hardcode to C for Create", '"C"', "Hardcoded"),
          ("*4", "Order status", "1", "Hardcode to 1", '"1"', "Hardcoded")]
HF, FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
wb = Workbook()


def table(ws, head, data, widths, start=1):
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font, c.fill = HF, FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for d in data:
        ws.append(list(d))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


m = wb.active
m.title = "Mapping"
m.append(["YA004MFG  -  Open works orders"])
m["A1"].font = Font(bold=True, size=13)
m.append(["Built from works_orders_extract.sql, dbo.ASC_WOR_TBL, status blank (open). Columns exactly as "
          "Ethan's template, same order."])
m.append([])
table(m, ["#", "Sage field", "X3 column", "Template example", "Guide note", "Access source", "Status"],
      [[i, f[1], f[0], f[2], f[3], f[4], f[5]] for i, f in enumerate(FIELDS, 1)], [4, 16, 10, 18, 34, 60, 12])

wos = [r["w"] for r in load]
n = wb.create_sheet("Notes")
for r in [
    ["YA004MFG  -  what the data says, and what is still open"], [],
    ["Open works orders in Access", len(wor)],
    ["In the load (clean run)", len(clean)],
    ["In the load (UAT now)", len(now), f"{len(not_in_uat)} more once the 137 added products are in"],
    ["Held back", len(held)],
] + [[f"   {k}", v] for k, v in why.most_common()] + [
    [],
    ["Part made", sum(1 for w in wos if float(w["good_qty"]) > 0), "quantity is what is still to make"],
    ["With cost already booked", sum(1 for w in wos if float(w["booked_mat"]) + float(w["booked_lab"]) + float(w["booked_subcon"]) > 0),
     f"£{sum(float(w['booked_mat']) + float(w['booked_lab']) + float(w['booked_subcon']) for w in wos):,.2f} booked - not carried by this template"],
    ["Scheduled before today", sum(1 for w in wos if w["scheduled"] < "20261027")],
    ["Sub-assembly orders (have a parent order)", sum(1 for w in wos if w["parent_order"])],
    ["Prototype orders (WP)", sum(1 for w in wos if w["order_type"] == "WP")],
    [],
    ["Open with Ethan"],
    ["1. Order number", "Left blank for X3 to number. Or should it carry the Access number (e.g. 70021093/02)?"],
    ["2. Quantity", "Planned = still to make (ordered less made). Confirm."],
    ["3. Work in progress", "Part-made orders have material and labour booked in Access. This template carries no cost: how is WIP brought across?"],
    ["4. Past dates", "Orders scheduled before go-live keep their Access dates. Move them to go-live?"],
    ["5. Sub-assemblies", "Child works orders (e.g. 70021093/01-/03 under 70021093) load as orders in their own right. "
                          "Will X3 also create them from the parent's BOM, so they come in twice?"],
]:
    n.append(r)
n["A1"].font = Font(bold=True, size=13)
n.column_dimensions["A"].width = 44
n.column_dimensions["B"].width = 100

d = wb.create_sheet("Data With Headings")          # name as in Ethan's template - Marc's macro reads it
d.append([f[0] for f in FIELDS])
d.append([f[1] for f in FIELDS])
for c in d[2]:
    c.font = Font(bold=True)
for r in now:
    d.append(r)
for i in range(1, 12):
    d.column_dimensions[get_column_letter(i)].width = 16
d.freeze_panes = "A3"

g = wb.create_sheet("Guide")
table(g, ["Field Description", "Example", "Notes"], [[f[1], f[2], f[3]] for f in FIELDS], [20, 18, 40])

a = wb.create_sheet("Access reference")
table(a, ["X3 product", "Access works order", "Parent order", "Sales order", "Type", "Ordered", "Made",
          "To make", "Created", "Scheduled", "Cost booked"],
      [[r["code"], r["w"]["works_order"], r["w"]["parent_order"], r["w"]["sales_order"], r["w"]["order_type"],
        float(r["w"]["orig_qty"]), float(r["w"]["good_qty"]), r["todo"], r["start"], r["w"]["scheduled"],
        round(float(r["w"]["booked_mat"]) + float(r["w"]["booked_lab"]) + float(r["w"]["booked_subcon"]), 2)]
       for r in load], [18, 14, 12, 18, 6, 8, 8, 8, 10, 10, 11])
h = wb.create_sheet("Held back")
table(h, ["Access product", "Access works order", "Type", "Ordered", "Made", "Scheduled", "Reason"],
      [[r["w"]["product"], r["w"]["works_order"], r["w"]["order_type"], float(r["w"]["orig_qty"]),
        float(r["w"]["good_qty"]), r["w"]["scheduled"], why_] for r, why_ in held], [20, 14, 6, 8, 8, 10, 34])
wb.save(f"{OUT}/YA004MFG_guide_MAPPED.xlsx")

print("open", len(wor), "load clean", len(clean), "UAT now", len(now), "held", len(held), dict(why))
print("past", sum(1 for w in wos if w["scheduled"] < "20261027"), "parent", sum(1 for w in wos if w["parent_order"]),
      "WP", sum(1 for w in wos if w["order_type"] == "WP"), "cats", Counter(cat.get(r["code"]) for r in load))
print("frac qty", sum(1 for r in load if r["todo"] != int(r["todo"])),
      "rejected/scrap", sum(1 for w in wor if float(w["rejected_qty"]) or float(w["scrap_qty"])),
      "end<start", sum(1 for r in load if r["end"] < r["start"]))
