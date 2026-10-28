"""YA004BICY and YA004BISY: the customer and supplier open-item files in their new layout.

What changed from BICX / BISX (James, 28 Oct):
  * the templates are renamed BICY and BISY
  * BISY loses the CURLED column
  * both get a new column, Company, between C (Customer / Supplier) and D (Site). Its value is
    the X3 company code that site S01A belongs to - COMPANY below, to confirm with Marc. Until
    then it is left blank.
Everything else is the BICX / BISX file exactly as built (populate_guides.py): same rows,
same values, same order. The build reads every data row back and checks that.

Usage:
    COMPANY=<code> python build_bicy_bisy.py
"""
import os

from openpyxl import load_workbook
from openpyxl.styles import Font

BASE = "/tmp/claude-0/-home-user-ClaudeVision/e9ffe9d9-c7dd-5e34-8a16-6cb3c65dadea/scratchpad"
SRC = os.environ.get("SRC_DIR", f"{BASE}/out_mapped")
OUT = os.environ.get("OUT_DIR", f"{BASE}/out_mapped")
COMPANY = os.environ.get("COMPANY", "")
NOTE = ("New in this layout (28 Oct): between {bp} and Site. The X3 company code that site S01A "
        "belongs to" + (f": {COMPANY}" if COMPANY else " - to confirm with Marc, blank until then"))


def rows_of(ws):
    return [[c.value for c in r] for r in ws.iter_rows(min_row=2) if any(c.value is not None for c in r)]


def find_row(ws, col, value, start=1):
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, col).value == value:
            return r
    raise KeyError(value)


def build(old, new, bp, drop_curled):
    wb = load_workbook(f"{SRC}/{old}_guide_MAPPED.xlsx")
    d = wb["Data with Headings"]
    before = rows_of(d)
    heads = [c.value for c in d[1]]
    if drop_curled:
        d.delete_cols(heads.index("CURLED") + 1)
    d.insert_cols(4)
    d.cell(1, 4, "Company").font = Font(bold=True)
    for r in range(2, d.max_row + 1):
        if d.cell(r, 1).value is not None:
            d.cell(r, 4, COMPANY or None)

    # read back: the same rows, with Company in D and (BISY) no CURLED
    after = rows_of(d)
    expect = []
    for r in before:
        r = list(r)
        if drop_curled:
            del r[heads.index("CURLED")]
        expect.append(r[:3] + [COMPANY or None] + r[3:])
    assert [a[:len(e)] for a, e in zip(after, expect)] == expect and len(after) == len(before), new

    m = wb["Mapping"]
    m.cell(1, 1, str(m.cell(1, 1).value).replace(old, new))
    if drop_curled:
        m.delete_rows(find_row(m, 2, "CURLED"))
    at = find_row(m, 2, bp, 4) + 1
    m.insert_rows(at)
    for c, v in enumerate([None, "Company", None, "", "(new column)", "-", NOTE.format(bp=bp), "To confirm"], 1):
        m.cell(at, c, v)
    n = 1
    for r in range(find_row(m, 1, "#") + 1, m.max_row + 1):
        if m.cell(r, 2).value:
            m.cell(r, 1, n)
            n += 1

    g = wb["Guide"]
    if drop_curled:                      # the CURLED row has no field name, example GBP
        cur = find_row(g, 1, "Currency")
        if g.cell(cur + 1, 1).value is None and g.cell(cur + 1, 2).value == "GBP":
            g.delete_rows(cur + 1)
    at = find_row(g, 1, "Customer") + 1   # the supplier guide labels its code "Customer" too
    g.insert_rows(at)
    for c, v in enumerate(["Company", COMPANY or None, COMPANY or None, "Sage X3 company of the site - new column"], 1):
        g.cell(at, c, v)

    nt = wb["Notes"]
    nt.cell(1, 1, str(nt.cell(1, 1).value).replace(old, new))
    nt.insert_rows(3, 3)
    nt.cell(3, 1, f"{new} layout").font = Font(bold=True)
    nt.cell(4, 1, f"Was {old}. New column Company (D), between {bp} and Site"
                  + ("; the CURLED column is removed" if drop_curled else "")
                  + ". " + ("Company = " + COMPANY if COMPANY else "Company is blank until Marc confirms the code")
                  + ". Every row and value is otherwise as in the " + old + " file.")
    path = f"{OUT}/{new}_guide_MAPPED.xlsx"
    wb.save(path)
    print(path.split("/")[-1], len(after), "rows checked,", len(after[0]) if after else 0, "columns")


build("YA004BICX", "YA004BICY", "Customer", drop_curled=False)
build("YA004BISX", "YA004BISY", "Supplier", drop_curled=True)
