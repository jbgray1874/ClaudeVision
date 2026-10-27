"""Customers_to_add.xlsx: live Access customers that are not in Sage X3, with a proposed new code each.

Input is the vw_DQ_BPC result for customers not graded OLD and not in migration.XREF_SAGE_BPC,
pasted as tab-separated: code, name, grade, reason, last invoice.

New codes follow the X3 pattern: first 3 letters of the name (THE ignored), then the next number
from 100 not already used in the customer load file. Within each 3-letter group the accounts that
Customer Services have already been told about go first (PROMISED), then the most recently invoiced,
so the busiest accounts get the lowest numbers and the codes don't move between reruns unless the
list itself changes.

Usage:
    python build_customers_to_add.py <missing .tsv> <customer load .txt (V7)> <output .xlsx>
"""
import collections
import csv
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Codes already given to Sue Rogers (27 Oct 2026) - keep them
PROMISED = ["TIE02E", "TIE02POL"]
ALREADY_IN_X3 = {"HCL02": "HOT100"}       # X3 spells it "Chocolate"; xref v3 has the pair
NOT_CUSTOMERS = {"ESD01", "FSC01"}        # discount and fuel scale charge ledger accounts


def base(name):
    n = re.sub(r"^THE\s+", "", name.upper().strip())
    return (re.sub(r"[^A-Z]", "", n) + "XXX")[:3]


def order(r):
    code, name, _, _, last = r
    promised = PROMISED.index(code) if code in PROMISED else len(PROMISED)
    return (base(name), promised, -int((last or "0000-00-00").replace("-", "")), code)


def main(src, load, dest):
    rows = [(l.split("\t") + [""] * 5)[:5] for l in open(src, encoding="utf-8").read().splitlines() if l.strip()]
    used = {r[1] for r in csv.reader(open(load, encoding="cp850"), delimiter=";") if len(r) > 1}
    out = []
    for code, name, grade, reason, last in sorted(rows, key=order):
        new = ""
        if code in ALREADY_IN_X3:
            act, new = "Already in X3", ALREADY_IN_X3[code]
            fix = 'None. Only the cross-reference needs version 3 (it misses HCL02 because X3 spells it "Chocolate")'
        elif grade == "REVIEW":
            act, fix = "Leave out (intercompany)", "Finance to confirm if intercompany accounts are wanted"
        elif code in NOT_CUSTOMERS:
            act, fix = "Leave out (not a customer)", "A ledger account for discounts / fuel scale charge, not a customer. Finance to confirm"
        else:
            b, i = base(name), 100
            while f"{b}{i}" in used:
                i += 1
            new = f"{b}{i}"
            used.add(new)
            if grade == "GOOD":
                act, fix = "Add", "None - newer than the customer file"
            elif reason.startswith("Category"):
                act, fix = "Add", "Overseas customer: category ROW, tax rule ROWC (as UNI100)"
            elif reason.startswith("UK customer"):
                act, fix = "Add", "Check the country. Overseas: set it, with ROW / ROWC. UK: add the postcode"
            elif reason == "No address":
                act, fix = "Add when it has an address", "Needs an address"
            else:
                act, fix = "Add", "Separate Access account with the same name as another - its own new code"
        out.append([code, name, grade, reason, last or None, act, new, fix])

    hf, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers to add"
    ws.append(["Old Access code", "Name", "Access grade", "Why it was left out", "Last invoice",
               "Action", "Proposed new code", "Fix before it loads"])
    for r in out:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, fill, Alignment(wrap_text=True)
    for i, w in enumerate([14, 40, 20, 34, 12, 26, 14, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    totals = collections.Counter(r[5] for r in out)
    n = wb.create_sheet("Notes", 0)
    for line in [
        ["Live customers missing from Sage X3"], [],
        ["Source", "migration.vw_DQ_BPC: every Access customer not graded OLD (dormant) and not in the cross-reference"],
        ["Why they were missing", 'The customer import took only accounts graded GOOD. Accounts sharing a name were set aside as '
                                  '"possible duplicates", overseas accounts as "incomplete", and accounts newer than the customer file were not in it'],
        ["Totals", ", ".join(f"{k}: {v}" for k, v in totals.most_common())],
        ["Proposed codes", "First three letters of the name (THE ignored), then the next free number. Within a group, accounts Customer "
                           "Services already have codes for go first (TIE02E = TEC103, TIE02POL = TEC104), then the most recently invoiced"],
        ["How they get into X3", "New rows in the customer import (YA004BPC) and delivery addresses (YA004BPCDL) - the import creates the accounts"],
        ["Next", "Their details from Access (address, currency, VAT, terms) go into the customer spreadsheet as new rows; the "
                 "cross-reference gets the old-to-new pairs; BICX and SOHPJ are rebuilt so their invoices and orders come in"],
    ]:
        n.append(line)
    n["A1"].font = Font(bold=True, size=13)
    n.column_dimensions["A"].width = 22
    n.column_dimensions["B"].width = 120
    for r in n.iter_rows(min_row=3):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(dest)
    print(dest, dict(totals))
    return out


if __name__ == "__main__":
    for r in main(*sys.argv[1:4]):
        if r[6]:
            print(f"  {r[0]:10} {r[6]:8} {r[1]}")
