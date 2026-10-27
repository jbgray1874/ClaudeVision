"""Map every Access bin holding stock to an X3 location from Marc's list.

Each mapping says how it was made, because they are not equally safe:
  exact            the Access bin is an X3 location already
  rule             the same location written another way - AA2 is AA02, A10-UP is A10U
  suggested        a reasonable reading that someone should confirm - HARDX6 is the
                   Hardings external store, TRAIL7 is trailer T007
  default          no X3 equivalent: a person's name, a work area, a subcontractor.
                   Sent to the single general location A001 for UAT
"""
import csv
import json
import re
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "/tmp/claude-0/-home-user-ClaudeVision/e9ffe9d9-c7dd-5e34-8a16-6cb3c65dadea/scratchpad"
DEFAULT = "A001"

loc = json.load(open(f"{BASE}/loclist/loc.json"))["xl/worksheets/sheet2.xml"][1:]
X3 = {r[1].strip().upper(): r[0].strip() for r in loc if len(r) > 1 and r[1].strip()}

# external stores and trailers, read off the names - these are suggestions, not facts
EXTERNAL = [
    (r"^(HARD|HAR\d)", "HARD", "Hardings external store"),
    (r"^SYST", "SYST", "Syston external store"),
    (r"^(ALDER|ABG)", "ALDE", "Alderham Barford"),
    (r"^PION", "PION", "Pionier"),
    (r"^LOG", "LOGX", "Logix"),
    (r"^TPSM", "TPSM", "TPSM"),
    (r"^TRAIL(\d)$", None, "trailer"),
]


def map_bin(b):
    u = b.strip().upper()
    if u == "(NO BIN)":
        return DEFAULT, "default", "no bin in Access"
    if u in X3:
        return u, "exact", ""
    m = re.fullmatch(r"([A-Z]{2})-?(\d{1,2})", u)               # racking: AA2, AA-12
    if m and f"{m.group(1)}{int(m.group(2)):02d}" in X3:
        return f"{m.group(1)}{int(m.group(2)):02d}", "rule", "racking number padded to two digits"
    m = re.fullmatch(r"A(\d{1,2})[-\s]?(UP|U|LO|L)", u)        # goods-in: A10-UP, A34 LO
    if m:
        c = f"A{int(m.group(1)):02d}{m.group(2)[0]}"
        if c in X3:
            return c, "rule", "goods-in bay with upper/lower level"
    m = re.fullmatch(r"([BC])(\d{1,3})", u)                     # goods-in: B5, C12
    if m and f"{m.group(1)}{int(m.group(2)):03d}" in X3:
        return f"{m.group(1)}{int(m.group(2)):03d}", "rule", "goods-in bay padded to three digits"
    for pat, target, what in EXTERNAL:
        m = re.match(pat, u)
        if m:
            if what == "trailer":
                t = f"T{int(m.group(1)):03d}"
                if t in X3:
                    return t, "suggested", f"trailer {m.group(1)} read as despatch location {t}"
            elif target in X3:
                return target, "suggested", f"read as the {what}"
    return DEFAULT, "default", "no X3 equivalent"


# ---- Access bins, by warehouse -------------------------------------------------
blocks, cur = [], []
for line in open(f"{BASE}/bins/out.txt", encoding="utf-8", errors="replace").read().split("\n"):
    if line.strip() == "":
        if cur:
            blocks.append(cur)
            cur = []
    else:
        cur.append(line)
if cur:
    blocks.append(cur)
acc = [r.split("\t") for r in blocks[0]]

rows = []
for wh, b, lots, prods, qty, val in acc:
    x3, how, note = map_bin(b)
    rows.append([wh, b, x3, X3.get(x3, ""), how, note, int(lots), int(prods), float(qty), float(val)])

# ---- the workbook --------------------------------------------------------------
HF = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="44546A")
COL = {"exact": "D5E8D4", "rule": "DEEBF7", "suggested": "FFF2CC", "default": "F8CBAD"}
wb = Workbook()
s = wb.active
s.title = "Summary"
s.append(["Access bins mapped to the X3 location list"])
s["A1"].font = Font(bold=True, size=13)
s.append([f"{len(X3)} X3 locations on Marc's list. Every Access bin holding stock is mapped below; "
          f"the 'How' column says how safe each mapping is."])
s.append([])
s.append(["How", "Meaning", "Bins", "Lots", "Stock value"])
for c in s[4]:
    c.font, c.fill = HF, FILL
MEAN = {"exact": "the Access bin is already an X3 location",
        "rule": "the same location written differently (AA2 = AA02, A10-UP = A10U)",
        "suggested": "a reading of the name - external store or trailer - to confirm",
        "default": "no X3 equivalent - sent to A001 for UAT"}
tot = Counter()
for how in ("exact", "rule", "suggested", "default"):
    rs = [r for r in rows if r[4] == how]
    s.append([how, MEAN[how], len(rs), sum(r[6] for r in rs), round(sum(r[9] for r in rs), 2)])
    s.cell(s.max_row, 1).fill = PatternFill("solid", fgColor=COL[how])
    s.cell(s.max_row, 5).number_format = "#,##0.00"
s.append(["Total", "", len(rows), sum(r[6] for r in rows), round(sum(r[9] for r in rows), 2)])
s.cell(s.max_row, 5).number_format = "#,##0.00"
for c in s[s.max_row]:
    c.font = Font(bold=True)
for i, w in enumerate([12, 70, 8, 8, 14], 1):
    s.column_dimensions[get_column_letter(i)].width = w

m = wb.create_sheet("Bin mapping")
m.append(["Access warehouse", "Access bin", "X3 location", "X3 location type", "How", "Note",
          "Lots", "Products", "Quantity", "Stock value"])
for c in m[1]:
    c.font, c.fill = HF, FILL
    c.alignment = Alignment(wrap_text=True, vertical="top")
for r in sorted(rows, key=lambda r: ("exact rule suggested default".split().index(r[4]), r[0], r[1])):
    m.append(r)
    m.cell(m.max_row, 5).fill = PatternFill("solid", fgColor=COL[r[4]])
    m.cell(m.max_row, 10).number_format = "#,##0.00"
for i, w in enumerate([10, 14, 12, 10, 11, 44, 7, 9, 11, 13], 1):
    m.column_dimensions[get_column_letter(i)].width = w
m.freeze_panes = "A2"
m.auto_filter.ref = m.dimensions

d = wb.create_sheet("Unmatched - biggest first")
d.append(["Access warehouse", "Access bin", "Lots", "Stock value", "What it looks like"])
for c in d[1]:
    c.font, c.fill = HF, FILL
for r in sorted([r for r in rows if r[4] == "default"], key=lambda r: -r[9]):
    d.append([r[0], r[1], r[6], r[9], r[5]])
    d.cell(d.max_row, 4).number_format = "#,##0.00"
for i, w in enumerate([10, 14, 8, 14, 30], 1):
    d.column_dimensions[get_column_letter(i)].width = w
d.freeze_panes = "A2"

wb.save(f"{BASE}/out_mapped/Stock_bin_mapping_to_X3.xlsx")
json.dump({r[1]: r[2] for r in rows}, open(f"{BASE}/loclist/bin_map.json", "w"))

for how in ("exact", "rule", "suggested", "default"):
    rs = [r for r in rows if r[4] == how]
    print(f"{how:10} bins {len(rs):5}  lots {sum(r[6] for r in rs):5}  value {sum(r[9] for r in rs):14,.2f}")
same = defaultdict(set)
for r in rows:
    same[r[1].upper()].add(r[2])
print("bin names mapping to different X3 locations in different warehouses:",
      sum(1 for v in same.values() if len(v) > 1))
