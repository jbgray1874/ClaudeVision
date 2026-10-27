"""Product-supplier list, built from product_suppliers_v3.sql grid 2.

Every Sage supplier each product has been bought from, ranked by priority, with price,
currency, purchase unit and lead time. Matched to the product codes exactly as they are
spelled in the UAT product file (302 of them are mixed case, and the V6 set keeps the
same spelling), and to the supplier codes in the UAT supplier file.

The columns are mapped to what each field means. The X3 field names wait for Ethan's
product-supplier template: no template, no guessing.
"""
import collections
import csv
import io
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "/tmp/claude-0/-home-user-ClaudeVision/e9ffe9d9-c7dd-5e34-8a16-6cb3c65dadea/scratchpad"
OUT = f"{BASE}/out_psup"
SHOT = ["BUMPERRAIL4", "BUMPERRAIL41", "FIXING150", "FIXING2924", "FIXING403", "SUB-12422-42-SA0",
        "EXTR19", "FIXING16", "FIXING48", "FIXING56", "FIXING1151", "FIXING405", "FIXING909",
        "POWDER312", "SUBMETAL33"]

lines = [l.split("\t") for l in open(f"{BASE}/psup3/out.txt", encoding="utf-8").read().splitlines() if l.strip()]
g2 = [r for r in lines if len(r) == 16]
g3 = [r for r in lines if len(r) == 3]


def read(path):
    return [r for r in csv.reader(open(path, encoding="cp850"), delimiter=";") if len(r) > 1]


uat = {r[1].upper(): r for r in read(f"{BASE}/uploads_final/ZSDIITM_Cleansing.txt")}
z = zipfile.ZipFile(f"{BASE}/out_std6/Standardised_load_files_V6.zip")
v6 = {r[1].upper(): r for r in csv.reader(io.TextIOWrapper(z.open("ZSDIITM_load6_products.txt"), encoding="cp850"),
                                         delimiter=";") if len(r) > 1}
bps = {r[1]: r[2] for r in read(f"{BASE}/upl2/bps.txt")}
old = {r[0].upper(): r[1] for r in list(csv.reader(open(f"{BASE}/out_lists/Preferred_supplier_by_product.csv",
                                                        encoding="utf-8-sig")))[1:]}


def desc(r):
    return next((d.strip() for d in (r[4], r[3], r[2]) if d.strip()), "")


UNIT = {"EA": "EA", "EACH": "EA"}
rows = {"uat": [], "v6": [], "held": []}
flags = collections.Counter()
for (prod, sup, prio, _sname, _d, acc, cur, uom, price, uop, per, lead, n, qty, first, last) in g2:
    key = prod.upper()
    src = uat.get(key) or v6.get(key)
    code = src[1] if src else prod
    note = []
    if uom not in UNIT:
        note.append(f"Access unit '{uom}' - loaded as EA, check the pack size")
    p, u, pp = float(price), float(uop), float(per)
    if p == 0:
        note.append("no price on the last order")
    elif pp > 0 and u > 0 and abs(u / pp - p) > 0.05 * p:
        note.append(f"price basis unclear: {p:g} each, or {u:g} per {pp:g}")
    if sup not in bps:
        note.append("supplier not in the UAT supplier file")
    out = [code, desc(src) if src else "", sup, bps.get(sup, ""), int(prio), UNIT.get(uom, "EA"),
           round(p, 5), cur, int(lead) if lead not in ("", "NULL") else "", int(n), float(qty),
           first, last, acc, "; ".join(note)]
    for k in note:
        flags[k.split(":")[0].split(" - ")[0]] += 1
    rows["uat" if key in uat else "v6" if key in v6 else "held"].append(out)
for k in rows:
    rows[k].sort(key=lambda r: (r[0].upper(), r[4]))

# products in UAT that Access has never bought, by product type
bought = {r[0].upper() for r in rows["uat"]}
never = sorted([[r[1], desc(r), r[0]] for k, r in uat.items() if k not in bought and r[0] in ("RAWML", "FGSL")],
               key=lambda r: (r[2], r[0].upper()))

# ---- the workbook --------------------------------------------------------------
HF, FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
HEAD = ["Product code", "Product description", "Supplier code", "Supplier name", "Priority",
        "Purchase unit", "Purchase price", "Currency", "Lead time (days)", "Order lines",
        "Quantity ordered", "First ordered", "Last ordered", "Access supplier code", "Check"]
WID = [22, 32, 10, 34, 8, 9, 12, 8, 9, 8, 11, 11, 11, 11, 44]


def sheet(ws, head, data, widths):
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = HF, FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for d in data:
        ws.append(d)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


u = rows["uat"]
u_prods = {r[0] for r in u}
tot = collections.Counter(r[0] for r in uat.values())
cov = collections.Counter(uat[p.upper()][0] for p in u_prods)
gained = sorted(p for p in u_prods if p.upper() not in old)

wb = Workbook()
s = wb.active
s.title = "Summary"
for line in [
    ["Product-supplier list - Sage codes, matched to UAT"],
    [],
    ["Built from", "Access purchase orders, every one since history starts (1 September 2022). "
                   "Access holds no supplier against the part itself - STK_STOCK_2's ten supplier columns are empty."],
    ["Product codes", "Exactly as spelled in the UAT product file, which the V6 products, routes, prices and BOMs also use."],
    ["Supplier codes", "Sage codes from the UAT supplier file. Every one is in it."],
    ["Priority", "1 is the supplier last bought from; 2, 3 ... the others, most recent first."],
    [],
    ["", "Rows", "Products"],
    ["For UAT now - product is in UAT", len(u), len(u_prods)],
    ["   of which priority 1", sum(1 for r in u if r[4] == 1), ""],
    ["   of which second and later suppliers", sum(1 for r in u if r[4] > 1), ""],
    ["Next run - product is in the V6 file, not yet in UAT", len(rows["v6"]), len({r[0] for r in rows["v6"]})],
    ["Held back - product is in neither", len(rows["held"]), len({r[0] for r in rows["held"]})],
    [],
    ["Against the first list (7,552 products, priority 1 only)"],
    ["Products gained", len(gained)],
    ["Products lost", sum(1 for p in old if p not in {x.upper() for x in u_prods})],
    ["Priority-1 supplier changed", sum(1 for r in u if r[4] == 1 and old.get(r[0].upper()) not in (None, r[2]))],
    [],
    ["UAT products with a supplier, by type", "With supplier", "In UAT"],
]:
    s.append(line)
for k in ("RAWML", "FGSL", "FGML", "PHANT"):
    s.append([k, cov[k], tot[k]])
s.append(["FGML and PHANT are made in-house, so most need no supplier. RAWML and FGSL without one have "
          "not been bought since September 2022 - see 'No purchase history'."])
s.append([])
s.append(["Checks raised", "Rows"])
for k, v in flags.most_common():
    s.append([k, v])
s.append([])
s.append(["Products in the planning screenshot"])
s.append(["Product", "Suppliers (priority order)", "Shown in X3"])
shown = {"FIXING1151", "FIXING405", "FIXING909", "POWDER312", "SUBMETAL33", "EXTR19"}
for p in SHOT:
    s.append([p, ", ".join(r[2] for r in u if r[0].upper() == p), "yes" if p in shown else "blank"])
s["A1"].font = Font(bold=True, size=13)
s.column_dimensions["A"].width = 52
s.column_dimensions["B"].width = 60
s.column_dimensions["C"].width = 12

sheet(wb.create_sheet("For UAT now"), HEAD, u, WID)
sheet(wb.create_sheet("Next run (V6)"), HEAD, rows["v6"], WID)
sheet(wb.create_sheet("Held back"), HEAD, rows["held"], WID)
sheet(wb.create_sheet("No purchase history"), ["Product code", "Product description", "Product type"],
      never, [24, 40, 10])
sheet(wb.create_sheet("Suppliers without Sage code"),
      ["Access supplier", "Products only they supply", "Last ordered"],
      sorted([[a, int(n), d] for a, n, d in g3], key=lambda r: -r[1]), [14, 12, 12])

m = wb.create_sheet("Mapping")
sheet(m, ["Column", "What it is", "Access source", "X3 field"], [
    ["Product code", "Sage product code", "ASC_POR_TBL.POR_PART_ONLY, cleansed, spelled as in the UAT product file", "Ethan's template"],
    ["Supplier code", "Sage supplier code", "ASC_POR_TBL.POR_SUP_CODE via migration.XREF_SAGE_BPS", "Ethan's template"],
    ["Priority", "1 = main supplier", "Most recent order first", "Ethan's template"],
    ["Purchase unit", "Unit bought in", "POR_UOM; EACH and EA both loaded as EA, the products' unit", "Ethan's template"],
    ["Purchase price", "Price on the last order", "POR_FORN_PRICE, in the order currency", "Ethan's template"],
    ["Currency", "Order currency", "POR_CURR_CODE, £ as GBP", "Ethan's template"],
    ["Lead time (days)", "Average order-to-due days", "POR_DUE_DATE - POR_ORDER_DATE, all orders", "Ethan's template"],
    ["Product description, supplier name", "For reading only", "UAT product and supplier files", "not loaded"],
    ["Order lines, quantity, first/last ordered", "For reading only", "ASC_POR_TBL", "not loaded"],
], [30, 28, 60, 18])
wb.save(f"{OUT}/Product_supplier_list_V2.xlsx")

with open(f"{OUT}/Product_supplier_list_V2_for_UAT.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(HEAD[:1] + HEAD[2:3] + HEAD[4:10])
    for r in u:
        w.writerow([r[0], r[2], r[4], r[5], r[6], r[7], r[8], r[9]])

print("uat", len(u), len(u_prods), "v6", len(rows["v6"]), "held", len(rows["held"]), "gained", len(gained),
      "never", len(never), collections.Counter(r[2] for r in never))
print(flags)
print(cov, tot)
