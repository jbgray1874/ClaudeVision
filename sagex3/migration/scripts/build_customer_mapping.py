"""The customer mapping table: every live Access customer and its Sage X3 code, existing or new.

Puts together
  * the 217 customers already in X3 ("Old to new" tab of Customer_codes_old_to_new.xlsx,
    which is the XREF_SAGE_BPC paste with WCP02 -> WCP101 and HCL02 -> HOT100 corrected), and
  * the new accounts from Customers_to_add.xlsx (build_customers_to_add.py),
and writes
  * the mapping workbook (old code, name, new code, status, open items) and
  * sql/xref_new_customer_accounts.sql, which adds the new pairs to migration.XREF_SAGE_BPC so
    the views stop dropping those customers' open invoices (BICX) and open orders (SOHPJ).
    Codes only - the names come from Access when it runs.

Usage:
    python build_customer_mapping.py <Customer_codes_old_to_new.xlsx> <Customers_to_add.xlsx>
                                     <output .xlsx> <output .sql>
"""
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

IN_X3 = "In Sage X3 UAT"
NEW = "New - created by the next customer import"
NEW_LATER = "New - created once it has an address"


def rows(path, sheet):
    ws = load_workbook(path, read_only=True)[sheet]
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]


def main(old_new, to_add, dest, sql_dest):
    # Open items per account, from the earlier "No new code yet" tab where there is one
    open_items = {r[0]: (r[3], r[4]) for r in rows(old_new, "No new code yet") if isinstance(r[3], (int, float))}

    mapping, left_out = [], []
    for code, name, new, new_name, note in rows(old_new, "Old to new"):
        mapping.append([code, name, new, IN_X3, note or ""])
    in_x3 = {m[0] for m in mapping}
    new_pairs = []
    for code, name, grade, reason, last, act, new, fix in rows(to_add, "Customers to add"):
        if code in in_x3:
            continue
        if act.startswith("Add"):
            mapping.append([code, name, new, NEW_LATER if "address" in act else NEW, fix])
            if act == "Add":
                new_pairs.append((code, new))
        else:
            left_out.append([code, name, act, fix])
    mapping.sort(key=lambda m: m[0].upper())
    for m in mapping:
        m += list(open_items.get(m[0], (None, None)))

    hf, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="44546A")
    newfill = PatternFill("solid", fgColor="FFF2CC")

    def sheet(ws, head, data, widths):
        ws.append(head)
        for r in data:
            ws.append(r)
        for c in ws[1]:
            c.font, c.fill, c.alignment = hf, fill, Alignment(wrap_text=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    n = wb.active
    n.title = "Notes"
    counts = {s: sum(1 for m in mapping if m[3] == s) for s in (IN_X3, NEW, NEW_LATER)}
    for line in [
        ["Customer mapping - old Access code to Sage X3 code"], [],
        ["Mapping", f"{len(mapping)} live customers, one Sage X3 code each: {counts[IN_X3]} already in Sage X3 UAT, "
                    f"{counts[NEW]} new accounts (highlighted) created by the next customer import, "
                    f"{counts[NEW_LATER]} to follow once it has an address."],
        ["New codes", "First three letters of the name (THE ignored), then the next free number. TIE02E = TEC103 and "
                      "TIE02POL = TEC104 are fixed; the other new codes are proposed and only final once loaded."],
        ["Not migrated", f"{len(left_out)} accounts deliberately left out (intercompany and ledger accounts) - Finance to confirm. "
                         "Customers with no open items that have not traded for 2 years are not migrated and are not listed."],
        ["Same name, different code", "Several companies have more than one Access account (per country or currency, e.g. "
                                      "Techtronic, Boots, Kurt Geiger). Each keeps its own new code - check the old code, not the name."],
        ["Open items", "Number and value of open invoices held against the old code, where known (27 Oct 2026)."],
    ]:
        n.append(line)
    n["A1"].font = Font(bold=True, size=13)
    n.column_dimensions["A"].width = 24
    n.column_dimensions["B"].width = 120
    for r in n.iter_rows(min_row=3):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Old to new")
    sheet(ws, ["Old Access code", "Customer name (Access)", "New Sage X3 code", "Status", "Note",
               "Open items", "Open value"], mapping, [14, 40, 14, 40, 70, 11, 14])
    for row in ws.iter_rows(min_row=2):
        if row[3].value != IN_X3:
            for c in row:
                c.fill = newfill
        row[6].number_format = "#,##0.00"
    sheet(wb.create_sheet("Not migrated"), ["Old Access code", "Name", "Why", "To confirm"], left_out, [14, 40, 28, 80])
    wb.save(dest)

    sql = [
        "-- xref_new_customer_accounts.sql  -  generated by scripts/build_customer_mapping.py",
        "--",
        "-- Live customers that were not in the first customer import, with their new Sage X3 codes.",
        "-- Adds them to migration.XREF_SAGE_BPC so vw_BICX_V2 and vw_SOH_V2 stop dropping their open",
        "-- invoices and open orders. The accounts themselves are created by the customer import",
        "-- (new rows in YA004BPC), which has to load before BICX and SOHPJ.",
        "--",
        "-- Run AFTER xref_sage_codes.sql (it drops and rebuilds the table) and BEFORE the extract.",
        "-- Re-runnable.",
        "USE SDILive;",
        "GO",
        "SET NOCOUNT ON;",
        "-- Only our own migration tables change. The table may be an older build whose text columns",
        "-- were sized to what it held then ('new account' and some names didn't fit), so widen them.",
        "ALTER TABLE migration.XREF_SAGE_BPC ALTER COLUMN SAGE_CODE varchar(10) COLLATE DATABASE_DEFAULT NULL;",
        "ALTER TABLE migration.XREF_SAGE_BPC ALTER COLUMN ACCESS_NAME nvarchar(100) COLLATE DATABASE_DEFAULT NULL;",
        "ALTER TABLE migration.XREF_SAGE_BPC ALTER COLUMN SAGE_NAME nvarchar(100) COLLATE DATABASE_DEFAULT NULL;",
        "ALTER TABLE migration.XREF_SAGE_BPC ALTER COLUMN MATCHED_BY varchar(20) COLLATE DATABASE_DEFAULT NULL;",
        "GO",
        "DELETE FROM migration.XREF_SAGE_BPC WHERE MATCHED_BY = 'new account';",
        "INSERT INTO migration.XREF_SAGE_BPC (ACCESS_CODE, SAGE_CODE, ACCESS_NAME, SAGE_NAME, MATCHED_BY)",
        "SELECT p.ACCESS_CODE, p.SAGE_CODE, LTRIM(RTRIM(a.CUNAME)), LTRIM(RTRIM(a.CUNAME)), 'new account'",
        "FROM (VALUES",
    ]
    sql += [f"    ('{c}', '{s}')" + ("," if i < len(new_pairs) - 1 else "") for i, (c, s) in enumerate(new_pairs)]
    sql += [
        ") p (ACCESS_CODE, SAGE_CODE)",
        "JOIN dbo.SL_ACCOUNTS a ON LTRIM(RTRIM(a.CUCODE)) COLLATE DATABASE_DEFAULT = p.ACCESS_CODE COLLATE DATABASE_DEFAULT",
        "WHERE NOT EXISTS (SELECT 1 FROM migration.XREF_SAGE_BPC x",
        "                   WHERE x.ACCESS_CODE = p.ACCESS_CODE COLLATE DATABASE_DEFAULT);",
        "",
        "-- Check: added = expected, and a new code must never already belong to another account",
        f"SELECT COUNT(*) AS new_accounts_added, {len(new_pairs)} AS expected,",
        "       (SELECT COUNT(*) FROM migration.XREF_SAGE_BPC n JOIN migration.XREF_SAGE_BPC o",
        "          ON o.SAGE_CODE = n.SAGE_CODE AND o.ACCESS_CODE <> n.ACCESS_CODE",
        "        WHERE n.MATCHED_BY = 'new account') AS new_codes_clashing",
        "FROM migration.XREF_SAGE_BPC WHERE MATCHED_BY = 'new account';",
        "GO",
        "",
    ]
    open(sql_dest, "w", encoding="utf-8", newline="\n").write("\n".join(sql))
    print(dest, counts, "left out", len(left_out))
    print(sql_dest, len(new_pairs), "pairs")


if __name__ == "__main__":
    main(*sys.argv[1:5])
