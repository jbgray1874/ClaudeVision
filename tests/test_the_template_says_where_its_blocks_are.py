"""The engine reads each block's rows from the template, so widening a block in Excel moves
every write with it (D-303). On 27 Sep 2026 the BOM block was widened again; with fixed rows
every run would have stopped at the layout guard until someone edited code."""
from __future__ import annotations

import copy
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import wb_populate as wb  # noqa: E402


def _template(extra_bom_rows: int):
    """The Estimate sheet's labels and totals formula, with the BOM widened by N rows."""
    d = extra_bom_rows
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Estimate"
    ws.cell(10, 3, "Bill of Materials (Per Unit)")
    ws.cell(51 + d, 3, "Wire")
    ws.cell(52 + d, 3, "Part Description")
    ws.cell(61 + d, 3, "Sheet Steel")
    ws.cell(62 + d, 3, "Part Description")
    ws.cell(82 + d, 3, "Other Sheet Material")
    ws.cell(83 + d, 3, "Part Description")
    ws.cell(92 + d, 3, "Total Material Cost")
    ws.cell(92 + d, 13, f"=(SUM(M11:M{50 + d})+SUM(M{53 + d}:M{60 + d})+SUM(M{63 + d}:M{81 + d})"
                        f"+SUM(M{84 + d}:M{91 + d})+AF{83 + d})")
    ws.cell(94 + d, 3, "Labour")
    ws.cell(95 + d, 3, "Operation")
    ws.cell(168 + d, 3, "Total Labour Cost (Including  Downtime)")
    return ws


def test_a_widened_bom_moves_every_block_below_it():
    cm = copy.deepcopy(wb.CELL_MAP)
    moved = wb.derive_cellmap_from_template(_template(30), cm, [])
    assert (cm["bom"]["first_row"], cm["bom"]["last_row"]) == (11, 80)
    assert (cm["tube"]["first_row"], cm["tube"]["last_row"]) == (83, 90)
    assert (cm["steel"]["first_row"], cm["steel"]["last_row"]) == (93, 111)
    assert (cm["other_sheet"]["first_row"], cm["other_sheet"]["last_row"]) == (114, 121)
    assert (cm["labour"]["first_row"], cm["labour"]["last_row"]) == (126, 197)
    assert cm["powder"] == {"rate_cell": "AF112", "total_cell": "AF113"}
    assert set(moved) == {"bom", "tube", "steel", "other_sheet", "labour"}
    wb._verify_template_matches_cellmap(_template(30), cm, [])     # the guard now agrees


def test_the_current_template_changes_nothing():
    cm = copy.deepcopy(wb.CELL_MAP)
    assert wb.derive_cellmap_from_template(_template(0), cm, []) == {}
    assert cm["powder"] == {"rate_cell": "AF82", "total_cell": "AF83"}


def test_a_formula_the_labels_contradict_moves_nothing():
    ws = _template(30)
    ws.cell(51 + 30, 3, "Something else")          # "Wire" is not where the formula says
    cm = copy.deepcopy(wb.CELL_MAP)
    assert wb.derive_cellmap_from_template(ws, cm, []) == {}
    assert cm["bom"]["last_row"] == wb.CELL_MAP["bom"]["last_row"]


def test_a_block_row_with_no_total_formula_is_filled_from_the_first_row():
    """The widened template had no Total Value formula on BOM rows 32-61 (D-304)."""
    ws = _template(30)
    for r in range(11, 81):
        if not 32 <= r <= 61:
            ws.cell(r, 13, f"=(J{r}*K{r})*(100%+L{r})")
    cm = copy.deepcopy(wb.CELL_MAP)
    wb.derive_cellmap_from_template(ws, cm, [])
    flags = []
    filled = wb.fill_missing_block_totals(ws, cm, flags)
    assert filled["bom"] == list(range(32, 62))
    assert ws.cell(45, 13).value == "=(J45*K45)*(100%+L45)"
    assert any("fill the formula down" in f for f in flags)

