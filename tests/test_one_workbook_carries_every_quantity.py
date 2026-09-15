"""One estimate, and a _qty10 filed beside it that nobody asked for any more.

    "Let's look at collapsing all the s/sheets into one when we have multiple unit
     quantities."                                           — James Gray, SDI, 15 Sep 2026
    "Still generating multiple s/sheets"                     — James Gray, same day, on the
                                                               run that produced both files

BOTH MECHANISMS WERE RUNNING AT ONCE. The quantity sweep predates the Material Price Break
tab: it recalculates the estimate at each quantity and saves each as its own workbook, which
was the only way to see 10 off before the tab existed. The tab now answers the same question
inside one workbook, on the estimators' own template, with the sheet's own LOOKUP against
$D$6 — and the sweep carried on filing copies underneath it.

WHY THE COPIES ARE THE WORSE ANSWER, not merely the redundant one. A variant looks exactly
like a finished estimate: right drawings, right blanks, plausible unit cost, SDI template
around it. That is why every one of them has to open on a READ THIS FIRST page disclaiming
itself — freight priced at a quantity nobody is quoting, bought-ins that never took their
price break. The tab needs no disclaimer, because nothing was recalculated behind anyone's
back.

THE SWEEP STILL RUNS. Its figures are what the Quantity Breaks tab and the report read. What
stops is the filing of copies — and only where the one sheet can actually carry the breaks,
because a machine whose template has not been widened yet must still get its other
quantities rather than silently getting none.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import config                                                       # noqa: E402
import quantity_sweep as QS                                         # noqa: E402


# ── the decision ─────────────────────────────────────────────────────────────────────────

def test_the_one_sheet_being_able_to_carry_them_is_what_settles_it(monkeypatch):
    monkeypatch.setattr(config, "QUANTITY_VARIANT_WORKBOOKS", None, raising=False)
    monkeypatch.setattr(config, "MATERIAL_PRICE_BREAK", {"enabled": True}, raising=False)
    assert QS.file_a_workbook_per_quantity() is False


def test_a_template_that_cannot_carry_them_still_gets_its_quantities(monkeypatch):
    """THE CASE THAT MAKES A FLAT False WRONG. Turn the break tab off — an un-widened
    template on some machine — and the copies are the only way to see 10 off at all."""
    monkeypatch.setattr(config, "QUANTITY_VARIANT_WORKBOOKS", None, raising=False)
    monkeypatch.setattr(config, "MATERIAL_PRICE_BREAK", {"enabled": False}, raising=False)
    assert QS.file_a_workbook_per_quantity() is True


def test_it_can_be_forced_either_way(monkeypatch):
    """A setting somebody can find and change, which is the point of it being in config."""
    monkeypatch.setattr(config, "MATERIAL_PRICE_BREAK", {"enabled": True}, raising=False)
    monkeypatch.setattr(config, "QUANTITY_VARIANT_WORKBOOKS", True, raising=False)
    assert QS.file_a_workbook_per_quantity() is True
    monkeypatch.setattr(config, "MATERIAL_PRICE_BREAK", {"enabled": False}, raising=False)
    monkeypatch.setattr(config, "QUANTITY_VARIANT_WORKBOOKS", False, raising=False)
    assert QS.file_a_workbook_per_quantity() is False


def test_the_shipped_default_is_one_workbook():
    """What an estimator actually gets on the current build, asserted on the real config
    rather than on a patched one."""
    assert config.QUANTITY_VARIANT_WORKBOOKS is None
    assert config.MATERIAL_PRICE_BREAK["enabled"] is True
    assert QS.file_a_workbook_per_quantity() is False


# ── the run obeys it ─────────────────────────────────────────────────────────────────────

def test_the_run_passes_the_decision_to_the_sweep():
    """`save_variants=True` was hard-wired at the call site, so the setting could exist and
    change nothing."""
    src = (ROOT / "src" / "main.py").read_text(encoding="utf-8")
    assert "save_variants=_save" in src
    assert "save_variants=True" not in src, "nothing may file copies unconditionally"
    assert "file_a_workbook_per_quantity" in src


def test_the_sweep_itself_is_not_switched_off():
    """The figures are still needed — the Quantity Breaks tab and the report read them. Only
    the filing of copies stops."""
    src = (ROOT / "src" / "main.py").read_text(encoding="utf-8")
    assert "_swept = _sweep(" in src
    assert "write_quantity_breaks_tab" in src


# ── and the report does not go quiet about it ────────────────────────────────────────────

def test_no_copies_filed_is_said_out_loud_not_left_blank():
    """The sentence naming the filed workbooks used to vanish when there were none, and an
    absent sentence about other quantities reads as "there are none" — the opposite of what
    one workbook carrying all of them means."""
    src = (ROOT / "src" / "estimate_explained.py").read_text(encoding="utf-8")
    assert "Every quantity is in THIS workbook" in src
    assert "No separate copies are filed." in src


def test_the_break_table_is_filled_before_the_sweep_measures_it():
    """The 19:02 book: the SHEET was perfectly right — break row 8 stepped exactly as
    predicted — and the Quantity Breaks comparison was flat at £4.09 material, because the
    sweep set D6 to each quantity while the table was still empty: J14's guard fell to the
    1-off literal and every quantity carried the whole £1.92 of packing. The table is the
    thing the sweep measures, so it goes in first."""
    src = (ROOT / "src" / "main.py").read_text(encoding="utf-8")
    fill = src.index("write_price_breaks(")
    sweep = src.index("_swept = _sweep(")
    tab = src.index("write_quantity_breaks_tab(")
    assert fill < sweep < tab, "fill the table, then recalc against it, then compare"


# ── the mechanism itself, exercised on real cells rather than on source text ─────────────
# The ordering pin above proves the calls stand in the right order; nothing in it proves a
# single cell gets a value. These do — the 19:02 packing line, written and read back.

def _a_book_like_the_template():
    import openpyxl
    wb = openpyxl.Workbook()
    est = wb.active
    est.title = "Estimate"
    wb.create_sheet("Material Price Break")
    return wb


_PACKING_LINE = {
    "code": "PACKAGING", "description": "Bagged and boxed, Howard's stated method",
    "sheet_row": 14,                                        # J14 -> break row 8
    "order_gbp": 1.92,
    "order_gbp_at": {1: 1.92, 10: 2.19, 50: 3.37, 250: 13.09, 1000: 46.69},
}


def test_the_break_row_carries_the_stepped_values_cell_by_cell():
    """Exactly the figures the 19:02 sheet proved right: 1.92 / 0.219 / 0.0674 /
    0.05236 / 0.04669 per unit, read back off the written cells."""
    import material_price_break as MPB
    wb = _a_book_like_the_template()
    res = MPB.write_price_breaks(wb, [dict(_PACKING_LINE)], [10, 50, 250, 1000],
                                 dict(config.MATERIAL_PRICE_BREAK))
    assert res["rows"] == 1, res
    ws = wb["Material Price Break"]
    target = 14 + int(config.MATERIAL_PRICE_BREAK.get("row_offset", -6))
    got = [ws.cell(row=target, column=c).value for c in range(4, 9)]      # D..H
    assert got == [1.92, 0.219, 0.0674, 0.05236, 0.04669], got
    # and padded to the table's full width so LOOKUP's last column is never empty
    assert ws.cell(row=target, column=14).value == 0.04669                # N


def test_the_quantity_vector_lands_on_the_estimate_and_never_descends():
    import material_price_break as MPB
    wb = _a_book_like_the_template()
    MPB.write_price_breaks(wb, [dict(_PACKING_LINE)], [10, 50, 250, 1000],
                           dict(config.MATERIAL_PRICE_BREAK))
    est = wb["Estimate"]
    import re as _re
    _m = _re.match(r"([A-Z]+)(\d+)",
                   str(config.MATERIAL_PRICE_BREAK.get("qty_vector_first_cell", "F180")))
    col, row0 = _m.group(1), int(_m.group(2))
    vec = [est[f"{col}{row0 + i}"].value for i in range(11)]
    assert vec[:5] == [1, 10, 50, 250, 1000], vec
    assert vec == sorted(vec), "LOOKUP requires a non-descending vector"
    assert vec[-1] == 1000, "padded with the last break, not left as formula zeros"


def test_an_empty_fill_puts_the_warning_on_the_tab_itself(tmp_path):
    """The OTHER way the 19:02 book happens: the fill fails or writes nothing and the
    sweep measures an empty table. The comparison still prints — its figures are honest
    reads of the sheet as it stands — but the caution is on the sheet's own face, not
    only in a run log nobody re-opens."""
    import openpyxl
    from quantity_breaks_tab import write_quantity_breaks_tab, SHEET_NAME
    p = tmp_path / "book.xlsx"
    _a_book_like_the_template().save(p)
    swept = {"rows": [{"quantity": 1, "material": 4.09, "labour": 2.0, "unit": 9.57},
                      {"quantity": 10, "material": 4.09, "labour": 1.0, "unit": 6.79}]}
    warning = "The Material Price Break table did NOT fill on this run"
    out = write_quantity_breaks_tab(p, swept, requested=[10], warning=warning)
    assert out == SHEET_NAME
    ws = openpyxl.load_workbook(p)[SHEET_NAME]
    assert warning in str(ws["A4"].value), ws["A4"].value


def test_a_healthy_run_carries_no_warning(tmp_path):
    import openpyxl
    from quantity_breaks_tab import write_quantity_breaks_tab, SHEET_NAME
    p = tmp_path / "book.xlsx"
    _a_book_like_the_template().save(p)
    swept = {"rows": [{"quantity": 1, "material": 2.39, "labour": 2.0, "unit": 7.74}]}
    write_quantity_breaks_tab(p, swept, requested=[1])
    ws = openpyxl.load_workbook(p)[SHEET_NAME]
    assert ws["A4"].value in (None, ""), ws["A4"].value


def test_the_run_log_and_the_tab_warning_come_from_one_place():
    """main computes the warning once, prints it, and hands THE SAME OBJECT to the tab —
    two separately-worded warnings would drift into contradiction."""
    src = (ROOT / "src" / "main.py").read_text(encoding="utf-8")
    assert "_mpb_warning = (" in src
    assert "warning=_mpb_warning" in src
    guard = src.index("_mpb_warning = None")
    sweep = src.index("_swept = _sweep(")
    assert guard < sweep, "the health of the table is judged before the sweep measures it"
