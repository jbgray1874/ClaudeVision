"""A drawing sheet is read as a drawing, and its bill is held against what the sheet states (D-415).

12675-01, 8 Oct 2026. The model under the V1 GA was one undetailed block, correctly refused
(D-414), and the vision read took the GA sheet alone — under a prompt written for a photograph:
"estimate in mm from visible human-scale cues". The sheet states 600 × 400 with R25 corners, a
1250 body on adjustable feet, 2mm STEEL CONSTRUCTION, POWDER COATED RAL 7021, a CENTRAL SOLID
DIVIDER, 20 x CUSTOMER BAGS at an estimated 1.5 kg, and 67.44 kg. The read returned four posts
1,580 long (the top of the bags), twenty bars (the bag count) 530 long (the bag's width), a
1,250 × 300 divider and two 600 × 40 base bars: about 15 kg of steel against 37 kg of stand.
Nothing after the answer checked it, and it was costed at £242.63 a unit.

Now: the sheet's own printed text goes in with the image under a preamble that makes the
dimensions and notes facts, the goods held not parts and a fitting a line only where it is
drawn; the bill is weighed and measured against the sheet's stated body and weight, put back
once with what broke, and what still disagrees is asked on the unit. A render is read exactly
as before — same prompt, same cache key.
"""
from __future__ import annotations

import hashlib
import json
import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import concept_scan as cs  # noqa: E402
import model_takeoff as mt  # noqa: E402

# The words the V1 GA's text layer carries, in the order pdfplumber gives them (abridged).
GA_WORDS = """100 PITCH FOR STACKED BAGS ESTIMATED - TO BE CONFIRMED
1250 16.5 1580 ADJUSTABLE FEET 2mm STEEL CONSTRUCTION WELDED AND FINISHED FLUSH
POWDER COATED - RAL 7021 BLACK GREY 17.82° 400 600 R25 SECTION A-A 2 x STACKS OF 10 BAGS
CENTRAL SOLID DIVIDER 380 556.5 530 662 186.5 ESTIMATED BAG WEIGHT - 1.5kg
20 x CUSTOMER BAGS 67.44kg WEIGHT: SEE PART DRAWINGS MATERIAL: SEE PART DRAWINGS"""

FACTS = {"body_mm": {"height": 1250, "width": 600, "depth": 400},
         "body_basis": "front view 1250 body; plan 600 x 400",
         "material": "MILD STEEL", "thickness_mm": 2, "finish": "POWDER COATED RAL 7021",
         "stated_weight_kg": 67.44, "weight_includes_goods": "yes",
         "goods": "20 customer bags at 1.5 kg", "goods_count": 20, "goods_unit_weight_kg": 1.5,
         "goods_weight_kg": 30}


def _made(name, l, w, qty=1, ops=("laser_cutting",), t=2, drawn="front view"):
    return {"name": name, "kind": "fabricated", "material_guess": "MILD STEEL",
            "assumed_blank_mm": {"length": l, "width": w, "thickness": t},
            "quantity": qty, "operations": list(ops), "drawn": drawn,
            "why_size": "from the sheet"}


def _feet(drawn="front view, labelled ADJUSTABLE FEET"):
    return {"name": "ADJUSTABLE FOOT", "kind": "bought_in", "quantity": 4,
            "assumed_blank_mm": {"length": 0, "width": 0, "thickness": 0},
            "operations": [], "drawn": drawn}


def _the_1205_bill():
    """What the read returned at £242.63: posts, bars, a divider and base bars, ~15 kg."""
    return {"product": {"name": "BAG STAND"}, "sheet_facts": dict(FACTS),
            "parts": [_made("OUTER VERTICAL POST", 1580, 40, 4),
                      _made("HORIZONTAL BAG SUPPORT", 530, 25, 20),
                      _made("CENTRAL SOLID DIVIDER", 1250, 300),
                      _made("BASE FRAME MEMBER", 600, 40, 2),
                      _feet()]}


def _the_sheets_bill():
    """The stand the sheet draws: four 2 mm faces, a base and the divider, on four feet."""
    return {"product": {"name": "BAG STAND"}, "sheet_facts": dict(FACTS),
            "parts": [_made("FRONT / BACK PANEL", 1250, 600, 2),
                      _made("SIDE PANEL", 1250, 400, 2),
                      _made("BASE", 600, 400),
                      _made("CENTRAL SOLID DIVIDER", 1250, 600, drawn="section A-A"),
                      _feet()],
            "unit_operations": ["welding", "powder_coating", "assembly"]}


# ── the prompt: a render is untouched, a sheet is read as a sheet ─────────────────────────

def test_a_render_keeps_its_prompt_and_its_cache_key():
    """No render pack re-asks the model because of this: the render prompt, its version and
    its cache key are exactly what they were."""
    assert cs.CONCEPT_PROMPT_VERSION == "c4"
    assert hashlib.sha256(cs._PROMPT.encode()).hexdigest()[:12] == cs._PROMPT_FINGERPRINT
    render = cs.concept_prompt_text([b"png"])
    assert "THESE ARE DRAWING SHEETS" not in render
    assert cs._cache_key([b"png"], "m") == cs._cache_key([b"png"], "m", "", sheet=False)
    assert cs._cache_key([b"png"], "m", sheet=True, sheet_words=GA_WORDS) != cs._cache_key([b"png"], "m")


def test_the_sheet_prompt_cannot_change_without_its_version():
    got = hashlib.sha256((cs._SHEET_PREAMBLE + cs._SHEET_TEXT_SECTION + cs._RECHECK_SECTION
                          + cs._ENVELOPE_SECTION).encode()).hexdigest()[:12]
    assert got == cs._SHEET_PROMPT_FINGERPRINT, (
        f"the sheet prompt changed. Bump SHEET_PROMPT_VERSION (now {cs.SHEET_PROMPT_VERSION!r}) "
        f"and set _SHEET_PROMPT_FINGERPRINT = {got!r}")


def test_a_sheet_read_is_told_the_sheet_is_the_facts_and_given_its_words():
    text = cs.concept_prompt_text([b"png"], sheet=True, sheet_words=GA_WORDS)
    assert text.startswith("THESE ARE DRAWING SHEETS, NOT PHOTOGRAPHS")
    for rule in ("A DIMENSION ON THE SHEET IS A FACT", "WHAT THE PRODUCT HOLDS IS NOT A PART",
                 "A dimension to the top of the goods\n   is not the height of the product",
                 "A FITTING IS A LINE ONLY WHERE THE SHEET DRAWS OR LABELS IT",
                 '"sheet_facts"', '"drawn"', '"weight_includes_goods"'):
        assert rule in text, rule
    assert "2mm STEEL CONSTRUCTION" in text and "20 x CUSTOMER BAGS" in text
    # The JSON in the preamble is literal braces once formatted, not a format error.
    assert '{"body_mm": {"height": 0' in text
    again = cs.concept_prompt_text([b"png"], sheet=True, sheet_words=GA_WORDS,
                                   recheck="- OUTER VERTICAL POST is sized 1580 × 40 mm")
    assert "DID NOT AGREE WITH WHAT THE SHEET STATES" in again and "1580 × 40" in again


def test_the_sheets_words_come_from_its_own_text_layer(tmp_path):
    import pymupdf

    pdf = tmp_path / "12675-01-GA.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page(width=600, height=400)
        page.insert_text((40, 60), "2mm STEEL CONSTRUCTION")
        page.insert_text((40, 90), "WEIGHT: 67.44kg")
        doc.save(str(pdf))
    words = cs.sheet_text([str(pdf)])
    assert "2mm STEEL CONSTRUCTION" in words and "67.44kg" in words
    assert cs.sheet_text([str(tmp_path / "x-ENQUIRY_BRIEF.pdf")]) == ""
    assert len(cs.sheet_text([str(pdf)] * 2000)) <= cs.SHEET_TEXT_MAX_CHARS


# ── the check: the bill weighed and measured against the sheet ───────────────────────────

def test_the_1205_bill_does_not_agree_with_the_sheet():
    check = cs.sheet_check(_the_1205_bill(), GA_WORDS)
    assert check["checked"] and not check["agrees"]
    assert check["net_weight_kg"] == pytest.approx(37.44)
    assert 13 < check["sighted_weight_kg"] < 17
    words = " ".join(check["failures"])
    assert "OUTER VERTICAL POST is sized 1580 × 40 mm, larger than the product body" in words
    assert "(1250 × 600 × 400 mm, from the sheet)" in words
    assert "parts are missing or undersized" in words and "less 30 kg of 20 customer bags" in words
    assert [b["part"] for b in check["breaches"]] == ["OUTER VERTICAL POST"]


def test_the_stand_the_sheet_draws_contradicts_nothing_but_is_unverified_without_a_model():
    """D-417: the body figures are the read's own when no model gave one — 1,580 is printed on
    the sheet too — so nothing independent says they are the body. Not "agrees"."""
    check = cs.sheet_check(_the_sheets_bill(), GA_WORDS)
    assert check["checked"] and not check["failures"], check["failures"]
    assert check["verdict"] == "unverified" and not check["agrees"]
    assert any("the body size is the read's own figures" in u for u in check["unverified"])
    assert 1.2 < check["ratio"] < 1.7, "gross blanks weigh more than the finished stand"
    assert cs.sheet_check_sentence(check).startswith("CONCEPT BILL UNVERIFIED AGAINST THE SHEET")


def test_a_figure_the_sheet_does_not_print_is_not_checked_against():
    """A depth of 490 the model says the sheet states, and the sheet does not print, is not a
    fact: it is named and left out of the envelope."""
    bill = _the_sheets_bill()
    bill["sheet_facts"] = dict(FACTS, body_mm={"height": 1580, "width": 600, "depth": 490})
    check = cs.sheet_check(bill, GA_WORDS)
    assert "body depth 490" in check["not_on_sheet"]
    assert check["body_mm"]["depth"] == 0.0 and check["body_mm"]["height"] == 1580


def test_a_folded_blank_is_its_unfolded_shape_and_may_exceed_the_body():
    bill = _the_sheets_bill()
    bill["parts"] = [_made("WRAPPED CARCASS", 2000, 1250, ops=("laser_cutting", "folding"))]
    assert not cs.sheet_check(bill, GA_WORDS)["breaches"]


def test_a_sheet_that_states_nothing_is_not_called_sound():
    bill = {"parts": [_made("PANEL", 500, 300)], "sheet_facts": {}}
    check = cs.sheet_check(bill, "")
    assert not check["checked"] and not check["agrees"] and not check["failures"]
    assert cs.sheet_check_sentence(check).startswith("CONCEPT BILL NOT CHECKED")


def test_an_unsettled_goods_split_is_unverified_never_agrees():
    """D-417, the reviewer's case: the read cannot say whether 67.44 kg includes the bags and
    gives no bag weight. A light bill must not come back "agrees"."""
    bill = _the_1205_bill()
    bill["sheet_facts"] = dict(FACTS, weight_includes_goods="not stated", goods_count=0,
                               goods_unit_weight_kg=0, goods_weight_kg=0)
    bill["parts"] = [p for p in bill["parts"] if p["name"] not in ("OUTER VERTICAL POST",
                                                                    "HORIZONTAL BAG SUPPORT")]
    bill["unit_operations"] = ["powder_coating"]
    check = cs.sheet_check(bill, GA_WORDS, envelope=[600, 1250, 400])
    assert not check["agrees"] and check["verdict"] == "unverified"
    assert any("a bill too light could not be caught" in u for u in check["unverified"])
    bill["parts"].append(_made("SLAB", 1250, 600, 20))
    assert "parts are oversized or counted twice" in " ".join(
        cs.sheet_check(bill, GA_WORDS)["failures"])


def test_with_the_bags_weight_printed_both_readings_are_weighed():
    """Count × unit weight, both printed: 30 kg. Not saying whether 67.44 includes them, a bill
    too light under BOTH readings (37.44 and 67.44) fails; one that fits only one is unverified."""
    facts = dict(FACTS, weight_includes_goods="not stated", goods_weight_kg=0)
    light = dict(_the_1205_bill(), sheet_facts=facts)
    light["parts"] = [p for p in light["parts"] if p["name"] != "OUTER VERTICAL POST"]
    check = cs.sheet_check(light, GA_WORDS)
    assert check["weight_readings_kg"] == [67.44, 37.44] and check["goods_weight_kg"] == 30
    assert "parts are missing or undersized" in " ".join(check["failures"])
    both = dict(_the_sheets_bill(), sheet_facts=facts)            # ~55 kg fits both readings
    check = cs.sheet_check(both, GA_WORDS, envelope=[600, 1250, 400])
    assert check["agrees"], check
    heavier = dict(both, parts=both["parts"] + [_made("INNER LINER", 1250, 600, 2)])  # ~78 kg
    check = cs.sheet_check(heavier, GA_WORDS, envelope=[600, 1250, 400])
    assert check["verdict"] == "unverified", check
    assert any("only if that weight excludes the goods" in u for u in check["unverified"])


# ── undrawn fittings are asked, not costed ───────────────────────────────────────────────

def test_a_fitting_the_sheet_does_not_draw_is_set_aside():
    bill = _the_sheets_bill()
    bill["parts"].append(dict(_feet(drawn=""), name="CASTOR"))
    kept, aside = cs.set_aside_undrawn(bill)
    assert [a["name"] for a in aside] == ["CASTOR"]
    assert "ADJUSTABLE FOOT" in [p["name"] for p in kept["parts"]], "a labelled fitting stays"
    # A read that did not answer the sheet's schema says nothing either way: nothing dropped.
    silent = {"parts": [{"name": "CASTOR", "kind": "bought_in", "quantity": 4}]}
    assert cs.set_aside_undrawn(silent) == (silent, [])


def test_a_sheet_that_draws_a_lid_and_no_hinge_asks_rather_than_mints_one():
    answer = {"parts": [_made("LID", 600, 400), _made("BODY PANEL", 600, 1000)]}
    render_parts = cs.parts_from_concept(answer, "JOB")
    assert any(p["description"] == "HINGE" for p in render_parts), "a render still nets the hinge"
    sheet_parts = cs.parts_from_concept(answer, "JOB", sheet=True)
    assert not any(p["description"] == "HINGE" for p in sheet_parts)
    lid = next(p for p in sheet_parts if p["description"] == "LID")
    assert any("draws LID and no hinge" in q["issue"] for q in lid["manufacturing_questions"])
    assert "read off the design-intent sheet" in " ".join(lid["review_flags"])


def test_what_the_check_could_not_settle_is_asked_on_the_unit():
    parts = cs.parts_from_concept(_the_1205_bill(), "JOB", sheet=True)
    unit = cs.unit_assembly_part(parts, _the_1205_bill(), "JOB")
    check = cs.sheet_check(_the_1205_bill(), GA_WORDS)
    n = cs.raise_sheet_questions([unit] + parts, check, [{"name": "CASTOR", "quantity": 4}])
    assert n == 2
    issues = [q["issue"] for q in unit["manufacturing_questions"]]
    assert issues[0].startswith("The sighted bill does not agree with the sheet")
    assert "CASTOR x4 was listed by the read but the sheet does not draw" in issues[1]
    assert cs.raise_sheet_questions([unit] + parts, check, []) == 0, "asked once"


# ── the re-ask: put back once with what broke, the better of two kept ───────────────────

def _sheet_pdf(tmp_path):
    import pymupdf

    pdf = tmp_path / "12675-01-GA.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page(width=600, height=400)
        for n, line in enumerate(GA_WORDS.splitlines()):
            page.insert_text((10, 30 + 14 * n), line, fontsize=5)
        doc.save(str(pdf))
    return pdf


def test_a_bill_that_breaks_the_sheet_is_put_back_once_and_the_better_kept(tmp_path, monkeypatch):
    pdf = _sheet_pdf(tmp_path)
    monkeypatch.setattr(cs, "_cache_dir", lambda: tmp_path / "cache")
    asked = []

    def fake(pngs, model, brief="", **kw):
        asked.append(kw)
        return json.dumps(_the_1205_bill() if not kw.get("recheck") else _the_sheets_bill())

    monkeypatch.setattr(cs, "_call_vision_llm", fake)
    out = cs.read_sheet_concept([str(pdf)])
    assert len(asked) == 2 and asked[0]["sheet"] and "1250" in asked[0]["sheet_words"]
    assert "OUTER VERTICAL POST is sized 1580" in asked[1]["recheck"]
    # No model envelope here, so the best a sheet's own figures can earn is "unverified" (D-417).
    assert not out["check"]["failures"] and out["check"]["second_read_taken"]
    assert out["check"]["verdict"] == "unverified"
    assert out["check"]["first_read_failures"]
    assert "after a second read" in cs.sheet_check_sentence(out["check"])
    # The same pack again replays both answers from the cache: the model is not asked a third time.
    out2 = cs.read_sheet_concept([str(pdf)])
    assert len(asked) == 2 and out2["check"]["verdict"] == "unverified"


def test_a_second_read_that_is_no_better_is_not_taken(tmp_path, monkeypatch):
    pdf = _sheet_pdf(tmp_path)
    monkeypatch.setattr(cs, "_cache_dir", lambda: tmp_path / "cache")
    monkeypatch.setattr(cs, "_call_vision_llm",
                        lambda pngs, model, brief="", **kw: json.dumps(_the_1205_bill()))
    out = cs.read_sheet_concept([str(pdf)])
    assert not out["check"]["agrees"] and out["check"]["rechecked"]
    assert out["check"]["second_read_taken"] is False
    assert "did no better" in cs.sheet_check_sentence(out["check"])


# ── the book says what the check found ───────────────────────────────────────────────────

def test_the_take_off_sentence_says_what_the_check_found_and_names_a_brief_only_when_used():
    base = {"source": "vision_concept", "design": "12675-01-GA", "brief_used": False}
    bad = dict(base, sheet_check={"failures": ["the made parts weigh about 14.8 kg as sized"]})
    s = mt.sentence(bad)
    assert "with the enquiry brief" not in s
    assert "The bill does not agree with the sheet: the made parts weigh about 14.8 kg" in s
    assert mt.banner(bad).startswith(
        "UNCHECKED CONCEPT FIGURE — CONCEPT READ OF A DESIGN-INTENT SHEET — 12675-01-GA")
    assert "not a checked budget" in s
    unsure = dict(base, sheet_check={"agrees": False, "failures": [],
                                     "unverified": ["the body size is the read's own figures"]})
    assert "it is unverified: the body size is the read's own figures" in mt.sentence(unsure)
    assert mt.banner(unsure).startswith("UNCHECKED CONCEPT FIGURE")
    good = dict(base, brief_used=True, sheet_check={"agrees": True, "failures": []})
    assert "with the enquiry brief" in mt.sentence(good) and "and agrees" in mt.sentence(good)
    assert mt.banner(good).startswith("CONCEPT READ OF A DESIGN-INTENT SHEET"), \
        "only a bill that passed every check loses the UNCHECKED label"


def test_the_report_callout_carries_the_check():
    import job_report_html as jrh

    t = {"source": "vision_concept", "design": "12675-01-GA", "design_title": "GA.PDF",
         "chosen_by": "the GA", "parts": 5, "brief_used": False,
         "sheet_check": {"checked": True, "agrees": False,
                         "failures": ["OUTER VERTICAL POST is sized 1580 × 40 mm"]}}
    html = jrh._render_concept_takeoff({"concept_takeoff": t})
    assert "UNCHECKED CONCEPT FIGURE — the bill does not agree with the sheet:" in html
    assert "1580 × 40" in html and "not a checked budget" in html
    assert "with the enquiry brief" not in html and "off the sheet" in html


def test_the_report_concept_section_says_what_the_sheet_stated():
    import job_report_html as jrh

    concept = {"sheet_read": True, "sheet_facts": FACTS,
               "sheet_check": cs.sheet_check(_the_1205_bill(), GA_WORDS),
               "set_aside_undrawn": [{"name": "CASTOR", "quantity": 4}],
               "assumptions": [{"part_number": "JOB-CPT01", "description": "POST",
                                "field": "blank_length_mm", "value": 1580, "cue": "sheet"}]}
    html = jrh._concept_assumptions_section({"concept_read": concept})
    assert "read off the sheet by the vision model" in html
    assert "1250 × 600 × 400 mm" in html and "67.44 kg" in html
    assert "CONCEPT BILL DOES NOT AGREE WITH THE SHEET" in html
    assert "CASTOR x4" in html


def test_the_scan_reads_a_design_intent_sheet_as_a_sheet():
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    assert "_sheet_read = bool(_design_intent_pages)" in src
    assert "concept_scan.read_sheet_concept(\n                    _pack, refresh=_fresh, brief=_brief," in src
    assert "concept_scan.parts_from_concept(_answer, _job_name, sheet=_sheet_read)" in src
    assert "concept_scan.raise_sheet_questions(" in src


# ── D-415c: a lasered or cut-to-length line is a part SDI cuts ───────────────────────────

def test_a_line_whose_route_cuts_it_is_a_part_sdi_cuts():
    import estimate_explained as ee

    assert ee._sdi_cuts("A", {}, {}, {"route_operations": ["laser_cutting", "powder_coating"]})
    assert ee._sdi_cuts("A", {}, {}, {"length": 1580})
    assert ee._sdi_cuts("A", {"A": {"x": 1}}, {}, None)
    assert not ee._sdi_cuts("FOOT", {}, {}, {"route_operations": ["assembly"]})
    rows = ee._missing_drawings(
        [{"code": "JOB-CPT01", "price": 2.0, "qty": 4}, {"code": "FOOT", "price": 11.0, "qty": 4}],
        {}, {}, {},
        record={"JOB-CPT01": {"route_operations": ["laser_cutting"]},
                "FOOT": {"route_operations": []}})
    assert [r["cut"] for r in rows] == [True, False]


# ── D-416: the model's envelope is the body; the goods' numbers are never the steel's ────
#
# James Gray's review of the same run: give the read the extracted GA facts and the model
# envelope as constraints, have it propose the box's panels and route, and refuse an answer
# that spends the bag count, the bag size or the stack height on the steel.

FULL_FACTS = dict(FACTS, goods_count=20, goods_dimensions_mm=[556.5, 530, 662, 186.5, 380],
                  dimensions_to_goods_mm=[1580])
ENVELOPE = [600, 1250, 400]


def _with(bill, facts=FULL_FACTS):
    return dict(bill, sheet_facts=dict(facts))


def test_the_1205_bill_spends_the_bags_on_the_steel_and_fails_for_it():
    check = cs.sheet_check(_with(_the_1205_bill()), GA_WORDS, envelope=ENVELOPE)
    words = " ".join(check["failures"])
    assert "HORIZONTAL BAG SUPPORT x20 takes the goods' count as its quantity" in words
    assert "HORIZONTAL BAG SUPPORT is sized off a dimension of the goods (530 mm)" in words
    assert "OUTER VERTICAL POST is sized off a dimension to the top of the goods (1580 mm)" in words
    assert "(1250 × 600 × 400 mm, from the SolidWorks model's envelope)" in words
    assert check["body_source"] == "the SolidWorks model's envelope"
    took = {(g["part"], g["took"]) for g in check["goods_spent"]}
    assert ("HORIZONTAL BAG SUPPORT", "count") in took


def test_the_stand_the_sheet_draws_passes_every_check_against_the_model_envelope():
    check = cs.sheet_check(_with(_the_sheets_bill()), GA_WORDS, envelope=ENVELOPE)
    assert check["agrees"], check["failures"]
    said = cs.sheet_check_sentence(check)
    assert "from the SolidWorks model's envelope" in said
    assert "no part takes the goods' count or size" in said


def test_a_goods_figure_that_is_also_the_bodys_is_the_bodys():
    """A panel the body's width is not sized off the goods because a bag is that wide too."""
    facts = dict(FULL_FACTS, goods_dimensions_mm=[600, 530])
    check = cs.sheet_check(_with(_the_sheets_bill(), facts), GA_WORDS, envelope=ENVELOPE)
    assert check["agrees"], check["failures"]


def test_one_hook_per_item_held_may_take_the_goods_count_when_it_says_so():
    bill = _with({"parts": [dict(_made("BAG HOOK", 150, 20, 20), quantity_basis="seen")]})
    assert not cs.sheet_check(bill, GA_WORDS)["agrees"]
    bill["parts"][0]["quantity_basis"] = "one per bag, drawn on the front view"
    assert not any("count" in f for f in cs.sheet_check(bill, GA_WORDS)["failures"])


def test_goods_figures_the_sheet_does_not_print_are_not_used():
    facts = dict(FULL_FACTS, goods_count=24, goods_dimensions_mm=[531], dimensions_to_goods_mm=[])
    bill = _with({"parts": [_made("BAR", 531, 25, 24)]}, facts)
    check = cs.sheet_check(bill, GA_WORDS)
    assert not check["goods_spent"]
    assert "goods count 24" in check["not_on_sheet"] and "goods dimension 531" in check["not_on_sheet"]


def test_the_envelope_goes_to_the_model_and_into_the_key():
    text = cs.concept_prompt_text([b"png"], sheet=True, sheet_words=GA_WORDS,
                                  envelope=ENVELOPE, envelope_of="12675-01-Stacking Holder Block")
    assert "12675-01-Stacking Holder Block is one undetailed body 600 × 1250 × 400 mm" in text
    assert "NEVER SPEND THE GOODS' NUMBERS ON THE STEEL" in text
    assert "MAKE THE BODY FROM ITS PANELS" in text
    assert "a general legend or\n   specification block" in text
    assert '"goods_count": 0' in text and '"dimensions_to_goods_mm"' in text
    plain = cs.concept_prompt_text([b"png"], sheet=True, sheet_words=GA_WORDS)
    assert "THE DESIGN'S MODEL" not in plain
    k = cs._cache_key([b"png"], "m", sheet=True, sheet_words=GA_WORDS)
    assert cs._cache_key([b"png"], "m", sheet=True, sheet_words=GA_WORDS, envelope=[]) == k
    assert cs._cache_key([b"png"], "m", sheet=True, sheet_words=GA_WORDS, envelope=ENVELOPE) != k


def test_the_re_ask_carries_the_envelope_both_times(tmp_path, monkeypatch):
    pdf = _sheet_pdf(tmp_path)
    monkeypatch.setattr(cs, "_cache_dir", lambda: tmp_path / "cache")
    asked = []

    def fake(pngs, model, brief="", **kw):
        asked.append(kw)
        bill = _the_1205_bill() if not kw.get("recheck") else _the_sheets_bill()
        return json.dumps(_with(bill))

    monkeypatch.setattr(cs, "_call_vision_llm", fake)
    out = cs.read_sheet_concept([str(pdf)], envelope=ENVELOPE, envelope_of="BLOCK")
    assert [a.get("envelope") for a in asked] == [ENVELOPE, ENVELOPE]
    assert "takes the goods' count" in asked[1]["recheck"]
    assert out["check"]["agrees"] and out["check"]["body_source"] == "the SolidWorks model's envelope"


# ── D-417: "agrees" only when every check ran; labels, gauge and finish are on the bill ──
#
# James Gray's review of ba02402: a figure printed anywhere was accepted (1,580 as the body),
# an unsettled goods split skipped the light-side check, so the old bars-and-posts bill could
# return agrees=True; the labelled divider and feet, the 2 mm gauge and the finish were not
# required; and a failed bill's figure did not say it was unchecked.

# The V1 GA's own text layer as pdfplumber gives it (the views' part, abridged after the legend).
REAL_TEXT = """100 PITCH FOR STACKED BAGS
ESTIMATED - TO BE CONFIRMED
 1250 
 16.5 
 1580 
ADJUSTABLE FEET
2mm STEEL CONSTRUCTION
WELDED AND FINISHED FLUSH
POWDER COATED - RAL 7021 BLACK GREY
 400 
 600 
 R25 
SECTION A-A
2 x STACKS OF 10 BAGS
EMPTY HOLDER
CENTRAL SOLID DIVIDER
STOPS HANDLES BECOMING
TANGLED BETWEEN STACKS
 380 
 556.5 
 530 
 662 
 186.5 
ESTIMATED BAG WEIGHT - 1.5kg
WEIGHT CAN BE ADDED TO BASE IF REQUIRED
20 x CUSTOMER BAGS
67.44kg
SEE PART DRAWINGS
FINISH SPECIFICATIONS:
• POWDERCOATING: BETWEEN 80 - 120 MICRON
• CHROME PLATING: NICKEL LAYER = 8 - 12 MICRON.
• Q195 UP TO 3mm THICK FOR POWDER COATED STEEL"""


def test_the_sheet_labels_its_feet_and_divider_and_nothing_from_the_legend():
    labels = cs.sheet_labels(REAL_TEXT, {"goods": "20 customer bags at 1.5 kg"})
    assert [(l["label"], l["head"]) for l in labels] == [
        ("ADJUSTABLE FEET", "FOOT"), ("CENTRAL SOLID DIVIDER", "DIVIDER")]
    # A label the read lists counts only where the sheet prints it; the goods are never one.
    more = cs.sheet_labels(REAL_TEXT, {"goods": "customer bags",
                                       "labelled_parts": ["EMPTY HOLDER", "BACK PANEL",
                                                          "20 x CUSTOMER BAGS"]})
    assert "EMPTY HOLDER" in [l["label"] for l in more]
    assert "BACK PANEL" not in [l["label"] for l in more], "not printed on the sheet"
    assert not any(l["head"] == "BAG" for l in more)


def test_a_bill_without_the_labelled_divider_or_feet_disagrees():
    bill = _with(_the_sheets_bill())
    bill["parts"] = [p for p in bill["parts"] if p["name"] not in ("CENTRAL SOLID DIVIDER",
                                                                    "ADJUSTABLE FOOT")]
    words = " ".join(cs.sheet_check(bill, REAL_TEXT, envelope=ENVELOPE)["failures"])
    assert "the sheet labels 'ADJUSTABLE FEET' and the bill has no part for it" in words
    assert "the sheet labels 'CENTRAL SOLID DIVIDER' and the bill has no part for it" in words


def test_a_part_off_the_stated_gauge_or_a_bill_without_the_stated_finish_disagrees():
    bill = _with(_the_sheets_bill())
    bill["parts"][0] = _made("FRONT / BACK PANEL", 1250, 600, 2, t=3)
    words = " ".join(cs.sheet_check(bill, REAL_TEXT, envelope=ENVELOPE)["failures"])
    assert "FRONT / BACK PANEL is 3 mm, and the sheet states 2 mm mild steel" in words
    bare = _with(_the_sheets_bill())
    bare["unit_operations"] = ["welding", "assembly"]
    words = " ".join(cs.sheet_check(bare, REAL_TEXT, envelope=ENVELOPE)["failures"])
    assert "nothing in the bill carries powder coating" in words


def test_the_reviewers_case_the_old_bill_with_1580_as_body_and_no_split_is_never_agrees():
    """The read takes 1,580 (printed) as the body, says nothing of the goods, and cannot tell
    whether 67.44 kg includes them. In ba02402 the old bars-and-posts bill came back agreeing."""
    facts = {"body_mm": {"height": 1580, "width": 600, "depth": 400}, "material": "MILD STEEL",
             "thickness_mm": 2, "finish": "POWDER COATED - RAL 7021", "stated_weight_kg": 67.44,
             "weight_includes_goods": "not stated"}
    bill = dict(_the_1205_bill(), sheet_facts=facts, unit_operations=["powder_coating"])
    check = cs.sheet_check(bill, REAL_TEXT)
    assert not check["agrees"] and check["verdict"] in ("unverified", "disagrees")
    assert any("the body size is the read's own figures" in u for u in check["unverified"])
    # The same read with the model's envelope: the 1,580 posts break the 1,250 body outright.
    assert cs.sheet_check(bill, REAL_TEXT, envelope=ENVELOPE)["verdict"] == "disagrees"
    # And a read that names 1,580 both as the body and as a height to the goods contradicts itself.
    said_both = dict(bill, sheet_facts=dict(facts, dimensions_to_goods_mm=[1580]))
    assert ("the read gives 1580 mm as the product body and as a dimension to the top of the goods"
            in " ".join(cs.sheet_check(said_both, REAL_TEXT)["failures"]))


def test_only_a_bill_that_passes_everything_agrees_and_says_what_it_passed():
    check = cs.sheet_check(_with(_the_sheets_bill()), REAL_TEXT, envelope=ENVELOPE)
    assert check["verdict"] == "agrees" and check["agrees"], (check["failures"], check["unverified"])
    said = cs.sheet_check_sentence(check)
    for bit in ("from the SolidWorks model's envelope", "no part takes the goods' count or size",
                "every component the sheet labels is on the bill (ADJUSTABLE FEET, CENTRAL SOLID DIVIDER)",
                "every part of the stated material is at the stated gauge",
                "the stated finish is on the route"):
        assert bit in said, bit


def test_an_unverified_or_failed_figure_is_labelled_unchecked_beside_the_price():
    import openpyxl
    import wb_populate as wp

    for sc in ({"agrees": False, "verdict": "unverified", "failures": [],
                "unverified": ["the body size is the read's own figures"]},
               {"agrees": False, "verdict": "disagrees", "failures": ["a post is 1580"]}):
        ws = openpyxl.Workbook().active
        ws["A1"], ws["B1"] = "Unit Cost", 242.63
        s = {"concept_takeoff": {"source": "vision_concept", "design": "12675-01-GA",
                                 "sheet_check": sc}}
        flags = []
        assert wp._write_concept_takeoff_banner(ws, s, flags) is True
        written = [c.value for c in ws[1] if isinstance(c.value, str) and "CONCEPT" in c.value]
        assert written and written[0].startswith("UNCHECKED CONCEPT FIGURE")
    st = {"concept_takeoff": {"source": "vision_concept", "design": "12675-01-GA",
                              "sheet_check": {"agrees": True, "verdict": "agrees", "failures": []}}}
    ws = openpyxl.Workbook().active
    ws["A1"], ws["B1"] = "Unit Cost", 199.0
    wp._write_concept_takeoff_banner(ws, st, [])
    assert not any(isinstance(c.value, str) and c.value.startswith("UNCHECKED") for c in ws[1])


def test_an_unverified_bill_is_asked_on_the_unit_too():
    parts = cs.parts_from_concept(_the_sheets_bill(), "JOB", sheet=True)
    unit = cs.unit_assembly_part(parts, _the_sheets_bill(), "JOB")
    check = cs.sheet_check(_with(_the_sheets_bill()), REAL_TEXT)
    assert check["verdict"] == "unverified"
    assert cs.raise_sheet_questions([unit] + parts, check, []) == 1
    q = unit["manufacturing_questions"][0]
    assert q["issue"].startswith("The sighted bill could not be verified against the sheet")
    assert "not a checked budget" in q["assumption"]
