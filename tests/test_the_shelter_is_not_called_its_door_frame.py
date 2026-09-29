"""The header named the shelter after one of its doors.

12645, DRS External Shelter for Tesco, 19:17 book of 28 Sep 2026: Estimate!D5 read
"12645-DRS EXTERNAL SHELTER V2" and D4, the Description box beside it, read

    Description    DOOR FRAME

which is 12645-02GA's title, one of three sub-assemblies under the shelter. The customer
quotation was headed with the same words, because it uses the same resolver.

Two things went wrong, one after the other:

  * THE PRODUCT'S OWN SHEET WAS NOT FOUND. The graph's product root is the top sheet's
    whole identity, "12645-DRS EXTERNAL SHELTER V2". Its file "12645 - DRS External Shelter
    V2_REVA.PDF" splits into the number "12645" and the title "DRS External Shelter V2", and
    title_from_files compared the product with the number alone, so the product's own file
    never named the product.
  * A SIBLING WAS TAKEN INSTEAD. With no title found, the last resort picked the shortest
    part number the job owns that is an assembly. That was 12645-02GA.

The fix is the first: the top sheet is recognised by the resolver that recognised it as the
product, so its title is found and the fallback is never reached (D-319). Gating the
fallback itself was measured and left out: on declared runs it blanked the box on 7332-01
and 12349-02, which have no titled product file.
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import product_identity as pi                                            # noqa: E402
import route_compiler as rc                                              # noqa: E402
from client_quote_html import _drawing_identity                          # noqa: E402
from wb_populate import write_job_identity_header                        # noqa: E402

FILES = ["12645 - DRS External Shelter V2_REVA.PDF", "12645-01GA V2_REVA.PDF",
         "12645-02GA_REVA.pdf", "12645-03GA_REVA.pdf"]
ROOT_ID = "12645-DRS EXTERNAL SHELTER V2"
NOTE = "assembly (from the SolidWorks model's own tree)"


def _rec(pn, desc, **over):
    r = {"part_number": pn, "description": desc}
    r.update(over)
    return r


RECORDS = [
    _rec("12645-01-01M", "BACK PANEL"),
    _rec("12645-01GA", "QTY."),
    _rec("12645-01GAV2", "BODY FRAME"),
    _rec("12645-02-01M", "Heel post"),
    _rec("12645-02GA", "DOOR FRAME", is_assembly_parent=True, is_sub_assembly=True),
    _rec("12645-03-01M", "Door panel"),
    _rec("12645-03GA", "DOOR", is_assembly_parent=True, is_sub_assembly=True),
]


def _summary(product=ROOT_ID, files=FILES, declared="12645"):
    return {
        "declared_product": declared,
        "job_source_pdfs": [{"name": f} for f in files],
        "parts": list(RECORDS),
        "estimate_summary": {"canonical_route_shadow": {
            "declared_product": declared, "product_root": product,
            "top_assembly": product, "top_assemblies": [product] if product else [],
            "issues": [],
            "nodes": [{"part_number": product, "description": NOTE}] if product else []}},
    }


def _sheet():
    ws = openpyxl.Workbook().active
    ws["C4"] = "Description"
    ws["C5"] = "Drawing No."
    return ws


# ── the product's own sheet names it ─────────────────────────────────────────────────────

def test_the_top_sheet_names_the_root_the_graph_minted_from_it():
    """Every spelling the graph gives the top sheet finds the title on its file."""
    for root in (ROOT_ID, "12645 - DRS EXTERNAL SHELTER V2",
                 rc.job_drawing_numbers({"job_source_pdfs": FILES})[0]):
        assert pi.title_from_files(root, FILES) == "DRS External Shelter V2", root


def test_a_second_top_sheet_does_not_lend_the_product_its_title():
    """The whole identity is compared, not the bare job number, so another titled sheet
    filed under the same job number is not the product's."""
    files = FILES + ["12645 - Plinth For The Big Shelter_REVA.PDF"]
    assert pi.title_from_files(ROOT_ID, files) == "DRS External Shelter V2"


def test_a_sub_assembly_is_never_titled_from_the_top_sheet():
    assert pi.title_from_files("12645-02GA", FILES) == ""


# ── the header, the quote and the scope line say the same thing ──────────────────────────

def test_the_description_box_is_the_shelter_not_its_door_frame():
    s = _summary()
    num, rev, title = _drawing_identity(s, "12645")
    assert (num, rev, title) == (ROOT_ID, "Rev A", "DRS External Shelter V2")
    ws = _sheet()
    write_job_identity_header(ws, s, "12645")
    assert ws["D4"].value == "DRS External Shelter V2"


def test_the_scope_line_names_the_same_title():
    line = rc.product_scope_sentences(_summary())[0]
    assert line.startswith(f"Priced as {ROOT_ID} (DRS External Shelter V2)"), line


def test_the_quotation_is_headed_by_the_same_title():
    from client_quote_html import build_quote_html
    page = build_quote_html(_summary(), "12645")
    assert re.search(r"<h1>DRS External Shelter V2</h1>", page)
    assert "DOOR FRAME</h1>" not in page


def test_the_products_own_record_still_names_it_under_another_spelling():
    """12349-02 must not regress. The product root is the minted 12349-02-69-GA, and its
    own record 12349-02-69 (the same sheet under the one resolver) carries the title."""
    s = {"parts": [_rec("12349-02-69-GA", NOTE), _rec("12349-02-69", "GRAVITY FEEDER MODULES"),
                   _rec("12349-02-69-04M", "LID")],
         "estimate_summary": {"canonical_route_shadow": {
             "product_root": "12349-02-69-GA", "top_assembly": "12349-02-69-GA",
             "top_assemblies": ["12349-02-69-GA"], "issues": [],
             "nodes": [{"part_number": "12349-02-69-GA", "description": NOTE}]}}}
    assert _drawing_identity(s, "12349-02")[2] == "GRAVITY FEEDER MODULES"


def test_with_no_product_named_the_old_rules_stand():
    """An undeclared hand run keeps the assembly-owns-parts rule (7332-01, 12349-02)."""
    s = {"parts": [_rec("7332-01-101", "FRAME WELDMENT", is_assembly_parent=True),
                   _rec("7332-01-001", "BASE")]}
    assert _drawing_identity(s, "7332-01")[2] == "FRAME WELDMENT"


def test_the_rev_box_is_the_top_sheets_own_revision():
    """Same blind spot, next box along. The Rev box looked for the product's file by its
    first token, found none, and took the first REV in the pack. On 12645 every file is
    REVA, so the 19:17 book's "A" was right by luck. A top sheet at C under a body at A
    must read C."""
    s = _summary(files=["12645-01GA V2_REVA.PDF", "12645 - DRS External Shelter V2_REVC.PDF"])
    assert _drawing_identity(s, "12645")[1] == "Rev C"
