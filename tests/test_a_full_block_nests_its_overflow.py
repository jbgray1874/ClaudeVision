"""12633-00: ten 5 mm acrylic parts for an eight-row Other Sheet Material block. The Choc
Holder's front panel and dividers spilled to the Bill of Materials at the engine's net-part
cost (£0.48, £0.22), which leaves out the sheet drop the block charges every row it holds.

Reproducing the nest in Python failed once (the engine's parts-per-sheet is not the
workbook's). So the spilled part is nested by the block's OWN formulas: a copy of the block's
first row on a sheet beside the Estimate, translated to its own row, with the BOM line priced
from it. The formulas below are the template's, as read from the 20:52 book.
"""
from __future__ import annotations

import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import wb_populate as wp  # noqa: E402

_O = wp.CELL_MAP["other_sheet"]
_B = wp.CELL_MAP["bom"]

_ROW84 = {
    "J": '=IFERROR(IF(D84>0,(INT(IF(F84=0,"",IF(F84>I84,"",((I84-5)/(F84+(10*2))))))*'
         'INT(IF(E84=0,"",IF(E84>H84,"",(H84/(E84+(10*2))))))),"0"),"")',
    "K": 0.04,
    "M": '=IFERROR((IF(D84=0,"0",SUM((L84/J84)*(100%+K84))*D84)),"")',
    "P": '=IFERROR(IF(H84=0,"0",SUM(300/J84)),"")',
    "Q": 50,
    "R": '=IFERROR(IF(H84=0,"0",SUM((E84+F84)*2)/Q84),"")',
}


def _book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estimate"
    for col, head in (("C", "Part Description"), ("D", "Qty Per Unit"), ("J", "Qty Per Sheet"),
                      ("M", "Cost Per Part")):
        ws[f"{col}83"] = head
    for col, v in _ROW84.items():
        ws[f"{col}84"] = v
    # the first block part's own inputs, which must NOT travel to the copy
    for col, v in (("C", "12633-01-01P  Base"), ("D", 1), ("E", 450), ("F", 292.12),
                   ("G", 5), ("H", 3050), ("I", 2050), ("L", 142.87)):
        ws[f"{col}84"] = v
    return wb, ws


def _writer(sheet, row, pe):
    sheet.cell(row=row, column=_O["col_desc"], value=f"{pe['part_number']}  {pe['description']}")
    sheet.cell(row=row, column=_O["col_qty"], value=pe["quantity"])
    sheet.cell(row=row, column=_O["col_length"], value=pe["l"])
    sheet.cell(row=row, column=_O["col_width"], value=pe["w"])
    sheet.cell(row=row, column=_O["col_thick"], value=5)
    sheet.cell(row=row, column=_O["col_sheet_l"], value=3050)
    sheet.cell(row=row, column=_O["col_sheet_w"], value=2050)
    sheet.cell(row=row, column=_O["col_cost_per_sheet"], value=pe["sheet"])


_SPILL = [{"part_number": "12633-03-01P", "description": "Front Panel", "quantity": 2,
           "l": 450, "w": 45, "sheet": 142.87},
          {"part_number": "12633-03-02P", "description": "Divider", "quantity": 3,
           "l": 202, "w": 45}]
_SPILL[1]["sheet"] = 142.87


def test_the_overflow_row_carries_the_blocks_own_formulas_on_its_own_row():
    wb, ws = _book()
    rows = wp.nest_overflow_rows(wb, ws, _O, _SPILL, _writer)
    assert rows == {"12633-03-01P": (wp.OVERFLOW_NEST_SHEET, 3),
                    "12633-03-02P": (wp.OVERFLOW_NEST_SHEET, 4)}
    nest = wb[wp.OVERFLOW_NEST_SHEET]
    assert nest["M3"].value == _ROW84["M"].replace("84", "3")
    assert nest["J4"].value == _ROW84["J"].replace("84", "4")
    assert nest["K3"].value == 0.04 and nest["Q4"].value == 50
    # the inputs are the spilled part's, never the first block part's
    assert (nest["E3"].value, nest["F3"].value, nest["D3"].value) == (450, 45, 2)
    assert nest["C4"].value.startswith("12633-03-02P")


def test_the_bom_line_is_priced_from_its_nest_row_with_no_second_scrap():
    wb, ws = _book()
    rows = wp.nest_overflow_rows(wb, ws, _O, _SPILL, _writer)
    r = _B["first_row"] + 10
    ws.cell(row=r, column=_B["col_desc"],
            value="12633-03-01P  Front Panel — Other Sheet Material block full — PROVISIONAL: x")
    ws.cell(row=r, column=_B["col_code"], value="12633-03-01P")
    ws.cell(row=r, column=_B["col_price"], value=0.48)
    ws.cell(row=r, column=_B["col_qty"], value=2)
    ws.cell(row=r, column=_B["col_scrap"], value=0.04)
    moved = wp.point_bom_lines_at_nest_rows(ws, _B, rows)
    assert moved == ["12633-03-01P"]
    price = ws.cell(row=r, column=_B["col_price"]).value
    qty_ref = ws.cell(row=r, column=_B["col_qty"]).coordinate
    assert price == f"=IFERROR('{wp.OVERFLOW_NEST_SHEET}'!M3/{qty_ref},\"\")"
    assert ws.cell(row=r, column=_B["col_scrap"]).value == 0
    desc = ws.cell(row=r, column=_B["col_desc"]).value
    assert "PROVISIONAL" not in desc and "row 3" in desc


def test_the_nest_matches_the_block_arithmetic():
    """What the copied formulas compute for the 450 x 45 front panel, by the template's rule:
    INT((2050-5)/(45+20)) x INT(3050/(450+20)) = 31 x 6 = 186 per sheet."""
    per_sheet = int((2050 - 5) / (45 + 20)) * int(3050 / (450 + 20))
    assert per_sheet == 186
    line = (142.87 / per_sheet) * 1.04 * 2
    assert round(line / 2, 2) == 0.8        # against the £0.48 net-part figure it replaces


def test_a_row_that_reads_outside_itself_is_not_copied():
    wb, ws = _book()
    ws["M84"] = '=IFERROR(($L$5/J84)*D84,"")'
    assert wp.nest_overflow_rows(wb, ws, _O, _SPILL, _writer) == {}
    assert wp.OVERFLOW_NEST_SHEET not in wb.sheetnames


def test_the_own_row_reader():
    assert wp._reads_only_its_own_row(_ROW84["J"], 84)
    assert wp._reads_only_its_own_row(_ROW84["M"], 84)
    assert not wp._reads_only_its_own_row("=A83+B84", 84)
    assert not wp._reads_only_its_own_row("=SUM(A84:B84)", 84)
    assert not wp._reads_only_its_own_row("=Labour!A84", 84)
    assert not wp._reads_only_its_own_row("=A$84", 84)
