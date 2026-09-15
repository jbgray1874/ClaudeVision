"""A break table that is right at 1, right at 250, and £0.00 at the quantity being quoted.

The 15:40 10975-02 book, measured:

    header   1   10   50   250   1000   1000   1000   1000   1000   1000   1000
    prices  .09  .09  .09   .09    .09      -      -      -      -      -      -
                                                                          ^ £0.00 at 1000

THE HEADER WAS PADDED AND THE PRICES WERE NOT. The padding of the header is deliberate and
correct: the break tab's row 4 is =Estimate!F180..F190, a FORMULA, so an empty source cell
renders as 0 and the vector descends — and LOOKUP over a descending vector does not error,
it returns the wrong column. Repeating the last break keeps it ascending. That was settled
by a run and written down.

What was not carried across is that LOOKUP resolves to the LAST cell holding the largest
value not above $D$6. With 1000 repeated to column N, an order of 1000 lands on N — and N
held nothing. So every break-driven line priced at zero at the top of the table.

It was invisible at 1 and at 250, which are the two quantities anyone checks.

Padding one row of a paired table and not the other is the whole failure. They are one
mechanism and they are now written by one loop.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from material_price_break import write_price_breaks                  # noqa: E402

CFG = {"sheet": "Material Price Break", "estimate_sheet": "Estimate", "row_offset": -6,
       "first_bom_row": 11, "last_bom_row": 50, "qty_vector_first_cell": "F180",
       "first_price_col": 4, "last_price_col": 14}


def _book():
    wb = openpyxl.Workbook()
    wb.active.title = "Estimate"
    wb.create_sheet("Material Price Break")
    return wb


def _line(row=13, unit=0.09):
    return {"sheet_row": row, "code": "TAPE113C", "description": "EPDM TAPE", "unit_gbp": unit}


# ── the failure ──────────────────────────────────────────────────────────────────────────

def test_the_price_row_is_as_wide_as_the_quantity_row():
    wb = _book()
    res = write_price_breaks(wb, [_line()], [1, 10, 50, 250, 1000], CFG)
    assert res["rows"] == 1
    ws, est = wb["Material Price Break"], wb["Estimate"]
    header = [est.cell(180 + i, 6).value for i in range(11)]
    prices = [ws.cell(7, 4 + i).value for i in range(11)]
    assert None not in header, header
    assert None not in prices, f"the price row stops short of the header: {prices}"
    assert len(header) == len(prices)


def test_the_top_quantity_resolves_to_a_price_and_not_a_blank():
    """LOOKUP lands on the LAST column holding the largest value not above $D$6. With the
    last break repeated to the end of the table, that is the last column — so the last
    column has to carry the price."""
    wb = _book()
    write_price_breaks(wb, [_line()], [1, 10, 50, 250, 1000], CFG)
    ws, est = wb["Material Price Break"], wb["Estimate"]
    _last = 14                                                     # column N
    assert est.cell(190, 6).value == 1000, "the header's last column is the top break"
    assert ws.cell(7, _last).value == 0.09, "and so is the price beneath it"


def test_the_padding_is_the_last_break_not_a_repeat_of_the_first():
    """Padding with anything other than the top break would price the largest order at a
    smaller order's rate, which is the error the table exists to prevent."""
    wb = _book()
    write_price_breaks(wb, [{"sheet_row": 13, "code": "BOX481", "order_gbp": 1.89}],
                       [1, 10, 250], CFG)
    ws = wb["Material Price Break"]
    _at_250 = round(1.89 / 250, 5)
    assert ws.cell(7, 6).value == _at_250, "the last asked-for break"
    for col in range(7, 15):
        assert ws.cell(7, col).value == _at_250, f"and every padded column after it ({col})"


def test_a_table_already_full_of_quantities_is_not_padded_over():
    """AN ESTIMATOR'S OWN FIGURE OUTRANKS ANYTHING DERIVED, in the padded columns exactly as
    in the asked-for ones — the padding must not become a way to overwrite typed work."""
    wb = _book()
    ws = wb["Material Price Break"]
    ws.cell(7, 14).value = 0.05                                     # somebody typed it
    res = write_price_breaks(wb, [_line()], [1, 10, 50, 250, 1000], CFG)
    assert ws.cell(7, 14).value == 0.05
    assert res["skipped_occupied"] >= 1


# ── the template's own broken references ─────────────────────────────────────────────────

def test_a_broken_reference_on_an_empty_row_is_cleared():
    """J19 and J20 came out of the widened template as

        =LOOKUP($D$6,'Material Price Break'!$D$4:$N$4,'Material Price Break'!#REF!)

    #REF! propagates through =(J19*K19)*(100%+L19) into M, and M is what the block total
    sums — two untouched empty rows can take Total Material out."""
    wb = _book()
    est = wb["Estimate"]
    est.cell(19, 10).value = "=LOOKUP($D$6,'x'!$D$4:$N$4,'x'!#REF!)"
    res = write_price_breaks(wb, [_line()], [1, 10], CFG)
    assert est.cell(19, 10).value is None
    assert "J19" in (res.get("cleared_broken_refs") or [])


def test_a_broken_reference_on_a_row_with_a_part_is_reported_not_tidied_away():
    """A row carrying a part is the estimator's line. A broken formula there is a fact to
    report — silently deleting it would remove a price cell somebody is relying on."""
    wb = _book()
    est = wb["Estimate"]
    est.cell(21, 8).value = "10975-02-A01"
    est.cell(21, 10).value = "=LOOKUP($D$6,'x'!$D$4:$N$4,'x'!#REF!)"
    res = write_price_breaks(wb, [_line()], [1, 10], CFG)
    assert est.cell(21, 10).value is not None, "left exactly as it is"
    assert any("J21" in r and "#REF!" in r for r in res["refused"]), res["refused"]


def test_a_healthy_formula_is_never_touched():
    wb = _book()
    est = wb["Estimate"]
    _ok = "=LOOKUP($D$6,'Material Price Break'!$D$4:$N$4,'Material Price Break'!D10:N10)"
    est.cell(16, 10).value = _ok
    write_price_breaks(wb, [_line()], [1, 10], CFG)
    assert est.cell(16, 10).value == _ok
