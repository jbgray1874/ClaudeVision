"""Two general arrangements of ONE stand (a colourway pair) collapse to one BOM.

7332-01 carries 7332-01-GA (revK) and 7332-01-GA2 (Rev[A]) — the same Harrods stand in a
different colour. Each GA prints its own weldment number (101, 102), so both became graph roots
and the cascade booked every leaf twice (base 1->2, legs 2->4) and minted two weldments' worth
of weld/dress/powder. That is NOT the 12392 case, where two DIFFERENT arrangements genuinely both
ship; the discriminator is the child set. Same assembly + (near-)identical children -> keep the
highest revision, fold the other in as a colourway variant. Different children -> keep both.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import route_compiler as rc  # noqa: E402


_LEAVES = ["7332-01-001", "7332-01-002", "7332-01-003", "7332-01-008"]


def _assembly(parent, children):
    return {"part_number": parent,
            "children": [{"part_number": c, "qty": 2 if c.endswith("-002") else 1}
                         for c in children]}


def _nodes_by_pn(graph):
    return {n.part_number: n for n in graph["nodes"]}


def test_two_colourway_gas_collapse_keeping_the_higher_revision():
    parts = [
        {"part_number": "7332-01-101", "description": "FRAME WELDMENT", "revision": "K"},
        {"part_number": "7332-01-102", "description": "FRAME WELDMENT (GOLD)", "revision": "A"},
    ] + [{"part_number": c, "description": c} for c in _LEAVES]
    llm = {"assemblies": [_assembly("7332-01-101", _LEAVES),
                          _assembly("7332-01-102", _LEAVES)]}
    g = rc.build_part_graph(parts, llm_extract=llm)
    nodes = _nodes_by_pn(g)

    # the lower-revision twin is gone; the revK structure survives
    assert "7332-01-101" in nodes
    assert "7332-01-102" not in nodes
    # leaves are single-stand, not doubled
    assert nodes["7332-01-001"].qty_per_unit == 1
    assert nodes["7332-01-002"].qty_per_unit == 2      # 2 per stand, not 4
    # the dropped colourway is recorded, not lost
    variants = (g["records"].get("7332-01-101") or {}).get("colourway_variants") or []
    assert any(v.get("part_number") == "7332-01-102" for v in variants)
    # and it rides on the surviving NODE's evidence, where the deliverables read it
    node_variants = (nodes["7332-01-101"].evidence or {}).get("colourway_variants") or []
    assert any(v.get("part_number") == "7332-01-102" for v in node_variants)


def test_a_collapsed_colourway_is_named_on_the_provenance_tab():
    """The dropped GA2 is shown on the AI Provenance tab — in the hierarchy block and in the
    identity block — not lost to a console log. (The Canonical BOM tab it used to sit on was
    folded into AI Provenance.)"""
    import openpyxl
    import wb_populate as wp
    import estimation_report as er

    node = {
        "part_number": "7332-01-101", "description": "FRAME WELDMENT", "kind": "assembly",
        "qty_per_unit": 1, "parents": [], "children": [],
        "evidence": {"colourway_variants": [
            {"part_number": "7332-01-102", "description": "FRAME WELDMENT (GOLD)",
             "revision": "A"}]},
    }
    summary = {"estimate_summary": {
        "part_estimates": [{"part_number": "7332-01-101", "description": "FRAME WELDMENT",
                            "quantity": 1}],
        "canonical_route_shadow": {"nodes": [node], "decisions": []},
    }}
    rows = wp.canonical_bom_rows(summary)
    note = str(rows[0][wp.CANONICAL_BOM_HEADERS.index("Colourway variant")] or "")
    assert "7332-01-102" in note and "colourway variant" in note.lower()

    wb = openpyxl.Workbook()
    wb.active.title = "Estimate"
    er.add_provenance_sheet(wb, summary, {"job_number": "7332"})
    assert "Canonical BOM" not in wb.sheetnames
    text = " ".join(str(c.value) for r in wb["AI Provenance"].iter_rows() for c in r if c.value)
    assert "Colourway collapsed" in text and "7332-01-102" in text


def test_two_different_arrangements_are_both_kept():
    """12392 shape: a panel GA and a bracket-set GA — different child sets, both ship."""
    panel_kids = ["12392-04-001", "12392-04-002"]
    bracket_kids = ["12392-02-010", "12392-02-011"]
    parts = [
        {"part_number": "12392-04-GA", "description": "PANEL GA", "revision": "B"},
        {"part_number": "12392-02-GA", "description": "BRACKET SET GA", "revision": "B"},
    ] + [{"part_number": c, "description": c} for c in panel_kids + bracket_kids]
    llm = {"assemblies": [_assembly("12392-04-GA", panel_kids),
                          _assembly("12392-02-GA", bracket_kids)]}
    g = rc.build_part_graph(parts, llm_extract=llm)
    nodes = _nodes_by_pn(g)
    assert "12392-04-GA" in nodes and "12392-02-GA" in nodes   # neither collapsed


def test_a_single_ga_is_unchanged():
    parts = [{"part_number": "7332-01-101", "description": "FRAME WELDMENT", "revision": "K"}] \
        + [{"part_number": c, "description": c} for c in _LEAVES]
    llm = {"assemblies": [_assembly("7332-01-101", _LEAVES)]}
    g = rc.build_part_graph(parts, llm_extract=llm)
    nodes = _nodes_by_pn(g)
    assert "7332-01-101" in nodes
    assert nodes["7332-01-001"].qty_per_unit == 1
    assert nodes["7332-01-002"].qty_per_unit == 2
