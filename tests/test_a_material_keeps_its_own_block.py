"""A part is charged as the material it was stated as, in a block that can price it (D-602).

Material preservation, Phase 1 (SCRUM-100/107, end to end under SCRUM-134/136). The 30 Oct
sweep found the workbook's Sheet Steel block taking any non-board sheet part, brass, glass
or a material nobody holds a rate for, and charging it at the one mild-steel £/tonne; and
"Aluminium, 3mm" priced under ALUMINIUM but weighed under its own name at steel's density,
2.9 times too heavy, then sent to Sheet Steel because the own-rate route missed the name.

These tests run the real workbook writer on a template fixture: the estimating template is
not in the repo, so the fixture carries the block labels and totals formula the writer
reads, at the template's own rows. Each part's charged cells are read back from the book.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import concept_scan  # noqa: E402
import estimator  # noqa: E402
import material_preservation as mp  # noqa: E402
import wb_populate  # noqa: E402


# ── the template fixture ─────────────────────────────────────────────────────────

def _template(path: Path) -> Path:
    """The Estimate sheet's block labels and totals formula at the template's rows."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Estimate"
    ws.cell(10, 3, "Bill of Materials (Per Unit)")
    ws.cell(51, 3, "Wire")
    ws.cell(52, 3, "Part Description")
    ws.cell(61, 3, "Sheet Steel")
    ws.cell(62, 3, "Part Description")
    ws.cell(82, 3, "Other Sheet Material")
    ws.cell(83, 3, "Part Description")
    ws.cell(92, 3, "Total Material Cost")
    ws.cell(92, 13, "=(SUM(M11:M50)+SUM(M53:M60)+SUM(M63:M81)+SUM(M84:M91)+AF83)")
    ws.cell(94, 3, "Labour")
    ws.cell(95, 3, "Operation")
    ws.cell(168, 3, "Total Labour Cost (Including  Downtime)")
    book.save(path)
    return path


def _concept_part(name, form, section, material, wall=0.0):
    answer = {"product": {"name": "X"}, "unit_operations": ["assembly"], "parts": [{
        "name": name, "kind": "fabricated", "sighted_material": material.lower(),
        "material_guess": material,
        "assumed_blank_mm": {"length": 600, "width": section[0], "thickness": section[1]},
        "stock": {"form": form, "section_mm": section, "wall_mm": wall, "length_mm": 600,
                  "aperture_mm": [], "wire_mm": 0},
        "quantity": 1, "operations": ["saw"], "why_size": "drawing"}]}
    parts = concept_scan.parts_from_concept(answer, "JOB")
    parts = parts if isinstance(parts, list) else parts.get("parts", parts)
    p = [q for q in parts if name in str(q.get("description", ""))][0]
    if wall:
        p["section_stock"].update({"a": float(section[0]), "b": float(section[1]), "t": float(wall)})
    p["part_number"] = name.replace(" ", "-")
    return p


def _sheet_part(pn, material, thickness=3.0):
    return {"part_number": pn, "description": f"{pn} PANEL", "normalized_material": material,
            "normalized_thickness_mm": thickness, "quantity": 1,
            "blank_length_mm": 600.0, "blank_width_mm": 300.0}


def _costed(part):
    part["material_estimate"] = estimator.estimate_material(part)
    me = part["material_estimate"]
    part["unit_cost_gbp"] = me.get("cost_per_part_gbp")
    part["extended_total_cost_gbp"] = me.get("extended_material_cost_gbp")
    return part


@pytest.fixture(scope="module")
def book(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("d602")
    os.environ["SDI_WB_TEMPLATE"] = str(_template(tmp / "template.xlsx"))
    old_out = wb_populate.CELL_MAP["output_dir"]
    wb_populate.CELL_MAP["output_dir"] = str(tmp / "out")
    try:
        parts = [_costed(p) for p in (
            _sheet_part("P-MS", "MILD_STEEL", 1.5),
            _sheet_part("P-ALU", "Aluminium, 3mm"),
            _concept_part("TIMBER BATTEN", "flat_bar", [50, 38], "TIMBER"),
            _concept_part("STEEL BAR", "flat_bar", [50, 38], "MILD STEEL"),
            _concept_part("STEEL TUBE", "box_section", [25, 25], "MILD STEEL", wall=1.5),
            _concept_part("PVC SECTION", "flat_bar", [50, 10], "PVC"),
            _sheet_part("P-BRASS", "BRASS", 1.5))]
        # A reader stamped the brass as sheet stock: the route that sent it to Sheet Steel.
        parts[-1]["material_estimate"].update({"stock_form": "sheet", "thickness_mm": 1.5,
                                               "blank_length_mm": 600, "blank_width_mm": 300})
        summary = {"part_estimates": parts, "job_folder": str(tmp),
                   "estimate_summary": {"part_estimates": parts, "canonical_route_shadow": {
                       "decisions": [], "nodes": [], "mode": "cutover"}}}
        out = wb_populate.populate_workbook(summary, "D602")
    finally:
        wb_populate.CELL_MAP["output_dir"] = old_out
        os.environ.pop("SDI_WB_TEMPLATE", None)
    assert out, "the workbook writer produced no book"
    ws = openpyxl.load_workbook(out)["Estimate"]
    cm = wb_populate.CELL_MAP
    rows = {}
    for block in ("bom", "steel", "other_sheet"):
        for r in range(cm[block]["first_row"], cm[block]["last_row"] + 1):
            desc = str(ws.cell(r, 3).value or "")
            if desc:
                rows[desc.split()[0]] = (block, r)
    return {"ws": ws, "rows": rows, "parts": {p["part_number"]: p for p in parts}}


def _block(book, pn):
    assert pn in book["rows"], f"{pn} is not on the sheet at all"
    return book["rows"][pn][0]


def _cell(book, pn, block_key, col_key):
    r = book["rows"][pn][1]
    return book["ws"].cell(r, wb_populate.CELL_MAP[block_key][col_key]).value


# ── the reviewer's four cases, read back from the charged cells ───────────────────

def test_a_timber_batten_is_charged_as_timber(book):
    assert _block(book, "TIMBER-BATTEN") == "bom"
    assert _cell(book, "TIMBER-BATTEN", "bom", "col_price") < 2.0, "0.68 kg of timber"
    assert book["parts"]["TIMBER-BATTEN"]["material_estimate"]["material"] == "TIMBER"


def test_a_solid_metal_bar_is_charged_at_its_own_weight(book):
    assert _block(book, "STEEL-BAR") == "bom"
    assert 8.5 < book["parts"]["STEEL-BAR"]["material_estimate"]["unit_material_mass_kg"] < 9.5
    assert _cell(book, "STEEL-BAR", "bom", "col_price") > 20.0


def test_a_hollow_tube_is_charged_as_a_tube(book):
    assert _block(book, "STEEL-TUBE") == "bom"
    assert 0.5 < book["parts"]["STEEL-TUBE"]["material_estimate"]["unit_material_mass_kg"] < 0.8
    assert 0 < _cell(book, "STEEL-TUBE", "bom", "col_price") < 10.0


def test_a_plastic_section_is_never_weighed_or_charged_as_steel(book):
    assert _block(book, "PVC-SECTION") == "bom"
    assert book["parts"]["PVC-SECTION"]["material_estimate"]["material"] == "PVC"
    assert _cell(book, "PVC-SECTION", "bom", "col_price") in (None, 0, ""), (
        "PVC has no rate: an honest gap, not a steel price")


# ── sheet parts: the block that can price them ─────────────────────────────────────

def test_mild_steel_stays_in_sheet_steel(book):
    assert _block(book, "P-MS") == "steel"


def test_brass_never_reaches_sheet_steel_and_is_asked(book):
    assert _block(book, "P-BRASS") == "bom"
    flags = book["parts"]["P-BRASS"].get("review_flags") or []
    ask = [f for f in flags if f.startswith("NOT SHEET STEEL")]
    assert ask and "Confirm a price for BRASS" in ask[0]


def test_aluminium_with_its_gauge_in_the_name_is_aluminium(book):
    me = book["parts"]["P-ALU"]["material_estimate"]
    assert 1.4 < me["unit_material_mass_kg"] < 1.5, "600 x 300 x 3 aluminium, not 4.24 kg of steel"
    assert _block(book, "P-ALU") == "other_sheet", "at aluminium's own rate, not mild steel's"
    per_sheet = _cell(book, "P-ALU", "other_sheet", "col_cost_per_sheet")
    assert per_sheet and 60 < per_sheet < 80, "2500 x 1250 x 3 at 2700 kg/m3 and £2.75/kg"


# ── the preservation check itself ─────────────────────────────────────────────────

def test_the_same_material_under_the_rate_lookups_resolution():
    assert mp.same_material("Aluminium, 3mm", "ALUMINIUM")
    assert mp.same_material("MILD_STEEL", "MILD STEEL")
    assert not mp.same_material("TIMBER", "MILD STEEL")


def test_an_unrecorded_change_of_material_is_asked():
    p = {"part_number": "B1", "normalized_material": "TIMBER",
         "material_estimate": {"material": "MILD STEEL"}}
    q = mp.material_question(p)
    assert q and q.startswith("MATERIAL NOT PRESERVED: B1 is stated as TIMBER and charged as "
                              "MILD STEEL") and "confirm" in q


def test_a_recorded_change_is_not_asked_twice():
    p = {"part_number": "B1", "normalized_material": "TIMBER",
         "material_estimate": {"material": "MILD STEEL"},
         "review_flags": ["MATERIAL CONFLICT: the reading says 'TIMBER', but the part is a "
                          "25x25x1.5 hollow section ... confirm the material"]}
    assert mp.material_question(p) is None


def test_only_steel_the_block_prices_is_sheet_steel():
    for ok in ("", "MILD_STEEL", "2MM MS", "ZINTEC", "GALVANISED STEEL", "CR4", "MILDSTEEL"):
        assert mp.sheet_steel_can_price(ok), ok
    for no in ("STAINLESS STEEL", "ALUMINIUM", "Aluminium, 3mm", "BRASS", "GLASS", "COPPER"):
        assert not mp.sheet_steel_can_price(no), no


def test_the_writer_asks_every_unrecorded_change_where_the_money_is_placed():
    src = (ROOT / "src" / "wb_populate.py").read_text(encoding="utf-8")
    body = src.split("THE MATERIAL IS PRESERVED WHERE THE MONEY IS PLACED")[1][:1500]
    assert "sheet_steel_question" in body and "material_question" in body
    assert src.index("THE MATERIAL IS PRESERVED WHERE THE MONEY IS PLACED") < src.index(
        "# ── BOM block: desc, code, PRICE (engine), qty, scrap")


def test_density_follows_the_key_the_price_resolves_to():
    assert estimator._density_for_material("Aluminium, 3mm") == 2700
    assert estimator._density_for_material("ALUMINIUM") == 2700
    assert estimator._density_for_material("GLASS") is None
