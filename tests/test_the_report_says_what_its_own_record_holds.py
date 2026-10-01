"""The report renders what its own record holds — shares that add up, names the template
uses, reasons that were recorded, counts from one list.

12173-02 Card Spinner (M&S), report of 1 Oct 2026, hierarchy section and sections 3, 5, 11:

  * a tab under two frames printed "16 / £0.33" under EACH, so siblings did not add up to
    their parent (and the ticket strip did the same under the pocket ×8 and the rack ×1);
  * frame weld 12173-03-201 was titled "FRONT FRAME ASSEMBLY 1" — row 1 of its own parts
    list, because on SDI's template DESCRIPTION is only the table's column header;
  * the riser and the hook arm were "costed by length on the Tube block" — the template's
    WIRE block (Tube is a labour department);
  * "treat any £0 on this line as MISSING, not free" beside £0.31 charged at Estimate!84;
  * the meshes' Dimensions cell printed a bare "1.0 mm", the pack's document-level figure;
  * "a difference of £2.88 on the sheet and on no line here" — the sheet's POWDER row;
  * section 11 guessed "a gauge, a blank size or a labour rate" for an MFC panel whose gauge
    and blank were known and whose labour was charged — its board had no price;
  * section 3 opened "None of the following change the arithmetic" over weld cues that
    drove 16 Weld and Dress rows;
  * 29 parts listed "Welding on the drawing" over a tally of 27, from two records.

Synthetic inputs only; the job's facts are in the docstrings.
"""
from __future__ import annotations

import copy
import html
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import costed_facts as CF                                                 # noqa: E402
import document_builder as DB                                             # noqa: E402
import estimator                                                          # noqa: E402
import estimator_inputs as EI                                             # noqa: E402
import extractor_patterns as EP                                           # noqa: E402
import job_report_html as J                                               # noqa: E402
import plain_english as PE                                                # noqa: E402
import route_compiler as RC                                               # noqa: E402
import wb_populate as WB                                                  # noqa: E402


def _text(page: str) -> str:
    return html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", page or "")))


# ── 2.0 a line under two parents is shared, not printed twice ─────────────────────────

def _node(pn, kind, parents, q, kids=(), byp=None):
    d = {"part_number": pn, "kind": kind, "parents": list(parents), "qty_per_unit": q,
         "children": [{"part_number": c, "qty": n} for c, n in kids]}
    if byp:
        d["qty_by_parent"] = byp
    return d


def _line(pn, q, gbp, row):
    return {"part_number": pn, "identity": pn, "kind": "leaf", "qty_per_unit": q,
            "charged_ext_gbp": gbp, "charged_cell": f"Estimate!M{row}", "price_origin": {}}


_BASE = [_node("A", "assembly", [], 1, [("B", 1), ("C", 1)]),
         _node("B", "assembly", ["A"], 1, [("X", 1), ("T", 8)]),
         _node("C", "assembly", ["A"], 1, [("Y", 1), ("T", 8)]),
         _node("X", "leaf", ["B"], 1), _node("Y", "leaf", ["C"], 1)]
_REC = {"lines": [_line("X", 1, 6.968, 11), _line("Y", 1, 6.968, 12), _line("T", 16, 0.3328, 93)],
        "run": {"material_gbp": 14.2688}}


def _tree(nodes, rec=_REC):
    return J._render_bom_tree({"canonical_route_shadow": {"nodes": nodes}}, rec)


def _qty_cells(html, pn):
    return re.findall(rf'<td><code>{pn}</code>.*?<td class="n">([\d.]+)</td>', html)


def test_a_shared_leaf_is_split_between_its_parents_and_siblings_add_up():
    """12173-03-06M under 202 and 203: 8 + 8, £0.17 + £0.17, and 202 + 203 = 201."""
    h = _tree(_BASE + [_node("T", "leaf", ["B", "C"], 16)])
    assert h.count("members charged £7.13") == 2 and "members charged £14.27" in h
    assert _qty_cells(h, "T") == ["8", "8"]
    assert "the sheet charges this line once" in h and "of 16 / £0.33 at Estimate!M93" in h
    assert "£7.30" not in h, "the whole line printed under each parent is the defect"


def test_the_cascades_own_split_wins_over_the_raw_edge():
    """The cascade can keep a part's own count over an edge; the tree follows the cascade,
    the same figures the BOMs & Routes trail prints."""
    nodes = _BASE + [_node("T", "leaf", ["B", "C"], 12, byp={"B": 4, "C": 8})]
    rec = {**_REC, "lines": _REC["lines"][:2] + [_line("T", 12, 0.30, 93)],
           "run": {"material_gbp": 14.236}}
    assert _qty_cells(_tree(nodes, rec), "T") == ["4", "8"]


def test_with_no_counts_the_line_is_charged_once_under_its_first_parent():
    """No count from any reader: money is not split on a guess."""
    nodes = [_node(n["part_number"], n["kind"], n["parents"], n["qty_per_unit"],
                   [(c["part_number"], 0 if c["part_number"] == "T" else c["qty"])
                    for c in n["children"]]) for n in _BASE]
    h = _tree(nodes + [_node("T", "leaf", ["B", "C"], 16)])
    assert "members charged £7.30" in h and "members charged £6.97" in h
    assert "charged there" in h


def test_a_shared_sub_assembly_is_drawn_once_with_a_reference_row():
    """A sub-assembly under two parents: drawn in full under one, a share row under the
    other, and the root still adds up once."""
    nodes = [_node("A", "assembly", [], 1, [("B", 1), ("C", 1)]),
             _node("B", "assembly", ["A"], 1, [("S", 1)]),
             _node("C", "assembly", ["A"], 1, [("S", 1)]),
             _node("S", "assembly", ["B", "C"], 2, [("Z", 1)]),
             _node("Z", "leaf", ["S"], 2)]
    rec = {"lines": [_line("Z", 2, 4.0, 20)], "run": {"material_gbp": 4.0}}
    h = _tree(nodes, rec)
    t = _text(h)
    assert h.count("<summary><code>S</code>") == 1, "drawn once"
    assert "rendered in full under B; this parent's share £2.00" in t
    assert t.count("members charged £4.00") == 2, "the root and S itself"
    assert t.count("members charged £2.00") == 2, "B and C, each its share"
    assert "shared: £2.00 under B, £2.00 under C" in t


def test_the_cascade_records_each_parent():
    """route_compiler: PartNode.qty_by_parent is the per-parent figure the trail prints."""
    parts = [{"part_number": "A", "quantity": 1, "is_assembly_parent": True},
             {"part_number": "B", "quantity": 1, "is_sub_assembly": True},
             {"part_number": "C", "quantity": 1, "is_sub_assembly": True},
             {"part_number": "T", "quantity": 8}]
    extract = {"assemblies": [
        {"part_number": "A", "children": [{"part_number": "B", "qty": 1},
                                          {"part_number": "C", "qty": 1}]},
        {"part_number": "B", "children": [{"part_number": "T", "qty": 8}]},
        {"part_number": "C", "children": [{"part_number": "T", "qty": 8}]}]}
    g = RC.build_part_graph(parts, extract, [], ["A"], declared_product="A")
    t = {n.part_number: n for n in g["nodes"]}["T"]
    assert t.qty_by_parent == {"B": 8.0, "C": 8.0}
    assert sum(t.qty_by_parent.values()) == t.qty_per_unit == 16


# ── 2.1 a parts-table heading is not a title ─────────────────────────────────────────

_PAGE = ("ITEM DWG NO.\nDESCRIPTION\nQTY\n1\nA-100-202\nFRONT FRAME ASSEMBLY\n1\n2\nA-100-203\n"
         "SIDE FRAME ASSEMBLY\n1\nA-100-201\nDRAWING No\nFRAME WELD ASSEMBLY\nDRAWING")


def test_a_parts_table_heading_is_not_read_as_a_title():
    tb = EP.extract_title_block_fields(_PAGE)
    assert tb["descriptions"] == []
    assert tb["descriptions_refused"] and tb["descriptions_refused"][0].startswith("QTY 1")
    assert EP._extract_description_candidates(
        "DESCRIPTION LENGTH QTY 1 30 x 30 x 2mm TUBE 1532 1") == []


def test_a_real_title_field_survives():
    assert EP._extract_description_candidates(
        "DESCRIPTION: SHELF BRACKET\nMATERIAL: MILD STEEL") == ["SHELF BRACKET"]


def test_the_gate_every_reader_passes():
    assert not DB._is_good_description("QTY 1 A-100-202 FRONT FRAME ASSEMBLY 1 2 A-100-203")
    assert DB._is_good_description("FRAME WELD ASSEMBLY")


def _index(descriptions, bom_rows=()):
    summary = {"document_analysis": {"bom_rows": list(bom_rows)}, "pages": [{
        "page_number": 6, "page_role": {"primary_role": "detail"}, "pattern_summary": {},
        "page_analysis": {"title_block": {"drawing_numbers": ["A-100-201"],
                                          "descriptions": list(descriptions)}}}]}
    return next(p for p in DB.build_part_index(summary) if p["part_number"] == "A-100-201")


def test_the_part_record_carries_no_bled_title():
    p = _index(["QTY 1 A-100-202 FRONT FRAME ASSEMBLY 1 2 A-100-203 SIDE FRAME ASSEMBLY 1"])
    assert p["description"] in (None, "")
    assert p["description_refused"], "what was refused is kept, so the gap can be said"


def test_a_row_of_its_own_page_with_its_count_is_not_its_title():
    """The cleaned signature from any reader: '<row words> <row qty>' off the same page."""
    row = {"part_number": "A-100-202", "description": "FRONT FRAME ASSEMBLY", "quantity": 1,
           "source_page": 6}
    assert _index(["FRONT FRAME ASSEMBLY 1"], [row])["description"] in (None, "")
    # A weldment titled like one of its members, with no count after it, is untouched.
    assert _index(["FRONT FRAME ASSEMBLY"], [row])["description"] == "FRONT FRAME ASSEMBLY"


def test_an_untitled_assembly_is_said_not_invented():
    parts = [{"part_number": "GA", "quantity": 1, "is_assembly_parent": True},
             {"part_number": "SA", "quantity": 1, "is_sub_assembly": True,
              "description_refused": ["QTY 1 P-2 FRONT FRAME 1"]},
             {"part_number": "SB", "quantity": 1, "is_sub_assembly": True},
             {"part_number": "P-1", "quantity": 1}, {"part_number": "P-2", "quantity": 1}]
    extract = {"assemblies": [
        {"part_number": "GA", "children": [{"part_number": "SA", "qty": 1},
                                           {"part_number": "SB", "qty": 1}]},
        {"part_number": "SA", "children": [{"part_number": "P-1", "qty": 1}]},
        {"part_number": "SB", "children": [{"part_number": "P-2", "qty": 1}]}]}
    g = RC.build_part_graph(parts, extract, [], ["GA"], declared_product="GA")
    hits = [i for i in g["issues"] if i.get("code") == "no_description_from_any_reader"]
    assert [i["part_number"] for i in hits] == ["SA"], "only where text was read and refused"
    assert {n.part_number: n.description for n in g["nodes"]}["SA"] == ""


# ── 2.2 a block is named as the template names it ─────────────────────────────────────

def _wire_summary(rows, part_extra=None):
    part = {"part_number": "P-1", "description": "ARM", "quantity": 7,
            "normalized_material": "MILD STEEL"}
    part.update(part_extra or {})
    return part, {"estimate_summary": {"part_estimates": [part], "final_estimate": {
        "material_rows": rows, "totals": {"material_gbp": 0.31, "labour_gbp": 0,
                                          "unit_gbp": 0.31}}},
        "workbook_totals": {"source": "excel_calculated"}}


_WIRE_ROW = [{"block": "tube", "description": "P-1  ARM", "qty_per_unit": 7,
              "total_value_gbp": 0.31, "workbook_row": 84, "charged_cell": "Estimate!M84"}]


def test_a_wire_row_names_the_wire_block():
    _, s = _wire_summary(copy.deepcopy(_WIRE_ROW))
    lab = CF.costed_job(s)["lines"][0]["price_origin"]["label"]
    assert "on the Wire block — Estimate!84" in lab and "Tube" not in lab
    assert WB.block_title("tube") == "Wire" and WB.block_title("steel") == "Sheet Steel"


def test_the_block_title_is_read_from_the_template():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, label in ((10, "Bill of Materials"), (51, "Wire (cut lengths)"), (61, "Sheet Steel"),
                     (82, "Other Sheet Material"), (92, "Total Material Cost")):
        ws.cell(r, 3).value = label
    ws.cell(92, 13).value = "=(SUM(M11:M50)+SUM(M53:M60)+SUM(M63:M81)+SUM(M84:M91)+AF83)"
    cm = copy.deepcopy(WB.CELL_MAP)
    WB.derive_cellmap_from_template(ws, cm)
    assert cm["tube"]["title"] == "Wire (cut lengths)"
    assert cm["steel"]["title"] == "Sheet Steel"


# ── 2.3 "treat any £0 as missing" only while the line is £0 ───────────────────────────

_NO_GEO = PE.NO_GEOMETRY_SENTENCE


def test_a_charged_line_states_what_its_money_rests_on():
    part, s = _wire_summary(copy.deepcopy(_WIRE_ROW), {
        "native_material_without_geometry": True, "review_flags": [_NO_GEO],
        "material_estimate": {"stock_estimate": {"wire_length_mm": 120}}})
    rec = CF.costed_job(s)
    flags = rec["lines"][0]["review_flags"]
    assert _NO_GEO not in flags
    assert any("£0.31" in f and "Estimate!M84" in f and "an unrecorded reader" in f
               for f in flags), flags
    notes = J.part_review_notes([part], rec["lines"])
    assert not any("MISSING, not free" in n["note"] for n in notes)


def test_an_uncharged_line_keeps_the_warning():
    part, s = _wire_summary([], {"native_material_without_geometry": True,
                                 "review_flags": [_NO_GEO]})
    rec = CF.costed_job(s)
    assert _NO_GEO in rec["lines"][0]["review_flags"]
    assert any("MISSING, not free" in n["note"] for n in J.part_review_notes([part], rec["lines"]))


# ── 2.4 a gauge alone states its basis ───────────────────────────────────────────────

def _gauge_summary():
    refused = [{"part_number": p, "normalized_thickness_mm": 2.0, "_displaced": {
        "normalized_thickness_mm": [{"value": 1.0, "source": "drawing_deterministic",
                                     "applied": False}]}} for p in ("A", "B", "C")]
    kept = {"part_number": "M", "normalized_thickness_mm": 1.0,
            "thickness_source": "drawing_deterministic"}
    return {"estimate_summary": {"part_estimates": refused + [kept]}}


def test_the_census_names_where_a_document_figure_stayed_in_force():
    s = _gauge_summary()
    assert CF.boilerplate_thickness_values(s) == {1.0: ["A", "B", "C"]}, "shape unchanged"
    assert CF.document_level_gauges(s)[1.0] == {"refused_on": ["A", "B", "C"], "kept_on": ["M"]}
    rec = CF.costed_job(s)
    m = next(l for l in rec["lines"] if l["part_number"] == "M")
    assert m["thickness_basis"]["document_level"] is True
    assert m["thickness_basis"]["also_on"] == ["A", "B", "C"]
    assert any("gauge in force on M" in d.get("issue", "") for d in rec["decisions_required"])


def test_the_dimensions_cell_names_a_document_level_gauge():
    cell = J._line_dimensions({"kind": "leaf", "thickness_mm": 1.0, "thickness_basis": {
        "mm": 1.0, "source": "drawing_deterministic", "document_level": True,
        "also_on": ["A", "B", "C"]}}, {})
    assert cell.startswith("1 mm") and "3 other parts" in cell and "confirm" in cell


def test_a_gauge_with_its_own_source_names_it():
    cell = J._line_dimensions({"kind": "leaf", "thickness_mm": 1.5, "thickness_basis": {
        "mm": 1.5, "source": "dxf", "document_level": False, "also_on": []}}, {})
    assert cell.startswith("1.5 mm") and "from " in cell and "confirm" not in cell


# ── 2.5 every sheet row that carries money is on the record ──────────────────────────

_PANEL = {"part_number": "P-1", "description": "PANEL", "quantity": 1,
          "normalized_material": "MILD STEEL"}
_PANEL_ROW = {"block": "steel", "description": "P-1  PANEL", "qty_per_unit": 1,
              "total_value_gbp": 5.0, "workbook_row": 63, "charged_cell": "Estimate!M63"}


def _rows_summary(rows, gbp, parts=(_PANEL,)):
    rows = [dict(_PANEL_ROW)] + list(rows) if parts else list(rows)
    gbp = gbp + (5.0 if parts else 0.0)
    return {"estimate_summary": {"part_estimates": [dict(p) for p in parts], "final_estimate": {
        "material_rows": rows, "totals": {"material_gbp": gbp, "labour_gbp": 0, "unit_gbp": gbp}}},
        "workbook_totals": {"source": "excel_calculated"}}


_POWDER = {"block": "bom", "part_code": "POWDER",
           "description": "Powder — computed from coated surface area (1 m2) at £4.00/kg",
           "total_value_gbp": 2.88, "workbook_row": 28, "charged_cell": "Estimate!M28"}


def test_the_powder_row_is_a_line_and_the_hierarchy_reconciles():
    rec = CF.costed_job(_rows_summary([dict(_POWDER)], 2.88))
    assert [l["part_number"] for l in rec["lines"]] == ["P-1", "POWDER"]
    line = rec["lines"][1]
    assert line["charged_ext_gbp"] == 2.88 and line["price_origin"]["class"] == "sheet_row"
    assert line["kind"] == "sheet_row" and line["material_label"] == "Powder"
    assert not rec["gaps"]["unpriced"] and not rec["gaps"]["indicative_market"]
    h = J._render_bom_tree({}, rec)
    assert "on no line here" not in h and "they agree" in h


def test_a_second_row_under_one_key_is_not_lost():
    a = dict(_POWDER, part_code="X1", description="X1 A", total_value_gbp=1.0, workbook_row=20)
    b = dict(a, workbook_row=21)
    rec = CF.costed_job(_rows_summary([a, b], 2.0))
    assert round(sum(l["charged_ext_gbp"] for l in rec["lines"]), 2) == 7.0
    assert len({l["identity"] for l in rec["lines"]}) == 3


def test_a_row_the_sheet_holds_at_a_market_figure_says_so():
    ai = dict(_POWDER, part_code="FIXING9", description="CONCRETE SLAB", supplier="xAI",
              total_value_gbp=171.24, workbook_row=12)
    rec = CF.costed_job(_rows_summary([ai], 171.24))
    assert rec["gaps"]["indicative_market"] == ["FIXING9"]


def test_with_no_part_list_the_sheet_is_not_minted_as_lines():
    assert CF.costed_job(_rows_summary([dict(_POWDER)], 2.88, parts=()))["lines"] == []


def test_the_material_breakdown_has_no_residual():
    assert CF.RESIDUAL_LABEL not in dict(
        CF.charged_breakdown_by_material(_rows_summary([dict(_POWDER)], 2.88)))


# ── 4.0 section 11 names the input that is missing ───────────────────────────────────

def _mfc(monkeypatch):
    monkeypatch.setattr(estimator, "_researched_board_rate_m2", lambda *a, **k: None)
    part = {"part_number": "P1", "normalized_material": "MFC", "normalized_thickness_mm": 18,
            "quantity": 2, "blank_length_mm": 1470, "blank_width_mm": 288,
            "normalized_geometry": {"blank_length_mm": 1470, "blank_width_mm": 288},
            "manufacturing_interpretation": {"stock_form": "sheet"}}
    return {"part_number": "P1", "normalized_material": "MFC",
            "material_estimate": estimator.estimate_material(part)}


def test_a_board_with_no_price_is_blank_for_its_price_not_its_geometry(monkeypatch):
    pe = _mfc(monkeypatch)
    assert pe["material_estimate"]["cost_per_part_gbp"] is None
    assert pe["material_estimate"]["unpriced_inputs"] == ["density", "rate_per_kg"]
    r = EI.unpriced_reason_for_row(pe)
    assert (r["category"], r["owner"]) == ("no_price_source", "estimator"), "owner unchanged"
    assert "MFC" in r["detail"] and "1470" in r["detail"] and "catalogue row" not in r["detail"]
    cat, why, sup = PE.why_no_price(r["category"], part_is_fabricated=True, reason=r)
    assert sup is False and cat == "Its material has no price"
    assert "labour rate" not in why and "section 5" not in why


def test_section_11_prints_the_computed_reason(monkeypatch):
    r = EI.unpriced_reason_for_row(_mfc(monkeypatch))
    s = {"final_estimate": {"material_rows": [
        {"description": "P1  BACK PANEL", "price_gbp": 0, "unpriced_reason": r}]}}
    t = _text(J._unpriced_section(s))
    assert "MFC" in t and "1470" in t
    assert "a gauge, a blank size or a labour rate" not in t and "section 5" not in t


def test_a_missing_gauge_is_named_as_the_gauge():
    me = estimator.estimate_material({
        "part_number": "P2", "normalized_material": "MILD STEEL", "quantity": 1,
        "blank_length_mm": 500, "blank_width_mm": 300,
        "normalized_geometry": {"blank_length_mm": 500, "blank_width_mm": 300}})
    r = EI.unpriced_reason_for_row({"part_number": "P2", "material_estimate": me})
    assert r["missing"] == ["thickness"] and r["category"] == "not_measured"
    assert "gauge" in r["detail"]


def test_a_bought_in_is_never_told_its_gauge_is_missing():
    r = EI.unpriced_reason_for_row({"part_number": "BI-BOLT",
                                    "material_estimate": {"unpriced_inputs": ["thickness"]}})
    assert "missing" not in r and r["category"] == "no_price_source"


def test_with_no_recorded_input_nothing_is_guessed():
    cat, why, sup = PE.why_no_price("no_price_source", part_is_fabricated=True)
    assert sup and "material plus labour" in why
    assert "labour rate" not in why and "section 5" not in why and "almost always" not in why


# ── 4.1 a cue that became a charge says so; the lead makes no blanket claim ──────────

def _charged_job():
    return {"estimate_summary": {
        "part_estimates": [{"part_number": "P-PANEL", "risk_flags": ["weld_required"]},
                           {"part_number": "P-BEARING", "risk_flags": ["weld_required"]}],
        "estimate_review_signals": {"parts_flagged": [
            {"part_number": "P-PANEL", "reasons": [{"code": "risk_flag", "detail": "weld_required"}]},
            {"part_number": "P-BEARING", "reasons": [{"code": "risk_flag", "detail": "weld_required"}]}]},
        "canonical_route_shadow": {"decisions": [
            {"decision_id": "d1", "operation": "welding", "status": "required",
             "target_id": "P-PANEL", "participants": [], "source": "drawing_notes"},
            {"decision_id": "d2", "operation": "dress_welds", "status": "required",
             "target_id": "P-PANEL", "participants": [], "source": "override_rule"}]},
        "final_estimate": {"labour_rows": [
            {"operation": "Weld (CO2)", "workbook_row": 194, "total_value_gbp": 23.77, "qty_per_unit": 2},
            {"operation": "Dress Welds", "workbook_row": 200, "total_value_gbp": 15.30, "qty_per_unit": 2},
            {"operation": "Assemble/pack (Metal)", "workbook_row": 209, "total_value_gbp": 9.04,
             "qty_per_unit": 1}]},
        "workbook_labour": {"rows": [
            {"workbook_row": 194, "wb_operation": "Weld (CO2)", "engine_operations": ["welding"],
             "part_numbers": ["P-PANEL"], "decision_ids": ["d1"]},
            {"workbook_row": 200, "wb_operation": "Dress Welds", "engine_operations": ["dress_welds"],
             "part_numbers": ["P-PANEL"], "decision_ids": ["d2"]},
            {"workbook_row": 209, "wb_operation": "Assemble/pack (Metal)",
             "engine_operations": ["assembly"], "part_numbers": ["P-BEARING"], "decision_ids": []}]}}}


def test_a_cue_that_became_a_charge_is_not_called_arithmetic_free():
    s = _charged_job()
    assert [c["workbook_row"] for c in CF.charges_behind_flag(s, "P-PANEL", "weld_required")] == [194, 200]
    assert CF.charges_behind_flag(s, "P-BEARING", "weld_required") == []
    assert CF.charges_behind_flag(s, "P-PANEL", "low_part_confidence") == []
    t = _text(J._render_review_items(J._extract_review_items(s)))
    assert "change the arithmetic" not in t
    assert "1 of them drove a charge on the sheet" in t and "2 Estimate rows" in t
    assert "Estimate row 194" in t
    assert "a note on the drawing" in t and "an SDI override rule" in t
    assert "drawing_notes" not in t and "override_rule" not in t, "English, not codes"
    assert "Charged on the sheet on 1 of these part(s)" in t


def test_an_uncharged_cue_keeps_its_plain_impact():
    assert "Review against the drawing." in _text(
        J._render_review_items(J._extract_review_items(_charged_job())))


def test_a_bare_review_makes_no_claim_either_way():
    t = _text(J._render_review_items({"flagged_parts": [], "risk_flag_tally": {}, "provisional": [],
                                      "part_notes": [{"note": "fold count: 1 ... is CHARGED",
                                                      "parts": ["P1"]}]}))
    assert "change the arithmetic" not in t and "drove a charge" not in t


# ── 4.2 one flag, one count, from one reconciled list ────────────────────────────────

def _flag_job():
    pe = [{"part_number": "P-PANEL", "risk_flags": ["weld_required"]},
          {"part_number": "P-BEARING", "risk_flags": ["weld_required"]},
          {"part_number": None, "description": "STICKER", "risk_flags": ["weld_required"]}]
    return {"estimate_summary": {
        "part_estimates": pe,
        # As wb_populate builds them: shallow copies, the nameless record dropped.
        "canonical_part_estimates": [dict(p) for p in pe if p["part_number"]],
        "estimate_review_signals": CF.review_signals(pe)},
        "workbook_labour": {"rows": [
            {"workbook_row": 190, "wb_operation": "Weld (CO2)", "engine_operations": ["welding"],
             "part_numbers": ["P-PANEL"]},
            {"workbook_row": 209, "wb_operation": "Assemble/pack (Metal)",
             "engine_operations": ["assembly"], "part_numbers": ["P-BEARING"]}]}}


def test_reconcile_reaches_the_list_the_report_reads():
    s = _flag_job()
    CF.reconcile_risk_flags(s)
    jp = {p["part_number"]: p for p in CF.job_parts(s)}
    assert "weld_required" not in jp["P-BEARING"]["risk_flags"]
    assert [x["flag"] for x in jp["P-BEARING"]["superseded_risk_flags"]] == ["weld_required"]
    CF.reconcile_risk_flags(s)                                      # idempotent
    jp = {p["part_number"]: p for p in CF.job_parts(s)}
    assert len(jp["P-BEARING"]["superseded_risk_flags"]) == 1


def test_one_flag_one_count_on_one_page():
    s = _flag_job()
    CF.reconcile_risk_flags(s)
    review = J._extract_review_items(s)
    listed = [fp["part"] for fp in review["flagged_parts"]
              if any(f["code"] == "weld_required" for f in fp["findings"])]
    assert listed == ["P-PANEL"] and review["risk_flag_tally"].get("weld_required") == 1
    assert "P-BEARING" not in _text(J._render_checklist(review, {}))
    assert review["not_on_sheet"] == ["STICKER"], "named, not dropped"
    t = _text(J._render_review_items(review))
    assert "Flags on lines that are not on the sheet" in t and "STICKER" in t
    assert "×1" in t


def test_before_the_workbook_the_frozen_signals_still_render():
    s = _flag_job()
    del s["workbook_labour"]
    del s["estimate_summary"]["canonical_part_estimates"]
    assert len(J._extract_review_items(s)["flagged_parts"]) == 3


def test_a_bought_in_is_not_asked_for_a_gauge():
    """The lazy-susan bearing was told to 'ask the drawing office for the gauge'."""
    out = estimator.estimate_part({"part_number": "BI-BEARING", "description": "LAZY SUSAN BEARING",
                                   "quantity": 1, "page_roles": ["bought_in"]}, job_quantity=1)
    assert not [f for f in (out.get("risk_flags") or []) if str(f).startswith("missing_material_")]
    made = estimator.estimate_part({"part_number": "P-1", "description": "BRACKET",
                                    "normalized_material": "MILD STEEL", "quantity": 1,
                                    "blank_length_mm": 100, "blank_width_mm": 50,
                                    "normalized_geometry": {"blank_length_mm": 100,
                                                            "blank_width_mm": 50}}, job_quantity=1)
    assert "missing_material_thickness" in (made.get("risk_flags") or [])
