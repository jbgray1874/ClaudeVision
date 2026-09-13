"""The check that catches a re-shaped template read the one cell a reshape moves.

12349-02 needs the Other Sheet Material block widened: it holds eight lines, the job has
nine acrylic and board parts, and the ninth — 01A-07 — was dumped on the bill of materials
at a net-part price instead of being nested. Widening it means inserting rows, and inserting
rows moves everything below, including the Total Material Cost formula this guard reads.

It read `M92` by address. After a widening the total sits at M94, the read finds no formula,
the guard returns quietly, and the run writes a stale CELL_MAP into a re-shaped sheet —
part rows on top of the totals. The one occasion the check exists for is the one occasion it
would have been blind.

The label does not move. That is what it is found by now.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import wb_populate                                                      # noqa: E402

CM = {"bom": {"first_row": 11, "last_row": 50},
      "tube": {"first_row": 53, "last_row": 60},
      "steel": {"first_row": 63, "last_row": 81},
      "other_sheet": {"first_row": 84, "last_row": 91}}


def _sheet(total_row: int, other_last: int = 91):
    """A template whose Total Material Cost sits where a reshape would have put it."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=total_row, column=3, value="Total Material Cost")
    ws.cell(row=total_row, column=13,
            value=f"=(SUM(M11:M50)+SUM(M53:M60)+SUM(M63:M81)+SUM(M84:M{other_last})+AF83)")
    return ws


def test_the_guard_agrees_with_the_map_it_was_written_for():
    wb_populate._verify_template_matches_cellmap(_sheet(92), {k: dict(v) for k, v in CM.items()})


def test_a_widened_block_is_caught_even_though_the_total_moved():
    """The real case: two rows inserted, so Other Sheet runs 84..93 and the total is at 94."""
    with pytest.raises(RuntimeError) as err:
        wb_populate._verify_template_matches_cellmap(
            _sheet(94, other_last=93), {k: dict(v) for k, v in CM.items()})
    assert "other_sheet" in str(err.value)
    assert "84..91" in str(err.value) and "M84:M93" in str(err.value)


def test_the_guard_still_catches_a_reshape_that_did_not_move_the_total():
    with pytest.raises(RuntimeError):
        wb_populate._verify_template_matches_cellmap(
            _sheet(92, other_last=95), {k: dict(v) for k, v in CM.items()})


def test_a_sheet_with_no_total_row_is_not_a_failure():
    """Non-fatal by design: a missing formula falls back to the constants rather than
    stopping a run that may be perfectly sound."""
    wb = openpyxl.Workbook()
    wb_populate._verify_template_matches_cellmap(wb.active, {k: dict(v) for k, v in CM.items()})


def test_a_label_without_a_formula_beside_it_is_not_a_failure():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=92, column=3, value="Total Material Cost")
    wb_populate._verify_template_matches_cellmap(ws, {k: dict(v) for k, v in CM.items()})
