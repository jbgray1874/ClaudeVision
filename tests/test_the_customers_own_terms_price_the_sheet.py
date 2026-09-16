"""One rebate default is right for at most one customer.

The Estimate's own totals formula is M170 = ((material + labour)/(100% − M172))/0.93 —
M172 is the customer rebate and the trailing divisor the overhead absorption. The blank
template ships rebate 0 and /0.93 and nothing ever set them, so 11908-21 (M&S) went out
missing the 1.8% uplift and on the wrong divisor. Tony Ford's own estimate for 0359967
prints the office's table on its Labour tab: Tesco 2.7%/0.93 · TTI 6.6%/0.93 ·
M&S 1.8%/0.92 · Boots 2.7%/0.93. "No rebate charge allowed" — his letter, 16 Sep 2026.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")


# ── the table resolves by customer, bounded, never guessed ───────────────────────────────

def test_the_office_table_resolves_the_spellings_one_customer_uses():
    from config import customer_commercial_terms as terms
    for spelling in ("M&S", "Marks & Spencer PLC", "MARKS AND SPENCER"):
        t = terms(spelling)
        assert t and t["rebate_fraction"] == 0.018 and t["absorption_divisor"] == 0.92, \
            (spelling, t)
    t = terms("Tesco Stores Ltd")
    assert t and t["rebate_fraction"] == 0.027 and t["absorption_divisor"] == 0.93


def test_an_unlisted_customer_gets_no_terms_not_the_nearest_neighbours():
    from config import customer_commercial_terms as terms
    assert terms("PITTI DESIGN") is None, "TTI must not match inside another word"
    assert terms("Somebody New Ltd") is None
    assert terms("") is None and terms(None) is None


# ── the JSON equivalent uses the same table ──────────────────────────────────────────────

def test_the_equivalent_pricing_books_the_ms_terms():
    from estimator import _build_workbook_equivalent_pricing as build
    out = build([], material_total=23.356, labour_total=28.219, customer="M&S")
    assert out["m107_rebate_fraction"] == 0.018
    assert out["overhead_absorption_factor"] == 0.92
    # (23.356 + 28.219) / 0.982 / 0.92 — Tony's own £57.09 arithmetic
    assert abs(out["m105_total_unit_cost_gbp"] - 57.0875) < 0.01, out
    assert out["assumptions"]["source"] == "customer_commercial_terms"


def test_no_customer_keeps_the_config_default():
    from estimator import _build_workbook_equivalent_pricing as build
    out = build([], material_total=10.0, labour_total=10.0)
    assert out["m107_rebate_fraction"] == 0.066, "the TTI default stands when nobody is named"


# ── and the workbook's own cells carry them ──────────────────────────────────────────────

def _sheet_like_the_template():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["K170"] = "Total Unit Cost Price"
    ws["M170"] = "=((M92+M168)/(100%-M172))/0.93"
    ws["K172"] = "Rebate Calculator"
    ws["M172"] = 0
    return ws


def test_the_rebate_and_divisor_land_on_the_sheet():
    from wb_populate import _apply_customer_terms
    ws, flags = _sheet_like_the_template(), []
    _apply_customer_terms(ws, "M&S", flags)
    assert ws["M172"].value == 0.018, ws["M172"].value
    assert ws["M170"].value.endswith("/0.92"), ws["M170"].value
    assert any("CUSTOMER TERMS" in f for f in flags)


def test_a_typed_rebate_is_never_overwritten():
    from wb_populate import _apply_customer_terms
    ws, flags = _sheet_like_the_template(), []
    ws["M172"] = 0.05                       # the estimator's own figure
    _apply_customer_terms(ws, "M&S", flags)
    assert ws["M172"].value == 0.05
    assert any("estimator's figure stands" in f for f in flags)


def test_an_unlisted_customer_leaves_the_template_untouched():
    from wb_populate import _apply_customer_terms
    ws, flags = _sheet_like_the_template(), []
    _apply_customer_terms(ws, "Somebody New Ltd", flags)
    assert ws["M172"].value == 0
    assert ws["M170"].value.endswith("/0.93")
    assert not flags
