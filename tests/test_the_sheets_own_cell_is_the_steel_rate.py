r"""The sheet's own cell is the steel rate. The engine reads it; it does not have its own.

James Gray, 18 September 2026, on 401912-02:

    "the spreadsheet rate is the controlling rate. The engine should read the steel
     £/tonne input from the workbook, use that same figure in its calculation/report, and
     let the estimator amend it in the sheet when required. It should not independently
     substitute £1.45/kg, £950/t or a config fallback."

FOUR RATES FOR ONE MATERIAL, AND ONE BOOK CARRIED TWO OF THEM AT ONCE. 401912-02's divider
charged £3.07 — the workbook's own L5 at £900/tonne — and published £3.88 beside it from
another path. Three people read that as the engine disagreeing with the workbook about the
METHOD. It was not: it was a rate written in several places.

    £900/tonne     the sheet's own cell, which is what charged the job
    £950/tonne     WORKBOOK_INPUT_DEFAULTS, which is what the engine used
    £800/tonne     MATERIAL_PRICE_GBP_PER_KG
    £1.45/kg       a live per-kilo rate from the price service

The duplication is admitted in config's own comment — "Blank sheet = £900" written directly
above a default of 950 — and `test_a_rate_lives_in_one_place` froze it as one of six known
pairs awaiting an estimator's ruling. This is that ruling.

So the ESTIMATING TEMPLATE's cell is read and used, and the config default becomes what it
should always have been: the answer when the template cannot be reached, never a second
opinion. An estimator changes the rate where they already change it — in the sheet — and
nobody edits Python. Where a per-kilo rate is also available it is REFUSED and the refusal
is stated, because a number quietly not used is indistinguishable from a number nobody had.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import config  # noqa: E402
import estimator  # noqa: E402

_DIVIDER = {"part_number": "401912-02-01M", "normalized_material": "MILD_STEEL",
            "normalized_thickness_mm": 2.0, "blank_length_mm": 460.0,
            "blank_width_mm": 356.61, "quantity": 1,
            "geometry_source": "dxf_flat_pattern"}


def _divider(**over):
    return dict(_DIVIDER, **over)


def _fresh(monkeypatch):
    """The cache is per process; a test that reads a different template needs it empty."""
    monkeypatch.setattr(config, "_WORKBOOK_INPUT_CACHE", {})


# ── the rate comes from the sheet ────────────────────────────────────────────────────

def test_the_template_cell_is_read_and_named(monkeypatch, tmp_path):
    """£900 in Estimate!L5 is the rate, and the record says which cell it came from —
    the estimator can then change it where they already work."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estimate"
    ws["L5"] = 900.0
    path = tmp_path / "Blank Estimate Sheet 2026.xlsx"
    wb.save(path)
    _fresh(monkeypatch)
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE", path)
    value, source = config.workbook_input_value("sheet_steel_cost_per_tonne_gbp")
    assert value == 900.0
    assert "Estimate!L5" in source


def test_an_unreachable_template_falls_back_and_says_so(monkeypatch, tmp_path):
    """A missing or locked template is not a crisis — but the sheet must not claim a rate
    came from a cell nobody could open."""
    _fresh(monkeypatch)
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE", tmp_path / "nope.xlsx")
    value, source = config.workbook_input_value("sheet_steel_cost_per_tonne_gbp")
    assert value == config.WORKBOOK_INPUT_DEFAULTS["sheet_steel_cost_per_tonne_gbp"]
    assert "WORKBOOK_INPUT_DEFAULTS" in source


def test_an_empty_or_nonsense_cell_does_not_become_a_rate(monkeypatch, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estimate"
    ws["L5"] = "see purchasing"
    path = tmp_path / "Blank Estimate Sheet 2026.xlsx"
    wb.save(path)
    _fresh(monkeypatch)
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE", path)
    value, source = config.workbook_input_value("sheet_steel_cost_per_tonne_gbp")
    assert value == config.WORKBOOK_INPUT_DEFAULTS["sheet_steel_cost_per_tonne_gbp"]
    assert "WORKBOOK_INPUT_DEFAULTS" in source


# ── and the engine costs from it, and says it did ───────────────────────────────────

def test_the_steel_line_is_costed_from_the_sheets_rate(monkeypatch):
    monkeypatch.setattr(config, "workbook_input_value",
                        lambda _k: (900.0, "the estimating template's Estimate!L5"))
    part = _divider()
    out = estimator.estimate_material(part)
    assert out["cost_method"] == "workbook_sheet_steel_formula"
    # 2500 x 1250 at 2 mm = 49.06 kg; £900/t / 15 per sheet x 1.04 = £3.06
    assert 3.0 <= out["unit_material_cost_gbp"] <= 3.12, out["unit_material_cost_gbp"]
    assert part["steel_rate_used"]["gbp_per_tonne"] == 900.0
    assert "Estimate!L5" in part["steel_rate_used"]["source"]


def test_the_record_carries_the_nest_the_rate_was_spread_over():
    """A rate alone does not make a figure checkable — the sheet it was divided by does."""
    part = _divider()
    estimator.estimate_material(part)
    rec = part["steel_rate_used"]
    assert rec["parts_per_sheet"] == 15
    assert rec["sheet_mm"] == [2500, 1250]


# ── no parallel resolver ─────────────────────────────────────────────────────────────

def test_a_live_per_kilo_rate_does_not_displace_the_sheet(monkeypatch):
    """£1.45/kg x 2.575 kg x 1.04 = £3.88, the figure that appeared beside the sheet's own
    £3.07. The sheet's cell is the controlling rate and the other is not used."""
    monkeypatch.setattr(config, "workbook_input_value",
                        lambda _k: (900.0, "the estimating template's Estimate!L5"))
    monkeypatch.setattr(estimator, "_resolve_material_price",
                        lambda *_a, **_k: {"result": {}, "applied_price_per_kg": 1.45,
                                           "applied_basis": "GBP_per_kg"})
    part = _divider()
    out = estimator.estimate_material(part)
    assert out["cost_method"] == "workbook_sheet_steel_formula"
    assert out["unit_material_cost_gbp"] < 3.5, "the per-kilo rate priced the line"


def test_the_refused_rate_is_stated_rather_than_dropped(monkeypatch):
    """A number quietly not used is indistinguishable from a number nobody had."""
    monkeypatch.setattr(config, "workbook_input_value",
                        lambda _k: (900.0, "the estimating template's Estimate!L5"))
    monkeypatch.setattr(estimator, "_resolve_material_price",
                        lambda *_a, **_k: {"result": {}, "applied_price_per_kg": 1.45,
                                           "applied_basis": "GBP_per_kg"})
    part = _divider()
    estimator.estimate_material(part)
    said = " ".join(str(f) for f in part.get("review_flags", []))
    assert "STEEL RATE" in said
    assert "£1.45/kg" in said
    assert "NOT used" in said


def test_the_ordinary_steel_part_is_told_nothing():
    """The config per-kilo fallback exists for every steel on the books and was never going
    to price this line. Flagging it would put a sentence on every steel part in the shop,
    which is how a warning stops being read — the same lesson three other flags have paid
    for this week."""
    part = _divider()
    estimator.estimate_material(part)
    assert not any("STEEL RATE" in str(f) for f in part.get("review_flags", []))


def test_a_steel_sheet_part_prices_even_with_no_per_kilo_rate_at_all(monkeypatch):
    """The formula does not use one, so a missing per-kg rate is no longer a reason to
    abandon the route and read as unpriced."""
    monkeypatch.setattr(estimator, "_resolve_material_price",
                        lambda *_a, **_k: {"result": {}, "applied_price_per_kg": None,
                                           "applied_basis": None})
    monkeypatch.setattr(estimator, "_price_per_kg_for_material", lambda *_a, **_k: None)
    out = estimator.estimate_material(_divider())
    assert out["cost_method"] == "workbook_sheet_steel_formula"
    assert out["unit_material_cost_gbp"] > 0


# ── the controls ─────────────────────────────────────────────────────────────────────

def test_a_part_with_no_blank_is_not_forced_down_the_sheet_route():
    """No area, no share of a sheet. It must not invent one."""
    out = estimator.estimate_material(
        _divider(blank_length_mm=None, blank_width_mm=None))
    assert out.get("cost_method") != "workbook_sheet_steel_formula"


def test_a_non_steel_material_is_untouched():
    """The ruling is about steel sheet. Acrylic, board and timber keep their own routes."""
    out = estimator.estimate_material({
        "part_number": "X-01", "normalized_material": "ACRYLIC",
        "normalized_thickness_mm": 3.0, "blank_length_mm": 300.0,
        "blank_width_mm": 200.0, "quantity": 1})
    assert out.get("cost_method") != "workbook_sheet_steel_formula"
    assert "steel_rate_used" not in out


def test_a_rate_the_sheet_states_but_nobody_could_mean_is_refused(monkeypatch):
    """£5/tonne or £40,000/tonne is a typo, not a ruling. The clamp stands, and says so."""
    monkeypatch.setattr(config, "workbook_input_value",
                        lambda _k: (5.0, "the estimating template's Estimate!L5"))
    part = _divider()
    estimator.estimate_material(part)
    assert part["steel_rate_used"]["gbp_per_tonne"] != 5.0
    assert "was not used" in part["steel_rate_used"]["source"]
