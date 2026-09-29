"""Total Labour Hours By Dept. must add up to the labour rows' hours.

12645 DRS External Shelter, book of 28 Sep 2026 (12645_20260928_191745.xlsx): Excel's cached
D244 read 9.85 h, FOLD and LASM only, while J126:J197 summed 12.60 h. The template's hidden
Labour sheet pairs column A (department) with column B (hours) by single-cell reference; when
the labour block was widened from 40 to 72 rows, Excel shifted A41:B100 to the rows below it
(Estimate!G198:J257) instead of extending them, so rows 166-197 were never read. P.Coat,
Assemble/pack (Metal), Weld (CO2), Dress Welds, P.Coat and Assemble/pack — 2.75 h — were
missing from the department table (D-325). Labour money was unaffected: M198 sums the rows.

The guard reads the Labour sheet's own references and the labour block the template's labels
bound — never fixed row numbers — and repairs the produced book when they disagree.
Recalculated in LibreOffice after the repair, the 12645 book gives D244 = 12.6042 h, equal to
J126:J197.
"""
from __future__ import annotations

import copy
import os
import sys

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import wb_populate as wp  # noqa: E402

FIRST, LAST, CAP = 30, 41, 20          # a synthetic layout: 12 labour rows, a 20-row scan


def _book(refs):
    """An Estimate sheet whose labour block is bounded by its labels, and a Labour sheet whose
    column A/B point at `refs` (one Estimate row per Labour row, None for blank)."""
    wb = openpyxl.Workbook()
    est = wb.active
    est.title = "Estimate"
    est.cell(FIRST - 1, 3, "Operation")
    est.cell(LAST + 1, 3, "Total Labour Cost (Including  Downtime)")
    lab = wb.create_sheet("Labour")
    for i, r in enumerate(refs, start=1):
        if r is not None:
            lab.cell(i, 1, f"=Estimate!G{r}")
            lab.cell(i, 2, f"=Estimate!J{r}")
    lab["C1"] = "=Estimate!C60"
    lab["C2"] = ArrayFormula("C2", f'=IF(ISERROR(INDEX($A$1:$B${CAP},SMALL(IF($A$1:$A${CAP}=$C$1,'
                                   f'ROW($A$1:$A${CAP})),ROW(1:1)),2)),"",1)')
    return wb


def _cm():
    cm = copy.deepcopy(wp.CELL_MAP)
    cm["labour"]["first_row"], cm["labour"]["last_row"] = FIRST, LAST
    return cm


# The 12645 shape: the first rows of the block, then a jump past it.
SKIPPING = list(range(FIRST, FIRST + 8)) + list(range(LAST + 1, LAST + 1 + CAP - 8))


def test_a_labour_sheet_that_skips_rows_is_caught():
    cov = wp.labour_sheet_coverage(_book(SKIPPING), _cm())
    assert cov["ok"] is False
    assert cov["missing"] == list(range(FIRST + 8, LAST + 1))
    assert cov["stray"] and min(cov["stray"]) == LAST + 1
    assert (cov["dept_col"], cov["hours_col"], cov["capacity"]) == ("G", "J", CAP)


def test_the_repair_points_it_at_exactly_the_labour_block():
    wb, flags = _book(SKIPPING), []
    wp.repair_labour_sheet_references(wb, _cm(), flags)
    lab = wb["Labour"]
    assert [lab.cell(i, 1).value for i in range(1, LAST - FIRST + 2)] == \
        [f"=Estimate!G{r}" for r in range(FIRST, LAST + 1)]
    assert lab.cell(LAST - FIRST + 1, 2).value == f"=Estimate!J{LAST}"
    assert all(lab.cell(i, 1).value is None for i in range(LAST - FIRST + 2, CAP + 1))
    assert wp.labour_sheet_coverage(wb, _cm())["ok"] is True
    assert any("TEMPLATE FAULT REPAIRED" in f for f in flags)


def test_a_sheet_that_already_covers_the_block_is_left_alone():
    wb, flags = _book(list(range(FIRST, LAST + 1))), []
    assert wp.repair_labour_sheet_references(wb, _cm(), flags)["ok"] is True
    assert not any("TEMPLATE FAULT" in f for f in flags)


def test_a_block_the_labels_do_not_bound_is_not_repaired():
    """The first try of this guard ran on a book whose layout could not be read, kept the
    map's stale constants, and pointed the department table at the wrong rows. An unconfirmed
    block is flagged and nothing is written."""
    wb, flags = _book(SKIPPING), []
    cm = _cm()
    cm["labour"]["first_row"] -= 5
    cov = wp.repair_labour_sheet_references(wb, cm, flags)
    assert cov.get("unconfirmed") is True
    assert wb["Labour"]["A9"].value == f"=Estimate!G{LAST + 1}", "nothing may be rewritten"
    assert any("NOT checked or repaired" in f for f in flags)


def test_a_block_larger_than_the_scan_is_a_template_fault_not_a_repair():
    wb, flags = _book(SKIPPING), []
    cm = _cm()
    wb["Estimate"].cell(LAST + 1, 3, None)
    cm["labour"]["last_row"] = FIRST + CAP + 4
    wb["Estimate"].cell(FIRST + CAP + 5, 3, "Total Labour Cost")
    wp.repair_labour_sheet_references(wb, cm, flags)
    assert any("scans 20 rows but the labour block has" in f for f in flags)


def test_the_populate_path_runs_the_guard_after_the_layout_is_read():
    import inspect
    src = inspect.getsource(wp.populate_workbook)
    d = src.index("derive_cellmap_from_template(ws, cm, flags)")
    r = src.index("repair_labour_sheet_references(wb, cm, flags)")
    assert d < r
