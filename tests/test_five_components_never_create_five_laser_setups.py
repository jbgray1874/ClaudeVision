"""One row per component on the laser; one set-up per nest, however many rows share it.

    "Labour rates should be separate line each component as reflect the different time per
     component" ... "if multiple components used but don't exceed a single sheet, laser
     rate would be reduced to not exceed the given set up time"
                                        — Howard Thurley, SDI estimating, 7332-01

His own sheet carries FIVE separate LASM rows. The engine keyed laser rows by (operation,
material, gauge), so 7332-01-003 (a 441x10 strap) and -004 (a 15.88 square cap) shared one
row at a blended 441/hr — his figures for the two are 235 and 900, nearly four to one, and
a single number cannot be checked against either.

THE TWO HALVES SHIP TOGETHER, which is what D-066 recorded and why the split waited.
Splitting the key alone would book five 10-minute set-ups where the floor loads one program
and cuts one sheet — precisely the error he warned about, worse than the blended row. So:

    the split       canonical_labour_groups keys a laser decision by its component, and
                    stamps every row with the nest it belongs to — same laser, same
                    material, same gauge, the tuple that decides whether parts can share
                    a sheet program.

    the set-up      allocate_laser_nest_setup wraps the template's OWN set-up lookup as
                    (lookup)/N on each row of the nest. The workbook stays authoritative —
                    the minutes still come from THEIR table, every share moves if the
                    estimators change it, and the nest's sum is exactly one allowance BY
                    CONSTRUCTION, not by rounding. Touched only after checking, per row,
                    that the row's hours formula (J) really reads its set-up cell (L);
                    anything else is left alone and said out loud.

These tests EXECUTE both halves — the grouping function and the allocator, on a real
worksheet — they do not read source text.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

from openpyxl import Workbook                                          # noqa: E402

import wb_populate as wb                                               # noqa: E402

LB = {"col_hours": 10, "col_setup": 12}


def _pe(pn, thk=2.5, material="MILD STEEL"):
    return {"part_number": pn, "normalized_material": material,
            "normalized_thickness_mm": thk,
            "material_estimate": {"material": material, "thickness_mm": thk,
                                  "stock_form": "sheet"},
            "labour_estimate": {"batch_hours": {"laser_cutting": 0.2}}}


def _summary(decisions, pns):
    return {"canonical_route_shadow": {
        "nodes": [{"part_number": p, "qty_per_unit": 1} for p in pns],
        "decisions": decisions}}


def _laser_decision(i, pn, thk=None):
    d = {"decision_id": f"d{i}", "operation": "laser_cutting", "status": "required",
         "target_id": pn, "participants": [pn], "scope": "part", "sequence": 10,
         "qty_per_unit": 1}
    return d


def _five_part_groups(thicknesses=(2.5, 2.5, 2.5, 2.5, 2.5)):
    pns = [f"7332-01-00{i}" for i in range(1, 6)]
    decisions = [_laser_decision(i, p) for i, p in enumerate(pns)]
    estimates = [_pe(p, thk) for p, thk in zip(pns, thicknesses)]
    return wb.canonical_labour_groups(_summary(decisions, pns), estimates, 6)


def _template_rows(ws, rows):
    """The labour block the way the WB 2026 template builds it: J reads L/60, L looks the
    set-up minutes up from the operation name."""
    for r in rows:
        ws.cell(row=r, column=LB["col_hours"],
                value=f'=IF(C{r}="","",((60/I{r})*H{r}/60)*$D$6)+((L{r}/60))')
        ws.cell(row=r, column=LB["col_setup"],
                value=f'=IF(C{r}="","",VLOOKUP(C{r},RATES!A:D,4,0))')


# ── the split ────────────────────────────────────────────────────────────────────────────

def test_five_components_get_five_laser_rows():
    groups = _five_part_groups()
    lasers = [g for g in groups.values() if g["wb_op"] == "Laser (Metal)"]
    assert len(lasers) == 5, [g["parts"] for g in groups.values()]
    for g in lasers:
        assert len(g["parts"]) == 1, "one row, one component — no blended rate"


def test_one_gauge_is_one_nest():
    groups = _five_part_groups()
    nests = {g["laser_nest_id"] for g in groups.values()}
    assert len(nests) == 1, nests


def test_a_different_gauge_is_a_different_nest():
    """The nesting constraint is per gauge — a 1.5mm part cannot share a 2.5mm sheet."""
    groups = _five_part_groups(thicknesses=(2.5, 2.5, 2.5, 1.5, 1.5))
    nests = {g["laser_nest_id"] for g in groups.values()}
    assert len(nests) == 2, nests


def test_non_laser_grouping_is_untouched():
    """Fold still groups by tooling — the split is the laser's, not the shop's."""
    pns = ["A1", "A2"]
    decisions = [
        dict(_laser_decision(0, "A1"), operation="folding"),
        dict(_laser_decision(1, "A2"), operation="folding"),
    ]
    groups = wb.canonical_labour_groups(
        _summary(decisions, pns), [_pe(p) for p in pns], 6)
    assert len(groups) == 1, list(groups)
    (g,) = groups.values()
    assert sorted(g["parts"]) == ["A1", "A2"]
    assert not g.get("laser_nest_id")


# ── the set-up: five rows, ONE allowance ─────────────────────────────────────────────────

def test_five_rows_carry_one_department_setup_between_them():
    """THE REGRESSION D-066 EXISTS TO PREVENT. Each row's set-up cell becomes the
    template's own lookup divided by 5 — so the nest sums to exactly one allowance, and
    the figure still moves if the estimators change their table."""
    groups = _five_part_groups()
    ws = Workbook().active
    rows = list(range(96, 101))
    _template_rows(ws, rows)
    for g, r in zip(groups.values(), rows):
        g["workbook_row"] = r
    flags = []
    wb.allocate_laser_nest_setup(ws, groups, LB, flags)
    for r in rows:
        v = str(ws.cell(row=r, column=LB["col_setup"]).value)
        assert v.startswith("=(") and v.endswith(")/5"), v
        assert "VLOOKUP" in v, "the template's own lookup survives inside the share"
    assert any("ONE department set-up" in f for f in flags)
    for g in groups.values():
        assert g["laser_nest_size"] == 5
        assert g["setup_share"].startswith("1/5 ")


def test_a_lone_laser_row_keeps_the_templates_own_setup():
    pns = ["7332-01-001"]
    groups = wb.canonical_labour_groups(
        _summary([_laser_decision(0, pns[0])], pns), [_pe(pns[0])], 6)
    ws = Workbook().active
    _template_rows(ws, [96])
    original = ws.cell(row=96, column=LB["col_setup"]).value
    for g in groups.values():
        g["workbook_row"] = 96
    wb.allocate_laser_nest_setup(ws, groups, LB, [])
    assert ws.cell(row=96, column=LB["col_setup"]).value == original


def test_two_nests_are_allocated_separately():
    groups = _five_part_groups(thicknesses=(2.5, 2.5, 2.5, 1.5, 1.5))
    ws = Workbook().active
    rows = list(range(96, 101))
    _template_rows(ws, rows)
    ordered = sorted(groups.values(), key=lambda g: g["laser_nest_id"])
    for g, r in zip(ordered, rows):
        g["workbook_row"] = r
    wb.allocate_laser_nest_setup(ws, groups, LB, [])
    shares = sorted(str(ws.cell(row=r, column=LB["col_setup"]).value)[-2:]
                    for r in rows)
    assert shares == ["/2", "/2", "/3", "/3", "/3"], shares


def test_an_unrecognised_template_is_not_touched_and_says_so():
    """No silent gaps, in either direction. If the row's hours formula does not read its
    set-up cell, dividing L would change nothing it claims to change — so NOTHING in the
    nest is written (a half-allocated nest under-charges invisibly) and the sheet says the
    set-up is overstated."""
    groups = _five_part_groups()
    ws = Workbook().active
    rows = list(range(96, 101))
    _template_rows(ws, rows)
    # One row of the five has a numeric J — a revised template, a moved column.
    ws.cell(row=98, column=LB["col_hours"], value=1.25)
    originals = {r: ws.cell(row=r, column=LB["col_setup"]).value for r in rows}
    for g, r in zip(groups.values(), rows):
        g["workbook_row"] = r
    flags = []
    wb.allocate_laser_nest_setup(ws, groups, LB, flags)
    for r in rows:
        assert ws.cell(row=r, column=LB["col_setup"]).value == originals[r], r
    assert any("OVERSTATED" in f for f in flags)
    assert not any(g.get("setup_share") for g in groups.values())


# ── the allocation is row DATA, auditable off the record ─────────────────────────────────

def test_the_route_record_carries_the_nest_and_the_share():
    groups = _five_part_groups()
    ws = Workbook().active
    rows = list(range(96, 101))
    _template_rows(ws, rows)
    for g, r in zip(groups.values(), rows):
        g["workbook_row"] = r
    wb.allocate_laser_nest_setup(ws, groups, LB, [])
    record = wb.build_workbook_labour(groups, canonical_mode=True)
    for row in record["rows"]:
        assert row["laser_nest_id"], row
        assert row["laser_nest_size"] == 5
        assert "1/5" in row["setup_share"]
