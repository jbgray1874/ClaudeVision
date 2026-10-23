"""The K-Wellness 18:17 book, charge by charge: each thing once, at its own scope (D-495..D-499).

The book came back at £1,433.81 a unit. £830 of it was two test allowances charged on each of
200 units though the brief says testing is "NOT in the unit price"; the Assemble/pack row
charged 165.9 minutes beside an explanation of 82.95; the labels were counted 79 and charged
one; the fold row told the reader a bracket ran at 5807 an hour; and the read's own
construction — brackets WELDED TO the trays, strips "one per shelf" across trays of two sizes,
no fixings named, a "flat" blank the size of the unit's own footprint — reached the rows
unexamined. Each rule here is generic: it reads the records' own words, counts and sources,
never a part code or a job.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import concept_scan as cs                                                # noqa: E402
import config                                                            # noqa: E402

BRIEF = (
    "Assembly, inspection, protective packaging and palletised delivery included. 200 off. "
    "The 5% spares and the design, development,\nprototype and testing costs are NOT in the "
    "unit price - the estimator lists them\nseparately in the response. QR component labels "
    "(20 x 20 mm) are a per-unit requirement: production, application and verification. QC "
    "is batch setup plus sampled and unit checks, not full inspection per unit. Design, "
    "prototypes, testing and documentation are per-order project costs, shown separately.")


# ── D-495: the unit's stated assembly time is charged once ─────────────────────────────

def test_a_stated_time_found_under_the_departments_other_name_is_not_added_twice():
    import wb_populate as W
    hours = {"KW-CPT00": {"handling": 1.3825}}
    claims = {"KW-CPT00": "the per-part handling allowance"}
    ops = {"KW-CPT00": {"handling", "assembly"}}
    g = {"parts": ["KW-CPT00"], "engine_ops": ["assembly"], "run_hours_per_unit": 1.3825,
         "hours_by_part": {"KW-CPT00": {"bh": 0.0, "qty_per_unit": 1, "rh": 1.3825}}}
    h, who = W.stated_row_hours(g, claims, hours, ops)
    assert who == ["KW-CPT00"] and abs(h - 1.3825) < 1e-9, h       # 82.95 min, not 165.9


def test_the_grouping_records_the_share_it_summed_for_the_assemblys_own_time():
    import wb_populate as W
    summary = {"estimate_summary": {"canonical_route_shadow": {"nodes": [], "decisions": [
        {"decision_id": "d1", "operation": "assembly", "status": "required",
         "target_id": "KW-CPT00", "participants": ["KW-CPT01", "KW-CPT02"],
         "scope": "assembly", "qty_per_unit": 1}]}}}
    estimates = [{"part_number": "KW-CPT00", "quantity": 1,
                  "labour_estimate": {"batch_hours": {"handling": 276.75},
                                      "run_hours_per_unit": {"handling": 1.3825}}},
                 {"part_number": "KW-CPT01", "quantity": 6, "labour_estimate": {}},
                 {"part_number": "KW-CPT02", "quantity": 1, "labour_estimate": {}}]
    groups = W.canonical_labour_groups(summary, estimates, 200)
    (g,) = [v for v in groups.values() if v.get("assembly_own_time")]
    assert abs(g["run_hours_per_unit"] - 1.3825) < 1e-9
    assert abs(g["hours_by_part"]["KW-CPT00"]["rh"] - 1.3825) < 1e-9
    h, _ = W.stated_row_hours(g, {"KW-CPT00": "x"}, {"KW-CPT00": {"handling": 1.3825}},
                              {"KW-CPT00": {"handling", "assembly"}})
    assert abs(h - 1.3825) < 1e-9, f"charged {h * 60:.1f} min for 82.95 stated"


# ── D-496: a count the engine worked out is kept ───────────────────────────────────────

def test_a_counted_quantity_over_the_cap_is_kept_and_an_unexplained_one_is_not():
    import estimator as E
    counted = {"part_number": "J-NV01", "quantity": 79,
               "quantity_counted": {"count": 79, "basis": "1 per component × 79 component(s)"}}
    assert E._sanitise_part_quantity(counted) == 79
    assert any(isinstance(f, dict) and f.get("flag") == "quantity_over_cap_kept"
               for f in counted.get("review_flags") or [])
    noise = {"part_number": "J-NV01", "quantity": 79}
    assert E._sanitise_part_quantity(noise) == 1


def test_per_component_labels_carry_their_count_and_its_working():
    (labels,) = cs.consumable_lines(["QR component labels 20x20mm"], "J", pieces_fitted=79)
    assert labels["quantity"] == 79
    assert labels["quantity_counted"]["count"] == 79
    assert "79 component" in labels["quantity_counted"]["basis"]


# ── D-497: a project cost's scope comes from evidence ──────────────────────────────────

def _req(name, scope, requirement="", basis=""):
    return {"name": name, "scope": scope, "requirement": requirement,
            "provisional_gbp": 450.0, "provisional_basis": basis}


def test_the_brief_moves_a_test_the_read_put_on_every_unit_to_the_order():
    (r,) = cs.standards_requirements({"applicable_requirements": [
        _req("BTC topple/tilt/bump tests", "per_unit",
             basis="3 hours test rig + report at UK trade rate")]}, brief=BRIEF)
    assert r["scope"] == "per_order" and r["scope_source"] == "brief"
    assert r["read_scope"] == "per_unit"
    # The brief's own words, quoted — either of its two per-order sentences.
    assert r["scope_basis"].startswith("the brief:")
    assert ("NOT in the unit price" in r["scope_basis"]
            or "per-order project costs" in r["scope_basis"])
    (p,) = cs.project_cost_plan([r], {}, 200)
    assert p["unit_gbp"] == 0.0 and p["order_gbp"] == 450.0
    assert "Scope per order: the brief" in p["flag"]
    (f,) = cs.requirement_flags([r])
    assert "the read had it per unit" in f


def test_a_line_whose_name_says_nothing_is_classed_by_its_allowance_basis():
    """'BTC006/007 custom shelf deflection' names no test; its basis does."""
    (r,) = cs.standards_requirements({"applicable_requirements": [
        _req("BTC006/007 custom shelf deflection", "per_unit",
             basis="2 hours test + report at UK trade rate")]}, brief=BRIEF)
    assert r["scope"] == "per_order"


def test_the_brief_keeps_labels_on_the_unit_and_qc_per_batch_despite_a_negated_per_unit():
    labels, qc = cs.standards_requirements({"applicable_requirements": [
        _req("QR component labels", "per_order"),
        _req("QC inspection", "per_unit")]}, brief=BRIEF)
    assert labels["scope"] == "per_unit" and labels["scope_source"] == "brief"
    # "not full inspection per unit" does not say per unit.
    assert qc["scope"] == "per_batch" and qc["scope_source"] == "brief"


def test_without_the_brief_the_requirements_own_words_decide():
    reqs = cs.standards_requirements({"applicable_requirements": [
        _req("Stability test", "per_unit", requirement="a type test on the design"),
        _req("Weld inspection", "per_order", requirement="each unit is inspected before despatch"),
        _req("Load test", "per_unit", requirement="re-test after any design change"),
        _req("Mystery", "per_unit")]})
    by = {r["name"]: r for r in reqs}
    assert by["Stability test"]["scope"] == "per_order"
    assert by["Weld inspection"]["scope"] == "per_unit"
    assert by["Load test"]["scope"] == "per_order"
    assert "a design change brings it back" in by["Load test"]["scope_basis"]
    # No evidence: the read's call stands, labelled as an assumption to confirm.
    assert by["Mystery"]["scope"] == "per_unit" and by["Mystery"]["scope_source"] == "read"
    assert "assumption to confirm" in by["Mystery"]["scope_basis"]


def test_the_estimators_scope_outranks_everything_and_batches_multiply_a_batch_cost():
    (r,) = cs.standards_requirements({"applicable_requirements": [
        _req("BTC topple/tilt/bump tests", "per_unit")]}, brief=BRIEF)
    (p,) = cs.project_cost_plan([r], {"project_scope": {"topple": "per_unit"}}, 200)
    assert p["unit_gbp"] == 450.0 and "the estimator's answer" in p["flag"]
    (qc,) = cs.standards_requirements({"applicable_requirements": [
        _req("QC inspection", "per_unit")]}, brief=BRIEF)
    (one,) = cs.project_cost_plan([qc], {}, 200)
    assert one["order_gbp"] == 450.0 and "ONE batch is taken at every quantity" in one["flag"]
    (three,) = cs.project_cost_plan([qc], {"project_batches": 3}, 200)
    assert three["order_gbp"] == 1350.0 and "ONE batch" not in three["flag"]
    assert three["batch_size"] == 67 and three["per_batch_gbp"] == 450.0


def test_a_per_batch_cost_is_recounted_at_every_quantity_break():
    (qc,) = cs.standards_requirements({"applicable_requirements": [
        _req("QC inspection", "per_unit")]}, brief=BRIEF)
    (om,) = cs.standards_requirements({"applicable_requirements": [
        dict(_req("O&M manual", "per_order"), provisional_gbp=650.0)]}, brief=BRIEF)
    plan = cs.project_cost_plan([qc, om], {"project_batch_size": 50}, 200)
    pc = {"total_gbp": sum(p["order_gbp"] for p in plan),
          "lines": [{k: p.get(k) for k in ("scope", "state", "order_gbp", "per_batch_gbp",
                                           "batch_size")} for p in plan]}
    assert pc["total_gbp"] == 4 * 450 + 650                       # 200 off = 4 batches of 50
    assert cs.project_costs_at(pc, 1) == 450 + 650                # one batch, the manual once
    assert cs.project_costs_at(pc, 10) == 450 + 650
    assert cs.project_costs_at(pc, 51) == 2 * 450 + 650
    assert cs.project_costs_at(pc, 200) == 4 * 450 + 650
    # No batch size stated: one batch at every quantity, never the order's count reused.
    plan1 = cs.project_cost_plan([qc], {}, 200)
    pc1 = {"lines": [{k: p.get(k) for k in ("scope", "state", "order_gbp", "per_batch_gbp",
                                            "batch_size")} for p in plan1]}
    assert cs.project_costs_at(pc1, 1) == cs.project_costs_at(pc1, 200) == 450


def test_the_quantity_breaks_tab_recalculates_project_costs_per_column(tmp_path):
    import openpyxl
    from quantity_breaks_tab import write_quantity_breaks_tab, SHEET_NAME
    p = tmp_path / "book.xlsx"
    openpyxl.Workbook().save(p)
    pc = {"total_gbp": 2450.0, "lines": [
        {"scope": "per_batch", "state": "provisional", "order_gbp": 1800.0,
         "per_batch_gbp": 450.0, "batch_size": 50},
        {"scope": "per_order", "state": "provisional", "order_gbp": 650.0}]}
    swept = {"rows": [{"quantity": 1, "material": 10, "labour": 5, "unit": 20.0},
                      {"quantity": 10, "material": 10, "labour": 4, "unit": 18.0},
                      {"quantity": 200, "material": 10, "labour": 3, "unit": 16.0}]}
    write_quantity_breaks_tab(p, swept, requested=[200], project_costs=pc)
    ws = openpyxl.load_workbook(p)[SHEET_NAME]
    (row,) = [r for r in range(1, ws.max_row + 1)
              if str(ws.cell(r, 1).value or "").startswith("Project costs £")]
    assert "per batch each batch" in ws.cell(row, 1).value
    assert [ws.cell(row, c).value for c in (2, 3, 4)] == [1100.0, 1100.0, 2450.0]
    assert ws.cell(row + 1, 4).value == round(16.0 * 200 + 2450.0, 2)


def test_a_saved_requirement_without_a_scope_source_is_resolved_from_the_brief():
    (p,) = cs.project_cost_plan([_req("BTC topple/tilt/bump tests", "per_unit")], {}, 200,
                                brief=BRIEF)
    assert p["scope"] == "per_order" and p["unit_gbp"] == 0.0


def test_a_per_order_project_line_is_charged_once_not_unpriced():
    import costed_facts as cf
    import estimator_inputs as ei
    line = {"part_number": "PROJECT-03", "description": "O&M manual",
            "project_cost_line": {"scope": "per_order", "state": "provisional",
                                  "order_gbp": 650.0}}
    o = cf._price_origin(line, "commercial", "bom", 0.0, 0.0, 17, False, row_text="")
    assert o["class"] == "per_order_project" and o["firmness"] == cf.NIL
    assert "£650.00 charged ONCE" in o["label"]
    r = ei.unpriced_reason_for_row(line)
    assert r["category"] == "not_applicable" and "£650.00" in r["detail"]
    # With no figure at all it is still a price to add.
    open_line = dict(line, project_cost_line={"scope": "per_order", "state": "to_price",
                                              "order_gbp": None})
    assert cf._price_origin(open_line, "commercial", "bom", 0.0, 0.0, 17, False,
                            row_text="")["class"] == "unpriced_commercial"


# ── D-498: a shared row's members are shown at their own rates ─────────────────────────

def test_a_members_rate_is_its_own_piece_time_not_inflated_by_its_count():
    import wb_populate as W
    hbp = {"KW-CPT10": {"bh": 1.0333, "qty_per_unit": 30, "rh": 0.005},
           "KW-CPT01": {"bh": 2.0333, "qty_per_unit": 6, "rh": 0.01},
           "LEGACY": {"bh": 2.0, "qty_per_unit": 4}}
    rates = {p: round(r, 1) for p, r, _ in W.member_rates(hbp, 200)}
    assert rates["KW-CPT10"] == 200.0          # was shown as 5807/hr
    assert rates["KW-CPT01"] == 100.0          # was shown as 590/hr
    assert rates["LEGACY"] == 100.0            # batch fallback: order ÷ batch hours, no × count


def test_the_rows_charge_is_unchanged_by_the_shown_rates():
    import wb_populate as W
    summary = {"estimate_summary": {"canonical_route_shadow": {"nodes": [], "decisions": [
        {"decision_id": f"d{i}", "operation": "folding", "status": "required",
         "target_id": pn, "participants": [pn], "scope": "part", "qty_per_unit": q}
        for i, (pn, q) in enumerate((("KW-CPT10", 30), ("KW-CPT01", 6)))]}}}
    est = [{"part_number": pn, "quantity": q, "normalized_thickness_mm": 2.0,
            "material_estimate": {"stock_form": "sheet", "material": "MILD STEEL",
                                  "blank_length_mm": 100, "blank_width_mm": 50},
            "labour_estimate": {"batch_hours": {"folding": bh},
                                "run_hours_per_unit": {"folding": rh}}}
           for pn, q, bh, rh in (("KW-CPT10", 30, 1.0333, 0.005),
                                 ("KW-CPT01", 6, 2.0333, 0.01))]
    groups = W.canonical_labour_groups(summary, est, 200)
    folds = [g for g in groups.values() if "folding" in (g.get("engine_ops") or [])]
    total = sum(g["run_hours_per_unit"] for g in folds)
    assert abs(total - (30 * 0.005 + 6 * 0.01)) < 1e-9      # 0.21 h a unit, as before


# ── D-499: the sighted construction, applied to its own lines ──────────────────────────

LF = ["laser_cutting", "folding", "powder_coating", "handling"]


def _p(name, q, L, W, ops, folds=0, basis="developed", seen="", why="", flange=0):
    return {"name": name, "kind": "fabricated", "quantity": q, "material_guess": "MILD STEEL",
            "assumed_blank_mm": {"length": L, "width": W, "thickness": 1.5},
            "operations": ops, "folds": folds, "blank_basis": basis, "seen": seen,
            "why_size": why, "flange_mm": flange}


def _fsdu(**over):
    parts = [
        _p("CORNER POST", 6, 1500, 80, LF + ["welding"], 2),
        _p("BASE / PLINTH", 1, 500, 500, LF + ["welding"], 4, why="500x500 footprint",
           flange=40),
        _p("TOP CAP", 1, 500, 500, LF, 4, why="500x500 footprint"),
        _p("HEADER", 6, 460, 80, LF, 1),
        _p("FRONT SHELF TRAY", 5, 460, 300, LF, 3),
        _p("SIDE SHELF TRAY (LEFT)", 5, 300, 250, LF, 3),
        _p("SIDE SHELF TRAY (RIGHT)", 5, 300, 250, LF, 3),
        _p("HOOKED BRACKET (WELDED TO TRAY)", 30, 40, 30,
           ["laser_cutting", "welding", "folding"], 1),
        _p("35MM SHELF EDGE STRIP", 15, 460, 35, LF, 1, seen="one per shelf x15 shelves",
           why="length from shelf width"),
    ]
    ans = {"product": {"name": "FSDU", "assumed_overall_mm": {"height": 1500, "width": 500,
                                                             "depth": 500}},
           "parts": parts}
    ans.update(over)
    return ans


def _by(parts):
    return {p["description"] if "STRIP" not in p["description"] else p["part_number"]: p
            for p in parts}


def test_a_piece_welded_to_a_named_host_joins_that_hosts_component_and_its_coat():
    parts = cs.parts_from_concept(_fsdu(), "J")
    br = _by(parts)["HOOKED BRACKET (WELDED TO TRAY)"]
    trays = {p["part_number"] for p in parts if "TRAY" in p["description"]
             and "BRACKET" not in p["description"]}
    assert set(br["welded_to_hosts"]) == trays
    assert set(br["coated_with_host"]) == trays          # hosts coated, bracket not
    assert any("welded to its host" in str(f) for f in br["review_flags"])
    unit = cs.unit_assembly_part(parts, _fsdu(), "J")
    total = sum(int(p["quantity"]) for p in parts)
    assert unit["pieces_fitted"] == total - 30            # the brackets fit with their trays
    assert unit["pieces_joined_to_hosts"] == {br["part_number"]: 30}


def test_a_piece_welded_to_an_uncoated_host_is_not_coated_and_no_host_no_change():
    ans = _fsdu()
    for p in ans["parts"]:
        if "TRAY" in p["name"] and "BRACKET" not in p["name"]:
            p["operations"] = ["laser_cutting", "folding"]
    br = _by(cs.parts_from_concept(ans, "J"))["HOOKED BRACKET (WELDED TO TRAY)"]
    assert br["welded_to_hosts"] and not br.get("coated_with_host")
    ans2 = _fsdu()
    ans2["parts"][7]["name"] = "HOOKED BRACKET (WELDED TO GANTRY)"
    br2 = _by(cs.parts_from_concept(ans2, "J"))["HOOKED BRACKET (WELDED TO GANTRY)"]
    assert not br2.get("welded_to_hosts")


def test_the_powder_area_includes_a_piece_coated_with_its_host_and_no_more():
    import wb_populate as W
    host = {"part_number": "T1", "quantity": 1,
            "material_estimate": {"stock_form": "sheet", "blank_length_mm": 460,
                                  "blank_width_mm": 300, "extended_material_cost_gbp": 2}}
    bracket = {"part_number": "B1", "quantity": 30, "coated_with_host": ["T1"],
               "material_estimate": {"stock_form": "sheet", "blank_length_mm": 40,
                                     "blank_width_mm": 30, "extended_material_cost_gbp": 1}}
    coated = W.route_coated_membership({"T1"})
    with_b = W.coated_sheet_area_m2([host, bracket], coated)
    alone = W.coated_sheet_area_m2([host], coated)
    assert abs((with_b - alone) - 0.04 * 0.03 * 2 * 30) < 1e-9
    nobody = W.route_coated_membership(set())
    assert W.coated_sheet_area_m2([bracket], nobody) == 0.0   # host not coated: no coat


def test_the_part_estimate_carries_the_coat_with_host_mark():
    import estimator as E
    pe = E.estimate_part({"part_number": "B1", "description": "bracket", "quantity": 2,
                          "normalized_material": "MILD STEEL", "normalized_thickness_mm": 2,
                          "blank_length_mm": 40, "blank_width_mm": 30,
                          "coated_with_host": ["T1"]}, job_quantity=10)
    assert pe.get("coated_with_host") == ["T1"]


def test_an_edging_counted_across_hosts_of_different_sizes_is_one_line_per_host():
    parts = cs.parts_from_concept(_fsdu(), "J")
    strips = [p for p in parts if "STRIP" in p["description"]]
    got = sorted((int(p["quantity"]), float(p["blank_length_mm"])) for p in strips)
    assert got == [(5, 300.0), (5, 300.0), (5, 460.0)]
    codes = {p["part_number"] for p in strips}
    assert "J-CPT09" in codes                              # the read's own code is kept
    assert all(p.get("strip_host") for p in strips)
    assert all(any("edging split per host" in str(f) for f in p["review_flags"])
               for p in strips)
    unit = cs.unit_assembly_part(parts, _fsdu(), "J")
    assert sum(int(p["quantity"]) for p in strips) == 15   # never more strips than shelves
    assert unit["pieces_fitted"] == 6 + 1 + 1 + 6 + 15 + 15


def test_an_edging_is_left_as_read_when_the_hosts_do_not_account_for_it():
    ans = _fsdu()
    ans["parts"][8]["quantity"] = 12                     # not the shelves' total
    strips = [p for p in cs.parts_from_concept(ans, "J") if "STRIP" in p["description"]]
    assert [(int(p["quantity"]), p["blank_length_mm"]) for p in strips] == [(12, 460.0)]


def _strips_after_approving(**approve):
    from source_precedence import apply_field, _source_key
    parts = cs.parts_from_concept(_fsdu(), "J")
    base = [p for p in parts if p.get("strip_host") and p["part_number"] == "J-CPT09"][0]
    for field, value in approve.items():
        base.pop(field, None)
        base.pop(_source_key(field), None)
        apply_field(base, field, value, "estimator_confirmed", note="estimator")
    dropped = cs.keep_approved_strip_lines(parts)
    strips = [p for p in parts if "STRIP" in p["description"]]
    return base, dropped, strips


def test_an_approved_count_for_the_whole_line_stands_whole():
    base, dropped, strips = _strips_after_approving(quantity=15)
    assert len(dropped) == 2 and strips == [base]
    assert sum(int(p["quantity"]) for p in strips) == 15        # the full count survives


def test_an_approval_short_of_the_whole_line_never_loses_the_other_hosts_strips():
    for approve in ({"quantity": 5}, {"blank_length_mm": 460.0}):
        base, dropped, strips = _strips_after_approving(**approve)
        assert dropped == [] and len(strips) == 3, approve
        assert sum(int(p["quantity"]) for p in strips) == 15, approve
        assert any("for this host only" in str(f) for f in base["review_flags"]), approve


def test_fixings_are_counted_per_mechanically_fitted_piece_when_the_read_names_none():
    parts = cs.parts_from_concept(_fsdu(), "J")
    mech = cs.mechanically_fitted(parts)
    # Counted: top cap 1, headers 6, strips 15. Welded post and plinth, the welded bracket,
    # and the trays that hang on it are not.
    assert mech["pieces"] == 1 + 6 + 15
    assert "hangs on" in " ".join(mech["excluded"].values())
    assert not cs.read_names_fixings(["QR labels"], parts)
    fx = cs.fixings_allowance_line(parts, "J", 2)
    per = config.CONCEPT_CONSUMABLE_ALLOWANCE["per_piece_fitted"]
    assert fx["quantity"] == per * 22 and fx["quantity_counted"]["count"] == per * 22
    assert fx["part_number"] == "J-NV02" and fx["unit_of_measure"] == "each"
    assert config.CONCEPT_CONSUMABLE_ALLOWANCE["reference_item"] in fx["description"]
    flag = fx["review_flags"][-1]
    assert "Counted:" in flag and "Not counted:" in flag and "Confirm" in flag


def test_handled_labelled_and_fastened_are_three_counts_each_with_its_basis():
    ans = _fsdu()
    ans["parts"].append({"name": "MAGNETIC GRAPHIC SET", "kind": "graphic", "quantity": 1})
    parts = cs.parts_from_concept(ans, "J")
    unit = cs.unit_assembly_part(parts, ans, "J")
    fc = unit["fitting_counts"]
    # Handled: every piece fitted at the unit, the 30 welded brackets with their trays.
    assert fc["handled"]["count"] == 6 + 1 + 1 + 6 + 15 + 15 + 1
    # Labelled: made components only — not the printed graphic, not the welded brackets.
    assert fc["labelled"]["count"] == fc["handled"]["count"] - 1
    # Fastened: never the welded post and plinth, the welded brackets or the hooked trays.
    assert fc["fastened"]["count"] == 1 + 6 + 15
    assert all(fc[k]["basis"] for k in fc)
    assert any("three counts" in str(f) for f in unit["review_flags"])
    (labels,) = cs.consumable_lines(["QR component labels"], "J",
                                    pieces_fitted=fc["handled"]["count"],
                                    components=fc["labelled"]["count"])
    assert labels["quantity"] == fc["labelled"]["count"]
    assert labels["items_to_apply"] == fc["labelled"]["count"]


def test_file_scan_counts_labels_on_labelled_components():
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    at = src[src.index("_nv_lines = concept_scan.consumable_lines("):][:600]
    assert 'components=((_unit or {}).get("components_labelled")' in at


def test_a_read_that_names_fixings_gets_no_second_line_and_its_set_counts_fixed_pieces():
    parts = cs.parts_from_concept(_fsdu(), "J")
    assert cs.read_names_fixings(["fixings and fasteners for assembly"], parts)
    (fset,) = cs.consumable_lines(["fixings and fasteners for assembly"], "J",
                                  pieces_fitted=49, fixed_pieces=22)
    per = config.CONCEPT_CONSUMABLE_ALLOWANCE["per_piece_fitted"]
    assert fset["consumable_allowance"]["count"] == per * 22


def test_file_scan_asks_for_unnamed_fixings_and_counts_the_set_on_fixed_pieces():
    """D-595 replaced D-499's automatic fixings line: fixings the read does not name are
    ASKED, never bought on no evidence; a set the read names still counts fixed pieces."""
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    at = src[src.index("_mech = concept_scan.mechanically_fitted(_sighted, unit_welded="):][:3200]
    assert "fixed_pieces=" in at and "read_names_fixings" in at
    assert "fixings_allowance_line" not in at, "no allowance is bought without evidence"
    assert "asked, not bought" in at
    assert "keep_approved_strip_lines(_sighted)" in src


def test_a_flat_claim_the_size_of_the_footprint_is_priced_with_its_flanges_and_kept():
    parts = _by(cs.parts_from_concept(_fsdu(), "J"))
    plinth, cap = parts["BASE / PLINTH"], parts["TOP CAP"]
    # The read's own 40 mm flange on the plinth; the house allowance on the cap.
    assert (plinth["blank_length_mm"], plinth["blank_width_mm"]) == (580.0, 580.0)
    house = config.CONCEPT_FLANGE_ALLOWANCE_MM
    assert cap["blank_length_mm"] == 500 + 2 * house
    for p in (plinth, cap):
        assert p["concept_read_blank_mm"]["length"] == 500.0    # the read's figure kept
        assert p["concept_blank_basis"] == "finished (by consistency)"
        assert any("CONCEPT CONSISTENCY" in str(f) for f in p["review_flags"])
    assert any("CONCEPT_FLANGE_ALLOWANCE_MM" in str(f) for f in cap["review_flags"])
    # A flat claim that is not the footprint is used as given.
    tray = parts["FRONT SHELF TRAY"]
    assert (tray["blank_length_mm"], tray["blank_width_mm"]) == (460.0, 300.0)
    assert not tray.get("concept_read_blank_mm")


def test_the_flange_allowance_reaches_material_cutting_and_coat_with_its_basis():
    """Not merely a flag: the developed blank is the blank the material is nested on, the
    perimeter the laser cuts (the Sheet Steel row's own length and width, plus any sighted
    cut-out) and the area the powder covers — and the working is on the assumption it sets."""
    import estimator as E
    import wb_populate as W
    ans = _fsdu()
    for p in ans["parts"]:
        if p["name"] == "TOP CAP":
            p["cutouts"] = [{"what": "opening", "count": 1, "length_mm": 200, "width_mm": 110}]
    parts = {p["description"]: p for p in cs.parts_from_concept(ans, "J")}
    plinth, cap = parts["BASE / PLINTH"], parts["TOP CAP"]
    pe = E.estimate_part(dict(plinth), job_quantity=200)
    me = pe["material_estimate"]
    assert (me["blank_length_mm"], me["blank_width_mm"]) == (580.0, 580.0)       # material
    area = W.coated_sheet_area_m2([pe], W.route_coated_membership(None))
    assert abs(area - 0.58 * 0.58 * 2) < 1e-9                                    # coat
    (a,) = [x for x in plinth["concept_assumptions"] if x["file_key"] == "blank_length_mm"]
    assert a["value"] == 580.0 and "40 mm each (the read's flange depth)" in a["cue"]
    house = config.CONCEPT_FLANGE_ALLOWANCE_MM
    side = 500 + 2 * house
    assert cap["manufacturing_features"]["cut_length_mm"] == round(
        2 * (side + side) + 2 * (200 + 110), 1)                                  # cutting
    (c,) = [x for x in cap["concept_assumptions"] if x["file_key"] == "blank_length_mm"]
    assert "CONCEPT_FLANGE_ALLOWANCE_MM" in c["cue"]


def test_a_finished_size_with_no_flange_takes_the_house_allowance_said():
    d = cs.developed_blank({"length": 400, "width": 300}, {"folds": 2, "blank_basis": "finished"})
    house = config.CONCEPT_FLANGE_ALLOWANCE_MM
    assert d["changed"] and d["width"] == 300 + 2 * house
    assert "CONCEPT_FLANGE_ALLOWANCE_MM" in d["working"]
    flat = cs.developed_blank({"length": 400, "width": 300},
                              {"folds": 2, "blank_basis": "developed"},
                              envelope={"width": 500, "depth": 500})
    assert not flat["changed"] and "used as given" in flat["working"]


def test_the_new_vocabulary_lives_in_config_and_holds_no_job():
    for name in ("CONCEPT_FIXING_WORDS", "CONCEPT_FIXING_FREE_WORDS",
                 "CONCEPT_WELDED_TO_WORDS", "CONCEPT_EDGE_STRIP_WORDS"):
        words = getattr(config, name)
        assert words and all(isinstance(w, str) for w in words)
    assert isinstance(config.CONCEPT_FLANGE_ALLOWANCE_MM, (int, float))
    rule = config.PROJECT_SCOPE_EVIDENCE
    assert {"categories", "per_order", "per_batch", "per_unit", "repeat_on_change"} <= set(rule)
    blob = repr([getattr(config, n) for n in (
        "CONCEPT_FIXING_WORDS", "CONCEPT_FIXING_FREE_WORDS", "CONCEPT_WELDED_TO_WORDS",
        "CONCEPT_EDGE_STRIP_WORDS", "PROJECT_SCOPE_EVIDENCE")]).lower()
    for job_word in ("boots", "wellness", "btc", "kwellness"):
        assert job_word not in blob


# ── D-500: the 20:49 book — volume pricing, the read's changes on the book, components named
#    inside another line, and the reviewer's four queued refinements ────────────────────────

def test_a_counted_item_is_asked_its_price_at_the_jobs_volume():
    import indicative_price as ip
    counted = {"description": "QR component label 20x20mm", "quantity": 48,
               "unit_of_measure": "each",
               "quantity_counted": {"count": 48, "basis": "1 per component × 48"}}
    b = ip.research_brief(counted, order_qty=200)
    assert b["volume_needed"] == 9600
    assert "= 9600 in all" in b["ask"] and "never a single-item retail listing" in b["ask"]
    plain = ip.research_brief({"description": "castor 75mm", "quantity": 4}, order_qty=200)
    assert "volume_needed" not in plain and "in all" not in plain["ask"]


def test_the_researcher_is_given_the_lines_count_basis():
    src = (ROOT / "src" / "estimator.py").read_text(encoding="utf-8")
    assert '"quantity_counted": part.get("quantity_counted"),' in src
    fx = cs.fixings_allowance_line(cs.parts_from_concept(_fsdu(), "J"), "J", 2)
    assert "at this job's volume" in fx["research_description"]


def test_the_reads_changes_are_on_the_book_and_each_dropped_line_is_a_decision():
    flags = cs.change_flags(["added: TRAY x5", "dropped: HOOKED BRACKET (WELDED TO TRAY) x30",
                             "HEADER: count 3 → 6 (no reason given)"],
                            baseline="new interpretation with the previous bill")
    assert flags[0].startswith("CONCEPT READ: new interpretation")
    assert "a different read" in flags[0]
    assert any(f.startswith("CONCEPT READ CHANGED") and "HEADER: count 3 → 6" in f
               and "dropped" not in f for f in flags)
    (drop,) = [f for f in flags if f.startswith("DROPPED FROM THE PREVIOUS READ")]
    assert "HOOKED BRACKET (WELDED TO TRAY) x30" in drop and "material" in drop
    same = cs.change_flags([], baseline="unchanged read reused")
    assert same == ["CONCEPT READ: unchanged read reused — the same read as the previous "
                    "book; differences are the engine's"]
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    at = src[src.index("_unit = concept_scan.unit_assembly_part(_sighted, _answer, _job_name)"):]
    assert "concept_scan.change_flags(" in at[:900]


def test_a_component_named_inside_another_line_is_found_unless_listed_or_a_feature():
    lines = [{"name": "FRONT SHELF TRAY (with welded hooked brackets)"},
             {"name": "TOP CAP with graphic pass-through cut-out"},
             {"name": "SIDE PANEL with welded seams"},
             {"name": "HEADER"}]
    (hit,) = cs.components_named_inside(lines)
    assert hit["component"] == "bracket" and "welded hooked brackets" in hit["phrase"]
    listed = lines + [{"name": "HOOKED BRACKET"}]
    assert cs.components_named_inside(listed) == []


def test_the_sheet_check_puts_the_read_back_naming_the_folded_component():
    ans = {"product": {"name": "FSDU"}, "parts": [
        {"name": "FRONT SHELF TRAY (with welded hooked brackets)", "kind": "fabricated",
         "quantity": 5, "assumed_blank_mm": {"length": 460, "width": 300, "thickness": 1.5}}]}
    chk = cs.sheet_check(ans)
    assert chk["components_inside"] and not chk["agrees"]
    assert any("names welded hooked brackets but no line lists them" in f
               for f in chk["failures"])


def _folded(*hosts):
    """The 20:49 shape: brackets named inside the trays' names, no line of their own."""
    ans = _fsdu()
    ans["parts"] = [p for p in ans["parts"] if "BRACKET" not in p["name"]]
    for p in ans["parts"]:
        if p["name"] in hosts:
            p["name"] = p["name"].replace("TRAY", "TRAY (with welded hooked brackets)", 1)
    return ans


_PREV_BRACKET = {"name": "HOOKED BRACKET (WELDED TO TRAY)", "kind": "fabricated",
                 "quantity": 30, "material": "MILD STEEL", "length": 40, "width": 30,
                 "thickness": 2.0, "folds": 1}


def test_a_component_still_folded_in_with_nothing_to_size_it_is_a_question():
    """D-500, kept by D-501 where no allowance can be reasoned: nothing sizes the component and
    it is not a small joined component the house allowance covers."""
    ans = _fsdu()
    for p in ans["parts"]:
        if p["name"] == "FRONT SHELF TRAY":
            p["name"] = "FRONT SHELF TRAY (with welded mesh panels)"
    parts = cs.parts_from_concept(ans, "J")
    res = cs.allow_components_inside(parts, {"parts": [{"name": "MESH PANEL", "quantity": 5}]},
                                     "J")
    assert res["minted"] == [] and len(res["asked"]) == 1
    tray = [p for p in parts if p["description"].startswith("FRONT SHELF TRAY")][0]
    (q,) = tray["manufacturing_questions"]
    assert "no line of the bill costs" in str(q) and "nothing on this read or the previous" in str(q)


def test_a_bonded_piece_takes_no_screws_and_hung_pieces_carry_the_retainer_note():
    ans = _fsdu()
    for p in ans["parts"]:
        if "STRIP" in p["name"]:
            p["name"] = "35MM SHELF EDGE STRIP (bonded)"
    parts = cs.parts_from_concept(ans, "J")
    mech = cs.mechanically_fitted(parts)
    assert mech["pieces"] == 1 + 6                       # top cap and headers only
    assert any("bonded" in why for why in mech["excluded"].values())
    fx = cs.fixings_allowance_line(parts, "J", 2)
    assert "retainers or anti-lift clips" in fx["review_flags"][-1]


def test_the_label_scope_states_its_too_small_exception(monkeypatch):
    parts = cs.parts_from_concept(_fsdu(), "J")
    unit = cs.unit_assembly_part(parts, _fsdu(), "J")
    assert "no too-small exception applied" in unit["fitting_counts"]["labelled"]["basis"]
    monkeypatch.setattr(config, "CONCEPT_LABEL_MIN_FACE_MM", 40, raising=False)
    unit2 = cs.unit_assembly_part(parts, _fsdu(), "J")
    assert unit2["components_labelled"] == unit["components_labelled"] - 15   # 35 mm strips
    assert "too small to carry one" in unit2["fitting_counts"]["labelled"]["basis"]


def test_the_quantity_breaks_tab_says_when_one_batch_is_assumed(tmp_path):
    import openpyxl
    from quantity_breaks_tab import write_quantity_breaks_tab, SHEET_NAME
    swept = {"rows": [{"quantity": 1, "material": 10, "labour": 5, "unit": 20.0},
                      {"quantity": 200, "material": 10, "labour": 3, "unit": 16.0}]}

    def _tab(batch_size):
        p = tmp_path / f"b{batch_size}.xlsx"
        openpyxl.Workbook().save(p)
        pc = {"total_gbp": 450.0, "lines": [
            {"code": "PROJECT-05", "scope": "per_batch", "state": "provisional",
             "order_gbp": 450.0, "per_batch_gbp": 450.0, "batch_size": batch_size}]}
        write_quantity_breaks_tab(p, swept, requested=[200], project_costs=pc)
        ws = openpyxl.load_workbook(p)[SHEET_NAME]
        return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "ONE batch is taken at every quantity" in _tab(None)
    assert "ONE batch is taken" not in _tab(50)


def test_the_coat_is_net_of_the_openings_cut_through_it():
    import estimator as E
    import wb_populate as W
    g = cs.cutout_geometry({"cutouts": [{"what": "opening", "count": 1, "length_mm": 200,
                                         "width_mm": 110},
                                        {"what": "hole", "count": 4, "length_mm": 10}]})
    assert abs(g["open_area_mm2"] - (200 * 110 + 4 * 3.14159 / 4 * 100)) < 0.2
    ans = _fsdu()
    for p in ans["parts"]:
        if p["name"] == "TOP CAP":
            p["cutouts"] = [{"what": "opening", "count": 1, "length_mm": 200, "width_mm": 110}]
    cap = [p for p in cs.parts_from_concept(ans, "J") if p["description"] == "TOP CAP"][0]
    assert cap["concept_open_area_mm2"] == 22000.0
    pe = E.estimate_part(dict(cap), job_quantity=10)
    assert pe["concept_open_area_mm2"] == 22000.0
    L, Wd = pe["material_estimate"]["blank_length_mm"], pe["material_estimate"]["blank_width_mm"]
    area = W.coated_sheet_area_m2([pe], W.route_coated_membership(None))
    assert abs(area - ((L / 1000) * (Wd / 1000) - 0.022) * 2) < 1e-9


# ── D-501: the 20:49 review — an answer reconciles before it is used, and every component the
#    read names reaches a charged line, an identified allowance, or a reasoned exclusion ──────

def _label_answer(**over):
    """The researched answer behind the 20:49 labels: 12p a label, found at the job's volume."""
    a = {"price_gbp": 0.12, "unit": "each", "source": "https://example-labels.co.uk/qr-20",
         "as_of": "2026-10-12", "quantity_basis": "9600 off, on rolls",
         "item_priced": "20 x 20 mm printed QR label, one label",
         "price_is_reproducible": True}
    a.update(over)
    return a


_LABEL = {"part_number": "J-NV02", "description": "QR component label 20x20mm",
          "is_bought_in": True, "concept": True, "concept_kind": "bought_in",
          "normalized_material": "BOUGHT_IN", "quantity": 48, "unit_of_measure": "each",
          "quantity_counted": {"count": 48, "basis": "1 per labelled component x 48"},
          "research_description": "20 x 20 mm QR label — the price of one item at this "
                                  "job's volume"}


def test_the_2049_label_answer_is_twelve_pence_each_and_its_working_is_the_line_total():
    import indicative_price as ip
    got = ip.resolve_indicative(dict(_LABEL), order_qty=200, ask=lambda b: _label_answer())
    assert got["calculation"]["working"] == "48 off x GBP 0.1200 each = GBP 5.7600 a unit"
    assert got["price_gbp"] == 5.76                 # the LINE's cost on one unit
    assert got["unit_price_gbp"] == 0.12            # the ITEM's price
    assert got["applies_per_item"] is True


def test_the_engine_charges_the_label_twelve_pence_each_once_not_twice(monkeypatch):
    import estimator as E
    asked = []
    monkeypatch.setattr(E, "_rung4_researcher",
                        lambda brief: (asked.append(brief), _label_answer())[1])
    pe = E.estimate_part(dict(_LABEL, review_flags=[]), job_quantity=200)
    me = pe["material_estimate"]
    assert me["unit_material_cost_gbp"] == 0.12     # was 5.76 — the line total taken as each
    assert me["extended_material_cost_gbp"] == 5.76  # was 276.48 — 48 × 5.76
    # Asked at THIS order's volume: estimate_part stamps assumed_job_quantity, and the
    # research rung used to read only job_quantity — every question went out at an order of 1.
    assert asked[0]["order_quantity"] == 200 and asked[0]["volume_needed"] == 9600
    (vol,) = [f for f in pe["review_flags"] if f.startswith("PRICED AT VOLUME")]
    assert "£0.1200 each was asked at 9600 items in all" in vol
    assert "every Quantity Breaks column re-uses that item price" in vol


def test_a_pack_price_is_divided_by_the_stated_pack_and_refused_without_one():
    import indicative_price as ip
    line = {"description": "M5 x 12 pan head screw", "quantity": 100, "unit_of_measure": "each"}
    per_box = _label_answer(price_gbp=9.00, unit="per box", quantity_basis="box of 200",
                            item_priced="box of 200 screws")
    got = ip.resolve_indicative(dict(line), ask=lambda b: per_box)
    assert got["unit_price_gbp"] == 0.045
    assert "GBP 9.0000 a box ÷ 200 = GBP 0.0450 each" in got["evidence"]["quantity_basis"]
    assert got["calculation"]["working"] == "100 off x GBP 0.0450 each = GBP 4.5000 a unit"
    no_size = ip.resolve_indicative(dict(line), ask=lambda b: dict(per_box, quantity_basis="trade"))
    assert no_size["price_gbp"] is None and "no pack size stated" in no_size["missing"]


def test_a_made_part_allowance_is_one_parts_price_not_the_lines(monkeypatch):
    import estimator as E
    monkeypatch.setattr(E, "_rung4_researcher", lambda brief: _label_answer(
        price_gbp=3.50, item_priced="one folded steel bracket"))
    pe = {"part_number": "J-CPT09", "description": "BRACKET", "quantity": 30,
          "assumed_job_quantity": 200, "review_flags": []}
    assert E._made_part_material_allowance(pe) == 3.50


def test_a_set_is_asked_as_the_whole_set_and_a_one_panel_answer_is_not_used(monkeypatch):
    import estimator as E
    import indicative_price as ip
    b = ip.research_brief({"description": "REPLACEABLE CARD / MAGNETIC GRAPHIC SET",
                           "quantity": 1, "sold_as_set": True})
    assert "COMPLETE set" in b["ask"] and b["sold_as_set"]
    gset = {"part_number": "J-CPT12", "description": "REPLACEABLE CARD / MAGNETIC GRAPHIC SET",
            "is_bought_in": True, "concept": True, "concept_kind": "graphic", "quantity": 1,
            "normalized_material": "PRINTED_PAPER", "assumed_job_quantity": 200,
            "review_flags": []}
    asked = []
    monkeypatch.setattr(E, "_rung4_researcher", lambda brief: (asked.append(brief), _label_answer(
        price_gbp=18.00, quantity_basis="200 off", item_priced="single panel magnetic graphic"))[1])
    res = E._resolve_part_system_cost(gset)
    assert asked and asked[0].get("sold_as_set") is True
    assert res["applied_unit_cost"] in (None, 0.0)
    assert any("PRICE SCOPE MISMATCH" in f and "£18.00" in f for f in gset["review_flags"])
    # The whole set, said so, is used.
    whole = dict(gset, review_flags=[])
    monkeypatch.setattr(E, "_rung4_researcher", lambda brief: _label_answer(
        price_gbp=64.00, quantity_basis="200 sets", item_priced="complete set of 5 panels"))
    assert E._resolve_part_system_cost(whole)["applied_unit_cost"] == 64.00
    # And a line that is not a set is never refused for a one-piece answer.
    panel = dict(gset, description="REAR GRAPHIC PANEL", review_flags=[])
    monkeypatch.setattr(E, "_rung4_researcher", lambda brief: _label_answer(
        price_gbp=18.00, quantity_basis="200 off", item_priced="single panel"))
    assert E._resolve_part_system_cost(panel)["applied_unit_cost"] == 18.00


def test_a_bracket_folded_into_its_trays_is_priced_once_from_the_previous_reads_line():
    parts = cs.parts_from_concept(_folded("FRONT SHELF TRAY", "SIDE SHELF TRAY (LEFT)"), "J")
    res = cs.allow_components_inside(parts, {"parts": [_PREV_BRACKET]}, "J")
    (cpa,) = res["minted"]                            # ONE line across both hosts
    assert cpa["part_number"] == "J-CPA01" and res["asked"] == []
    assert cpa["quantity"] == 30
    assert (cpa["blank_length_mm"], cpa["blank_width_mm"]) == (40.0, 30.0)
    assert cpa["normalized_thickness_mm"] == 2.0 and cpa["normalized_material"] == "MILD STEEL"
    hosts = [p["part_number"] for p in parts if "brackets" in p["description"]]
    assert cpa["welded_to_hosts"] == hosts and len(hosts) == 2
    assert cpa["coated_with_host"] == hosts           # the trays are powder coated
    # The allowance never welds (D-502): the trays' own names say the brackets are welded on,
    # so the trays carry the weld — added here, where this read left it off — charged once.
    assert cpa["inferred_operations"] == ["laser_cutting", "folding"]
    for p in parts:
        if p["part_number"] in hosts:
            assert "welding" in p["inferred_operations"]
    assert cpa["provisional_allowance"]["from_previous_read"]["name"] == _PREV_BRACKET["name"]
    flag = cpa["review_flags"][-1]
    assert flag.startswith("PROVISIONAL ALLOWANCE") and "Charged ONCE here" in flag
    assert "not confirmed by this read" in flag and "30 × 40 × 30 × 2 mm MILD STEEL" in flag
    assert "its weld and dress time sit on those lines' Weld and Dress rows, NOT on this line" in flag
    for p in parts:
        if p["part_number"] in hosts:
            assert "charged on J-CPA01 as a provisional allowance" in p["review_flags"][-1]
            assert any("welding added" in f for f in p["review_flags"])
            assert not p.get("manufacturing_questions")
    assert res["allowed"] == {_PREV_BRACKET["name"]: "J-CPA01"}
    assert not any("no size could be sighted" in f for f in cpa["review_flags"])


def test_a_host_row_in_the_previous_bill_never_sizes_the_component_it_names():
    """The 20:49 bill itself lists 'FRONT SHELF TRAY (with welded hooked brackets)' 460 × 300
    × 5 — a tray. It is not the brackets' size or count."""
    parts = cs.parts_from_concept(_folded("FRONT SHELF TRAY"), "J")
    host_row = {"name": "FRONT SHELF TRAY (with welded hooked brackets)", "kind": "fabricated",
                "quantity": 5, "material": "MILD STEEL", "length": 460, "width": 300,
                "thickness": 1.5, "folds": 3}
    (cpa,) = cs.allow_components_inside(parts, {"parts": [host_row]}, "J")["minted"]
    assert "house_allowance" in cpa["provisional_allowance"]
    assert cpa["blank_length_mm"] != 460.0 and cpa["quantity"] == 2 * 5


def test_the_joint_stays_on_a_host_that_already_welds():
    ans = _folded("FRONT SHELF TRAY")
    for p in ans["parts"]:
        if p["name"].startswith("FRONT SHELF TRAY"):
            p["operations"] = p["operations"] + ["welding"]
    parts = cs.parts_from_concept(ans, "J")
    (cpa,) = cs.allow_components_inside(parts, {"parts": [_PREV_BRACKET]}, "J")["minted"]
    tray = [p for p in parts if p["description"].startswith("FRONT SHELF TRAY")][0]
    assert "welding" not in cpa["inferred_operations"]
    assert tray["inferred_operations"].count("welding") == 1        # not added a second time
    assert not any("welding added" in f for f in tray["review_flags"])
    assert f"the joint is welded to {tray['part_number']}" in cpa["review_flags"][-1]
    assert cpa["provisional_allowance"]["joint_on"] == [tray["part_number"]]


def test_a_dropped_line_that_is_allowed_says_where_it_is_charged():
    flags = cs.change_flags(["dropped: HOOKED BRACKET (WELDED TO TRAY) x30",
                             "dropped: TOP CAP x1"],
                            baseline="new interpretation with the previous bill",
                            allowed={"HOOKED BRACKET (WELDED TO TRAY)": "J-CPA01"})
    drops = [f for f in flags if f.startswith("DROPPED FROM THE PREVIOUS READ")]
    assert "charged as a PROVISIONAL ALLOWANCE on J-CPA01" in drops[0]
    assert "not confirmed" in drops[0]
    assert "its material and work are missing" in drops[1]       # TOP CAP: still a gap


def test_two_hosts_naming_the_same_component_do_not_list_it_for_each_other():
    lines = [{"name": "FRONT SHELF TRAY (with welded hooked brackets)"},
             {"name": "SIDE SHELF TRAY (mirrored, with welded hooked brackets)"}]
    assert [h["host"] for h in cs.components_named_inside(lines)] == [l["name"] for l in lines]


def test_file_scan_prices_the_folded_component_and_asks_the_pitch():
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    assert "concept_scan.allow_components_inside(_sighted, _prev_bill, _job_name," in src
    assert '_sighted.extend(_inside["minted"])' in src
    assert "concept_scan.pitch_as_size_questions(_sighted, _brief" in src
    at = src.index("concept_scan.allow_components_inside(")
    assert at < src.index("concept_scan.save_bill(job_folder, _bill_job, _answer, records=_sighted")


def test_a_bonded_joint_is_priced_as_adhesive_not_left_out():
    ans = _fsdu()
    for p in ans["parts"]:
        if "STRIP" in p["name"]:
            p["name"] = "35MM SHELF EDGE STRIP (bonded)"
    parts = cs.parts_from_concept(ans, "J")
    adh = cs.adhesive_allowance_line(parts, "J", 3)
    assert adh["part_number"] == "J-NV03" and adh["unit_of_measure"] == "m"
    assert adh["quantity"] == round((5 * 460 + 5 * 300 + 5 * 300) / 1000.0, 2)   # 5.3 m
    assert adh["quantity_counted"]["count"] == adh["quantity"]
    tape = config.CONCEPT_ADHESIVE_ALLOWANCE["methods"]["tape"]
    assert adh["joining_method"] == "tape" and tape["reference_item"] in adh["description"]
    assert "price per m" in adh["research_description"]
    assert "provisional adhesive allowance" in adh["review_flags"][-1]
    # No bonded piece, no adhesive line.
    assert cs.adhesive_allowance_line(cs.parts_from_concept(_fsdu(), "J"), "J", 3) is None
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    assert "concept_scan.adhesive_allowance_lines(_sighted, _job_name" in src


def test_a_cut_out_read_at_the_stated_pitch_is_asked_and_still_priced():
    ans = _fsdu()
    for p in ans["parts"]:
        if p["name"] == "CORNER POST":
            p["cutouts"] = [{"what": "slot", "count": 40, "length_mm": 31.75, "width_mm": 5}]
    parts = cs.parts_from_concept(ans, "J")
    post = [p for p in parts if p["description"] == "CORNER POST"][0]
    cut_before = post["manufacturing_features"]["cut_length_mm"]
    brief = ("Shelving positions to be fixed to standard slotted upright vertical pitch of "
             "31.75mm. Body 500 x 500 mm.")
    asked = cs.pitch_as_size_questions(parts, brief)
    assert asked == ["J-CPT01: slot 31.75 × 5 mm = pitch 31.75 mm"]
    (q,) = post["manufacturing_questions"]
    assert "31.75 mm is stated as a pitch" in q["issue"] and "priced as read" in q["assumption"]
    assert post["manufacturing_features"]["cut_length_mm"] == cut_before      # no money moved
    # A slot that is not the pitch is not asked; a figure in another sentence is not a pitch.
    assert cs._pitch_figures(brief) == [{"mm": 31.75, "said":
                                         "standard slotted upright vertical pitch of 31.75mm"}]
    for p in ans["parts"]:
        if p["name"] == "CORNER POST":
            p["cutouts"] = [{"what": "slot", "count": 40, "length_mm": 20, "width_mm": 5}]
    assert cs.pitch_as_size_questions(cs.parts_from_concept(ans, "J"), brief) == []


def test_cutting_and_coating_take_the_same_opening_once():
    """The 20:49 cap's opening is 460 × 20 mm (960 mm of internal cut). The laser takes its
    perimeter and the coat its area — the same opening, deducted once, by the coat only."""
    import estimator as E
    import wb_populate as W
    ans = _fsdu()
    for p in ans["parts"]:
        if p["name"] == "TOP CAP":
            p["cutouts"] = [{"what": "opening", "count": 1, "length_mm": 460, "width_mm": 20}]
    cap = [p for p in cs.parts_from_concept(ans, "J") if p["description"] == "TOP CAP"][0]
    L, Wd = cap["blank_length_mm"], cap["blank_width_mm"]
    assert cap["manufacturing_features"]["cut_length_mm"] == round(2 * (L + Wd) + 960, 1)
    assert cap["concept_open_area_mm2"] == 9200.0
    pe = E.estimate_part(dict(cap), job_quantity=10)
    me = pe["material_estimate"]
    assert (me["blank_length_mm"], me["blank_width_mm"]) == (L, Wd)   # the blank is not cut down
    area = W.coated_sheet_area_m2([pe], W.route_coated_membership(None))
    assert abs(area - ((L / 1000) * (Wd / 1000) - 0.0092) * 2) < 1e-9


def test_a_project_cost_charged_once_is_not_an_unpriced_rate():
    import estimator_inputs as ei
    line = {"part_number": "PROJECT-02", "description": "Prototype",
            "project_cost_line": {"scope": "per_order", "state": "provisional",
                                  "order_gbp": 1200.0}}
    assert ei.canonical_pricing_status(line, 0.0) == ei.NOT_APPLICABLE
    open_line = dict(line, project_cost_line={"scope": "per_order", "order_gbp": None})
    assert ei.canonical_pricing_status(open_line, 0.0) == ei.UNPRICED


def test_the_banner_says_the_report_adds_the_checks_it_had_not_seen():
    import costed_facts as cf
    src = (ROOT / "src" / "costed_facts.py").read_text(encoding="utf-8")
    assert "a provisional count at workbook population — the consistency checks had not" in src
    assert "the report carries the final release status" in src
    assert "adds any check they raise" in src
    assert hasattr(cf, "outstanding_summary")


def test_the_quantity_breaks_tab_says_what_was_priced_at_this_orders_volume(tmp_path):
    import openpyxl
    from quantity_breaks_tab import write_quantity_breaks_tab, SHEET_NAME
    p = tmp_path / "v.xlsx"
    openpyxl.Workbook().save(p)
    swept = {"rows": [{"quantity": 1, "material": 10, "labour": 5, "unit": 20.0},
                      {"quantity": 200, "material": 10, "labour": 3, "unit": 16.0}]}
    write_quantity_breaks_tab(p, swept, requested=[1, 200],
                              volume_priced=["J-NV02 (£0.1200 each was asked at 9600 items "
                                             "in all (this order's quantity))"])
    ws = openpyxl.load_workbook(p)[SHEET_NAME]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Priced at this order's volume and re-used in every column: J-NV02" in text
    src = (ROOT / "src" / "main.py").read_text(encoding="utf-8")
    assert 'str(_f).startswith("PRICED AT VOLUME:")' in src


def test_the_d501_vocabulary_lives_in_config_and_holds_no_job():
    for name in ("CONCEPT_SET_WORDS", "PRICE_SINGLE_PIECE_WORDS", "CONCEPT_PITCH_WORDS"):
        words = getattr(config, name)
        assert words and all(isinstance(w, str) for w in words)
    assert config.CONCEPT_JOINED_COMPONENT_ALLOWANCE["components"]
    assert "pack" not in config.CONCEPT_SET_WORDS        # a pack is divided, not priced whole
    blob = repr([config.CONCEPT_SET_WORDS, config.PRICE_SINGLE_PIECE_WORDS,
                 config.CONCEPT_PITCH_WORDS, config.CONCEPT_ADHESIVE_ALLOWANCE,
                 config.CONCEPT_JOINED_COMPONENT_ALLOWANCE]).lower()
    for job_word in ("boots", "wellness", "kwellness", "31.75"):
        assert job_word not in blob


def test_with_no_read_sizing_it_a_small_component_is_a_house_allowance_counted_from_the_brief():
    """The 20:49 shape on the NEXT run: its saved bill no longer has the bracket line, so no
    read sizes them. They are not left unpriced: the count from the brief's own words, the size
    a house allowance at the host's gauge, every figure said, charged once."""
    parts = cs.parts_from_concept(_folded("FRONT SHELF TRAY", "SIDE SHELF TRAY (LEFT)"), "J")
    brief = ("Each shelf is a folded steel tray with a front lip, with two hooked\n"
             "brackets welded or riveted to the tray (not separate brackets).")
    res = cs.allow_components_inside(parts, {}, "J", brief=brief)
    (cpa,) = res["minted"]
    rule = config.CONCEPT_JOINED_COMPONENT_ALLOWANCE
    assert cpa["quantity"] == 2 * (5 + 5)
    assert (cpa["blank_length_mm"], cpa["blank_width_mm"]) == (
        float(rule["blank_mm"]["length"]), float(rule["blank_mm"]["width"]))
    assert cpa["normalized_thickness_mm"] == 1.5                  # the hosts' gauge
    assert cpa["inferred_operations"] == ["laser_cutting", "folding"]      # the hosts weld
    flag = cpa["review_flags"][-1]
    assert "2 per host from the enquiry brief ('two hooked brackets')" in flag
    assert rule["basis"] in flag and "Charged ONCE here" in flag
    assert "house_allowance" in cpa["provisional_allowance"]
    # No count anywhere: the plural is the assumption, and it says so.
    (cpa2,) = cs.allow_components_inside(
        cs.parts_from_concept(_folded("FRONT SHELF TRAY"), "J"), {}, "J")["minted"]
    assert cpa2["quantity"] == 2 * 5
    assert "the read names them in the plural and nothing states a count" in cpa2["review_flags"][-1]


def test_a_line_a_read_dropped_is_kept_as_history_and_sizes_a_later_allowance(tmp_path,
                                                                               monkeypatch):
    monkeypatch.setattr(cs, "_cache_dir", lambda: tmp_path / "engine" / "cache" / "vision_concept")
    job = tmp_path / "staged" / "Client" / "JOB"
    job.mkdir(parents=True)
    cs.save_bill(job, "JOB", _fsdu(), read_key="r1")                    # lists the brackets
    cs.save_bill(job, "JOB", _folded("FRONT SHELF TRAY"), read_key="r2")  # folds them in
    prev = cs.load_previous_bill(job, "JOB")
    (h,) = [r for r in prev["history"] if r["name"] == "HOOKED BRACKET (WELDED TO TRAY)"]
    assert h["dropped_by_read"] == "r2" and h["length"] == 40 and h["quantity"] == 30
    # History is a basis for an allowance, never a line of the read the model is shown.
    assert "HOOKED BRACKET (WELDED TO TRAY)" not in cs.previous_bill_text(prev)
    parts = cs.parts_from_concept(_folded("FRONT SHELF TRAY"), "J")
    res = cs.allow_components_inside(parts, prev, "J")
    (cpa,) = res["minted"]
    assert (cpa["blank_length_mm"], cpa["blank_width_mm"], cpa["quantity"]) == (40.0, 30.0, 30)
    assert "an earlier read's (dropped since, kept in the bill's history)" in cpa["review_flags"][-1]
    assert res["allowed"] == {"HOOKED BRACKET (WELDED TO TRAY)": cpa["part_number"]}
    # A third read that lists them again takes them out of history.
    cs.save_bill(job, "JOB", _fsdu(), read_key="r3")
    assert not [r for r in cs.load_previous_bill(job, "JOB")["history"]
                if r["name"] == "HOOKED BRACKET (WELDED TO TRAY)"]


def test_file_scan_hands_the_brief_to_the_component_allowance():
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    at = src[src.index("concept_scan.allow_components_inside("):][:200]
    assert 'brief=_brief or ""' in at


# ── D-502: the reviewer's four qualifications on D-501 — the set is never left at nothing, a
#    stated count beats an inferred one, the joint is charged once on its host, and every
#    allowance is construction-aware ─────────────────────────────────────────────────────────

_GSET = {"part_number": "J-CPT12", "description": "REPLACEABLE CARD / MAGNETIC GRAPHIC SET",
         "is_bought_in": True, "concept": True, "concept_kind": "graphic", "quantity": 1,
         "normalized_material": "PRINTED_PAPER", "assumed_job_quantity": 200}
_ONE_PANEL_2049 = "replaceable card/magnetic graphic set approx 1400x460 mm, one panel"


def _set_research(asked, reask=None):
    def research(brief):
        asked.append(brief)
        if reask is not None and str(brief.get("description")).startswith("COMPLETE SET"):
            return reask
        return _label_answer(price_gbp=18.00, quantity_basis="200 off",
                             item_priced=_ONE_PANEL_2049)
    return research


def test_the_2049_one_panel_answer_is_re_asked_once_as_the_complete_set(monkeypatch):
    import estimator as E
    asked = []
    monkeypatch.setattr(E, "_rung4_researcher", _set_research(asked, reask=_label_answer(
        price_gbp=64.00, quantity_basis="200 sets",
        item_priced="complete set: 3 header magnetics, rear panel graphic, 2 base cards")))
    part = dict(_GSET, review_flags=[])
    res = E._resolve_part_system_cost(part)
    assert len(asked) == 2
    # A NEW DESCRIPTION, so the price cache cannot return the refused one-panel answer.
    assert asked[1]["description"].startswith("COMPLETE SET: ")
    assert f"A previous answer priced only {_ONE_PANEL_2049} at GBP 18.00" in asked[1]["ask"]
    assert res["applied_unit_cost"] == 64.00
    assert res["result"]["selected"]["metadata"]["pricing_mode"] == "llm_indicative_set_reask"
    assert any(f.startswith("PRICE SCOPE: the first answer priced") and "£18.00" in f
               for f in part["review_flags"])


def test_a_set_the_re_ask_still_prices_one_piece_of_is_a_reasoned_allowance(monkeypatch):
    import estimator as E
    monkeypatch.setattr(E, "_rung4_researcher", _set_research([]))
    part = dict(_GSET, review_flags=[], set_pieces={"count": 4, "basis": "4 graphic positions"})
    res = E._resolve_part_system_cost(part)
    assert res["applied_unit_cost"] == 72.00                       # 4 × £18.00
    assert res["result"]["selected"]["metadata"]["pricing_mode"] == "set_allowance_from_one_piece"
    (flag,) = [f for f in part["review_flags"] if f.startswith("PROVISIONAL SET ALLOWANCE")]
    assert "4 piece(s) (4 graphic positions) × £18.00 = £72.00" in flag
    assert "an extrapolation from ONE piece, not a set price" in flag
    # With no count anywhere it is not priced, and says what is needed — never the £18 as the set.
    bare = dict(_GSET, review_flags=[])
    assert E._resolve_part_system_cost(bare)["applied_unit_cost"] in (None, 0.0)
    assert any("nothing states how many pieces the set holds" in f for f in bare["review_flags"])


def test_a_sets_pieces_come_from_its_own_words_or_the_positions_that_hold_them():
    ans = {"product": {"name": "FSDU"}, "parts": [
        _p("TOP CAP", 1, 500, 500, LF, 4),
        _p("HEADER", 3, 460, 80, LF, 1),
        _p("REAR GRAPHIC PANEL", 1, 1400, 460, LF),
        _p("BASE GRAPHIC HOLDER", 2, 500, 80, LF, 1),
        {"name": "REPLACEABLE CARD / MAGNETIC GRAPHIC SET", "kind": "graphic", "quantity": 1}]}
    parts = cs.parts_from_concept(ans, "J")
    gset = [p for p in parts if p["description"].endswith("SET")][0]
    # The brief this run used ties no graphic to the header ("metal header and rear graphic panel").
    cs.set_pieces_for(parts, "metal header and rear graphic panel; top cap with graphic "
                             "pass-through cut-out; 2 base graphic holders")
    assert gset["set_pieces"]["count"] == 1 + 1 + 2             # cap, rear panel, two holders
    assert "HEADER" not in gset["set_pieces"]["basis"]
    # A brief that does tie it counts the headers as positions too.
    cs.set_pieces_for(parts, "headers x3 (metal, to take a magnetic graphic or a card graphic "
                             "holder); rear graphic panel x1")
    assert gset["set_pieces"]["count"] == 3 + 1 + 2
    named = dict(gset, description="GRAPHIC SET OF 6")
    cs.set_pieces_for([named], "")
    assert named["set_pieces"] == {"count": 6, "basis": "the line's own words ('set of 6')"}


def test_a_count_the_brief_states_outranks_the_reads_and_the_read_is_disclosed():
    ans = _fsdu()                                               # the read gives HEADER x6
    parts = cs.parts_from_concept(ans, "J")
    res = cs.apply_brief_counts(parts, "Parts per unit: rear graphic panel x1; headers x3 "
                                       "(metal, to take a magnetic graphic); corner posts x6.")
    header = [p for p in parts if p["description"] == "HEADER"][0]
    assert header["quantity"] == 3 and res["applied"] == ["J-CPT04: 6 → 3 ('headers x3')"]
    from source_precedence import source_of
    assert source_of(header, "quantity") == cs.BRIEF_SOURCE
    assert any(f.startswith("BRIEF COUNT OVER THE READ: the enquiry brief states 'headers x3' — "
                            "costed at 3; the read gave 6") for f in header["review_flags"])
    post = [p for p in parts if p["description"] == "CORNER POST"][0]
    assert post["quantity"] == 6 and not any("BRIEF COUNT" in f for f in post["review_flags"])


def test_handed_lines_total_a_stated_count_and_a_split_that_does_not_is_asked():
    parts = cs.parts_from_concept(_fsdu(), "J")                 # side trays 5 left + 5 right
    assert cs.apply_brief_counts(parts, "10 side shelf trays (5 left + 5 right, mirrored)") == {
        "applied": [], "asked": []}
    parts2 = cs.parts_from_concept(_fsdu(), "J")
    res = cs.apply_brief_counts(parts2, "12 side shelf trays.")
    assert res["applied"] == [] and len(res["asked"]) == 1
    left = [p for p in parts2 if p["description"] == "SIDE SHELF TRAY (LEFT)"][0]
    assert left["quantity"] == 5 and "states '12 side shelf trays'" in str(
        left["manufacturing_questions"])


def test_a_per_host_or_each_sentence_is_not_a_parts_count():
    stated = cs.brief_stated_counts(
        "Each shelf is a folded steel tray with a front lip, with two hooked brackets welded to "
        "the tray. Make list: 6 slotted corner posts (31.75 mm slot pitch); 5 front shelf trays "
        "and 10 side shelf trays, each tray with its hooked brackets. Footprint 500 x 500 mm; "
        "200 off; breaks 1 and 10.")
    said = {s["said"] for s in stated}
    assert {"6 slotted corner posts", "5 front shelf trays", "10 side shelf trays"} <= said
    assert not any("bracket" in s for s in said)                  # per tray, not a count
    assert not any(s.startswith(("31", "500", "200", "75")) for s in said)


def test_an_estimators_count_stands_above_the_brief():
    import estimator_confirmed as ec
    parts = cs.parts_from_concept(_fsdu(), "J")
    header = [p for p in parts if p["description"] == "HEADER"][0]
    ec.apply_estimator_confirmed(parts, {"parts": {header["part_number"]: {"quantity": 4}}})
    cs.apply_brief_counts(parts, "headers x3")
    assert header["quantity"] == 4


def test_a_count_no_source_states_that_moved_between_reads_is_asked_on_its_line():
    parts = cs.parts_from_concept(_fsdu(), "J")
    asked = cs.question_unstated_count_changes(
        parts, ["HEADER: count 2 → 6 (no reason given)"], brief="metal header and rear panel")
    header = [p for p in parts if p["description"] == "HEADER"][0]
    assert asked == [f"{header['part_number']}: 2 → 6"]
    (q,) = header["manufacturing_questions"]
    assert "count 2 → 6 across reads, and neither the enquiry brief nor the pack states it" in \
        q["issue"] and "costed at this read's 6" in q["assumption"]
    # Stated by the brief, it is not a drift question — the brief decides it.
    parts2 = cs.parts_from_concept(_fsdu(), "J")
    assert cs.question_unstated_count_changes(
        parts2, ["HEADER: count 2 → 6 (no reason given)"], brief="headers x3") == []


def test_file_scan_settles_counts_before_anything_counts_from_the_lines():
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    a = src.index("concept_scan.apply_brief_counts(_sighted, _brief")
    assert src.index("_ec2.apply_estimator_confirmed(_sighted") < a
    assert a < src.index("concept_scan.allow_components_inside(_sighted")
    assert a < src.index("concept_scan.unit_assembly_part(_sighted, _answer, _job_name)")
    assert "concept_scan.question_unstated_count_changes(_sighted, _changes" in src
    assert "concept_scan.set_pieces_for(_sighted, _brief" in src


def test_an_earlier_reads_size_that_does_not_fit_its_host_is_not_used():
    parts = cs.parts_from_concept(_folded("FRONT SHELF TRAY"), "J")
    too_big = dict(_PREV_BRACKET, length=900, width=40)          # longer than the 460 mm tray
    (cpa,) = cs.allow_components_inside(parts, {"parts": [too_big]}, "J")["minted"]
    assert "house_allowance" in cpa["provisional_allowance"]
    assert "gave 900 × 40 mm, which does not fit its host (460 × 300 mm) — not used" in \
        cpa["review_flags"][-1]


def test_a_stated_count_beats_the_remembered_one_and_a_total_is_not_multiplied():
    parts = cs.parts_from_concept(_folded("FRONT SHELF TRAY"), "J")       # 5 trays
    (cpa,) = cs.allow_components_inside(parts, {"parts": [_PREV_BRACKET]}, "J",
                                        brief="each tray with its two hooked brackets")["minted"]
    assert cpa["quantity"] == 10 and (cpa["blank_length_mm"], cpa["blank_width_mm"]) == (40, 30)
    assert "the previous read's count 30 is not used: 2 per host" in cpa["review_flags"][-1]
    parts2 = cs.parts_from_concept(_folded("FRONT SHELF TRAY"), "J")
    (cpa2,) = cs.allow_components_inside(parts2, {}, "J",
                                         brief="Make list: 12 hooked brackets.")["minted"]
    assert cpa2["quantity"] == 12 and "12 in all from the enquiry brief" in cpa2["review_flags"][-1]


def test_the_adhesive_method_follows_the_words_and_the_shape():
    ans = _fsdu()
    for p in ans["parts"]:
        if "STRIP" in p["name"]:
            p["name"] = "35MM SHELF EDGE STRIP (bonded)"              # narrow: tape
        if p["name"] == "HEADER":
            p["name"] = "HEADER (glued)"                              # said: structural
        if p["name"] == "TOP CAP":
            p["name"] = "TOP CAP (bonded)"                            # wide, unsaid: structural
    lines = cs.adhesive_allowance_lines(cs.parts_from_concept(ans, "J"), "J", 3)
    by = {l["joining_method"]: l for l in lines}
    assert set(by) == {"tape", "structural adhesive"}
    assert by["tape"]["quantity"] == 5.3 and by["tape"]["unit_of_measure"] == "m"
    glue = by["structural adhesive"]
    assert glue["unit_of_measure"] == "each"
    flag = glue["review_flags"][-1]
    assert "('glued' in the read's words; round its perimeter)" in flag
    assert "(only 'bonded' said — PROVISIONAL, wider than a strip; round its perimeter)" in flag
    assert "a PROVISIONAL method, not a construction rule" in flag
    per = config.CONCEPT_ADHESIVE_ALLOWANCE["methods"]["structural adhesive"]["metres_per_unit"]
    cap = [p for p in cs.parts_from_concept(ans, "J") if p["description"] == "TOP CAP (bonded)"][0]
    hdr = [p for p in cs.parts_from_concept(ans, "J") if p["description"] == "HEADER (glued)"][0]
    metres = (2 * (cap["blank_length_mm"] + cap["blank_width_mm"])
              + 6 * 2 * (hdr["blank_length_mm"] + hdr["blank_width_mm"])) / 1000.0
    assert glue["quantity"] == round(metres / per, 2)
    assert [l["part_number"] for l in lines] == ["J-NV03", "J-NV04"]


def test_the_d502_vocabulary_lives_in_config_and_holds_no_job():
    rule = config.CONCEPT_ADHESIVE_ALLOWANCE
    assert {"methods", "strip_max_width_mm", "strip_method", "default_method"} <= set(rule)
    assert rule["strip_method"] in rule["methods"] and rule["default_method"] in rule["methods"]
    blob = repr([rule, config.CONCEPT_JOINED_COMPONENT_ALLOWANCE]).lower()
    for job_word in ("boots", "wellness", "kwellness"):
        assert job_word not in blob



# ── D-503: the reviewer's tightening of D-502 — prices to the tenth of a penny, counts that
#    keep their history, a set extrapolated position by position, the host's joint proven,
#    and an adhesive method that stays provisional ────────────────────────────────────────

def test_a_sub_penny_buy_price_reaches_the_sheet_unrounded(monkeypatch):
    import estimator as E
    import wb_populate as W
    monkeypatch.setattr(E, "_rung4_researcher", lambda brief: _label_answer(
        price_gbp=0.018, item_priced="20x20mm QR component label, one label"))
    pe = E.estimate_part(dict(_LABEL, review_flags=[]), job_quantity=200)
    sc = pe["cost_breakdown"]["system_cost"]
    assert sc["unit_cost_gbp"] == 0.018 and sc["applied_to_total"] is True
    assert W._bom_line_price(pe) == 0.018                    # was 0.02 — 11% over at 48 a unit


def test_a_count_that_moved_stays_a_question_on_a_replay(tmp_path, monkeypatch):
    monkeypatch.setattr(cs, "_cache_dir", lambda: tmp_path / "engine" / "cache" / "vision_concept")
    job = tmp_path / "staged" / "Client" / "JOB"
    job.mkdir(parents=True)
    a2, a6 = _fsdu(), _fsdu()
    for p in a2["parts"]:
        if p["name"] == "HEADER":
            p["quantity"] = 2
    cs.save_bill(job, "JOB", a2, read_key="r1")
    cs.save_bill(job, "JOB", a6, read_key="r2")
    cs.save_bill(job, "JOB", a6, read_key="r2")              # the same read, replayed
    prev = cs.load_previous_bill(job, "JOB")
    (row,) = [r for r in prev["parts"] if r["name"] == "HEADER"]
    assert row["counts_seen"] == [2.0, 6.0]
    parts = cs.parts_from_concept(a6, "J")
    asked = cs.question_unstated_count_changes(parts, [], "metal header", previous=prev)
    header = [p for p in parts if p["description"] == "HEADER"][0]
    assert asked == [f"{header['part_number']}: 2 → 6"]       # no change this read; still asked
    # A count the brief states is not asked — the brief decides it.
    parts2 = cs.parts_from_concept(a6, "J")
    assert cs.question_unstated_count_changes(parts2, [], "headers x3", previous=prev) == []
    src = (ROOT / "src" / "file_scan.py").read_text(encoding="utf-8")
    assert "previous=_prev_bill)" in src[src.index("question_unstated_count_changes(_sighted"):][:400]


def test_a_set_allowance_is_extrapolated_position_by_position_and_asks_its_contents(monkeypatch):
    import estimator as E
    ans = {"product": {"name": "FSDU"}, "parts": [
        _p("TOP CAP", 1, 560, 560, LF, 4),
        _p("REAR GRAPHIC PANEL", 1, 1400, 460, LF),
        _p("BASE GRAPHIC HOLDER", 2, 500, 80, LF, 1),
        {"name": "REPLACEABLE CARD / MAGNETIC GRAPHIC SET", "kind": "graphic", "quantity": 1}]}
    parts = cs.parts_from_concept(ans, "J")
    cs.set_pieces_for(parts, "top cap with graphic pass-through cut-out")
    gset = [p for p in parts if p["description"].endswith("SET")][0]
    pos = {p["line"]: p for p in gset["set_pieces"]["positions"]}
    assert pos["REAR GRAPHIC PANEL"]["size_mm"] == (1400.0, 460.0)
    assert pos["TOP CAP"]["size_mm"] is None and pos["TOP CAP"]["by"] == "the brief"
    (q,) = gset["manufacturing_questions"]
    assert "which graphics the set contains is not stated" in q["issue"]
    assert "A holder holds a graphic but is not one" in q["assumption"]
    monkeypatch.setattr(E, "_rung4_researcher", _set_research([]))
    part = dict(gset, assumed_job_quantity=200, review_flags=[], is_bought_in=True)
    res = E._resolve_part_system_cost(part)
    rate = 18.00 / (1.4 * 0.46)                                    # £/m² of the priced panel
    expected = 18.00 + rate * 1.4 * 0.46 + 2 * rate * 0.5 * 0.08   # cap at piece price + sized
    assert abs(res["applied_unit_cost"] - round(expected, 4)) < 1e-6
    (flag,) = [f for f in part["review_flags"] if f.startswith("PROVISIONAL SET ALLOWANCE")]
    assert "an extrapolation from ONE piece, not a set price" in flag
    assert "£27.95/m² over its 1400 × 460 mm" in flag
    assert "TOP CAP ×1, size not known (the brief) — at the priced piece's £18.00" in flag
    assert "BASE GRAPHIC HOLDER ×2, 500 × 80 mm (its face, an upper bound)" in flag


def test_the_joint_left_on_its_hosts_is_proven_charged_there_or_asked():
    import estimator as E
    a = _folded("FRONT SHELF TRAY")
    for p in a["parts"]:
        if p["name"].startswith("FRONT SHELF TRAY"):
            p["operations"] = p["operations"] + ["welding"]
    parts = cs.parts_from_concept(a, "J")
    res = cs.allow_components_inside(parts, {"parts": [_PREV_BRACKET]}, "J",
                                     brief="each tray with its two hooked brackets")
    parts += res["minted"]
    cpa = res["minted"][0]
    assert cpa["provisional_allowance"]["per_host_piece"] == 2.0
    pes = [E.estimate_part(dict(p), job_quantity=200) for p in parts]
    pe_cpa = [p for p in pes if p["part_number"] == cpa["part_number"]][0]
    assert pe_cpa["provisional_allowance"]["joint_on"] == ["J-CPT05"]
    (line,) = E.prove_joint_on_hosts(parts, pes, None)
    assert "JOINT ON THE HOST, PROVEN" in line and "min per bracket joint (2 a piece)" in line
    assert any("JOINT ON THE HOST, PROVEN" in f for f in pe_cpa["review_flags"])
    # The host's weld ruled out of the route: the joint costs nothing — a counted question.
    parts2 = [dict(p, review_flags=list(p.get("review_flags") or []),
                   manufacturing_questions=[]) for p in parts]
    route = {"decisions": [
        {"operation": "welding", "status": "ruled_out", "participants": ["J-CPT05"]},
        {"operation": "dress_welds", "status": "required", "participants": ["J-CPT05"]}]}
    (bad,) = E.prove_joint_on_hosts(parts2, pes, route)
    assert "JOINT NOT CHARGED ON ITS HOST" in bad and "J-CPT05 weld is not charged" in bad
    raw = [p for p in parts2 if p["part_number"] == cpa["part_number"]][0]
    assert any("JOINT NOT CHARGED" in q["issue"] for q in raw["manufacturing_questions"])
    src = (ROOT / "src" / "estimator.py").read_text(encoding="utf-8")
    assert "prove_joint_on_hosts(parts, part_estimates, canonical_route_shadow)" in src


def test_an_unstated_bonding_method_is_provisional_and_asked_with_material_and_load():
    ans = _fsdu()
    for p in ans["parts"]:
        if "STRIP" in p["name"]:
            p["name"] = "35MM SHELF EDGE STRIP (bonded)"
        if p["name"] == "FRONT SHELF TRAY":
            p["name"] = "FRONT SHELF TRAY (bonded)"
        if p["name"] == "HEADER":
            p["name"] = "HEADER (taped)"
    ans["parts"].append(dict(_p("ACRYLIC FACIA (bonded)", 1, 400, 200, ["laser_cutting"]),
                             material_guess="ACRYLIC"))
    parts = cs.parts_from_concept(ans, "J")
    lines = {l["joining_method"]: l for l in cs.adhesive_allowance_lines(parts, "J", 3)}
    tray = [p for p in parts if p["description"] == "FRONT SHELF TRAY (bonded)"][0]
    (q,) = tray["manufacturing_questions"]
    assert "bonded, and the joining method is not stated" in q["issue"]
    assert "LOAD-BEARING by its name" in q["assumption"]
    assert "provisionally structural adhesive" in q["assumption"]
    facia = [p for p in parts if p["description"] == "ACRYLIC FACIA (bonded)"][0]
    assert facia["normalized_material"].upper().startswith("ACRYLIC")
    assert "acrylic is usually taped — PROVISIONAL" in lines["tape"]["review_flags"][-1]
    assert "not load-bearing by its name" in facia["manufacturing_questions"][0]["assumption"]
    header = [p for p in parts if p["description"] == "HEADER (taped)"][0]
    assert not header.get("manufacturing_questions")            # the read stated the method
