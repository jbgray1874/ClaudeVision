"""The explanation belongs in the workbook, and must never cost a run to put it there.

A document that travels beside a spreadsheet arrives without it — forwarded on its own, saved
somewhere else, stale the moment a rate is edited. So it is written into the file as a tab.
Everything about that is a nicety next to the estimate itself, which is why the writer's most
important property is that it cannot break a run.
"""
from __future__ import annotations

import ast
import sys
from pathlib import Path

import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(SRC))

import estimate_explanation_tab                                         # noqa: E402


def test_a_file_that_is_not_a_workbook_declines_rather_than_raising(tmp_path):
    """A corrupt or half-written file must cost a printed line, never the run."""
    book = tmp_path / "12552-00.xlsx"
    book.write_bytes(b"not really a workbook")
    assert estimate_explanation_tab.write_tab(book) is None


def test_a_missing_workbook_is_reported_not_raised(tmp_path):
    assert estimate_explanation_tab.write_tab(tmp_path / "nothing.xlsx") is None


def _function(source: str, name: str) -> ast.FunctionDef:
    tree = ast.parse(source)
    return next(n for n in ast.walk(tree)
                if isinstance(n, ast.FunctionDef) and n.name == name)


def test_every_step_that_can_fail_is_caught():
    """Structurally, not by reading the prose — a comment saying so matches a grep and a
    missing try does not."""
    fn = _function((SRC / "estimate_explanation_tab.py").read_text(encoding="utf-8"),
                   "write_tab")
    handlers = [n for n in ast.walk(fn) if isinstance(n, ast.Try)]
    assert len(handlers) >= 2, (
        "building the document and writing the workbook are separate failures with "
        "separate reasons, and neither may reach the caller")
    assert any(isinstance(n, ast.Try) and n.finalbody for n in handlers), (
        "the workbook handle must be closed in a finally — the very next step of the run "
        "reopens this same file to add AI Provenance")


def test_the_run_never_stops_for_a_missing_tab():
    """main.py calls it inside its own try, so even an import error is survivable."""
    source = (SRC / "main.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    calls = [n for n in ast.walk(tree)
             if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)
             and n.func.id == "_write_explanation_tab"]
    assert calls, "the tab is written from main.py after the read-back"

    guarded = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Try):
            for call in calls:
                if any(call is inner for inner in ast.walk(node)):
                    guarded.append(call)
    assert len(guarded) == len(calls), (
        "an estimate that has taken an hour must not be lost to a tab that would not write")


def test_it_runs_after_the_read_back_and_not_before():
    """The tab prints Estimate!M and the sheet's own totals. Neither exists until Excel has
    calculated the template and the read-back has recorded what it found."""
    source = (SRC / "main.py").read_text(encoding="utf-8")
    assert source.index("stamp_real_totals_into_json") < source.index("_write_explanation_tab")


@pytest.mark.parametrize("name", ["build", "sections", "worksheet_rows"])
def test_the_tab_renders_the_document_rather_than_asking_again(name):
    """One builder. A tab that goes and reads the workbook itself is a second answer, and two
    answers to one question is the failure this document exists to expose."""
    source = (SRC / "estimate_explanation_tab.py").read_text(encoding="utf-8")
    assert f"estimate_explained.{name}" in source


def test_the_tab_wears_the_provenance_tabs_clothes(tmp_path):
    """One product, one look. Tim's reviewer read the flat COM rendering: "the rendered
    action table clips descriptions, assumptions and instructions." The tab now uses AI
    Provenance's own palette — banner, section fills, wrapped bordered tables, computed
    row heights — and a grid that fits a screen instead of seven columns at the cap."""
    import json
    import openpyxl
    from test_one_costed_record_for_every_deliverable import seventy_three_thirty_two

    job = seventy_three_thirty_two()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estimate"
    ws["D6"] = 6
    ws["C9"] = "Bill of Materials (Per Unit)"
    ws["H9"] = "Part code"
    r = 10
    for m in job["final_estimate"]["material_rows"]:
        if m["block"] != "bom":
            continue
        ws.cell(r, 3, m["description"])
        ws.cell(r, 8, m["part_code"])
        ws.cell(r, 10, m["unit_price_gbp"])
        ws.cell(r, 11, m["qty_per_unit"])
        r += 1
    xlsx = tmp_path / "7332-01_style.xlsx"
    wb.save(xlsx)
    jp = tmp_path / "7332-01.json"
    jp.write_text(json.dumps(job), encoding="utf-8")

    assert estimate_explanation_tab.write_tab(xlsx, jp) == "AI Explanation"
    out = openpyxl.load_workbook(xlsx)["AI Explanation"]

    title = out["A1"]
    assert title.value.startswith("SDI Intelligence")
    assert title.fill.fgColor.rgb.endswith("1F3864"), "the Provenance navy banner"
    section_rows = [c.row for row in out.iter_rows(max_col=1) for c in row
                    if c.fill and c.fill.fgColor and c.font.bold
                    and str(c.fill.fgColor.rgb).endswith("2F5496")
                    and str(c.value or "").isupper()]
    assert section_rows, "section titles carry the Provenance section blue"

    # No clipping: long table cells wrap, and their rows grow to hold the text.
    grown = 0
    for row in out.iter_rows(min_row=2):
        for c in row:
            if c.value and len(str(c.value)) > 60 and c.column <= 8:
                assert c.alignment.wrap_text, f"{c.coordinate} would clip"
                if (out.row_dimensions[c.row].height or 0) > 20:
                    grown += 1
    assert grown, "no row grew for its text — heights are not being computed"

    total_width = sum(d.width or 0 for d in out.column_dimensions.values())
    assert total_width <= 230, f"the grid is {total_width:.0f} units wide — a scroll, not a tab"
