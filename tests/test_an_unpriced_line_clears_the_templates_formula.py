"""An unpriced BOM line must leave a BLANK price cell, not the template's stale formula.

The 17:36 7332-01 book (16 Sep 2026) carried, on its plating line:

    J20  =LOOKUP($D$6,'Material Price Break'!$D$4:$N$4,'Material Price Break'!#REF!)

and on its freight line a LOOKUP misrouted into break row 45. Neither was written by the
engine: they are the blank template's own formulas, and the engine believed it had cleared
them. It had called ws.cell(row, col, value=None) — which in openpyxl is a READ. The value
argument is assigned only when it is not None, so an unpriced line kept whatever the
template held, and the broken reference propagated through M into the block total.

PACKAGING and DELIVERY escaped only because they write a literal 0.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

from openpyxl import Workbook                                          # noqa: E402

import wb_populate as wb                                               # noqa: E402

BROKEN = "=LOOKUP($D$6,'Material Price Break'!$D$4:$N$4,'Material Price Break'!#REF!)"


def test_openpyxl_does_not_clear_on_none_which_is_why_the_helper_exists():
    ws = Workbook().active
    ws.cell(row=20, column=10, value=BROKEN)
    ws.cell(row=20, column=10, value=None)            # the call the writer used to make
    assert ws.cell(row=20, column=10).value == BROKEN, "a read, not a write"


def test_the_helper_clears_the_templates_formula_for_an_unpriced_line():
    ws = Workbook().active
    ws.cell(row=20, column=10, value=BROKEN)
    wb._clear_or_set(ws, 20, 10, None)
    assert ws.cell(row=20, column=10).value is None


def test_the_helper_still_writes_a_price():
    ws = Workbook().active
    ws.cell(row=17, column=10, value=BROKEN)
    wb._clear_or_set(ws, 17, 10, 5.86)
    assert ws.cell(row=17, column=10).value == 5.86


def test_the_bom_writer_uses_the_helper_for_the_price_cell():
    """The one place the defect lived. Named so a refactor that reverts to ws.cell(...,
    value=None) fails here with the history rather than on a customer's book."""
    src = open(os.path.join(os.path.dirname(__file__), "..", "src", "wb_populate.py"),
               encoding="utf-8").read()
    assert '_clear_or_set(ws, row, b["col_price"], price)' in src
    assert 'value=price if price is not None else None' not in src
