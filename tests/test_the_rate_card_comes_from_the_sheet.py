r"""The rate card comes from the sheet, and it says so out loud.

James Gray, 18 September 2026 — the steel ruling, which is about labour too:

    "the spreadsheet rate is the controlling rate. The engine should read the ... input from
     the workbook, use that same figure in its calculation/report, and let the estimator
     amend it in the sheet when required. It should not independently substitute ... a
     config fallback."

and, 17 September, on where settings live:

    "config needs to be in config files. not json files lying around and being copied
     manually around."

EVERY DEPARTMENT RATE IN THE SHOP LIVED IN A GITIGNORED FILE. `rate_card_ingest` reads
the estimate template's own "Labour / Rate / Dept" block — the estimator's card, the one they
amend — and wrote it to a JSON BESIDE config.py, which `src/.gitignore`'s
blanket `*.json` rule has always excluded. config then overlaid it at import. So:

  * the numbers that price every route were outside version control
  * two machines could hold different cards and produce different money from one commit,
    with nothing on either book to compare
  * it was SILENT — the loaded-flag was assigned and read by NOTHING, so a run with
    a card and a run without looked identical on the page
  * and the overlay wrote INSIDE its loop, so a throw part-way left a rate card HALF
    APPLIED: some departments on this year's numbers and some on last year's

That last one is the reason this is not merely tidying. A half-applied card is the one
outcome nobody would choose and the one nothing would have reported.

The open £355.43-vs-£331.42 question is the same fault wearing a different hat: `_rate_table_find`
records a completed estimate booking P.Coat at 331.42 under dept POWDER while the engine
booked 355.43 under P/C. Neither of those settles anything -- a figure copied off one
estimator's sheet cannot validate or replace the current template rate -- and the answer was
always in the live template, which nothing read.
"""
from __future__ import annotations

import importlib
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import config  # noqa: E402


def _template(tmp_path, rows, header_row=10, sheet="Estimate"):
    """A workbook shaped like the estimate sheet's dept block: H/I/J/K."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws[f"H{header_row}"], ws[f"I{header_row}"] = "Labour", "Rate"
    ws[f"J{header_row}"], ws[f"K{header_row}"] = "Dept", "Setup"
    for n, (label, rate, dept, setup) in enumerate(rows, start=header_row + 1):
        ws[f"H{n}"], ws[f"I{n}"], ws[f"J{n}"], ws[f"K{n}"] = label, rate, dept, setup
    path = tmp_path / "Blank Estimate Sheet 2026.xlsx"
    wb.save(path)
    return path


_ORDINARY = [("Laser (Metal)", 71.50, "LASM", 10),
             ("Fold", 42.00, "FOLD", 15),
             ("P.Coat", 331.42, "POWDER", 15)]


@pytest.fixture
def card_from(tmp_path, monkeypatch):
    def _read(rows, **kw):
        monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE",
                            _template(tmp_path, rows, **kw))
        return config.read_template_rate_card()
    return _read


# ── the sheet's own block is the rate card ───────────────────────────────────────────

def test_the_templates_labour_block_is_read(card_from):
    card, source = card_from(_ORDINARY)
    assert card["laser_cutting"] == 71.50
    assert card["folding"] == 42.00
    assert "labour block" in source


def test_the_open_powder_question_is_answered_by_the_sheet(card_from):
    """£355.43 (P/C) against £331.42 (POWDER) was an open question, on a rate worth about
    sixty per cent of 401912-02's unit cost. It is not settled by either figure: it is settled
    by the TEMPLATE, which is the sheet an estimator amends. Whichever way the template's
    figure differs from the built-in default, the template's is taken -- that is the ruling."""
    card, _ = card_from(_ORDINARY)
    assert card["powder_coating"] == 331.42
    assert card["powder_coating"] != 355.43, "the config default is still deciding"


def test_the_block_is_found_wherever_it_sits(card_from):
    """The header is located, not assumed at a fixed row — a template whose block has moved
    is still read, and the alternative is a rate card that silently empties after an edit."""
    card, _ = card_from(_ORDINARY, header_row=40)
    assert card["folding"] == 42.00


def test_reading_stops_where_the_labour_block_stops(card_from):
    """The wire price-break rows below the block have a NUMBER in the dept column. Reading
    past the block turns a price break into an hourly rate."""
    card, _ = card_from(_ORDINARY + [("Wire break qty", 1.23, 5, None),
                                     ("Fold", 999.0, "FOLD", 1)])
    assert card["folding"] == 42.00, "read past the end of the labour block"


# ── and never a partial one ──────────────────────────────────────────────────────────

def test_an_unreadable_template_yields_nothing_rather_than_something(card_from, monkeypatch, tmp_path):
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE", tmp_path / "absent.xlsx")
    card, source = config.read_template_rate_card()
    assert card == {}
    assert "could not be opened" in source


def test_a_template_with_no_block_says_so_instead_of_guessing(card_from, tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    wb.active.title = "Estimate"
    path = tmp_path / "Blank Estimate Sheet 2026.xlsx"
    wb.save(path)
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE", path)
    card, source = config.read_template_rate_card()
    assert card == {}
    assert "no 'Labour / Rate / Dept' block" in source


def test_a_rate_nobody_could_mean_is_refused(card_from):
    """A decimal in the wrong place, or the setup column read by mistake. £355/hr is a real
    line rate, so the ceiling is generous and only catches nonsense."""
    card, _ = card_from([("Fold", 0.0, "FOLD", 15), ("Laser (Metal)", 71.5, "LASM", 10),
                         ("P.Coat", 99_000.0, "POWDER", 15)])
    assert "folding" not in card
    assert "powder_coating" not in card
    assert card["laser_cutting"] == 71.5


def test_a_label_the_engine_does_not_cost_is_ignored_quietly(card_from):
    card, _ = card_from(_ORDINARY + [("Sundry allowance", 12.0, "SUND", 0)])
    assert set(card) == {"laser_cutting", "folding", "powder_coating"}


# ── one mapping, one home ────────────────────────────────────────────────────────────

def test_the_ingester_and_the_engine_share_one_label_mapping():
    """Two copies of this mapping is two answers to 'what does p.coat cost' waiting to
    happen. The ingester imports config's."""
    import rate_card_ingest
    assert rate_card_ingest.LABEL_TO_OP is config.ESTIMATE_LABOUR_LABEL_TO_OP


def test_every_rate_the_engine_holds_can_name_its_source():
    for op in config.HOURLY_RATES_GBP:
        rate, source = config.hourly_rate(op)
        assert rate is not None
        assert source and source != "not in the rate card", op


def test_an_operation_with_no_rate_says_that_rather_than_inventing_one():
    rate, source = config.hourly_rate("teleportation")
    assert rate is None
    assert source == "not in the rate card"


# ── and the run says which card priced it ────────────────────────────────────────────

def test_the_run_states_where_its_rates_came_from():
    """The loaded-flag was assigned and read by nothing. A silent rate card is how
    two machines price one job differently and neither book can be blamed."""
    assert config.RATE_CARD_NOTES, "the run has nothing to say about its own rates"
    assert any("rate" in n.lower() for n in config.RATE_CARD_NOTES)


def test_the_banner_prints_the_rate_provenance(capsys, monkeypatch):
    import build_stamp
    monkeypatch.setattr(config, "RATE_CARD_NOTES",
                        ["7 department rate(s) read from the estimating template"])
    build_stamp.print_build_stamp()
    assert "[rates]" in capsys.readouterr().out


def test_a_disagreement_between_the_file_and_the_sheet_is_reported(tmp_path, monkeypatch):
    """The ingested file may have come off a DIFFERENT estimate. It does not overrule the
    sheet this job is costed on — and where the two differ, one of them is out of date and
    that is worth a sentence rather than a silent win."""
    import json
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE",
                        _template(tmp_path, _ORDINARY))
    monkeypatch.setattr(config, "RATE_CARD_NOTES", [])
    monkeypatch.setattr(config, "HOURLY_RATE_SOURCE", dict(config.HOURLY_RATE_SOURCE))
    monkeypatch.setattr(config, "HOURLY_RATES_GBP", dict(config.HOURLY_RATES_GBP))
    rc = tmp_path / "rate_card.json"
    rc.write_text(json.dumps({"by_op": {"folding": 38.0, "oven": 26.5}, "source": "1310"}))
    monkeypatch.setattr(config, "INGESTED_RATE_CARD_PATH", str(rc))
    config._apply_rate_cards()
    said = " ".join(config.RATE_CARD_NOTES)
    assert "DISAGREES" in said and "folding" in said
    assert config.HOURLY_RATES_GBP["folding"] == 42.00, "the file overruled the sheet"
    assert config.HOURLY_RATES_GBP["oven"] == 26.5, "the file was ignored where the sheet was silent"


# ── named for what it is, not for whose sheet it first read ─────────────────────────
#
# James Gray, 18 September 2026: "this sheet is dave not tim. we have a team of four
# estimators." The whole mechanism was named after one of them -- the file, the ingester,
# the loaded-flag. A shared office rate table named after one estimator is read as that
# person's personal figures: over-trusted by some, dismissed by others, and argued about by
# anybody who knows it is not their sheet. Per-figure attribution belongs in
# SHOP_STATED_PROVENANCE, where a name is attached to a ruling somebody actually made.

def test_no_estimator_name_is_baked_into_the_rate_card_mechanism():
    import inspect
    import rate_card_ingest
    for mod in (config, rate_card_ingest):
        src = inspect.getsource(mod)
        # The history may be recounted; the machinery may not carry the name.
        code = "\n".join(ln for ln in src.split("\n")
                          if not ln.lstrip().startswith("#") and '"""' not in ln)
        assert "TIM_RATE_CARD" not in code, f"{mod.__name__} still names an estimator"


def test_the_card_is_looked_for_under_its_own_name():
    assert config.INGESTED_RATE_CARD_PATH.endswith("rate_card.json")
    assert not config.INGESTED_RATE_CARD_PATH.endswith("tim_rate_card.json")


def test_a_runner_holding_the_old_filename_keeps_its_rates(tmp_path, monkeypatch):
    """A rename must not silently change a price. The old name is still read where it is the
    only one there -- and using it is SAID, so the legacy name cannot quietly outlive this."""
    import json
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE", tmp_path / "absent.xlsx")
    monkeypatch.setattr(config, "RATE_CARD_NOTES", [])
    monkeypatch.setattr(config, "HOURLY_RATES_GBP", dict(config.HOURLY_RATES_GBP))
    monkeypatch.setattr(config, "HOURLY_RATE_SOURCE", dict(config.HOURLY_RATE_SOURCE))
    monkeypatch.setattr(config, "INGESTED_RATE_CARD_PATH", str(tmp_path / "rate_card.json"))
    legacy = tmp_path / "tim_rate_card.json"
    legacy.write_text(json.dumps({"by_op": {"oven": 27.5}, "source": "an older estimate"}))
    monkeypatch.setattr(config, "LEGACY_RATE_CARD_PATH", str(legacy))
    config._apply_rate_cards()
    assert config.HOURLY_RATES_GBP["oven"] == 27.5
    assert any("legacy filename" in n for n in config.RATE_CARD_NOTES)


def test_the_new_name_is_preferred_where_both_exist(tmp_path, monkeypatch):
    import json
    monkeypatch.setattr(config, "AI_ESTIMATE_XLSX_TEMPLATE", tmp_path / "absent.xlsx")
    monkeypatch.setattr(config, "RATE_CARD_NOTES", [])
    monkeypatch.setattr(config, "HOURLY_RATES_GBP", dict(config.HOURLY_RATES_GBP))
    monkeypatch.setattr(config, "HOURLY_RATE_SOURCE", dict(config.HOURLY_RATE_SOURCE))
    current = tmp_path / "rate_card.json"
    current.write_text(json.dumps({"by_op": {"oven": 31.0}, "source": "current"}))
    legacy = tmp_path / "tim_rate_card.json"
    legacy.write_text(json.dumps({"by_op": {"oven": 27.5}, "source": "older"}))
    monkeypatch.setattr(config, "INGESTED_RATE_CARD_PATH", str(current))
    monkeypatch.setattr(config, "LEGACY_RATE_CARD_PATH", str(legacy))
    config._apply_rate_cards()
    assert config.HOURLY_RATES_GBP["oven"] == 31.0
    assert not any("legacy filename" in n for n in config.RATE_CARD_NOTES)
