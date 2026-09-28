"""Dave's pencil notes on 11650-06-GA's printed Estimate sheet (estimating, 28 Sep 2026).

  * "customer" — "Boots" and "11650-06-GA" were written over the template's own "Customer"
    and "Drawing No." labels in C3 and C5, with nothing in the value cells beside them.
  * "could the sub drawings come out of material and into labour" — the £0 "costed in Sheet
    Steel below" lines leave the bill of materials; each part is named where it is priced.
  * "2 columns" — the labour rows give the part numbers a column of their own.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import openpyxl  # noqa: E402

import config  # noqa: E402
import wb_populate as wp  # noqa: E402


def _template_top():
    """The 2026 template's header block and labour header, as the estimators laid it out."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C3"], ws["C4"], ws["C5"], ws["C6"], ws["C7"] = (
        "Customer", "Description", "Drawing No.", "Quantity", "Date")
    for rng in ("D3:G3", "D4:G4", "D5:E5", "D6:E6", "D7:E7"):
        ws.merge_cells(rng)
    ws["C125"], ws["D125"], ws["G125"], ws["H125"] = (
        "Operation", "Part Description", "Dept.", "Qty Per Unit")
    ws.merge_cells("D125:F125")
    ws.merge_cells("D126:F126")          # some labour rows merged D:F, some not
    return ws


LB = {"first_row": 126, "last_row": 197, "col_operation": 3, "col_desc": 4, "col_qty": 8}


def test_customer_and_drawing_number_go_beside_their_labels():
    ws = _template_top()
    assert wp.write_header_value_beside_label(ws, ("customer",), "D3", "Boots") == "D3"
    assert wp.write_header_value_beside_label(
        ws, ("drawing no.", "drawing no"), "D5", "11650-06-GA") == "D5"
    assert ws["C3"].value == "Customer" and ws["C5"].value == "Drawing No."
    assert ws["D3"].value == "Boots" and ws["D5"].value == "11650-06-GA"


def test_a_fallback_cell_holding_the_label_is_never_written_over():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A20"] = "Customer"                    # label outside the top rows: not found
    assert wp.write_header_value_beside_label(ws, ("customer",), "A20", "Boots") == ""
    assert ws["A20"].value == "Customer"


def test_the_header_map_names_value_cells_not_labels():
    assert wp.CELL_MAP["header"]["customer"] == "D3"
    assert wp.CELL_MAP["header"]["drawing_no"] == "D5"


def test_sub_drawings_are_off_the_bill_by_default():
    assert config.BOM_LISTS_PARTS_COSTED_IN_BLOCKS is False


def test_the_labour_header_gives_part_numbers_their_own_column():
    ws = _template_top()
    col = wp.labour_parts_column(ws, LB)
    assert col == 6
    assert ws.cell(125, 6).value == config.LABOUR_PARTS_COLUMN_LABEL
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "D125:E125" in merged and "D125:F125" not in merged
    assert ws["D125"].value == "Part Description"


def test_a_template_that_already_has_the_column_is_used_as_it_stands():
    ws = _template_top()
    ws.unmerge_cells("D125:F125")
    ws["E125"] = "Part No."
    assert wp.labour_parts_column(ws, LB) == 5
    assert ws["F125"].value is None


def test_a_merged_labour_row_is_narrowed_and_an_open_one_left_as_it_is():
    ws = _template_top()
    col = wp.labour_parts_column(ws, LB)
    c = wp.free_labour_parts_cell(ws, 126, 4, col)
    assert c is not None and c.coordinate == "F126"
    assert "D126:E126" in {str(m) for m in ws.merged_cells.ranges}
    c = wp.free_labour_parts_cell(ws, 127, 4, col)
    assert c is not None and c.coordinate == "F127"


def test_the_row_text_splits_into_description_and_parts():
    parts = ["11650-02-03M", "11650-03-01M"]
    rd = wp.labour_row_description("Fold", "MILD STEEL", 1.5, parts, bends=3)
    assert "11650-02-03M" in rd, "the combined text still names the parts"
    desc, pn = wp.split_labour_description(rd, "Fold", parts)
    assert pn == "11650-02-03M, 11650-03-01M"
    assert "11650" not in desc
    assert desc == "1.5mm MILD STEEL (3 bends)"


def test_more_than_six_parts_all_reach_the_column():
    parts = [f"P-{i}" for i in range(9)]
    rd = wp.labour_row_description("Assemble/pack (Metal)", "MILD STEEL", None, parts)
    desc, pn = wp.split_labour_description(rd, "Assemble/pack (Metal)", parts)
    assert pn.count("P-") == 9 and "more" not in desc


def test_no_label_configured_keeps_parts_in_the_text(monkeypatch):
    monkeypatch.setattr(config, "LABOUR_PARTS_COLUMN_LABEL", "")
    ws = _template_top()
    assert wp.labour_parts_column(ws, LB) is None
    assert "D125:F125" in {str(m) for m in ws.merged_cells.ranges}
