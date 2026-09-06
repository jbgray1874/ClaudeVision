"""The AI tabs read the same lines the Estimate sheet is built from.

The Estimate sheet is built from canonical_part_estimates — the list canonicalise produces,
which MINTS explicit bought-in BOM lines that never got a pricing record (11762-17's £1.20
class-word clip is one). The AI Material Detail and AI Price Provenance tabs annotate the
sheet's lines, so they must read that same list. They had been reading the pre-canonical
part_estimates, so the clip carried a price on the sheet and appeared on NEITHER tab — and
the report read the missing provenance row as "source not named" on a line that has a source.

This pins the writer half: a bought-in present only on canonical_part_estimates gets a row on
the AI Price Provenance tab, carrying its source.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import openpyxl  # noqa: E402
import wb_populate as wp  # noqa: E402


def _clip_pe():
    return {
        "part_number": "STD PART",
        "description": "PERFO PLASTIC LOCKING CLIP",
        "page_roles": ["bought_in"],
        "unit_cost_gbp": 1.20,
        "material_estimate": {"unit_material_cost_gbp": 1.20,
                              "cost_per_part_gbp": 1.20,
                              "cost_method": "standard_commodity_provisional"},
        "source": "standard_commodity_provisional",
        "review_flags": ["Provisional standard-commodity price — confirm against a supplier quote."],
    }


def _rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_a_canonical_only_boughtin_reaches_the_price_provenance_tab():
    wb = openpyxl.Workbook()
    summary = {"estimate_summary": {
        "part_estimates": [],                      # pre-canonical: clip absent
        "canonical_part_estimates": [_clip_pe()],  # canonical: clip present
    }}
    wp._append_ai_sheets(wb, summary, [])
    assert "AI Price Provenance" in wb.sheetnames
    rows = _rows(wb["AI Price Provenance"])
    parts = [str(r[0]) for r in rows[1:]]          # skip header
    assert "STD PART" in parts, f"clip missing from provenance tab: {parts}"
    # and the row names its source — the token, which the report then says in words
    clip = next(r for r in rows[1:] if str(r[0]) == "STD PART")
    header = rows[0]
    src = clip[header.index("Price Source")]
    assert "commodity" in str(src).lower()


def test_the_writer_falls_back_to_part_estimates_when_no_canonical_list():
    # the no-cutover / no-workbook path is unchanged: it still reads part_estimates
    wb = openpyxl.Workbook()
    summary = {"estimate_summary": {"part_estimates": [_clip_pe()]}}
    wp._append_ai_sheets(wb, summary, [])
    parts = [str(r[0]) for r in _rows(wb["AI Price Provenance"])[1:]]
    assert "STD PART" in parts
