"""A sheet that cannot be matched back to its job is a sheet somebody has to re-derive.

"Could you add description/date/drawing number at top please (easier to trace back when
requoting same job in future)" — the estimator, 12 Sep. The template already carries the
labels (C4 'Description', C7 'Date'); every pack went out with nothing beside them, and the
drawing-number box read "12349" on a job whose number is 12349-02, because the header write
took \\d+ of the folder name and stopped at the first hyphen.

The rules are the Rev box's rules: found by exact label (top rows only — 'Date' appears in
blocks further down the sheet), written only into an empty cell, values from the one
resolver that answers "what is this unit". A description that can say nothing better than
the drawing number is not written at all.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from wb_populate import full_drawing_number, write_job_identity_header   # noqa: E402


def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C4"] = "Description"
    ws["C7"] = "Date"
    ws["F7"] = "Prepared By"
    ws["C40"] = "Date"                     # a lower block's own label — must stay untouched
    return ws


def _summary(number="12349-02-69", title="RETAILER COUNTER UNIT"):
    return {"llm_full_extract": {"drawing_info": {"drawing_number": number, "title": title}}}


# ── the drawing number is the whole number ────────────────────────────────────────────────

def test_the_folder_name_keeps_its_number_groups():
    assert full_drawing_number({}, "12349-02 SolidWorks Pack") == "12349-02"


def test_the_title_block_beats_the_folder_name():
    assert full_drawing_number(_summary(), "12349-02 SolidWorks") == "12349-02-69"


def test_a_folder_with_no_number_is_not_mangled():
    assert full_drawing_number({}, "Boots Ladder Rack") == "Boots Ladder Rack"


# ── description and date land beside their own labels ────────────────────────────────────

def test_description_and_date_fill_their_labelled_cells():
    ws = _sheet()
    written = write_job_identity_header(ws, _summary(), "12349-02")
    assert sorted(written) == ["date", "description"]
    assert ws["D4"].value == "RETAILER COUNTER UNIT"
    assert re.fullmatch(r"\d{2}/\d{2}/\d{4}", str(ws["D7"].value))


def test_a_date_label_further_down_the_sheet_is_not_touched():
    ws = _sheet()
    write_job_identity_header(ws, _summary(), "12349-02")
    assert ws["D40"].value is None


def test_an_occupied_cell_is_never_displaced():
    ws = _sheet()
    ws["D4"] = "the estimator already wrote this"
    written = write_job_identity_header(ws, _summary(), "12349-02")
    assert "description" not in written
    assert ws["D4"].value == "the estimator already wrote this"


def test_a_description_that_is_only_the_number_is_not_written():
    """A number repeated under a Description label is noise wearing a label."""
    ws = _sheet()
    summary = {"llm_full_extract": {"drawing_info": {"drawing_number": "12349-02"}}}
    written = write_job_identity_header(ws, summary, "12349-02")
    assert "description" not in written
    assert ws["D4"].value is None


def test_a_sheet_with_no_labels_takes_no_writes():
    wb = openpyxl.Workbook()
    ws = wb.active
    assert write_job_identity_header(ws, _summary(), "12349-02") == []


# ── the unit's name, when the job has more than one root ─────────────────────────────────
# 12349-02's Description box came out EMPTY on the 13 Sep pack, on a job whose own provenance
# tab prints "GRAVITY FEEDER MODULES 3 A GRAVITY FEEDERS" against the node the graph calls
# the top. The compiler leaves `top_assembly` blank whenever a job has more than one root —
# it says so where it emits the field — and the resolver read only that one field, so the
# title fell through to the drawing NUMBER, which this module then correctly refuses to write
# under a Description label. The forest is `top_assemblies`.

def _shadow(roots, top="", nodes=()):
    return {"llm_full_extract": {"drawing_info": {"drawing_number": "12349-02"}},
            "estimate_summary": {"canonical_route_shadow": {
                "top_assembly": top, "top_assemblies": list(roots),
                "nodes": [{"part_number": p, "description": d} for p, d in nodes]}}}


def test_the_outermost_root_names_the_unit():
    """Several roots ship; the drawing number owns 12349-02-69, and -69 is outside -69-100."""
    from client_quote_html import _drawing_identity
    s = _shadow(["12349-02-69", "12349-02-69-100"],
                nodes=[("12349-02-69", "GRAVITY FEEDER MODULES"),
                       ("12349-02-69-100", "GRAVITY FEEDERS, 9 WIDE")])
    assert _drawing_identity(s, "12349-02")[2] == "GRAVITY FEEDER MODULES"


def test_a_single_root_names_the_unit_whatever_its_code():
    from client_quote_html import _drawing_identity
    s = _shadow(["ASSY-1"], nodes=[("ASSY-1", "COUNTER UNIT")])
    assert _drawing_identity(s, "12349-02")[2] == "COUNTER UNIT"


def test_two_unrelated_roots_claim_nothing():
    """When two different things ship and neither contains the other, there is no single
    answer to 'what is this unit called' — and inventing one is worse than a blank."""
    from client_quote_html import _drawing_identity
    s = _shadow(["8100-01", "9200-01"],
                nodes=[("8100-01", "LEFT STAND"), ("9200-01", "RIGHT STAND")])
    assert _drawing_identity(s, "12349-02")[2] == "12349-02"


def test_an_explicit_top_assembly_still_wins():
    from client_quote_html import _drawing_identity
    s = _shadow(["12349-02-69", "12349-02-69-100"], top="12349-02-69-100",
                nodes=[("12349-02-69", "GRAVITY FEEDER MODULES"),
                       ("12349-02-69-100", "GRAVITY FEEDERS, 9 WIDE")])
    assert _drawing_identity(s, "12349-02")[2] == "GRAVITY FEEDERS, 9 WIDE"


def test_the_header_then_writes_it():
    """End to end: the shape this pack had now fills the Description cell."""
    ws = _sheet()
    s = _shadow(["12349-02-69", "12349-02-69-100"],
                nodes=[("12349-02-69", "GRAVITY FEEDER MODULES"),
                       ("12349-02-69-100", "GRAVITY FEEDERS, 9 WIDE")])
    written = write_job_identity_header(ws, s, "12349-02")
    assert "description" in written
    assert ws["D4"].value == "GRAVITY FEEDER MODULES"


# ── the client, which was falling back to the job number ─────────────────────────────────
# "For traceability / easier identification can Client / Job Description / Date be populated
# for header" — the second estimator, the same week as the first asked for description, date
# and drawing number. The customer cell fell back to the JOB NUMBER, so the sheet read
# "7332-01" under a heading that means Harrods. The client is not on the drawing in any form
# we read; it is the folder the pack came from.

def _client(path):
    from wb_populate import client_from_job_folder
    return client_from_job_folder({"job_folder": path})


def test_the_client_is_read_from_the_enquiry_folder():
    assert _client(r"K:\Estimating\Completed\Live Enquiries\completed AI briefs"
                   r"\Harrods\7332-01-A3SignageStand") == "Harrods"


def test_the_other_live_pack_reads_too():
    assert _client(r"\\sdi-dc01\shareddata$\Shared\Estimating\Completed\AI Estimating"
                   r"\AISheets\SDIIntelligenceAISheet\fanatics\12349-02") == "fanatics"


def test_a_filing_directory_is_never_mistaken_for_a_customer():
    """The enquiry tree is deep and most of it is filing, not clients."""
    for path in (r"K:\Estimating\completed AI briefs\7332-01",
                 r"C:\ClaudeVision\output\12349-02",
                 r"K:\Estimating\Live Enquiries\7332-01"):
        assert _client(path) == "", path


def test_a_job_number_above_a_job_is_not_a_customer():
    assert _client(r"K:\jobs\7332\7332-01") == ""


def test_nothing_to_read_is_an_empty_answer_not_a_guess():
    assert _client("") == "" and _client("12349-02") == ""
