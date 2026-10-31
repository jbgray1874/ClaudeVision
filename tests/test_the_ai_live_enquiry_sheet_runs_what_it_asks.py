"""The AI Live Enquiry sheet runs what it asks for, from the top (D-608).

31 Oct 2026, James: "this s/sheet is read : C:\\ClaudeVision\\chatgptdropoff\\AI-Live-Enquiry.xlsx.
the columns AI Check ( column H ) is then checked. one by one from the top. if it states YES,
and there is a QTY ( column I ) and AICompleted ( column J ) is NO, or N, or Re-Run, we run the
estimate ... we need to look for the drawings in this location : K:\\Estimating\\Completed\\AI
Estimating\\Live Enquiry\\ matching off the customer in Column D with the directory below ...
then a pattern match for the Drawing No. Field F in the name of the folder in the directory
below the client name."

Then (D-608a): "Runs a row when [QTY] will be like the same that we put into the number of
units area like 1,5,10, and AICompleted (J) is NO, N or Re-Run or Blank. Also, when we complete
the run for that row, we populate the AICompleted column J with YES, so it is not re-ran
unless re-set." The YES is the one write to the workbook, made only while it is not open in
Excel; until it lands, the ledger stops a repeat. The workbook's tabs are Archived and Live,
and Live is the sheet read.
"""
from __future__ import annotations

import os
import sys
from datetime import datetime
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("SDI_OFFLINE", "1")

import config  # noqa: E402
import live_enquiry_runner as L  # noqa: E402

# The AI Live Enquiry sheet as it stands: A empty, the header on row 1, H AI CHECK, I QTY,
# J AICompleted, K Estimator.
HEADER = [None, "Enquiry Received", "Account Manager", "Customer", "Job Description",
          "Drawing No.", "Requested Estimate Completion Date", "AI CHECK", "QTY", "AICompleted",
          "Estimator"]


def _line(customer, drawing, qty="1,5,10", done="NO", ai="YES", desc="Editions Footwear Podium",
          due=datetime(2026, 11, 5)):
    return [None, datetime(2026, 10, 30), "Andy Morrans", customer, desc, drawing, due, ai, qty,
            done, "Tony"]


def _rows(*lines):
    return L.rows_from_table([HEADER, [None] * 11, *lines])


def _pack(root: Path, client: str, folder: str, *names: str) -> Path:
    d = root / client / folder
    d.mkdir(parents=True, exist_ok=True)
    for n in names or ("GA.pdf",):
        (d / n).write_bytes(b"%PDF-1.4 " + n.encode())
    return d


@pytest.fixture
def root(tmp_path):
    r = tmp_path / "Live Enquiry"
    r.mkdir()
    return r


def _plan(rows, root, ledger=None):
    ledger = ledger if ledger is not None else {"schema": L.LEDGER_SCHEMA, "jobs": {},
                                                "quantities": {}}
    return L.plan(rows, L._sources(root), ledger), ledger


class _Portal:
    def __init__(self):
        self.started = []

    def start(self, body):
        self.started.append(body["drawing_number"])
        return {"run_id": f"r{len(self.started)}", "output_path": "C:\\out"}

    def status(self, run_id):
        return {"status": "done", "engine_price_gbp": 10.0, "seconds": 1,
                "output_path": "C:\\out", "deliverables": []}


def _run(decisions, ledger, tmp_path):
    return L.run_ready(decisions, _Portal(), ledger, tmp_path / "ledger.json",
                       say=lambda s: None, poll_seconds=0, sleep=lambda s: None)


# ── the sheet ──────────────────────────────────────────────────────────────────────────

def test_the_sheet_is_read_by_its_headers_with_qty_and_aicompleted():
    row = _rows(_line("M&S", "12678-01-GA", qty="1,5,10", done="Re-Run"))[0]
    assert (row.customer, row.drawing, row.ai_check) == ("M&S", "12678-01-GA", "YES")
    assert row.quantities == "1,5,10" and row.ai_completed == "Re-Run"
    assert row.completed_column, "this sheet runs by its own rule"
    assert row.notes == "", "AICompleted is a headed column, not a note"


def test_the_workbook_defaults_to_the_one_james_named():
    """D-608b: K:\\Estimating\\Completed\\AI Estimating\\Live Enquiry\\AI-Live-Enquiry.xlsx, by
    its UNC path, beside the customer folders under the Live Enquiry root."""
    if os.getenv("SDI_LIVE_ENQUIRY_WORKBOOK"):
        pytest.skip("SDI_LIVE_ENQUIRY_WORKBOOK is set on this machine and wins")
    assert config.LIVE_ENQUIRY_WORKBOOK == (config.LIVE_ENQUIRY_ROOT.rstrip("\\/")
                                            + "\\AI-Live-Enquiry.xlsx")
    if not os.getenv("SDI_LIVE_ENQUIRY_ROOT"):
        assert config.LIVE_ENQUIRY_WORKBOOK.startswith("\\\\sdi-dc01\\"), "UNC, never K:"
    assert ":" not in config.LIVE_ENQUIRY_WORKBOOK[:3], "never a drive letter"


def test_the_workbook_beside_the_customer_folders_is_not_taken_for_one(root):
    """It sits in the Live Enquiry folder itself; only folders are customers."""
    (root / "AI-Live-Enquiry.xlsx").write_bytes(b"PK")
    _pack(root, "M&S", "12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA")), root)[0][0]
    assert d.status == "ready" and d.located.client_folder_name == "M&S"


# ── what runs ──────────────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("done", ["", "NO", "No", "N", "n", "Re-Run", "re run", "RERUN",
                                  "Re_Run"])
def test_yes_a_qty_and_aicompleted_no_n_or_rerun_runs(root, done):
    _pack(root, "M&S", "12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA", done=done)), root)[0][0]
    assert d.status == "ready", d.reason
    assert d.quantities == [1, 5, 10] and d.quantity_source == "the sheet's QTY column"


@pytest.mark.parametrize("done", ["YES", "Y", "Done", "Manual"])
def test_any_other_aicompleted_is_not_a_request_to_run(root, done):
    _pack(root, "M&S", "12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA", done=done)), root)[0][0]
    assert d.status == "not_asked"
    assert "AICompleted is" in d.reason and "NO, N or Re-Run runs it" in d.reason


def test_qty_reads_as_the_units_box_does(root):
    _pack(root, "M&S", "12676-01-GA2")
    d = _plan(_rows(_line("M&S", "12676-01-GA2", qty="1,10,20,30,50", done="")), root)[0][0]
    assert d.quantities == [1, 10, 20, 30, 50], "the first is the run, the rest are breaks"


def test_ai_check_not_yes_is_not_asked_whatever_aicompleted_says(root):
    _pack(root, "M&S", "12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA", ai="", done="NO")), root)[0][0]
    assert d.status == "not_asked" and "AI CHECK" in d.reason


def test_no_qty_on_the_sheet_holds_the_row_and_nothing_else_supplies_one(root):
    """The QTY column is the quantity: a QUANTITIES.txt in the folder or "100 off" in the
    description, which the older sheet accepted, do not count here."""
    _pack(root, "M&S", "12678-01-GA", "GA.pdf", "QUANTITIES.txt")
    (root / "M&S" / "12678-01-GA" / "QUANTITIES.txt").write_text("1,5", encoding="utf-8")
    rows = _rows(_line("M&S", "12678-01-GA", qty="", desc="Podium 100 off"))
    d, ledger = _plan(rows, root)
    ledger["quantities"][L.drawing_key("12678-01-GA")] = [7]
    d = L.plan(rows, L._sources(root), ledger)[0]
    assert d.status == "held" and "no QTY on the sheet" in d.reason


def test_ready_rows_run_from_the_top_of_the_sheet(root, tmp_path):
    """One by one from the top, not by the completion date."""
    for n in ("12678-01-GA", "12684-01-GA", "12676-01-GA2"):
        _pack(root, "M&S", n)
    rows = _rows(_line("M&S", "12678-01-GA", due=datetime(2026, 11, 9)),
                 _line("M&S", "12684-01-GA", due=datetime(2026, 11, 1)),
                 _line("M&S", "12676-01-GA2", due=datetime(2026, 11, 5)))
    decisions, ledger = _plan(rows, root)
    portal = _Portal()
    L.run_ready(decisions, portal, ledger, tmp_path / "ledger.json", say=lambda s: None,
                poll_seconds=0, sleep=lambda s: None)
    assert portal.started == ["12678-01-GA", "12684-01-GA", "12676-01-GA2"]


# ── not run twice, and Re-Run once per marking ─────────────────────────────────────────

def test_a_finished_row_is_not_run_again_while_its_yes_is_not_on_the_sheet(root, tmp_path):
    """The workbook was open, so the YES could not be written: the ledger stops a repeat."""
    _pack(root, "M&S", "12678-01-GA")
    rows = _rows(_line("M&S", "12678-01-GA", done=""))
    decisions, ledger = _plan(rows, root)
    assert len(_run(decisions, ledger, tmp_path)) == 1
    again = L.plan(rows, L._sources(root), ledger)[0]
    assert again.status == "done" and "AICompleted not yet written" in again.reason


def test_a_row_reset_after_its_yes_reached_the_sheet_runs_again(root, tmp_path):
    """"So it is not re-ran unless re-set": the YES landed, and a person cleared it."""
    _pack(root, "M&S", "12678-01-GA")
    rows = _rows(_line("M&S", "12678-01-GA", done=""))
    decisions, ledger = _plan(rows, root)
    _run(decisions, ledger, tmp_path)
    ledger["jobs"][L.job_key(rows[0])]["attempts"][-1]["sheet_marked"] = "2026-10-31T19:00:00"
    assert L.plan(_rows(_line("M&S", "12678-01-GA", done="YES")), L._sources(root),
                  ledger)[0].status == "not_asked"
    d = L.plan(rows, L._sources(root), ledger)[0]
    assert d.status == "ready" and "AICompleted reset" in d.reason


def test_rerun_runs_once_per_marking(root, tmp_path):
    _pack(root, "M&S", "12678-01-GA")
    no = _rows(_line("M&S", "12678-01-GA", done="NO"))
    rerun = _rows(_line("M&S", "12678-01-GA", done="Re-Run"))
    decisions, ledger = _plan(no, root)
    _run(decisions, ledger, tmp_path)                                 # first run
    d = L.plan(rerun, L._sources(root), ledger)[0]
    assert d.status == "ready" and d.reason == "Re-Run asked for on the sheet"
    assert len(_run([d], ledger, tmp_path)) == 1                      # the re-run
    assert L.plan(rerun, L._sources(root), ledger)[0].status == "done", \
        "the cell still says Re-Run, and that Re-Run has been answered"
    L.plan(no, L._sources(root), ledger)                              # the cell is changed back
    assert L.plan(rerun, L._sources(root), ledger)[0].status == "ready", \
        "marked Re-Run again, it runs again"


def test_rerun_lets_a_failed_job_go_again(root, tmp_path):
    _pack(root, "M&S", "12678-01-GA")
    rows = _rows(_line("M&S", "12678-01-GA", done="Re-Run"))
    _, ledger = _plan(rows, root)
    job = ledger["jobs"].setdefault(L.job_key(rows[0]), {"customer": "M&S",
                                                          "drawing": "12678-01-GA",
                                                          "attempts": []})
    fp = L.fingerprint(L.pack_files(root / "M&S" / "12678-01-GA"))
    job["attempts"].append({"status": "error", "fingerprint": fp, "quantities": [1, 5, 10],
                            "error": "boom"})
    assert L.plan(rows, L._sources(root), ledger)[0].status == "ready"


# ── the drawings: the customer's folder, then a folder carrying the number ────────────

def test_a_folder_carrying_the_drawing_number_inside_its_name_is_the_pack(root):
    _pack(root, "M&S", "0359974 Editions Footwear Podium 12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA")), root)[0][0]
    assert d.status == "ready", d.reason
    assert d.located.folder.name == "0359974 Editions Footwear Podium 12678-01-GA"


def test_a_number_inside_a_longer_number_or_another_drawing_is_not_a_match(root):
    _pack(root, "M&S", "Podium 112678-01-GA")
    _pack(root, "M&S", "Podium 12678-01-02-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA")), root)[0][0]
    assert d.status == "waiting" and "no 12678-01-GA folder under M&S" in d.reason


def test_two_folders_carrying_the_number_hold_the_row(root):
    _pack(root, "M&S", "0359974 Podium 12678-01-GA")
    _pack(root, "M&S", "Old Podium 12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA")), root)[0][0]
    assert d.status == "held" and "leave one" in d.reason


def test_a_folder_named_by_the_number_wins_over_one_carrying_it(root):
    _pack(root, "M&S", "12678-01-GA Podium")
    _pack(root, "M&S", "Copy of 12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA")), root)[0][0]
    assert d.status == "ready" and d.located.folder.name == "12678-01-GA Podium"


def test_the_customer_folder_is_matched_from_column_d(root):
    _pack(root, "TTI", "12754-01-GA Gen 5 Wall CDU")
    d = _plan(_rows(_line("TTI", "12754-01-GA", qty="1,364")), root)[0][0]
    assert d.status == "ready" and d.located.client_folder_name == "TTI"
    assert d.quantities == [1, 364]


def test_the_number_inside_is_looked_for_only_under_the_customers_folder(root):
    """At the root, the older layout still needs the exact number to lead the name."""
    _pack(root, "M&S", "unrelated")
    _pack(root, ".", "Copy of 12678-01-GA")
    d = _plan(_rows(_line("M&S", "12678-01-GA")), root)[0][0]
    assert d.status == "waiting"



# ── which sheet ────────────────────────────────────────────────────────────────────────

def test_the_live_tab_is_read_never_archived():
    assert L.pick_sheet(["Archived", "Live"]) == "Live"
    assert L.pick_sheet(["Sheet1", "Sheet2"]) == "Sheet1"
    assert L.pick_sheet(["Jobs"]) == "Jobs", "a workbook's only sheet"


def test_several_sheets_and_none_named_is_refused_not_read_at_the_first_tab():
    with pytest.raises(ValueError) as refused:
        L.pick_sheet(["Archived", "October"])
    assert "Archived, October" in str(refused.value) and "SDI_LIVE_ENQUIRY_SHEET" in str(refused.value)
    with pytest.raises(ValueError):
        L.pick_sheet(["Archived", "Live"], "Current")


# ── AICompleted written back ───────────────────────────────────────────────────────────

def _workbook(tmp_path, *lines, formula=True):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    arch = wb.active
    arch.title = "Archived"
    arch.append(HEADER)
    arch.append(_line("M&S", "12678-01-GA", done=""))        # an old row, never touched
    live = wb.create_sheet("Live")
    live.append(HEADER)
    live.append([None] * 11)
    for line in lines:
        live.append(line)
    if formula:
        live["M3"] = "=1+1"
    path = tmp_path / "chatgptdropoff" / "AI-Live-Enquiry.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _cell(path, sheet, ref):
    import openpyxl
    return openpyxl.load_workbook(path)[sheet][ref].value


def test_a_finished_rows_aicompleted_is_set_to_yes_and_nothing_else_changes(tmp_path):
    path = _workbook(tmp_path, _line("M&S", "12678-01-GA", done=""),
                     _line("M&S", "12684-01-GA", done="N"))
    row = [r for r in L.read_sheet(path) if r.drawing == "12684-01-GA"][0]
    ok, note = L.mark_completed(path, row)
    assert ok, note
    assert _cell(path, "Live", "J4") == "YES"
    assert _cell(path, "Live", "J3") is None, "only that row"
    assert _cell(path, "Live", "M3") == "=1+1", "formulas are kept, not their values"
    assert _cell(path, "Archived", "J2") is None, "the Archived tab is never touched"
    assert not list(path.parent.glob("*.sdi-writing*")), "no half-written file left behind"


def test_nothing_is_written_while_the_workbook_is_open_in_excel(tmp_path):
    path = _workbook(tmp_path, _line("M&S", "12678-01-GA", done=""))
    (path.parent / ("~$" + path.name)).write_bytes(b"lock")
    before = path.read_bytes()
    ok, note = L.mark_completed(path, L.read_sheet(path)[0])
    assert not ok and "open in Excel" in note
    assert path.read_bytes() == before


def test_a_row_moved_by_a_sort_is_found_by_its_job(tmp_path):
    path = _workbook(tmp_path, _line("M&S", "12684-01-GA"), _line("M&S", "12678-01-GA"))
    row = L.EnquiryRow(row=3, customer="M&S", drawing="12678-01-GA")    # it was row 3
    ok, note = L.mark_completed(path, row)
    assert ok and "row 4" in note
    assert _cell(path, "Live", "J4") == "YES" and _cell(path, "Live", "J3") == "NO", \
        "the row now at 3 is another job, left as it was"


def test_a_job_on_two_rows_or_none_is_left_alone(tmp_path):
    path = _workbook(tmp_path, _line("M&S", "12678-01-GA"), _line("M&S", "12678-01-GA"))
    ok, note = L.mark_completed(path, L.EnquiryRow(row=9, customer="M&S", drawing="12678-01-GA"))
    assert not ok and "2 other row(s)" in note
    ok, note = L.mark_completed(path, L.EnquiryRow(row=3, customer="TTI", drawing="99999"))
    assert not ok and "no other row(s)" in note


def test_self_running_writes_yes_after_the_run_and_then_the_row_is_not_asked(tmp_path, root,
                                                                            monkeypatch):
    import self_running as sr
    path = _workbook(tmp_path, _line("M&S", "12678-01-GA", done=""))
    _pack(root, "M&S", "0359974 Editions Footwear Podium 12678-01-GA")
    s = {"root": root, "ledger": tmp_path / "ledger.json", "workbook": path}

    class _On:
        def beat(self, rep):
            return True

    portal = _Portal()
    sr.cycle(s, portal, _On(), say=lambda m: None, wait={"poll_seconds": 0,
                                                          "sleep": lambda x: None})
    assert portal.started == ["12678-01-GA"]
    assert _cell(path, "Live", "J3") == "YES"
    after = L.plan(L.read_sheet(path), L._sources(root), L.load_ledger(s["ledger"]))[0]
    assert after.status == "not_asked"


def test_a_yes_that_could_not_be_written_is_written_on_a_later_read(tmp_path, root):
    import self_running as sr
    path = _workbook(tmp_path, _line("M&S", "12678-01-GA", done=""))
    _pack(root, "M&S", "12678-01-GA")
    s = {"root": root, "ledger": tmp_path / "ledger.json", "workbook": path}
    lock = path.parent / ("~$" + path.name)
    lock.write_bytes(b"lock")

    class _On:
        def beat(self, rep):
            return True

    portal = _Portal()
    wait = {"poll_seconds": 0, "sleep": lambda x: None}
    sr.cycle(s, portal, _On(), say=lambda m: None, wait=wait)
    assert _cell(path, "Live", "J3") is None, "open in Excel: not written"
    sr.cycle(s, portal, _On(), say=lambda m: None, wait=wait)
    assert portal.started == ["12678-01-GA"], "and not run a second time meanwhile"
    lock.unlink()                                                  # Excel is closed
    sr.cycle(s, portal, _On(), say=lambda m: None, wait=wait)
    assert _cell(path, "Live", "J3") == "YES" and portal.started == ["12678-01-GA"]
