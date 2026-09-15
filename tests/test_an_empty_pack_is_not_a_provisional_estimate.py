""""PROVISIONAL. not reported/unit at 1 off. 11908-21" — over nothing at all.

James, 15 Sep 2026, on the phone screenshot: "It looks like it didn't run" — and then:
"There were no attachments so no output was produced". The triggering email carried no
drawings; nothing was scanned, nothing was costed — and the covering note still walked its
seven sections and dressed the nothing as a draft price. A reader cannot tell that subject
line from a broken estimate, and either reading wastes an estimator's evening.

When there is no BOM line, no material row, no labour row and no unit figure, there is
exactly one thing worth saying, and it goes in the subject: NOTHING WAS ESTIMATED.
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
os.environ.setdefault("SDI_OFFLINE", "1")


def _empty_book(tmp: Path, name: str = "11908-21_empty.xlsx") -> Path:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Estimate"
    p = tmp / name
    wb.save(p)
    return p


def test_an_empty_pack_says_so_in_the_subject(tmp_path):
    from estimate_explained import covering_email
    out = covering_email(_empty_book(tmp_path))
    assert "NOTHING WAS ESTIMATED" in out["subject"], out["subject"]
    assert "PROVISIONAL" not in out["subject"]
    assert "not reported" not in out["subject"]


def test_the_note_names_the_usual_cause_and_the_remedy(tmp_path):
    from estimate_explained import covering_email
    out = covering_email(_empty_book(tmp_path))
    for text in (out["html"], out["text"]):
        assert "attachments" in text.lower()
        assert "re-send" in text.lower()
    assert "not" in out["html"].lower() and "price" in out["html"].lower(), \
        "the note must say nothing in it is a price"


def test_a_real_estimate_still_gets_the_full_note(tmp_path):
    """The guard keys on the estimate being EMPTY, never on its money being unreadable —
    a book with rows whose total failed to read keeps the existing behaviour and wording."""
    import openpyxl
    from estimate_explained import covering_email
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estimate"
    ws["C6"], ws["D6"] = "Quantity", 1
    ws["C10"], ws["H10"] = "Bill of Materials (Per Unit)", "Part code"
    ws["C11"], ws["H11"], ws["J11"], ws["K11"] = "WIDGET", "WID1", 2.5, 1
    p = tmp_path / "12345-01_real.xlsx"
    wb.save(p)
    out = covering_email(p)
    assert "NOTHING WAS ESTIMATED" not in out["subject"], out["subject"]
