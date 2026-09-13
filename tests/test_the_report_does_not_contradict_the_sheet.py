"""A report that contradicts the sheet it accompanies is worse than either alone.

12349-02's provenance tab said of the 6 mm MDF packer:

    Modelled but in no assembly BOM · 12349-02-69-08J · a fixture, jig or setup part in the
    model folder — not a component of the product, NOT COSTED; confirm

while the Estimate sheet beside it charged £0.95 for that part, with CNC Joinery and Wet
Spray against it. The line asserted "not costed" about every part the SolidWorks reader found
outside the assembly bill of materials, without once looking at the costed population.

Being absent from the assembly BOM is worth raising either way — it is how a jig gets priced
as a product part, and how a real component goes missing. What has to be true is the sentence
about the money.

Also pinned here: the Total Material Cost formula is found by LABEL. Widening the Other Sheet
Material block — which this job needs, it has nine board and acrylic parts for eight rows —
moves that cell, and the function that makes the total error-tolerant used to read it at M92.
It would have flagged "not the expected SUM formula" and left plain SUM in place, so one part
with no blank dimensions would #VALUE! Total Material, Unit Cost and Sell Price together.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import wb_populate                                                      # noqa: E402
from estimation_report import add_provenance_sheet                      # noqa: E402


def _summary(costed_gbp):
    """A job whose model folder holds 08J and whose sheet may or may not charge for it."""
    pe = {"part_number": "12349-02-69-08J", "description": "PACKER",
          "quantity": 1, "cost_breakdown": {}}
    if costed_gbp is not None:
        pe["extended_total_cost_gbp"] = costed_gbp
    return {
        "solidworks_native": {"applied": {"not_in_bom_parts": ["12349-02-69-08J"]}},
        "estimate_summary": {"part_estimates": [pe]},
    }


def _provenance_text(summary):
    wb = openpyxl.Workbook()
    add_provenance_sheet(wb, summary, {"pdf_name": "12349-02", "job_number": "12349",
                                       "scan_date": "13/09/2026"})
    ws = next(wb[n] for n in wb.sheetnames if "rovenance" in n)
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)


def test_a_costed_part_is_never_reported_as_not_costed():
    """The 08J case, exactly as it shipped."""
    text = _provenance_text(_summary(0.95))
    assert "12349-02-69-08J" in text
    assert "not costed" not in text.lower(), \
        "the report told an estimator a part was free while the sheet charged for it"
    assert "£0.95" in text


def test_it_still_says_which_of_the_two_things_happened():
    text = _provenance_text(_summary(0.95))
    assert "fixture or jig" in text and "short a line" in text, \
        "absent from the BOM and costed is worth raising — the row must say what to check"


def test_an_uncosted_fixture_keeps_the_original_wording():
    """A genuine jig, not in the price: the line that was right all along."""
    text = _provenance_text(_summary(None))
    assert "not costed" in text.lower()


def test_a_zero_cost_is_not_a_charge():
    text = _provenance_text(_summary(0.0))
    assert "not costed" in text.lower()


# ── the material total is found by label, so a widened block cannot blind it ──────────────

def _sheet_with_total_at(row, other_last=91):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=row, column=3, value="Total Material Cost")
    ws.cell(row=row, column=13,
            value=f"=(SUM(M11:M50)+SUM(M53:M60)+SUM(M63:M81)+SUM(M84:M{other_last})+AF83)")
    return ws


def test_the_material_total_is_made_error_tolerant_where_it_actually_sits():
    """Two rows inserted into Other Sheet Material: the total is at 94, not 92."""
    ws = _sheet_with_total_at(94, other_last=93)
    assert wb_populate._make_material_total_error_tolerant(ws) is True
    assert "_xlfn.AGGREGATE(9,6," in ws.cell(row=94, column=13).value
    assert "SUM(" not in re.sub(r"_xlfn\.AGGREGATE\(9,6,", "", ws.cell(row=94, column=13).value)


def test_it_still_works_where_the_template_has_always_had_it():
    ws = _sheet_with_total_at(92)
    assert wb_populate._make_material_total_error_tolerant(ws) is True


def test_a_sheet_with_no_total_row_is_reported_not_crashed():
    flags = []
    assert wb_populate._make_material_total_error_tolerant(
        openpyxl.Workbook().active, flags) is False
    assert flags and "Total Material Cost" in flags[0]
