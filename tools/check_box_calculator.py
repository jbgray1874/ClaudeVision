"""Compare BOTH box-calculator formulas against the Harley's order extract.

The BOX TEST SHEET is 327 real Harleys Corrugated Cases orders out of SDI Live: order
number, part code, description, quantity ordered and the price actually paid. Most
descriptions carry the box's outside size and its FEFCO style code. Two candidate totals
exist — the source sheet's own (running + set-up + the first TWO panel pairs, its
K10 = SUM(K7:K8)) and the three-pair sum — and NEITHER is assumed correct: this tool
reports both against the paid prices, overall and by quantity band, so the choice is made
on evidence. Only the plain six-panel 0201 carton is compared; every other style is
reported as out of scope rather than silently mispriced.

THE DIFFERENCE IS (paid − calculator) / paid. Negative means the calculator is HIGHER
than what Harleys charged; positive means lower. Say it that way round everywhere.

Usage (from the repo root; the user's copy of the sheets lives in
C:\\ClaudeVision\\chatgptdropoff\\AICalculator):

    python tools\\check_box_calculator.py "chatgptdropoff\\AICalculator\\BOX TEST SHEET.xlsx"
    python tools\\check_box_calculator.py <sheet.xlsx> --write-back out.xlsx

With --write-back the three-pair figure goes in column G and its difference in H, exactly
where the sheet reserves them; the original file is never touched.
"""
from __future__ import annotations

import argparse
import statistics
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import box_costing  # noqa: E402
import config  # noqa: E402

_BANDS = ((1, 9), (10, 49), (50, 199), (200, 10**9))


def _band(q) -> str:
    q = int(q or 1)
    for lo, hi in _BANDS:
        if lo <= q <= hi:
            return f"{lo}-{hi}" if hi < 10**9 else f"{lo}+"
    return "?"


def _report(label: str, rows) -> None:
    diffs = [d for *_x, d in rows if d is not None]
    if not diffs:
        print(f"\n{label}: no comparable rows")
        return
    within = sum(1 for d in diffs if abs(d) <= 0.25)
    print(f"\n{label} — (paid − calculator) / paid over {len(diffs)} rows "
          f"(negative = calculator higher than Harleys charged):")
    print(f"  median {statistics.median(diffs):+.1%}   mean {statistics.fmean(diffs):+.1%}   "
          f"within +/-25%: {within} ({within / len(diffs):.0%})")
    by_band: dict = {}
    for row, desc, qty, paid, calc, d in rows:
        if d is not None:
            by_band.setdefault(_band(qty), []).append(d)
    for band in sorted(by_band, key=lambda b: int(b.split("-")[0].rstrip("+"))):
        ds = by_band[band]
        print(f"    qty {band:>7}: median {statistics.median(ds):+.1%} over {len(ds)} rows")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("sheet", help="BOX TEST SHEET .xlsx (the SDI Live Harley's extract)")
    ap.add_argument("--write-back", default="",
                    help="write a copy with the three-pair figure in G and %% diff in H")
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(args.sheet, data_only=True)
    ws = wb.worksheets[0]

    styles_priced = set((getattr(config, "BOX_COSTING", {}) or {})
                        .get("styles_priced") or ("0201",))
    three_pair, two_pair, cells = [], [], []
    total, no_dims, declined = 0, 0, 0
    out_of_scope: dict = {}
    for r in ws.iter_rows(min_row=2):
        desc = r[3].value
        if not desc:
            continue
        total += 1
        qty, paid = r[4].value, r[5].value
        read = box_costing.parse_box_description(desc)
        style = read.get("style") or "(none)"
        if not read.get("dims_mm"):
            no_dims += 1
            continue
        if style not in styles_priced:
            out_of_scope[style] = out_of_scope.get(style, 0) + 1
            continue
        L, W, H = read["dims_mm"]
        res = box_costing.price_box(L, W, H, qty or 1)
        if not res.get("priced"):
            declined += 1
            continue
        for rows, calc in ((three_pair, res["unit_gbp"]),
                           (two_pair, res["sheet_total_gbp"])):
            diff = ((float(paid) - calc) / float(paid)) if paid else None
            rows.append((r[0].row, desc, qty, paid, calc, diff))
        if paid:
            cells.append((r[0].row, round(res["unit_gbp"], 2),
                          round((float(paid) - res["unit_gbp"]) / float(paid), 4)))

    print(f"rows read:            {total}")
    print(f"compared (0201):      {len(three_pair)}")
    print(f"no L x W x H in text: {no_dims}")
    print(f"out-of-scope styles:  {sum(out_of_scope.values())}  "
          f"{dict(sorted(out_of_scope.items()))}")
    print(f"declined (oversize/bad dims): {declined}")
    _report("THREE-PAIR sum (all six panels)", three_pair)
    _report("TWO-PAIR sum (the sheet's own K10 = SUM(K7:K8))", two_pair)

    worst = sorted(three_pair, key=lambda t: -abs(t[5] or 0))[:8]
    if worst:
        print("\n  largest three-pair differences:")
        for row, desc, qty, paid, calc, d in worst:
            print(f"   row {row}: paid £{float(paid):.2f}  calc £{calc:.2f}  "
                  f"({d:+.0%})  x{qty}  {str(desc)[:70]}")

    if args.write_back:
        for row, calc, diff in cells:
            ws.cell(row=row, column=7, value=calc)
            ws.cell(row=row, column=8, value=diff)
        wb.save(args.write_back)
        print(f"\nwritten: {args.write_back} (G = three-pair figure, H = (paid − calc)/paid; "
              f"the original is untouched)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
