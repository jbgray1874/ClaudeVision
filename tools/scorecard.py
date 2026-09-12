"""A small, brutal gate on one real pack.

WHAT THIS IS FOR. The regressions that reach an estimator are not "the model is 3% off" — they
are the same failure repeating: a fix lands, five thousand tests stay green, and a number that
was right last week is wrong again. The suite cannot catch that, because it tests the engine's
parts and the number lives in a finished workbook.

So: a handful of facts about ONE pack that have been proven and must never move. Not the
workbook frozen whole — that would teach the engine that everything on the sheet is right, and
on this pack it is not. Only what is settled.

    python tools\\scorecard.py tests\\scorecard\\12349-02.scorecard C:\\ClaudeVision\\output\\12349-02_20260912_144540.xlsx

Exit 0 = every pinned fact still holds. Exit 1 = something moved, named line by line.

Codes match on their SUFFIX so a line reads the way an estimator says it: `qty 04M = 1`, not
the full four-segment code. The format is in the scorecard file itself.
"""
from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:                                       # pragma: no cover
    print("openpyxl is needed: py -m pip install openpyxl")
    raise SystemExit(2)


def _rules(path: Path):
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.split("#", 1)[0].strip()
        if line:
            yield line


def _find_headers(ws, *wanted):
    """EVERY header row carrying these titles, not the first.

    The Estimate sheet has several nested blocks with the same column names — Wire, Sheet Steel
    and Other Sheet Material all head a column "Part Description". Taking the first match landed
    on the EMPTY Wire block and reported every steel gauge as missing: a checker crying wolf on a
    good pack, which is worse than no checker because the next person turns it off."""
    out = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400)):
        vals = {str(c.value).strip().lower(): c.column for c in row if c.value is not None}
        if all(any(w.lower() == k for k in vals) for w in wanted):
            out.append((row[0].row,
                        {w: vals[next(k for k in vals if k == w.lower())] for w in wanted}))
    return out


def _find_header(ws, *wanted):
    hits = _find_headers(ws, *wanted)
    return hits[0] if hits else (None, {})


def _rows_under(ws, hdr, code_col, limit=80):
    """Rows of a contiguous block, stopping where the block does."""
    out = []
    for r in range(hdr + 1, min(ws.max_row, hdr + limit)):
        v = ws.cell(row=r, column=code_col).value
        if v is None:
            break
        if str(v).strip().lower().startswith("total"):
            break
        out.append((r, str(v).strip()))
    return out


def _matches(code: str, want: str) -> bool:
    c, w = code.upper().replace(" ", ""), want.upper().replace(" ", "")
    return c == w or c.endswith("-" + w) or c.endswith(w) or w in c.split()


def read_facts(book: Path) -> dict:
    """Everything the scorecard can ask about, gathered once."""
    wb = openpyxl.load_workbook(book, data_only=True)
    ws = wb["Estimate"]
    facts = {"qty": {}, "gauge": {}, "dept": {}, "codes": []}

    hdr, cols = _find_header(ws, "Part code", "Qty Per Unit")
    if hdr:
        for r, code in _rows_under(ws, hdr, cols["Part code"]):
            facts["codes"].append(code)
            facts["qty"][code] = ws.cell(row=r, column=cols["Qty Per Unit"]).value

    # Both nested blocks carry a thickness under a different heading.
    for _hdr_words, _thk in ((("Part Description", "Gauge"), "Gauge"),
                             (("Part Description", "Thickness"), "Thickness")):
        for h, c in _find_headers(ws, *_hdr_words):
            for r, desc in _rows_under(ws, h, c["Part Description"]):
                code = desc.split()[0] if desc.split() else desc
                v = ws.cell(row=r, column=c[_thk]).value
                if v is not None:
                    facts["gauge"][code] = v
                    if code not in facts["codes"]:
                        facts["codes"].append(code)

    h, c = _find_header(ws, "Operation", "Part Description")
    if h:
        for r, op in _rows_under(ws, h, c["Operation"], limit=200):
            desc = str(ws.cell(row=r, column=c["Part Description"]).value or "")
            facts["dept"].setdefault(desc, []).append(op)
    return facts


def check(scorecard: Path, book: Path):
    facts = read_facts(book)
    fails = []
    for line in _rules(scorecard):
        parts = line.split()
        kind = parts[0]

        if kind in ("qty", "gauge") and "=" in line:
            want_code, want_val = [x.strip() for x in line[len(kind):].split("=", 1)]
            hit = [(k, v) for k, v in facts[kind].items() if _matches(k, want_code)]
            if not hit:
                fails.append(f"{line}   -> no {kind} row for '{want_code}' on this sheet")
                continue
            for k, v in hit:
                try:
                    if abs(float(v) - float(want_val)) > 1e-9:
                        fails.append(f"{line}   -> {k} is {v}")
                except (TypeError, ValueError):
                    fails.append(f"{line}   -> {k} is {v!r}, which is not a number")

        elif kind == "dept" and "=" in line:
            want_code, want_dept = [x.strip() for x in line[len(kind):].split("=", 1)]
            seen = [op for desc, ops in facts["dept"].items()
                    if any(_matches(t.strip("(),"), want_code) for t in desc.replace(",", " ").split())
                    for op in ops]
            if not seen:
                fails.append(f"{line}   -> no labour row mentions '{want_code}'")
            elif want_dept not in seen:
                fails.append(f"{line}   -> charged to {sorted(set(seen))}")

        elif kind == "no_part_number_contains":
            token = parts[1]
            bad = [c for c in facts["codes"] if token.upper() in c.upper()]
            if bad:
                fails.append(f"{line}   -> found {bad[:4]}")

        elif kind == "must_exist":
            if not any(_matches(c, parts[1]) for c in facts["codes"]):
                fails.append(f"{line}   -> not on the sheet")

        elif kind == "must_not_charge" and " on " in line:
            op, code = line[len("must_not_charge"):].split(" on ", 1)
            op, code = op.strip(), code.strip()
            for desc, ops in facts["dept"].items():
                if any(_matches(t.strip("(),"), code) for t in desc.replace(",", " ").split()) \
                        and any(op.lower() in o.lower() for o in ops):
                    fails.append(f"{line}   -> charged on '{desc[:60]}'")
    return fails


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    card, book = Path(sys.argv[1]), Path(sys.argv[2])
    if not card.is_file() or not book.is_file():
        print("usage: scorecard.py <scorecard file> <workbook>")
        return 2
    print(f"\n=== scorecard {card.name} against {book.name}\n")
    fails = check(card, book)
    for f in fails:
        print(f"   MOVED  {f}")
    print()
    if fails:
        print(f"   {len(fails)} pinned fact(s) moved. Something that was proven is no longer "
              f"true — fix it or, if the change is intended, change the scorecard on purpose.")
        return 1
    print("   every pinned fact still holds.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
