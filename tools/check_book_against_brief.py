"""Hold a run's saved record against an estimator's brief, and report EVERY pin.

    python tools/check_book_against_brief.py output/json/<job>.json docs/briefs/<job>.facts.json
                                             [--xlsx <book.xlsx>]

The replay harness (tests/replay) asserts a frozen, signed-off record and stops at the first
failure, which is right for a gate and wrong for a review: an estimator re-running a job wants
the whole list — what held, what failed, and the figure behind each. This reads the same
record every surface reads (costed_facts.costed_job) and the same facts format the replay layer
uses, so a brief written here can become a replay pin unchanged once a run is signed off.

GENERIC BY CONSTRUCTION. Nothing in this file knows a job: every number, code and word comes
from the facts file. Structure, never money — rates move.

Keys understood (all optional; the replay keys behave as they do there):
  product_root                   the declared product the record must be built from
  not_set_aside                  identities that must not be set aside outside the product
  quantities                     {code: qty per finished unit}
  exactly_one_line               {words: qty or null} — exactly one line whose code or
                                 description carries these words (a merged pair or a doubled
                                 line both fail)
  required_operations            {code: [op]} routed or charged on the line
  required_operations_incl_pieces{code: [op]} met by the line or its numbered pieces
  forbidden_operations           {code: [op]} not charged on the line
  forbidden_operations_anywhere  [op] charged on no line
  material_charged               [code] material money reaches the line
  no_cut_assemblies              true: no assembly / GA line carries a cutting operation
  forbidden_names_everywhere     [name] no line carries it
  unit_mass_kg                   {expected, tolerance_pct} — sum of the material mass the
                                 costing used x quantity, against the drawn mass
  max_unit_material_mass_kg      {code: kg} — a title-block mass used as material fails here
  min_unit_material_mass_kg      {code: kg}
  quantity_breaks                [q] the book's Quantity Breaks columns (needs --xlsx)
  provisional_material           {code: {min_unit_material_gbp}} — the line's material is priced
                                 (at least the floor) and, where its blank was not measured, the
                                 record says which dimensions were measured and which inferred
  commercial_basis               [code] each commercial line names the basis it was priced on, holds
                                 an order figure at every break, and — where that basis is SDI
                                 history — says the history is comparable on customer and quantity
  no_gauge_decision_on_assemblies true: no assembly record carries a gauge decision
  operation_times                {code: {op: {min_run_min, min_setup_min}}} — the op's charged
                                 time reaches the stated floor (proves the work, not only the
                                 word; the workbook row itself is verified only on the book)

Exit status 0 when every pin holds, 1 otherwise.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

_CUTTING = {"laser_cutting", "laser", "cnc_routing", "punching", "guillotine", "saw",
            "tube_cut"}


def _sq(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def _num(value: Any) -> Optional[float]:
    try:
        f = float(value)
    except (TypeError, ValueError):
        return None
    return f if f == f else None


class Result:
    def __init__(self) -> None:
        self.rows: List[Tuple[str, str, bool, str]] = []

    def add(self, check: str, item: str, ok: bool, detail: str = "") -> None:
        self.rows.append((check, item, bool(ok), detail))

    @property
    def failed(self) -> List[Tuple[str, str, bool, str]]:
        return [r for r in self.rows if not r[2]]


def _lines(summary: Dict[str, Any]) -> List[Dict[str, Any]]:
    import costed_facts as cf
    record = cf.costed_job(summary)
    return [l for l in (record.get("lines") or []) if isinstance(l, dict)]


def _ops(line: Optional[Dict[str, Any]], routed: bool = False) -> set:
    if not line:
        return set()
    out = {str(o).lower() for o in (line.get("operations") or [])}
    if routed:
        out |= {str(o).lower() for o in (line.get("route_operations") or [])}
    return out


def _product_root(summary: Dict[str, Any]) -> str:
    for holder in (summary.get("estimate_summary") or {}, summary):
        payload = holder.get("canonical_route_shadow") if isinstance(holder, dict) else None
        if isinstance(payload, dict) and payload.get("product_root"):
            return str(payload["product_root"])
    return ""


def _part_estimates(summary: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    es = summary.get("estimate_summary") or {}
    out: Dict[str, Dict[str, Any]] = {}
    for key in ("canonical_part_estimates", "part_estimates"):
        for p in es.get(key) or []:
            if isinstance(p, dict) and p.get("part_number"):
                out.setdefault(_sq(p["part_number"]), p)
    return out


def _unit_mass(pe: Optional[Dict[str, Any]]) -> Optional[float]:
    if not pe:
        return None
    me = pe.get("material_estimate") or {}
    return _num(me.get("unit_material_mass_kg"))


def _breaks_in_book(xlsx: Path) -> Tuple[List[int], Dict[int, float]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx), data_only=True, read_only=True)
    name = next((n for n in wb.sheetnames if n.strip().lower() == "quantity breaks"), None)
    if name is None:
        return [], {}
    qtys: List[int] = []
    unit: Dict[int, float] = {}
    for row in wb[name].iter_rows(values_only=True):
        head = str(row[0] or "").strip().lower() if row else ""
        if head == "quantity":
            qtys = [int(v) for v in row[1:] if _num(v) is not None]
        elif head.startswith("unit cost") and qtys:
            for q, v in zip(qtys, row[1:]):
                if _num(v) is not None:
                    unit[q] = float(v)
    return qtys, unit


def check(summary: Dict[str, Any], facts: Dict[str, Any],
          xlsx: Optional[Path] = None) -> Result:
    res = Result()
    lines = _lines(summary)
    by_id: Dict[str, Dict[str, Any]] = {}
    for l in lines:
        by_id.setdefault(_sq(l.get("part_number")), l)
    pes = _part_estimates(summary)

    def find(code: str) -> Optional[Dict[str, Any]]:
        return by_id.get(_sq(code))

    def pieces(code: str) -> List[Dict[str, Any]]:
        stem = _sq(code)
        return [l for k, l in by_id.items()
                if k != stem and re.fullmatch(re.escape(stem) + r"\d{1,2}", k)]

    # ── the product ─────────────────────────────────────────────────────────────
    want_root = facts.get("product_root")
    if want_root:
        got = _product_root(summary)
        res.add("product_root", want_root, _sq(got) == _sq(want_root),
                f"record is built from {got or 'nothing'}")
    _aside = {_sq(e.get("part_number")) for e in (summary.get("set_aside_outside_product")
                                                  or []) if isinstance(e, dict)}
    for code in facts.get("not_set_aside") or []:
        res.add("not_set_aside", code, _sq(code) not in _aside,
                "set aside outside the product" if _sq(code) in _aside else "inside the product")

    # ── quantities ──────────────────────────────────────────────────────────────
    for code, q in (facts.get("quantities") or {}).items():
        line = find(code)
        got = _num(line.get("qty_per_unit")) if line else None
        res.add("quantity", code, got is not None and abs(got - float(q)) < 0.01,
                f"book {got if got is not None else 'no line'} / brief {q:g}")
    for words, q in (facts.get("exactly_one_line") or {}).items():
        key = _sq(words)
        hits = [l for l in lines
                if key in _sq(l.get("part_number")) or key in _sq(l.get("description"))]
        ok = len(hits) == 1
        detail = f"{len(hits)} line(s): " + ", ".join(
            f"{l.get('part_number')} x{l.get('qty_per_unit')}" for l in hits[:4])
        if ok and q is not None:
            got = _num(hits[0].get("qty_per_unit"))
            ok = got is not None and abs(got - float(q)) < 0.01
            detail += f" / brief {q:g}"
        res.add("exactly_one_line", words, ok, detail)

    # ── routes ──────────────────────────────────────────────────────────────────
    for code, ops in (facts.get("required_operations") or {}).items():
        line = find(code)
        got = _ops(line, routed=True)
        for op in ops:
            res.add("required_op", f"{code} {op}", op.lower() in got,
                    "no line" if line is None else f"line carries {sorted(got)}")
    for code, ops in (facts.get("required_operations_incl_pieces") or {}).items():
        group = [l for l in [find(code)] + pieces(code) if l]
        got = set().union(*(_ops(l, routed=True) for l in group)) if group else set()
        for op in ops:
            res.add("required_op", f"{code} (or its pieces) {op}", op.lower() in got,
                    f"{[l.get('part_number') for l in group]} carry {sorted(got)}")
    for code, ops in (facts.get("forbidden_operations") or {}).items():
        line = find(code)
        got = _ops(line)
        for op in ops:
            res.add("forbidden_op", f"{code} {op}", op.lower() not in got,
                    "no line (see quantity)" if line is None else
                    ("charged" if op.lower() in got else "not charged"))
    for op in facts.get("forbidden_operations_anywhere") or []:
        guilty = [l.get("part_number") for l in lines if op.lower() in _ops(l)]
        res.add("forbidden_anywhere", op, not guilty, f"charged on {guilty}" if guilty else "")
    if facts.get("no_cut_assemblies"):
        from part_code_conventions import carries_assembly_role
        guilty = [l.get("part_number") for l in lines
                  if (str(l.get("kind") or "").lower() == "assembly"
                      or carries_assembly_role(str(l.get("part_number") or "")))
                  and _ops(l) & _CUTTING]
        res.add("no_cut_assemblies", "assemblies", not guilty,
                f"cut: {guilty}" if guilty else "no assembly is cut")

    # ── material ────────────────────────────────────────────────────────────────
    for code in facts.get("material_charged") or []:
        line = find(code)
        money = None
        if line:
            money = _num(line.get("charged_unit_gbp"))
            if money is None:
                money = _num(line.get("engine_unit_gbp"))
        res.add("material_charged", code, bool(money and money > 0),
                "no line" if line is None else f"£{money}")
    for name in facts.get("forbidden_names_everywhere") or []:
        hit = [l.get("part_number") for l in lines if _sq(name) == _sq(l.get("part_number"))]
        res.add("forbidden_name", name, not hit, f"on {hit}" if hit else "")
    # HOW a line was priced, not only whether (D-388). A frame priced at the global £/kg hold
    # and a frame priced from the catalogue's stock length both carry money; the brief can say
    # which methods are not acceptable for a line ("not") or which one it must be ("in").
    for code, want in (facts.get("cost_method") or {}).items():
        pe = pes.get(_sq(code)) or {}
        got = str(((pe.get("material_estimate") or {}).get("cost_method")) or "")
        ok = bool(got)
        if ok and want.get("not"):
            ok = got not in set(want["not"])
        if ok and want.get("in"):
            ok = got in set(want["in"])
        res.add("cost_method", code, ok, f"priced by {got or 'no method named'}"
                + (f" (not {want['not']})" if want.get("not") else "")
                + (f" (one of {want['in']})" if want.get("in") else ""))
    # WHAT the price names. The oak back is a bought price only if the row carries the
    # drawing's product code; a researched figure must name the decor it was asked with.
    for code, tokens in (facts.get("material_basis_names") or {}).items():
        pe = pes.get(_sq(code)) or {}
        blob = json.dumps(pe.get("material_estimate") or {}, default=str).upper()
        missing = [t for t in tokens if str(t).upper() not in blob]
        res.add("material_basis_names", code, not missing,
                f"missing {missing}" if missing else f"names {tokens}")

    # ── a provisional figure, said as one (D-454) ───────────────────────────────
    _wparts = {_sq(p.get("part_number")): p for p in
               ((summary.get("manufacturing_writeup") or {}).get("parts") or []) if isinstance(p, dict)}
    for code, want in (facts.get("provisional_material") or {}).items():
        pe = pes.get(_sq(code)) or {}
        me = pe.get("material_estimate") or {}
        money = _num(me.get("unit_material_cost_gbp"))
        floor = _num((want or {}).get("min_unit_material_gbp")) or 0.0
        _wpart = _wparts.get(_sq(code)) or {}
        rec = (_wpart or pe).get("_blank_provisional")
        # A CUT PATH IS NOT AN OUTLINE (D-456): "dxf_cut_length_only" carries "dxf" and the
        # first replay called 013's unmeasured blank a measured outline. One shared test.
        try:
            import blank_credibility as _bc
            _measured = _bc.blank_is_measured(_wpart)
        except Exception:                                        # noqa: BLE001
            _measured = False
        said = isinstance(rec, dict) and "measured_mm" in rec and "inferred_mm" in rec
        ok = money is not None and money >= floor and (said or _measured)
        res.add("provisional_material", code, ok,
                f"material £{money} (floor £{floor:g}); "
                + (f"basis {rec.get('basis')}, measured {rec.get('measured_mm')}, inferred "
                   f"{rec.get('inferred_mm')}" if said else
                   ("measured blank" if _measured else
                    "no provisional record and no measured blank — "
                    + str(_wpart.get('geometry_source') or 'no geometry source'))))
    if facts.get("commercial_basis"):
        cls = {str(c.get("code") or "").upper(): c for c in (summary.get("commercial_lines") or [])
               if isinstance(c, dict)}
        if not cls:
            try:
                from commercial_lines import collect_lines
                cls = {str(c.get("code") or "").upper(): c for c in collect_lines(summary)
                       if isinstance(c, dict)}
            except Exception:                                    # noqa: BLE001
                cls = {}
        want_breaks = [int(q) for q in (facts.get("commercial_breaks") or facts.get("quantity_breaks") or [])]
        for code in facts["commercial_basis"]:
            c = cls.get(str(code).upper())
            if not c:
                res.add("commercial_basis", code, False, "no commercial line")
                continue
            basis = str(c.get("basis_chosen") or "")
            breaks = {int(k): v for k, v in (c.get("order_gbp_at_breaks") or {}).items()}
            missing = [q for q in want_breaks if q not in breaks]
            # Weak history IS the justified working basis where the line records why the
            # counted shipment could not price (the market refused, nothing counted) — the
            # policy asks for the best-supported basis SAID, not for history never to win.
            # STRUCTURED EVIDENCE FIRST (D-457): the engine stamps shipment_refusal on the
            # line; wording (the parenthesis) is accepted only for records saved before the
            # field existed. "recorded no reason" is the engine saying its own evidence is
            # missing, and never passes.
            refusal = str(c.get("shipment_refusal") or "")
            researched = bool(c.get("shipment_working"))
            weak = "weak" in basis.lower()
            justified = bool(refusal) and "recorded no reason" not in refusal
            if not refusal and "(" in basis and "recorded no reason" not in basis:
                justified = True                    # a pre-field record, reason in the words
            hist_ok = ("history" not in basis.lower()) or ("comparable" in basis.lower()
                                                           and not weak) \
                or (weak and justified)
            klass = ("researched shipment" if researched else
                     "comparable history" if "comparable" in basis.lower() and not weak else
                     f"justified fallback — {refusal or 'reason in the basis words'}" if weak and hist_ok
                     else "unjustified fallback")
            ok = bool(basis) and not missing and hist_ok
            res.add("commercial_basis", code, ok,
                    (f"[{klass}] basis: {basis or 'NOT NAMED'}; " +
                     ("; ".join(f"{q}: £{float(breaks[q]) / max(q, 1):.2f}/unit" for q in sorted(breaks))
                      or "no breaks held") +
                     (f"; breaks missing {missing}" if missing else "") +
                     ("" if hist_ok else "; history won with no stated reason why the shipment "
                                        "basis did not price")))
    for code, ops_want in (facts.get("operation_times") or {}).items():
        pe = pes.get(_sq(code)) or {}
        proc = pe.get("process_estimate") or {}
        run = proc.get("run_times_min_per_unit") or {}
        setup = proc.get("setup_times_min") or {}
        hours = (pe.get("labour_estimate") or {}).get("run_hours_per_unit") or {}
        for op, want in (ops_want or {}).items():
            got_run = _num(run.get(op))
            if got_run is None and _num(hours.get(op)) is not None:
                got_run = round(float(hours[op]) * 60.0, 2)
            got_setup = _num(setup.get(op))
            ok = True
            if _num((want or {}).get("min_run_min")) is not None:
                ok = got_run is not None and got_run >= float(want["min_run_min"])
            if ok and _num((want or {}).get("min_setup_min")) is not None:
                ok = got_setup is not None and got_setup >= float(want["min_setup_min"])
            res.add("operation_time", f"{code} {op}", ok,
                    f"run {got_run} min (floor {(want or {}).get('min_run_min')}), "
                    f"setup {got_setup} min (floor {(want or {}).get('min_setup_min')}); "
                    f"the charged workbook row is proven only on the book")
    if facts.get("no_gauge_decision_on_assemblies"):
        import costed_facts as _cf
        from detail_page_geometry import not_cut_from_a_blank
        guilty = [p.get("part_number") for p in _wparts.values()
                  if not_cut_from_a_blank(p) and _cf.thickness_conflict(p)]
        res.add("no_gauge_decision", "assemblies", not guilty,
                f"gauge decision on {guilty}" if guilty else "no assembly carries a gauge decision")

    # ── mass ────────────────────────────────────────────────────────────────────
    for code, kg in (facts.get("max_unit_material_mass_kg") or {}).items():
        m = _unit_mass(pes.get(_sq(code)))
        res.add("max_mass", code, m is None or m <= float(kg),
                f"material mass {m} kg (limit {kg})")
    for code, kg in (facts.get("min_unit_material_mass_kg") or {}).items():
        m = _unit_mass(pes.get(_sq(code)))
        res.add("min_mass", code, m is not None and m >= float(kg),
                f"material mass {m} kg (floor {kg})")
    _um = facts.get("unit_mass_kg") or {}
    if _um.get("expected"):
        total, counted = 0.0, 0
        for l in lines:
            if str(l.get("kind") or "").lower() == "assembly":
                continue
            m = _unit_mass(pes.get(_sq(l.get("part_number"))))
            q = _num(l.get("qty_per_unit")) or 0.0
            if m:
                total += m * q
                counted += 1
        exp, tol = float(_um["expected"]), float(_um.get("tolerance_pct") or 25)
        res.add("unit_mass", f"{exp:g} kg", abs(total - exp) <= exp * tol / 100.0,
                f"book's material mass {total:.2f} kg over {counted} line(s) "
                f"(±{tol:g}%; blanks include the cut-outs, bought items weigh nothing here)")

    # ── breaks ──────────────────────────────────────────────────────────────────
    want = facts.get("quantity_breaks")
    if want:
        if xlsx is None:
            res.add("quantity_breaks", str(want), False, "pass --xlsx <book> to check")
        else:
            got, unit = _breaks_in_book(xlsx)
            res.add("quantity_breaks", str(want), sorted(got) == sorted(int(q) for q in want),
                    f"book columns {got}; unit £ " + ", ".join(
                        f"{q}: {unit[q]:.2f}" for q in got if q in unit))
    return res


def main(argv: Optional[Iterable[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("summary", type=Path)
    ap.add_argument("facts", type=Path)
    ap.add_argument("--xlsx", type=Path, default=None)
    a = ap.parse_args(list(argv) if argv is not None else None)
    summary = json.loads(a.summary.read_text(encoding="utf-8"))
    facts = json.loads(a.facts.read_text(encoding="utf-8"))
    res = check(summary, facts, a.xlsx)
    width = max((len(r[1]) for r in res.rows), default=10)
    for check_name, item, ok, detail in res.rows:
        print(f"{'PASS' if ok else 'FAIL'}  {check_name:<20} {item:<{width}}  {detail}")
    n_fail = len(res.failed)
    print(f"\n{len(res.rows) - n_fail} held, {n_fail} failed")
    return 1 if n_fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
